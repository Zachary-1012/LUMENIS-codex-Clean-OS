"""Offline generated-document regressions; no user E2E or production claims."""
import hashlib
import json
from datetime import datetime

from lumenis.importer.parser import parse_file


def test_xlsx_physical_rows_types_hidden_sheets_and_inert_formulas(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'First'
    ws.append(['Name', 'Value', 'Flag', 'Date', 'Formula', 'Empty'])
    ws.append(['item', 0, False, datetime(2026, 10, 18), '=1+1', None])
    ws.cell(4, 1, 'after blank')
    hidden = wb.create_sheet('Hidden')
    hidden.sheet_state = 'hidden'
    hidden.append(['=HYPERLINK("https://invalid.example/", "inert")', '#DIV/0!'])
    wb.create_sheet('Empty')
    path = tmp_path / 'typed.xlsx'
    wb.save(path)
    result = parse_file(str(path))
    assert result.metadata['sheets'] == ['First', 'Hidden', 'Empty']
    assert result.metadata['parser_version'] == 'typed-source-v2'
    assert result.metadata['formula_evaluation'] == 'never'
    rows = result.source_tables[0]['rows']
    assert [r['row_index'] for r in rows] == [1, 2, 3, 4]
    assert all(len(r['values']) == 6 for r in rows)
    values = rows[1]['values']
    assert values[1]['type'] == 'integer' and values[1]['value'] == 0
    assert values[2]['type'] == 'boolean' and values[2]['value'] is False
    assert values[3]['type'] == 'datetime' and values[3]['value'] == '2026-10-18T00:00:00'
    assert values[4] == {'type': 'formula', 'value': '=1+1', 'cell_type': 'f'}
    assert values[5]['type'] == 'blank' and values[5]['value'] is None
    assert all(c['value'] is None for c in rows[2]['values'])
    assert result.source_tables[1]['sheet_state'] == 'hidden'
    assert result.source_tables[1]['rows'][0]['values'][1]['type'] == 'error'
    assert result.source_tables[2]['rows'] == []
    assert len(result.tables[0]) == 4
    json.dumps(result.source_tables, allow_nan=False)
    for locator in result.locators:
        cell = result.tables[result.metadata['sheets'].index(locator['sheet'])][locator['row']-1][locator['column']-1]
        assert result.text[locator['start']:locator['end']] == cell
    assert len([l for l in result.locators if l['sheet'] == 'First']) == 24
    repeat = parse_file(str(path))
    assert repeat.source_tables == result.source_tables


def test_docx_order_blank_whitespace_and_cell_offsets(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph('')
    doc.add_paragraph('  ')
    doc.add_paragraph('repeated')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'cell'
    table.cell(1, 1).text = 'tail'
    doc.add_paragraph('repeated')
    path = tmp_path / 'ordered.docx'
    doc.save(path)
    result = parse_file(str(path))
    assert [b['source_category'] for b in result.source_blocks] == ['docx_paragraph']*3 + ['docx_table_row']*2 + ['docx_paragraph']
    assert [b['text'] for b in result.source_blocks[:2]] == ['', '  ']
    for block in result.source_blocks:
        loc = block['locator']
        assert result.text[loc['start']:loc['end']] == block['text']
    cells = [l for l in result.locators if l['unit'] == 'word_table_cell']
    assert len(cells) == 4
    for loc in cells:
        assert result.text[loc['start']:loc['end']] == result.tables[loc['table']-1][loc['row']-1][loc['column']-1]
    assert result.source_blocks[3]['source_fields'][1]['value'] == ''


def test_pdf_order_blank_page_repeat_and_nested_offsets(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    path = tmp_path / 'pages.pdf'
    canvas = Canvas(str(path))
    canvas.drawString(30, 750, 'repeat')
    canvas.drawString(30, 720, 'repeat')
    canvas.showPage()
    canvas.showPage()
    canvas.drawString(30, 750, 'final')
    canvas.save()
    result = parse_file(str(path))
    assert len(result.source_blocks) == 3
    assert result.source_blocks[1]['text'] == ''
    for index, block in enumerate(result.source_blocks, 1):
        loc = block['locator']
        assert loc['page_index'] == index
        assert result.text[loc['start']:loc['end']] == block['text']
        for nested in loc['text_blocks']:
            assert nested['offset_basis'] == 'page_text'
            if nested['start'] is not None:
                text = block['text'][nested['start']:nested['end']]
                assert hashlib.sha256(text.encode()).hexdigest() == nested['text_sha256']
    assert len(result.source_blocks[0]['locator']['text_blocks']) == 2


def test_current_csv_encoding_dialect_and_formula_retained(tmp_path):
    path = tmp_path / 'semicolon.csv'
    path.write_bytes('名称;值\n测试;0\n公式;=1+1\n'.encode('gb18030'))
    result = parse_file(str(path))
    assert result.metadata['encoding'] == 'gb18030'
    assert result.metadata['csv_delimiter'] == ';'
    assert result.tables == [[['名称', '值'], ['测试', '0'], ['公式', '=1+1']]]
    assert result.source_tables == [] and result.source_blocks == []


def test_xlsx_typed_time_duration_number_and_array_formula(tmp_path):
    from datetime import time, timedelta
    from openpyxl import Workbook
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    ws = wb.active
    ws.append([1.25, time(12, 34, 56), timedelta(days=2, seconds=3), '0'])
    ws['E1'] = ArrayFormula(ref='E1:E2', text='=SUM(A1:A2)')
    path = tmp_path / 'more-types.xlsx'
    wb.save(path)
    result = parse_file(str(path))
    values = result.source_tables[0]['rows'][0]['values']
    assert [(v['type'], v['value']) for v in values[:4]] == [
        ('number', 1.25), ('time', '12:34:56'), ('duration', '2 days, 0:00:03'), ('string', '0')]
    assert values[4]['type'] == 'formula' and values[4]['value'] == '=SUM(A1:A2)'
    assert result.tables[0][0][4] == '=SUM(A1:A2)'


def test_sparse_workbook_cannot_expand_unbounded_blank_typed_cells(tmp_path):
    import pytest
    from openpyxl import Workbook
    book = Workbook()
    book.active.cell(1_000_000, 100, "far corner")
    path = tmp_path / 'sparse.xlsx'
    book.save(path)
    with pytest.raises(ValueError, match='DOCUMENT_EXCEL_CELL_LIMIT'):
        parse_file(str(path))
