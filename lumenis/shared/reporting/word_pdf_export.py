"""
Word / PDF 报告导出模块
========================
将标准化报告数据导出为Word(.docx)和PDF文件
支持 CTH / 序境·商业运营与 BMP / 序境·品牌市场模板

依赖：python-docx（Word）、reportlab与已有嵌入字体（PDF）
"""
import hashlib
import json
import os
import re
import uuid
from contextlib import contextmanager
from datetime import datetime
from html import escape
from threading import RLock

from .standard_taxonomy import RiskLevel, Sentiment


_PDF_EXPORT_LOCK = RLock()
_PPTX_MAX_SLIDES = 512
_PPTX_PARAGRAPHS_PER_SLIDE = 24
_PPTX_CHARS_PER_PARAGRAPH = 64
_PPTX_MAX_SOURCE_CHARS = 1_000_000
_PPTX_LINE_BREAKS = ("\n", "\r", "\v", "\f", "\x1c", "\x1d", "\x1e", "\x85", "\u2028", "\u2029")
_PPTX_SECTION_LABELS = ("Title", "Report", "Lineage", "Sources", "Snapshot", "JSON")


def _display_value(value):
    return "—" if value is None else str(value)


def _saved_pptx_pages(sections: list[tuple[str, str]]) -> list[tuple[str, int, list[str]]]:
    total_chars = sum(len(text) for _, text in sections)
    if total_chars > _PPTX_MAX_SOURCE_CHARS:
        raise ValueError("WORK_EXPORT_PPTX_TOO_LARGE")

    pages = []
    for label, text in sections:
        lines = text.splitlines()
        if not lines:
            lines = [""]
        if text.endswith(_PPTX_LINE_BREAKS):
            lines.append("")

        paragraphs = []
        for line in lines:
            if not line:
                paragraphs.append("")
                continue
            paragraphs.extend(
                line[start:start + _PPTX_CHARS_PER_PARAGRAPH]
                for start in range(0, len(line), _PPTX_CHARS_PER_PARAGRAPH)
            )
        if not paragraphs:
            paragraphs = [""]

        for page_number, start in enumerate(range(0, len(paragraphs), _PPTX_PARAGRAPHS_PER_SLIDE), 1):
            pages.append((label, page_number, paragraphs[start:start + _PPTX_PARAGRAPHS_PER_SLIDE]))
            if len(pages) > _PPTX_MAX_SLIDES:
                raise ValueError("WORK_EXPORT_PPTX_TOO_LARGE")
    return pages


@contextmanager
def pdf_text_context():
    """Embed the existing product fonts and serialize ReportLab font subsets."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from lumenis.shared.database import resolve_path

    # ReportLab's TTFont objects must not build different documents concurrently.
    with _PDF_EXPORT_LOCK:
        fonts = []
        for name, filename in (("LumenisNotoSC", "NotoSansSC.ttf"),
                               ("LumenisNotoKR", "NotoSansKR.ttf")):
            path = resolve_path("app/assets/fonts/" + filename)
            if not os.path.isfile(path):
                raise RuntimeError("WORK_EXPORT_FONT_UNAVAILABLE")
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, path))
            fonts.append(pdfmetrics.getFont(name))

        def line_markup(text):
            runs = []
            for char in str(text).replace("\t", "    "):
                font = next((font for font in fonts
                             if font.face.charToGlyph.get(ord(char), 0)), None)
                if font is None:
                    raise RuntimeError("WORK_EXPORT_GLYPH_UNAVAILABLE")
                if runs and runs[-1][0] == font.fontName:
                    runs[-1][1].append(char)
                else:
                    runs.append((font.fontName, [char]))
            return "".join(f'<font name="{name}">{escape("".join(chars))}</font>'
                           for name, chars in runs) or "&#160;"

        def markup(text):
            return "<br/>".join(line_markup(line) for line in str(text).splitlines() or [""])

        yield fonts[0].fontName, markup


def delivery_markdown(document: dict, sources: list) -> str:
    text = document["markdown"]
    if not sources:
        return text
    language = "ko" if re.search(r"[가-힣]", document["title"]) else (
        "zh" if re.search(r"[\u3400-\u9fff]", document["title"]) else "en")
    labels = {"zh": ("来源", "资料"), "en": ("Sources", "Material"), "ko": ("출처", "자료")}[language]
    lines = [text, "", "## " + labels[0], ""]
    for index, source in enumerate(sources, 1):
        title = source.get("title") or f"{labels[1]} {index}"
        lines.append(f"{index}. {title}")
        for field in ("source_url", "quote"):
            if isinstance(source.get(field), str) and source[field]:
                lines.append(source[field])
        quoted = {source.get("quote")}
        for citation in document.get("citations", []):
            quote = citation.get("quote")
            if (source.get("evidence_id") is not None
                    and citation.get("evidence_id") == source["evidence_id"]
                    and isinstance(quote, str) and quote and quote not in quoted):
                lines.append(quote)
                quoted.add(quote)
        lines.append("")
    return "\n".join(lines)


def _delivery_blocks(markdown: str):
    lines = markdown.splitlines()
    index = 0
    fenced = False
    while index < len(lines):
        line = lines[index]
        if line.lstrip().startswith("```"):
            fenced = not fenced
            yield "paragraph", line
            index += 1
            continue
        heading = None if fenced else re.fullmatch(r"(#{1,6})\s+(.+)", line)
        if heading:
            yield "heading", (len(heading[1]), heading[2])
        elif (not fenced and index + 1 < len(lines) and "|" in line
              and re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*", lines[index + 1])
              and "\\|" not in line):
            header = [value.strip() for value in line.strip().strip("|").split("|")]
            rows = [header]
            index += 2
            while index < len(lines) and "|" in lines[index] and "\\|" not in lines[index]:
                row = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if len(row) != len(header):
                    break
                rows.append(row)
                index += 1
            yield "table", rows
            continue
        else:
            yield "paragraph", line
        index += 1


def _export_delivery(document: dict, *, sources: list, format: str) -> bytes:
    import io

    output = io.BytesIO()
    markdown = delivery_markdown(document, sources)
    blocks = list(_delivery_blocks(markdown))
    if format == "docx":
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        from lxml import etree

        doc = Document()
        doc.core_properties.title = document["title"]
        doc.core_properties.author = ""
        for style in doc.styles:
            if hasattr(style, "font"):
                style.font.name = "Arial"
                fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
                fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    fonts.attrib.pop(qn("w:" + attribute), None)
        for fonts in doc.styles.element.iter(qn("w:rFonts")):
            for name in ("ascii", "hAnsi", "cs"):
                fonts.set(qn("w:" + name), "Arial")
            fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + attribute), None)
        doc.styles["Normal"].font.size = Pt(11)
        font_table = next(part for part in doc.part.package.iter_parts()
                          if str(part.partname) == "/word/fontTable.xml")
        fonts = OxmlElement("w:fonts")
        for name in ("Arial", "Microsoft YaHei"):
            font = OxmlElement("w:font")
            font.set(qn("w:name"), name)
            fonts.append(font)
        font_table._blob = etree.tostring(fonts, xml_declaration=True, encoding="UTF-8", standalone=True)
        doc.add_heading(document["title"], level=0)
        for kind, value in blocks:
            if kind == "heading":
                doc.add_heading(value[1], level=value[0])
            elif kind == "table":
                table = doc.add_table(rows=1, cols=len(value[0]))
                table.style = "Light Shading Accent 1"
                for cell, text in zip(table.rows[0].cells, value[0]):
                    cell.text = text
                for row in value[1:]:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = text
            else:
                doc.add_paragraph(value)
        doc.save(output)
    elif format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import (
            LongTable,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            TableStyle,
        )

        with pdf_text_context() as (font_name, markup):
            normal = ParagraphStyle("Delivery", fontName=font_name, fontSize=10, leading=15,
                                    wordWrap="CJK", spaceAfter=6)
            title = ParagraphStyle("DeliveryTitle", parent=normal, fontSize=22, leading=30, spaceAfter=18)
            heading = ParagraphStyle("DeliveryHeading", parent=normal, fontSize=14, leading=20,
                                     spaceBefore=10, spaceAfter=8, keepWithNext=True)
            story = [Paragraph(markup(document["title"]), title)]
            for kind, value in blocks:
                if kind == "heading":
                    story.append(Paragraph(markup(value[1]), heading))
                elif kind == "table":
                    cells = [[Paragraph(markup(text), normal) for text in row] for row in value]
                    table = LongTable(cells, colWidths=[(A4[0] - 96) / len(value[0])] * len(value[0]), repeatRows=1)
                    table.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECE8")),
                        ("LINEBELOW", (0, 0), (-1, 0), .6, colors.HexColor("#728477")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ]))
                    story.extend([table, Spacer(1, 10)])
                elif value:
                    story.append(Paragraph(markup(value), normal))
                else:
                    story.append(Spacer(1, 6))
            SimpleDocTemplate(output, pagesize=A4, leftMargin=48, rightMargin=48,
                              title=document["title"], author="").build(story)
    elif format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        book = Workbook()
        sheet = book.active
        sheet.title = "Report"
        sheet.column_dimensions["A"].width = 90
        table_index = 0
        for kind, value in [("heading", (1, document["title"])), *blocks]:
            if kind == "table":
                table_index += 1
                table = book.create_sheet(f"Table {table_index}")
                for row in value:
                    if any(len(text) > 32767 for text in row):
                        raise ValueError("WORK_EXPORT_CELL_TOO_LARGE")
                    table.append(row)
                    for cell in table[table.max_row]:
                        cell.data_type = "s"
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                table.freeze_panes = "A2"
                table.auto_filter.ref = table.dimensions
                for cell in table[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="E6ECE8")
                    table.column_dimensions[cell.column_letter].width = 28
                text = table.title
            else:
                text = value[1] if kind == "heading" else value
            for start in range(0, max(1, len(text)), 30000):
                sheet.append([text[start:start + 30000]])
                cell = sheet.cell(sheet.max_row, 1)
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(size=14 if kind == "heading" else 11, bold=kind == "heading")
                if kind == "table":
                    cell.hyperlink = f"#'Table {table_index}'!A1"
        book.save(output)
    elif format == "pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt

        sections = []
        title, body = document["title"], []
        for kind, value in blocks:
            if kind == "heading":
                if body or title != document["title"]:
                    sections.append((title, "\n".join(body)))
                title, body = value[1], []
            elif kind == "table":
                body.extend(" | ".join(row) for row in value)
            else:
                body.append(value)
        sections.append((title, "\n".join(body)))
        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
        cover = presentation.slides.add_slide(presentation.slide_layouts[0])
        cover.shapes.title.text = document["title"]
        for title, page, lines in _saved_pptx_pages(sections):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            header = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12), Inches(.6))
            header.text_frame.text = title if page == 1 else f"{title} · {page}"
            header.text_frame.paragraphs[0].font.size = Pt(24)
            box = slide.shapes.add_textbox(Inches(.65), Inches(1.25), Inches(12), Inches(5.8))
            box.text_frame.word_wrap = True
            for index, text in enumerate(lines):
                paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
                paragraph.text = text
                paragraph.font.size = Pt(13)
                paragraph.space_after = Pt(1)
        presentation.save(output)
    else:
        raise ValueError("WORK_EXPORT_FORMAT_INVALID")
    return output.getvalue()


class ReportExporter:
    """报告导出器"""

    def __init__(self, output_dir: str = "data/exports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    @staticmethod
    def export_saved_document(document: dict, *, sources: list, format: str,
                              purpose: str = "audit_snapshot") -> bytes:
        """Export the complete saved snapshot in memory, including source lineage.

        Uses the existing Reporting exporter dependencies. A missing converter
        raises an error; a requested PDF never returns a Word file. No output
        directory or database is created by this method.
        """
        import io
        import json
        if purpose == "delivery":
            return _export_delivery(document, sources=sources, format=format)
        if purpose != "audit_snapshot":
            raise ValueError("WORK_EXPORT_PURPOSE_INVALID")
        snapshot = json.dumps(document, ensure_ascii=False, sort_keys=True, indent=2)
        sections = [document["title"], document["markdown"],
                    "Sources / original scope / rights", json.dumps(sources, ensure_ascii=False, indent=2),
                    "Complete saved snapshot", snapshot]
        output = io.BytesIO()
        if format == "docx":
            from docx import Document
            doc = Document()
            for section in sections:
                for line in section.splitlines():
                    doc.add_paragraph(line)
            doc.save(output)
        elif format == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            with pdf_text_context() as (font_name, markup):
                style = ParagraphStyle("SavedWork", fontName=font_name, fontSize=9,
                                       leading=13, wordWrap="CJK")
                story = []
                for section in sections:
                    for line in section.splitlines():
                        story.append(Paragraph(markup(line), style))
                    story.append(Spacer(1, 12))
                SimpleDocTemplate(output, pagesize=A4).build(story)
        elif format == "xlsx":
            from openpyxl import Workbook
            book = Workbook()
            sheet = book.active
            sheet.title = "Saved snapshot"
            for label, section in zip(("Title", "Report", "Lineage", "Sources", "Snapshot", "JSON"), sections):
                # Excel cells truncate at 32767 characters. Preserve every
                # character in ordered chunks, and force text to avoid formulas.
                for index, start in enumerate(range(0, len(section), 30000), 1):
                    sheet.append([label, index, section[start:start + 30000]])
                    sheet.cell(sheet.max_row, 3).data_type = "s"
            book.save(output)
        elif format == "pptx":
            from pptx import Presentation
            from pptx.util import Inches, Pt

            pages = _saved_pptx_pages(list(zip(_PPTX_SECTION_LABELS, sections)))
            presentation = Presentation()
            presentation.slide_width = Inches(13.333333)
            presentation.slide_height = Inches(7.5)
            blank_layout = presentation.slide_layouts[6]
            for label, page_number, paragraphs in pages:
                slide = presentation.slides.add_slide(blank_layout)
                continuation = f" (continued {page_number})" if page_number > 1 else ""
                heading = slide.shapes.add_textbox(Inches(0.55), Inches(0.30), Inches(12.2), Inches(0.55))
                heading.text_frame.text = f"Saved document | {label}{continuation}"
                title_paragraph = heading.text_frame.paragraphs[0]
                title_paragraph.font.size = Pt(16)
                title_paragraph.font.bold = True

                body = slide.shapes.add_textbox(Inches(0.62), Inches(1.02), Inches(12.05), Inches(6.0))
                body.text_frame.word_wrap = True
                body.text_frame.margin_left = Inches(0.04)
                body.text_frame.margin_right = Inches(0.04)
                body.text_frame.margin_top = Inches(0.02)
                body.text_frame.margin_bottom = Inches(0.02)
                for index, value in enumerate(paragraphs):
                    paragraph = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
                    paragraph.text = value
                    paragraph.font.size = Pt(11)
                    paragraph.space_after = Pt(0)
            presentation.save(output)
        else:
            raise ValueError("unsupported saved document export format")
        return output.getvalue()

    def export_word(self, report: dict, filename: str = None) -> str:
        """导出为Word文档"""
        try:
            from docx import Document
            from docx.shared import Pt, Inches, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return ""

        doc = Document()

        # 设置默认字体
        style = doc.styles["Normal"]
        style.font.name = "微软雅黑"
        style.font.size = Pt(10.5)

        # 标题
        brand_name = report.get("brand_name", "品牌")
        period_label = report.get("period_label", "报告")
        title = doc.add_heading(f"{brand_name} {period_label}", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 报告信息
        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info.add_run(f"报告周期：{report.get('start_date', '')} 至 {report.get('end_date', '')}\n").font.size = Pt(9)
        info.add_run(f"生成时间：{report.get('generated_at', '')}\n").font.size = Pt(9)
        info.add_run(f"报告ID：{report.get('report_id', '')}").font.size = Pt(9)

        doc.add_paragraph()  # 空行

        # 1. 执行摘要
        doc.add_heading("一、执行摘要", level=1)
        p = doc.add_paragraph()
        p.add_run(report.get("executive_summary", "")).bold = True

        # 2. 关键指标
        doc.add_heading("二、关键指标", level=1)
        metrics = report.get("key_metrics", [])
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "指标"
        hdr[1].text = "数值"
        hdr[2].text = "单位"
        for m in metrics:
            row = table.add_row().cells
            row[0].text = m.get("name", "")
            row[1].text = _display_value(m.get("value"))
            row[2].text = m.get("unit", "")

        # 3. 风险事件
        doc.add_heading("三、风险事件", level=1)
        risk_events = report.get("risk_events", [])
        if not risk_events:
            doc.add_paragraph("风险事件数据不可用。" if report.get("risk_events") is None else "本周期无风险事件。")
        for evt in risk_events:
            p = doc.add_paragraph()
            risk_label = evt.get("risk_label", evt.get("risk_level", ""))
            run = p.add_run(f"【{risk_label}】{evt.get('title', '')}")
            run.bold = True
            if evt.get("risk_level") in ("critical", "high"):
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            doc.add_paragraph(f"风险指数：{_display_value(evt.get('risk_score'))} | 提及数：{_display_value(evt.get('mention_count'))} | 来源数：{_display_value(evt.get('source_count'))}")
            doc.add_paragraph(f"预警：{evt.get('alert_message', '')}")
            doc.add_paragraph()

        # 4. 情感分析
        doc.add_heading("四、情感分析", level=1)
        sentiment = report.get("sentiment_analysis", {})
        dist = sentiment.get("distribution", {})
        doc.add_paragraph(
            f"整体情感：{sentiment.get('overall_label', '')} | "
            f"正面{_display_value(dist.get('positive'))}条（{_display_value(dist.get('positive_ratio'))}%） | "
            f"负面{_display_value(dist.get('negative'))}条（{_display_value(dist.get('negative_ratio'))}%） | "
            f"中性{_display_value(dist.get('neutral'))}条"
        )

        # 5. 传播分析
        doc.add_heading("五、传播分析", level=1)
        prop = report.get("propagation_analysis", {})
        doc.add_paragraph(f"互动水平：{prop.get('engagement_level', '')} | 总互动量：{_display_value(prop.get('total_engagement'))}")
        top_notes = prop.get("top_notes", [])
        if top_notes:
            doc.add_paragraph("热门笔记Top5：")
            for i, n in enumerate(top_notes, 1):
                doc.add_paragraph(
                    f"{i}. {n.get('title', '')}（{n.get('author', '')}）"
                    f" - 互动{_display_value(n.get('engagement'))} - {n.get('sentiment', '')}",
                    style="List Bullet"
                )

        # 6. 关键洞察
        doc.add_heading("六、关键洞察", level=1)
        for ins in report.get("key_insights", []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{ins.get('type', '')}] {ins.get('insight', '')}")
            if ins.get("metric") is not None:
                p.add_run(f"（{ins['metric']}）").font.size = Pt(9)

        # 7. 行动建议
        doc.add_heading("七、行动建议", level=1)
        for act in report.get("action_items", []):
            p = doc.add_paragraph()
            priority = act.get("priority", "")
            run = p.add_run(f"【{priority}】")
            if priority == "紧急":
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                run.bold = True
            p.add_run(f"{act.get('action', '')}")
            p_owner = doc.add_paragraph(f"负责：{act.get('owner', '')} | 时限：{act.get('deadline', '')}")
            p_owner.runs[0].font.size = Pt(9)

        # 8. 数据质量
        doc.add_heading("八、数据质量", level=1)
        dq = report.get("data_quality", {})
        doc.add_paragraph(
            f"采集总数：{_display_value(dq.get('total_collected'))}条 | "
            f"质量评分：{_display_value(dq.get('quality_score'))}/100 | "
            f"采集完整：{('是' if dq['collection_complete'] else '否') if dq.get('collection_complete') is not None else '—'}"
        )
        for w in dq.get("warnings", []):
            if w:
                doc.add_paragraph(f"⚠️ {w}", style="List Bullet")

        # 9. CTH专属（如果有）
        if (report.get("cth_specific", {}).get("reference_scope") is True):
            doc.add_heading("九、参考主体专项分析", level=1)
            cth = report["cth_specific"]
            tone = cth.get("tone_evaluation", {})
            doc.add_paragraph(
                f"调性匹配：{tone.get('overall', '')} | "
                f"温暖{_display_value(tone.get('warm_ratio'))}% | 精致{_display_value(tone.get('refined_ratio'))}% | 仪式感{_display_value(tone.get('ritual_ratio'))}%"
            )
            nodes = cth.get("active_activity_nodes", [])
            if nodes:
                doc.add_paragraph("当前活动节点：")
                for n in nodes:
                    doc.add_paragraph(f"{n.get('name', '')}（{n.get('month', '')}）- {n.get('tone', '')}", style="List Bullet")

        # 页脚
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(
            f"本报告由 {'LUMENIS Commerce Operations (CTH)' if report.get('brand_mode') == 'cth' else 'LUMENIS Brand & Market (BMP)'} 自动生成 | "
            f"数据来源可追溯 | {datetime.now().strftime('%Y-%m-%d')}"
        ).font.size = Pt(8)

        # 保存
        if not filename:
            brand = brand_name.replace(" ", "_")
            period = report.get("period", "report")
            date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{brand}_{period}_{date_str}.docx"

        filepath = os.path.join(self.output_dir, filename)
        doc.save(filepath)
        return filepath

    def export_pdf(self, report: dict, filename: str = None) -> str:
        """
        导出为PDF
        使用已有ReportLab与嵌入字体生成完整报告。
        """
        # 先生成Word
        word_path = self.export_word(report, filename.replace(".pdf", ".docx") if filename else None)
        if not word_path:
            return ""

        # All PDF paths share the verified glyph checks and embedded fonts.
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table
            from reportlab.lib.units import cm

            if not filename:
                brand = report.get("brand_name", "report").replace(" ", "_")
                period = report.get("period", "report")
                date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"{brand}_{period}_{date_str}.pdf"

            pdf_path = os.path.join(self.output_dir, filename)
            with pdf_text_context() as (font_name, markup):
                doc = SimpleDocTemplate(pdf_path, pagesize=A4)
                styles = getSampleStyleSheet()
                for style in styles.byName.values():
                    style.fontName = font_name
                    style.wordWrap = "CJK"
                story = []

                story.append(Paragraph(markup(f"{report.get('brand_name', '')} {report.get('period_label', '')}"), styles["Title"]))
                story.append(Spacer(1, 0.5 * cm))
                story.append(Paragraph(markup(f"周期：{report.get('start_date', '')} 至 {report.get('end_date', '')}"), styles["Normal"]))
                story.append(Paragraph(markup(f"生成：{report.get('generated_at', '')}"), styles["Normal"]))
                story.append(Paragraph(markup(f"报告ID：{report.get('report_id', '')}"), styles["Normal"]))
                story.append(Spacer(1, 0.5 * cm))

                story.append(Paragraph(markup("执行摘要"), styles["Heading2"]))
                story.append(Paragraph(markup(report.get("executive_summary", "")), styles["Normal"]))
                story.append(Spacer(1, 0.3 * cm))

                story.append(Paragraph(markup("关键指标"), styles["Heading2"]))
                metric_data = [[Paragraph(markup(value), styles["Normal"]) for value in ("指标", "数值", "单位")]]
                for metric in report.get("key_metrics", []):
                    display_value = "—" if metric.get("value") is None else str(metric["value"])
                    metric_data.append([Paragraph(markup(cell), styles["Normal"]) for cell in
                                        (metric.get("name", ""), display_value, metric.get("unit", ""))])
                story.append(Table(metric_data))
                story.append(Spacer(1, 0.3 * cm))

                story.append(Paragraph(markup("风险事件"), styles["Heading2"]))
                if not report.get("risk_events"):
                    story.append(Paragraph(markup(
                        "风险事件数据不可用。" if report.get("risk_events") is None else "本周期无风险事件。"), styles["Normal"]))
                for evt in report.get("risk_events", []):
                    score = "—" if evt.get("risk_score") is None else str(evt["risk_score"])
                    story.append(Paragraph(markup(
                        f"【{evt.get('risk_label', evt.get('risk_level', ''))}】{evt.get('title', '')} - 风险指数{score}"), styles["Normal"]))
                    story.append(Paragraph(markup(
                        f"提及数：{_display_value(evt.get('mention_count'))} | 来源数：{_display_value(evt.get('source_count'))}"), styles["Normal"]))
                    story.append(Paragraph(markup(f"预警：{evt.get('alert_message', '')}"), styles["Normal"]))
                story.append(Spacer(1, 0.3 * cm))

                story.append(Paragraph(markup("情感分析"), styles["Heading2"]))
                sentiment = report.get("sentiment_analysis", {})
                dist = sentiment.get("distribution", {})
                story.append(Paragraph(markup(
                    f"整体情感：{sentiment.get('overall_label', '')} | "
                    f"正面{_display_value(dist.get('positive'))}条（{_display_value(dist.get('positive_ratio'))}%） | "
                    f"负面{_display_value(dist.get('negative'))}条（{_display_value(dist.get('negative_ratio'))}%） | "
                    f"中性{_display_value(dist.get('neutral'))}条"), styles["Normal"]))

                story.append(Paragraph(markup("传播分析"), styles["Heading2"]))
                prop = report.get("propagation_analysis", {})
                story.append(Paragraph(markup(
                    f"互动水平：{prop.get('engagement_level', '')} | 总互动量：{_display_value(prop.get('total_engagement'))}"), styles["Normal"]))
                for index, note in enumerate(prop.get("top_notes", []), 1):
                    story.append(Paragraph(markup(
                        f"{index}. {note.get('title', '')}（{note.get('author', '')}） - "
                        f"互动{_display_value(note.get('engagement'))} - {note.get('sentiment', '')}"), styles["Normal"]))

                story.append(Paragraph(markup("关键洞察"), styles["Heading2"]))
                for insight in report.get("key_insights", []):
                    metric = f"（{insight['metric']}）" if insight.get("metric") is not None else ""
                    story.append(Paragraph(markup(
                        f"[{insight.get('type', '')}] {insight.get('insight', '')}{metric}"), styles["Normal"]))

                story.append(Paragraph(markup("行动建议"), styles["Heading2"]))
                for act in report.get("action_items", []):
                    story.append(Paragraph(markup(
                        f"【{act.get('priority', '')}】{act.get('action', '')}（{act.get('owner', '')}，{act.get('deadline', '')}）"),
                        styles["Normal"]))

                story.append(Paragraph(markup("数据质量"), styles["Heading2"]))
                quality = report.get("data_quality", {})
                complete = ('是' if quality['collection_complete'] else '否') if quality.get('collection_complete') is not None else '—'
                story.append(Paragraph(markup(
                    f"采集总数：{_display_value(quality.get('total_collected'))}条 | "
                    f"质量评分：{_display_value(quality.get('quality_score'))}/100 | 采集完整：{complete}"), styles["Normal"]))
                for warning in quality.get("warnings", []):
                    if warning:
                        story.append(Paragraph(markup(f"警告：{warning}"), styles["Normal"]))

                if report.get("cth_specific", {}).get("reference_scope") is True:
                    story.append(Paragraph(markup("参考主体专项分析"), styles["Heading2"]))
                    cth = report["cth_specific"]
                    tone = cth.get("tone_evaluation", {})
                    story.append(Paragraph(markup(
                        f"调性匹配：{tone.get('overall', '')} | 温暖{_display_value(tone.get('warm_ratio'))}% | "
                        f"精致{_display_value(tone.get('refined_ratio'))}% | 仪式感{_display_value(tone.get('ritual_ratio'))}%"), styles["Normal"]))
                    for node in cth.get("active_activity_nodes", []):
                        story.append(Paragraph(markup(
                            f"{node.get('name', '')}（{node.get('month', '')}）- {node.get('tone', '')}"), styles["Normal"]))

                doc.build(story)
            return pdf_path
        except Exception as e:
            print(f"PDF导出失败: {e}")
            return ""

    def generate_and_export(self, brand_mode: str = "cth", period: str = "daily",
                            start_date: str | None = None, end_date: str | None = None,
                            export_formats: list | None = None, brand_config: dict | None = None,
                            *, actor: dict | None = None, scope: dict | None = None,
                            scope_entry: dict | None = None) -> dict:
        """
        一键生成并导出报告

        Args:
            brand_mode: cth/zbmp
            period: daily/weekly/monthly/yearly
            export_formats: ["word", "pdf"]，默认两者都导出
            brand_config: BMP 模式品牌配置

        Returns:
            {"report": 报告数据, "word_path": ..., "pdf_path": ...}
        """
        if export_formats is None:
            export_formats = ["word", "pdf"]
        export_formats = tuple(dict.fromkeys(str(fmt).lower() for fmt in export_formats))
        if not export_formats or any(fmt not in {"word", "pdf"} for fmt in export_formats):
            raise ValueError("REPORT_EXPORT_FORMAT_INVALID")
        if actor is None or scope is None:
            raise PermissionError("REPORT_SCOPE_REQUIRED")

        import database
        from lumenis.shared.reporting.export_api import (
            _build_report,
            _verify_report_source_use,
        )
        from lumenis.shared.reporting.saved_work import _check_effect
        from lumenis.system.security.permission import OperationType

        exact_scope = database._normalize_event_scope(scope, allow_global=True)
        request_binding = {
            "scope": exact_scope, "period": period, "start": start_date,
            "end": end_date, "formats": list(export_formats),
        }
        request_resource = "report-export-request:v1:" + hashlib.sha256(
            json.dumps(request_binding, ensure_ascii=False, sort_keys=True,
                       separators=(",", ":")).encode("utf-8")
        ).hexdigest()
        _check_effect(actor, exact_scope, OperationType.DATA_EXPORT, request_resource)

        kind, start, end, report, _source_count, notes, events = _build_report(
            period, start_date, end_date, actor=actor, scope=exact_scope,
            scope_entry=scope_entry)
        result = {"report": report, "word_path": "", "pdf_path": ""}
        export_id = uuid.uuid4().hex
        word_target = os.path.join(self.output_dir, f"lumenis_{export_id}.docx")
        pdf_target = os.path.join(self.output_dir, f"lumenis_{export_id}_pdf.pdf")
        pdf_word_target = os.path.splitext(pdf_target)[0] + ".docx"
        invocation_paths = []
        if "word" in export_formats:
            invocation_paths.append(word_target)
        if "pdf" in export_formats:
            invocation_paths.extend((pdf_target, pdf_word_target))
        invocation_paths = list(dict.fromkeys(os.path.abspath(path) for path in invocation_paths))

        def remove_invocation_files(paths):
            for path in paths:
                try:
                    if os.path.isfile(path):
                        os.remove(path)
                except OSError:
                    pass

        try:
            if "word" in export_formats:
                word_path = self.export_word(report, os.path.basename(word_target))
                result["word_path"] = word_path
                if not word_path:
                    remove_invocation_files((word_target,))
            if "pdf" in export_formats:
                pdf_path = self.export_pdf(report, os.path.basename(pdf_target))
                result["pdf_path"] = pdf_path
                if not pdf_path:
                    remove_invocation_files((pdf_target,))
                if (os.path.isfile(pdf_word_target)
                        and os.path.abspath(pdf_word_target) != os.path.abspath(result["word_path"] or "")):
                    remove_invocation_files((pdf_word_target,))

            for fmt in export_formats:
                path = result.get(fmt + "_path")
                if not path:
                    continue
                with open(path, "rb") as exported:
                    digest = hashlib.sha256(exported.read()).hexdigest()
                resource = "report-export:v1:" + hashlib.sha256(
                    json.dumps({"report_id": report.get("report_id"), "scope": exact_scope,
                                "period": kind, "start": start, "end": end,
                                "sha256": digest, "format": fmt},
                               ensure_ascii=False, sort_keys=True,
                               separators=(",", ":")).encode("utf-8")
                ).hexdigest()
                _verify_report_source_use(actor, exact_scope, notes, events)
                _check_effect(actor, exact_scope, OperationType.DATA_EXPORT, resource)
            return result
        except BaseException:
            remove_invocation_files(invocation_paths)
            raise
