"""Actual domain/Word/Artifact owners in disposable DBs; not live acceptance."""
import copy
import io
import sqlite3

import pytest
from docx import Document

from tests.test_source_brief_job import job_db
from tests.test_source_brief_job_routes import client, GLOBAL, URL
from tests.test_internal_pitch_workflow import arguments, manuscript_arguments

BUSINESS = "/api/experience/business-works"
PITCH = "/api/experience/creation/internal-pitches"


@pytest.fixture(params=["CTH", "BMP"])
def owners(client, request):
    import ai_task_engine
    import lumenis_app_api
    import shared_artifacts
    import subject_registry
    from canonical_data import CanonicalRecord, persist_canonical
    from lumenis.shared.content_creation import internal_pitch as pitch
    from lumenis.shared.evidence import store
    from lumenis.system.security import permission

    api, app, db, entry, gate = client
    app.register_blueprint(lumenis_app_api.domain_api_bp)
    gate.grant_permission(permission.PermissionLevel.COMMIT)
    pitch.engine.cache_clear()
    with sqlite3.connect(db) as con:
        con.executescript(ai_task_engine._TASK_ENGINE_SCHEMA)
    shared_artifacts.ensure_artifact_table()
    if request.param == "CTH":
        headers = {**GLOBAL, "X-Lumenis-Context-Mode": "FOCUS", "X-Subject-ID": entry["subject_id"]}
        scope = {"domain_id": "cth", "subject_id": entry["subject_id"], "brand_focus_id": None, "global_scope": False}
        source_id = "brief-job-source"
        kind = "venue_operations"
    else:
        entry, _ = subject_registry.add_subject({"domain_id": "BMP", "subject_type": "brand",
                                               "name_en": "Owned route test brand", "name_zh": "隔离品牌输入"})
        headers = {"X-Domain-Id": "BMP", "X-Product-Id": "BMP", "X-Lumenis-Context-Mode": "FOCUS",
                   "X-Brand-Focus-ID": entry["brand_focus_id"]}
        scope = {"domain_id": "zbmp", "subject_id": None, "brand_focus_id": entry["brand_focus_id"], "global_scope": False}
        source_id = "owned-route-brand-source"
        kind = "brand_market"
        with sqlite3.connect(db) as con:
            con.row_factory = sqlite3.Row
            evidence = store.record_evidence(connection=con, source="isolated_test", project="LUMENIS",
                product="ZBMP", domain_id="zbmp", brand_focus_id=entry["brand_focus_id"], global_scope=False,
                feature_id="OWNED_ROUTE_TEST", result="PASS", status="verified", inputs={"source_ids": [source_id]},
                outputs={"synthetic_engineering_input": True})
            persist_canonical(con, CanonicalRecord(canonical_id="canonical-" + source_id, source_id=source_id,
                product_id="ZBMP", domain_id="zbmp", brand_focus_id=entry["brand_focus_id"],
                scope_status="VERIFIED", scope_evidence_ref=str(evidence), source_type="web_mention",
                source_platform="web", title="Attributed engineering statement", content="An isolated attributed statement.",
                url="https://example.org/owned-route", captured_at="2026-10-03T00:00:00Z"))
            con.commit()
    try:
        yield dict(api=api, app=app, db=db, gate=gate, headers=headers, scope=scope,
                   source_id=source_id, kind=kind, pitch=pitch)
    finally:
        pitch.engine.cache_clear()


def prepare_business(ws):
    from tests.test_growth_business_fixture import prepare_owner_source
    prepare_owner_source(ws)
    report = ws["api"].post(URL + "/generate", headers=ws["headers"], json={})
    assert report.status_code == 200, report.get_json()
    source_job = report.get_json()["result"]["job_id"]
    options = [{"option_id": label, "proposal": "Review a bounded " + label + " option.",
                "rationale": "Compare against the attributed material, without causal claims.",
                "source_ids": [ws["source_id"]], "assumptions": ["Needs operator review."],
                "unknowns": ["Business effect has not been observed."], "check": "Inspect the completed internal document."}
               for label in ("A", "B")]
    response = ws["api"].post(BUSINESS, headers=ws["headers"], json={"submission_id": "prepare-owner-route-01",
        "source_job_id": source_job, "kind": ws["kind"], "question": "Which reversible action should be reviewed?",
        "options": options})
    assert response.status_code == 200, response.get_json()
    result = response.get_json()["result"]
    assert result["action_service"]["state"] == "BOUND_READER"
    return result


def advance(ws, current, operation, payload, submission):
    response = ws["api"].post(BUSINESS + "/" + current["work_id"] + "/" + operation,
        headers=ws["headers"], json={"submission_id": submission, "expected_revision": current["revision"], "payload": payload})
    assert response.status_code == 200, response.get_json()
    return response.get_json()["result"]


@pytest.mark.parametrize("profile", ["legacy", "manuscript", "source"])
def test_normal_owner_action_unknown_outcome_reopen_and_real_word(owners, profile):
    ws = owners
    current = prepare_business(ws)
    current = advance(ws, current, "choose", {"option_id": "B", "reason": "Inspect the smaller internal deliverable.",
        "owner": "Consenting operator", "review_at": "2026-10-05T12:00:00+08:00"}, "choose-owner-route-01")
    values = (arguments if profile == "legacy" else manuscript_arguments)({"actor": None, "scope": ws["scope"]})
    values.pop("actor")
    if profile == "source":
        report_response = ws["api"].post(URL + "/generate", headers=ws["headers"], json={})
        assert report_response.status_code == 200
        report = report_response.get_json()["result"]
        values.pop("input_manifest", None)
        values.pop("shot_plan", None)
        values.update(profile="source-bound/v3", upstream_ref={
            "kind": "source_brief_job", "job_id": report["job_id"],
            "expected_artifact_hash": report["artifact_outcome"]["artifact_hash"]})
    pending = ws["api"].post(PITCH, headers=ws["headers"], json=values)
    assert pending.status_code == 202, pending.get_json()
    task_id = pending.get_json()["result"]["task_id"]
    produced = ws["api"].post(PITCH + "/" + task_id + "/run", headers=ws["headers"], json={"scope": ws["scope"]})
    assert produced.status_code == 200, produced.get_json()
    real = produced.get_json()["result"]
    current = advance(ws, current, "action", {"option_id": "B", "task_id": task_id}, "action-owner-route-01")
    assert current["action_state"] == "EXECUTED"
    assert current["action"]["receipt"]["sha256"] == real["action_receipt"]["sha256"]
    current = advance(ws, current, "outcome_unknown", {"limits": "No campaign or consumer result has been measured."},
                      "outcome-owner-route-01")
    assert current["business_outcome"] == "UNKNOWN"
    assert current["evaluation"]["state"] == "RECORDED"
    reopened = ws["api"].get(BUSINESS + "/" + current["work_id"], headers=ws["headers"])
    assert reopened.status_code == 200 and reopened.get_json()["result"] == current
    query = {key: (str(value).lower() if isinstance(value, bool) else value or "") for key, value in ws["scope"].items()}
    exported = ws["api"].get(PITCH + "/" + task_id + "/export", query_string=query, headers=ws["headers"])
    assert exported.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(exported.data)).paragraphs)
    assert "What changed?" in text
    assert exported.headers["X-Artifact-SHA256"] == real["action_receipt"]["output"]["sha256"]


def test_client_receipt_identity_and_global_scope_cannot_create_domain_work(owners):
    ws = owners
    with sqlite3.connect(ws["db"]) as con:
        before = con.execute("SELECT COUNT(*) FROM app_evidence_records WHERE feature_id='CTH_BMP_DECISION_WORK'").fetchone()[0]
    for injection in ({"actor": {"id": 1}}, {"action_receipt": {"execution_state": "EXECUTED"}}, {"scope": ws["scope"]}):
        denied = ws["api"].post(BUSINESS, headers=ws["headers"], json=injection)
        assert denied.status_code == 400
    global_headers = {"X-Domain-Id": "CTH", "X-Product-Id": "CTH", "X-Lumenis-Context-Mode": "GLOBAL_ALL"}
    denied = ws["api"].post(BUSINESS, headers=global_headers, json={"submission_id": "denied-global-01",
        "source_job_id": "job-" + "a" * 32, "kind": "venue_operations", "question": "Unbound question", "options": []})
    assert denied.status_code == 403
    with sqlite3.connect(ws["db"]) as con:
        after = con.execute("SELECT COUNT(*) FROM app_evidence_records WHERE feature_id='CTH_BMP_DECISION_WORK'").fetchone()[0]
    assert after == before


def test_authenticated_actor_consumes_actual_owner_identity_layout(owners, monkeypatch):
    """A canonical account actor is distinct from the open-workspace principal."""
    from account.models import ACCOUNT_TABLES_SQL
    import lumenis.auth.service
    ws = owners
    with sqlite3.connect(ws["db"]) as con:
        con.executescript(ACCOUNT_TABLES_SQL)
        con.execute("INSERT INTO v4_users(id,username,password_hash,role,status) VALUES(?,?,?,?,?)",
                    (921, "owned-route-actor", "isolated-test-only", "owner", "active"))
        con.commit()
    ws["app"].config["LUMENIS_OPEN_WORKSPACE"] = False
    # The transport returns an existing canonical actor. Permission reads its
    # actual role/status from the disposable account table, never this dict.
    monkeypatch.setattr(lumenis.auth.service, "get_current_user", lambda *a, **k: {"id": 921})
    current = prepare_business(ws)
    current = advance(ws, current, "choose", {"option_id": "B", "reason": "Review the internal artifact.",
        "owner": "Canonical test actor", "review_at": "2026-10-05T12:00:00+08:00"}, "choose-numeric-owner-01")
    values = arguments({"actor": None, "scope": ws["scope"]}, request="numeric-owner-pitch-01")
    values.pop("actor")
    pending = ws["api"].post(PITCH, headers=ws["headers"], json=values)
    assert pending.status_code == 202, pending.get_json()
    task_id = pending.get_json()["result"]["task_id"]
    produced = ws["api"].post(PITCH + "/" + task_id + "/run", headers=ws["headers"], json={"scope": ws["scope"]})
    assert produced.status_code == 200, produced.get_json()
    assert produced.get_json()["result"]["work"]["document"]["context"]["actor_id"] == 921
    current = advance(ws, current, "action", {"option_id": "B", "task_id": task_id}, "numeric-owner-action-01")
    assert current["action_state"] == "EXECUTED"
    with sqlite3.connect(ws["db"]) as con:
        con.execute("UPDATE v4_users SET status='inactive' WHERE id=921")
        con.commit()
    denied = ws["api"].get(BUSINESS + "/" + current["work_id"], headers=ws["headers"])
    assert denied.status_code in {401, 403}, denied.get_json()


def test_business_work_http_permission_mapping_keeps_shared_save_at_data_export():
    """Compile only the production resolver; no app startup or live services."""
    import ast
    import re
    from pathlib import Path
    from flask import request
    from lumenis.system.security.permission import OperationType

    tree = ast.parse(Path("app.py").read_text(encoding="utf-8-sig"))
    node = next(item for item in tree.body if isinstance(item, ast.FunctionDef)
                and item.name == "_pg_resolve_operation_type")
    resolver = {}
    exec(compile(ast.Module(body=[node], type_ignores=[]), "app.py", "exec"),
         {"re": re, "request": request, "OperationType": OperationType}, resolver)
    resolve = resolver["_pg_resolve_operation_type"]

    work_id = "work-" + "a" * 64
    assert resolve(f"{BUSINESS}/{work_id}/save-request", "POST") == OperationType.DATA_EXPORT
    assert resolve(f"{BUSINESS}/{work_id}/save", "POST") == OperationType.DATA_EXPORT
    assert resolve(f"{BUSINESS}/{work_id}/adopt", "POST") == OperationType.REPORT_REVIEW
