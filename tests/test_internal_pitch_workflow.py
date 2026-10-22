"""Real owner/file/transaction regressions in disposable databases."""
import copy
import hashlib
import io
import json

import pytest
from docx import Document
from flask import Flask

from lumenis.shared.database import DatabaseRepository
from lumenis.shared.content_creation import internal_pitch as pitch
from lumenis.shared.reporting import saved_work as work
from lumenis.system.security import permission


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    path = tmp_path / "internal-pitch.sqlite3"
    monkeypatch.setenv("ZBC_DB_PATH", str(path))
    monkeypatch.setenv("ZBC_ARTIFACT_DIR", str(tmp_path / "artifacts"))
    monkeypatch.setattr(permission, "_TESTING_BYPASS_ENABLED", False)
    monkeypatch.delenv("PERMISSION_GATE_TESTING", raising=False)
    import ai_task_engine
    import shared_artifacts
    import subject_registry
    repository = DatabaseRepository(str(path))
    for module in (ai_task_engine, shared_artifacts, work, pitch):
        monkeypatch.setattr(module, "get_repository", lambda: repository)
    with repository.connection() as con:
        con.executescript(ai_task_engine._TASK_ENGINE_SCHEMA)
        subject_registry.ensure_schema(con)
        con.execute("INSERT INTO lumenis_subject_registry (subject_id,brand_focus_id,domain_id,subject_type,name_en,name_zh,name_en_norm,name_zh_norm,status,created_at,updated_at) VALUES ('SUBJ-PITCH',NULL,'CTH','company','Pitch Center','Center','pitch center','center','configured','2026-10-03T00:00:00Z','2026-10-03T00:00:00Z')")
        con.commit()
    shared_artifacts.ensure_artifact_table()
    gate = permission.PermissionGate()
    gate.grant_permission(permission.PermissionLevel.COMMIT)
    monkeypatch.setattr(permission, "get_permission_gate", lambda: gate)
    pitch.engine.cache_clear()
    app = Flask(__name__)
    app.config["LUMENIS_OPEN_WORKSPACE"] = True
    with app.test_request_context("/", headers={"Host": "localhost"},
                                  environ_overrides={"REMOTE_ADDR": "127.0.0.1"}):
        actor = permission.resolve_local_workspace_principal(module="reports", action="write")
        scope = {"domain_id": "cth", "subject_id": "SUBJ-PITCH", "brand_focus_id": None, "global_scope": False}
        yield {"repo": repository, "actor": actor, "scope": scope, "app": app, "gate": gate, "root": tmp_path}
    pitch.engine.cache_clear()


def arguments(ws, request="request-0001"):
    from lumenis.shared.content_creation.contracts import Shot, ShotPlan
    return {"manuscript": {"title": "Internal questions", "summary": "Questions for internal review.",
             "sections": [{"heading": "Original text", "body": "What changed? No change is also a valid answer."}],
             "language": "en", "authorship": "user_supplied"},
            "shot_plan": ShotPlan((Shot(shot_purpose="Review the open question", duration_sec=8),)).to_dict(),
            "request_id": request, "actor": ws["actor"], "scope": ws["scope"]}


def prepare(ws, request="request-0001", **changes):
    values = arguments(ws, request)
    values.update(changes)
    return pitch.prepare_internal_pitch(**values)


def run(ws, task):
    return pitch.run_internal_pitch(task["task_id"], actor=ws["actor"], scope=ws["scope"])


def count(ws):
    return ws["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]


def test_normal_real_word_save_revision_reopen_and_download(workspace):
    ws = workspace
    pending = prepare(ws)
    assert pending["state"] == "pending" and pending["action_receipt"] is None
    first = run(ws, pending)
    assert first["state"] == "succeeded"
    receipt = first["action_receipt"]
    assert receipt["action_kind"] == "INTERNAL_TEXT_CREATION"
    assert receipt["external_effect"] == "NOT_PERFORMED" and receipt["business_outcome"] == "NOT_OBSERVED"
    output = pitch.download_internal_pitch(pending["task_id"], actor=ws["actor"], scope=ws["scope"])
    assert hashlib.sha256(output["content"]).hexdigest() == receipt["output"]["sha256"]
    assert "What changed? No change is also a valid answer." in [p.text for p in Document(io.BytesIO(output["content"])).paragraphs]
    args = arguments(ws, "revision-0002")
    args["manuscript"]["sections"][0]["body"] = "Which information is still missing?"
    parent = first["work"]["document"]
    second = run(ws, pitch.prepare_internal_pitch(**args, work_id=parent["work_id"],
        previous_artifact_id=receipt["artifact_id"], expected_draft_digest=parent["draft_digest"]))
    assert second["action_receipt"]["version"] == 2
    assert second["action_receipt"]["action_kind"] == "INTERNAL_TEXT_REVISION"
    assert second["action_receipt"]["parent_artifact_id"] == receipt["artifact_id"]
    assert second["action_receipt"]["parent_sha256"] == receipt["sha256"]
    assert pitch.read_internal_pitch(pending["task_id"], actor=ws["actor"], scope=ws["scope"])["action_receipt"] == receipt
    assert count(ws) == 4


def test_committed_retry_never_calls_word_owner_again(workspace, monkeypatch):
    task = prepare(workspace)
    first = run(workspace, task)
    import scripts.pitch_export
    monkeypatch.setattr(scripts.pitch_export, "export_pitch_docx", lambda *a, **k: pytest.fail("committed replay must not produce another Word"))
    assert run(workspace, task)["action_receipt"] == first["action_receipt"]
    assert prepare(workspace)["action_receipt"] == first["action_receipt"]
    assert count(workspace) == 2


def test_changed_input_under_same_request_is_conflict(workspace):
    prepare(workspace)
    values = arguments(workspace)
    values["manuscript"]["summary"] += " Changed."
    with pytest.raises(ValueError, match="INTERNAL_PITCH_REQUEST_REUSE_CONFLICT"):
        pitch.prepare_internal_pitch(**values)
    assert count(workspace) == 0


def test_pending_cancel_is_durable_and_not_executable(workspace):
    task = prepare(workspace)
    cancelled = pitch.cancel_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert cancelled["state"] == "cancelled" and cancelled["cancellation"] == "CANCELLED"
    assert cancelled["action_receipt"] is None
    with pytest.raises(ValueError, match="INTERNAL_PITCH_TASK_NOT_RUNNABLE"):
        run(workspace, task)
    assert count(workspace) == 0


def test_actual_export_stage_cancel_rolls_back_all_outputs(workspace, monkeypatch):
    task = prepare(workspace)
    import scripts.pitch_export
    owner = scripts.pitch_export.export_pitch_docx
    observed = []
    def cancelling(*args, **kwargs):
        callback = kwargs["check_continue"]
        calls = 0
        def check():
            nonlocal calls
            calls += 1
            if calls == 3:
                reply = pitch.cancel_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])
                observed.append(reply["state"])
            callback()
        return owner(*args, **{**kwargs, "check_continue": check})
    monkeypatch.setattr(scripts.pitch_export, "export_pitch_docx", cancelling)
    with pytest.raises(ValueError, match="INTERNAL_PITCH_CANCELLED"):
        run(workspace, task)
    assert observed == ["cancel_requested"]
    assert pitch.read_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])["state"] == "cancelled"
    assert count(workspace) == 0


def test_cancel_after_commit_preserves_exact_artifact(workspace):
    task = prepare(workspace)
    first = run(workspace, task)
    after = pitch.cancel_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert after["cancellation"] == "ALREADY_COMMITTED"
    assert after["state"] == "succeeded" and after["action_receipt"] == first["action_receipt"]
    assert count(workspace) == 2


def test_lost_ack_after_commit_is_reconciled_from_real_owner(workspace, monkeypatch):
    task = prepare(workspace)
    original = pitch.read_internal_pitch
    def lose(task_id, **kwargs):
        result = original(task_id, **kwargs)
        if result["state"] == "succeeded":
            raise RuntimeError("response connection lost")
        return result
    monkeypatch.setattr(pitch, "read_internal_pitch", lose)
    with pytest.raises(RuntimeError, match="connection lost"):
        run(workspace, task)
    monkeypatch.setattr(pitch, "read_internal_pitch", original)
    restored = run(workspace, task)
    assert restored["state"] == "succeeded" and count(workspace) == 2


def test_attachment_corruption_blocks_reopen(workspace):
    task = prepare(workspace)
    saved = run(workspace, task)
    attachment = saved["action_receipt"]["output"]["artifact_id"]
    path = workspace["root"] / "artifacts" / (attachment + ".docx")
    path.write_bytes(path.read_bytes() + b"corrupt")
    with pytest.raises(ValueError, match="WORK_ATTACHMENT_INTEGRITY_FAILED"):
        pitch.read_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])


def test_final_reopen_revocation_does_not_return_content(workspace, monkeypatch):
    task = prepare(workspace)
    run(workspace, task)
    original = work._read_attachment
    def revoke(*a, **k):
        raw = original(*a, **k)
        workspace["app"].config["LUMENIS_OPEN_WORKSPACE"] = False
        return raw
    monkeypatch.setattr(work, "_read_attachment", revoke)
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        pitch.read_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert count(workspace) == 2


def test_stale_revision_cannot_replace_committed_result(workspace):
    first = run(workspace, prepare(workspace))
    doc = first["work"]["document"]
    args = dict(work_id=doc["work_id"], previous_artifact_id=first["action_receipt"]["artifact_id"], expected_draft_digest=doc["draft_digest"])
    run(workspace, prepare(workspace, "revision-0002", **args))
    with pytest.raises(ValueError, match="WORK_REVISION_CONFLICT"):
        run(workspace, prepare(workspace, "revision-0003", **args))
    assert count(workspace) == 4


def test_wrong_scope_or_forged_actor_cannot_read_or_cancel(workspace):
    task = prepare(workspace)
    run(workspace, task)
    wrong = {**workspace["scope"], "subject_id": "SUBJ-MISSING"}
    for operation in (pitch.read_internal_pitch, pitch.cancel_internal_pitch):
        with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
            operation(task["task_id"], actor=workspace["actor"], scope=wrong)
        with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
            operation(task["task_id"], actor={"id": "workspace:local"}, scope=workspace["scope"])


def test_generic_task_cancel_cannot_bypass_internal_owner(workspace):
    task = prepare(workspace)
    with pytest.raises(PermissionError, match="managed owner"):
        pitch.engine().cancel_task(task["task_id"])
    assert pitch.read_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])["state"] == "pending"


@pytest.fixture
def client(workspace):
    # Execute the actual global auth/effect hooks without starting collectors,
    # legacy providers, shared Core processes or any formal database.
    import ast
    import re
    import secrets
    from pathlib import Path
    from flask import request, jsonify
    from lumenis_app_api import domain_api_bp
    app = workspace["app"]
    app.register_blueprint(domain_api_bp)
    names = {"_auth_enabled", "_is_cth_readonly_get", "_basic_auth", "_owner_authorized_root_request",
             "_pg_resolve_operation_type", "_permission_gate_write_check"}
    constants = {"_PG_WRITE_METHODS", "_PG_BYPASS_PATHS", "_CTH_READONLY_GET_PREFIXES"}
    nodes = [node for node in ast.parse(Path("app.py").read_text(encoding="utf-8-sig")).body if
        isinstance(node, ast.FunctionDef) and node.name in names or
        isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id in constants for t in node.targets)]
    assert {n.name for n in nodes if isinstance(n, ast.FunctionDef)} == names
    namespace = dict(app=app, request=request, jsonify=jsonify, re=re, secrets=secrets,
        AUTH_USER="isolated_basic_user", AUTH_PASSWORD="isolated_basic_password",
        _pg=workspace["gate"], OperationType=permission.OperationType)
    exec(compile(ast.Module(body=nodes, type_ignores=[]), "<actual-app-hooks>", "exec"), namespace)
    return app.test_client()


PREFIX = "/api/experience/creation/internal-pitches"


def http_input(ws):
    values = arguments(ws)
    values.pop("actor")
    return values


def query(ws):
    return {k: (str(v).lower() if k == "global_scope" else v or "") for k, v in ws["scope"].items()}


def test_actual_api_normal_create_read_export_cancel(workspace, client):
    response = client.post(PREFIX, json=http_input(workspace))
    assert response.status_code == 202
    task = response.get_json()["result"]["task_id"]
    response = client.post(PREFIX + "/" + task + "/run", json={"scope": workspace["scope"]})
    assert response.status_code == 200 and response.headers["Cache-Control"] == "no-store"
    saved = response.get_json()["result"]
    reopened = client.get(PREFIX + "/" + task, query_string=query(workspace))
    assert reopened.get_json()["result"]["action_receipt"] == saved["action_receipt"]
    output = client.get(PREFIX + "/" + task + "/export", query_string=query(workspace))
    assert output.status_code == 200
    assert hashlib.sha256(output.data).hexdigest() == output.headers["X-Artifact-SHA256"]
    assert output.headers["X-Artifact-ID"] == saved["action_receipt"]["output"]["artifact_id"]
    assert len(Document(io.BytesIO(output.data)).paragraphs) > 0
    cancelled = client.post(PREFIX + "/" + task + "/cancel", json={"scope": workspace["scope"]})
    assert cancelled.get_json()["result"]["cancellation"] == "ALREADY_COMMITTED"


def test_explicit_global_internal_creation_does_not_require_entity(workspace, client):
    values = http_input(workspace)
    values["scope"] = {"domain_id": "cth", "subject_id": None, "brand_focus_id": None, "global_scope": True}
    response = client.post(PREFIX, json=values)
    assert response.status_code == 202
    task = response.get_json()["result"]["task_id"]
    response = client.post(PREFIX + "/" + task + "/run", json={"scope": values["scope"]})
    assert response.status_code == 200
    assert response.get_json()["result"]["action_receipt"]["scope"] == values["scope"]


@pytest.mark.parametrize("change", ["actor", "duplicate", "nonfinite", "global", "language", "shot", "large"])
def test_api_rejects_bad_input_without_creating_task(workspace, client, change):
    values = http_input(workspace)
    if change == "actor":
        values["actor_id"] = 1
    elif change == "global":
        values["scope"]["global_scope"] = True
    elif change == "language":
        values["manuscript"]["language"] = []
    elif change == "shot":
        values["shot_plan"]["shots"] = [None]
    body = json.dumps(values)
    if change == "duplicate":
        body = '{"scope":{},"scope":{}}'
    if change == "nonfinite":
        body = '{"scope":{},"number":NaN}'
    if change == "large":
        body = " " * (1024 * 1024 + 1)
    response = client.post(PREFIX, data=body, content_type="application/json")
    # Mixed selected-Subject/global scope is rejected by the real owner guard.
    if change == "global":
        assert response.status_code == 403
        assert response.get_json() == {"ok": False, "code": "WORK_PERMISSION_DENIED"}
    else:
        assert response.status_code in {400, 401}
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0


def test_api_closed_workspace_and_forged_identity_do_not_bypass_auth(workspace, client):
    values = http_input(workspace)
    assert client.post(PREFIX, json=values, headers={"X-Actor-ID": "1"}).status_code == 401
    workspace["app"].config["LUMENIS_OPEN_WORKSPACE"] = False
    assert client.post(PREFIX, json=values).status_code == 401
    assert count(workspace) == 0


def test_api_effect_gate_denial_writes_nothing(workspace, client):
    workspace["gate"].revoke_permission(permission.PermissionLevel.VERIFY)
    response = client.post(PREFIX, json=http_input(workspace))
    assert response.status_code == 403
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0


def test_final_export_permission_revocation_rolls_back_task_and_artifact(workspace, monkeypatch):
    task = prepare(workspace)
    original = work._stage_bytes
    def revoke(*args, **kwargs):
        result = original(*args, **kwargs)
        workspace["gate"].revoke_permission(permission.PermissionLevel.VERIFY)
        return result
    monkeypatch.setattr(work, "_stage_bytes", revoke)
    with pytest.raises(ValueError, match="WORK_EFFECT_PERMISSION_DENIED"):
        run(workspace, task)
    assert count(workspace) == 0
    assert pitch.engine().get_pitch_task(task["task_id"], owner_context=work.authorize(
        workspace["actor"], workspace["scope"], action="read"))["status"] == "failed"


def test_actual_worker_process_exit_recovers_through_existing_task_owner(workspace):
    import os
    from pathlib import Path
    import subprocess
    import sys
    task = prepare(workspace)
    script = '''
import json, os, sys
from flask import Flask
from lumenis.system.security import permission
from lumenis.shared.content_creation import internal_pitch
from scripts import pitch_export
permission.set_testing_bypass(False)
app=Flask('real_worker_crash')
app.config['LUMENIS_OPEN_WORKSPACE']=True
pitch_export.export_pitch_docx=lambda *a, **k: os._exit(23)
with app.test_request_context('/', headers={'Host':'localhost'}, environ_overrides={'REMOTE_ADDR':'127.0.0.1'}):
    actor=permission.resolve_local_workspace_principal(module='reports', action='write')
    internal_pitch.run_internal_pitch(sys.argv[1], actor=actor, scope=json.loads(sys.argv[2]))
'''
    result = subprocess.run([sys.executable, "-c", script, task["task_id"], json.dumps(workspace["scope"])],
        cwd=Path.cwd(), env={**os.environ, "PYTHONPATH": os.pathsep.join(filter(None,
            (str(Path.cwd()), os.environ.get("PYTHONPATH", ""))))},
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=30,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    assert result.returncode == 23, result.stdout.decode(errors="replace")
    assert pitch.read_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])["state"] == "running"
    restored = run(workspace, task)
    assert restored["state"] == "succeeded" and restored["action_receipt"]["version"] == 1
    assert count(workspace) == 2


# New shape regressions are technical fixture evidence, never business acceptance.
def manuscript_arguments(ws, request="manuscript-request-0001"):
    args = arguments(ws, request)
    args.pop("shot_plan")
    materials = []
    for key, role, text in (("brief", "brief", "The complete fixture brief.\n"),
                            ("form", "application_form", "The separate complete fixture form.\n")):
        materials.append({"material_id": key, "role": role, "filename": key + ".docx",
            "size_bytes": 100, "sha256": "0" * 64, "text": text,
            "text_sha256": hashlib.sha256(text.encode()).hexdigest(),
            "source_url": "https://example.com/" + key + ".docx", "published_at": None,
            "captured_at": "2026-10-03T00:00:00Z", "source_ref": None,
            "evidence_ref": None, "source_use_ref": None})
    args.update(profile=pitch.MANUSCRIPT_PROFILE,
        input_manifest={"schema": pitch.INPUT_MANIFEST_SCHEMA, "state": "NOT_INGESTED", "materials": materials})
    return args


def test_manuscript_only_real_word_version_revision_full_manifest_and_reopen(workspace):
    ws = workspace
    args = manuscript_arguments(ws)
    first = run(ws, pitch.prepare_internal_pitch(**args))
    doc = first["work"]["document"]
    assert doc["contract"] == pitch.MANUSCRIPT_CONTRACT
    assert "shot_plan" not in doc["payload"] and "shot_timing" not in doc["payload"]["truth_state"]
    assert doc["payload"]["input_manifest"] == args["input_manifest"]
    assert doc["payload"]["truth_state"]["source_admission"] == "PENDING"
    assert doc["payload"]["truth_state"]["input_materials"] == "OPERATOR_DECLARED_NOT_INGESTED"
    assert doc["payload"]["truth_state"]["owned_model_execution"] == "NOT_RUN"
    second_args = manuscript_arguments(ws, "manuscript-revision-0002")
    second_args["manuscript"]["sections"][0]["body"] = "A revised original question."
    second_args.update(work_id=doc["work_id"], previous_artifact_id=first["action_receipt"]["artifact_id"],
                       expected_draft_digest=doc["draft_digest"])
    second = run(ws, pitch.prepare_internal_pitch(**second_args))
    assert second["action_receipt"]["version"] == 2
    assert second["action_receipt"]["parent_sha256"] == first["action_receipt"]["sha256"]
    assert second["work"]["document"]["payload"]["input_manifest"] == args["input_manifest"]
    output = pitch.download_internal_pitch(second["task_id"], actor=ws["actor"], scope=ws["scope"])
    assert "A revised original question." in [p.text for p in Document(io.BytesIO(output["content"])).paragraphs]
    assert hashlib.sha256(output["content"]).hexdigest() == second["action_receipt"]["output"]["sha256"]
    assert pitch.read_internal_pitch(first["task_id"], actor=ws["actor"], scope=ws["scope"])["work"]["document"] == doc
    assert count(ws) == 4


def test_exact_v1_task_created_with_old_owner_remains_byte_identical_under_new_owner(workspace, monkeypatch):
    import base64
    import types
    import zlib
    # Immutable pre-upgrade owner is a test-only compatibility golden, not a
    # runtime owner. Inline compressed bytes keep the test independent of W4
    # work paths, Git history depth and external source downloads.
    legacy = zlib.decompress(base64.b64decode("eJzNWutz47YR/66/AuV9oS463TlNM60TzlRnKxdNfbIry3lU0XBgErJwpkiVBG0rN/7fu7sA+BJlS/Flpvpg8wEsFvv87YKO47yXccjEg8yUjG/Yz0kavr1cJuoi4nGPTXl2O4xvZCwYh2GX/E6EMOSWJfexSDO2SFImYyXSmEfsHl70O51xwrJgKVacBangSiZAh0c0RgmWqSQVPVjQTFrn15HMlrg20FoloYhYmuRK9DsXIl3JLAMCTGZA4p5vMpaKLImAC3a9YWopSs6JI7ZIkxU9j+SdgMH/zUWm+h3HcTr0yvcXucpT4ftMrtZJqmBfcaKIzazTMc8+ZUlsr1Nhr5RYrRcyEprSIo8DlSRRZglFae4HHHau36+5Wkby2r69gNuOfsOlr0CuvtCCNQNKUetRUb4Sscz62ZKnIuwHCYg5Vr6VKT1IeaCK9VFrPWZ110ok5Ipf86xY80YoPxXrJJOglU3rFHydkoDNnAxtwEddM56RzhvzNhkIqp+JIE+l2vTXFS1qCufwiPYw3axFp3NyPp5OBidT5rHPjgydY+ZYWta03qylCpZOjzl3YHUwEwcdOb0Oq/wc1LYEnv0130QJx/9qmcHQmbPicZ4FqVwrJJKBjPw1CAlvVJqrpZ+BCQhn/tjxL0/OL4bAi+uEyYrL2AeWcE5+/UkEytxdp+AO/iIJ8sw8uYmSax75WZCshdPtdDr/LOzBXfGHTP4uvKNuJxQLpvXudo+J/VSAPcYV9bs8V4kvY6n88Nr7gUeZQHo40ycreFAu6D1Je4xWA+8KUJqG3is2fFhHMpCKfTg7fz84ezs4O0MHuuORDFs8lk2rblQ6naGm/SoVuHuwtRhYyGAs2CKj/bOFFFGYfcfGV7COfhQL0BOqO5ICPJfdgLjACyvbpZVho8skBcm07scz27J7Lw2pMRxDQp4Gwghgh5CcezBHVE2TCx9UFIA3Lhaww8akmqn2TwfTgT/85eJ8Mq2savhbg6eAx/hWuD7ZrPu6x0rjA7rW9Ho2OIH90I7LVWtW3fiBWnWs8sZJDAziDoCCuQMe7mQCNsnBZReczFW/eoqmeFjDzsFvwpQvwOjkDbBF04xEIXpeCjCXBNNAoHK4/PmbyraKhEFJQiqM0ndSuzxFrT7GX2Jfqwbcax8tgcXmAsZiOO6jR2eu1hg+cD9Xvfq4JuWKhx9XRP6UEGwQKQQMM4trTE2zOXKnxQ3vzNVeNNvUAiTaHu9Fr1VhQLD1+WNXC1NLLuMLYeOjS/LVb0vpgbzpeS1ozmmQXDAXkiUEExlDwIwDDG2l1EMZqC6KKhOq8qLL/oKRXUkVCR1JVyuebuhSBFrWcA0Kusn5DQ3RgQFQwdp5rAkEiO9kYGaWAD1lKu3ase/Y9ywScdvAPoyTa7fbZd977K/v3h2wlN1EY7HWIXaZA8gXwqjQb3+vqcQg4N8pQQpKarfJQYKryHvHgtUR5ZI8o1RACTDPROpn+RpjfnjQ6oUVwNqABVUhzSPUS1N35WhS27ffmCCFv5RLADc/ofkO0zRJXWc0ng4n48GZfzGanvzofxyMry5PJqOLqT8a/zQ4G52aULMgsyXKuLX2BcuFWjzBDGu4gXlqfGApeAhJlhBEEm4aUjKS4vFmB+3ZbcPcyufWxmgjt7gFd2u1brciqpeJC2NqGSjKaLszTmQETuuigUdGLniZaVwlxQIg6YJM+U6GAibjffasRSGNmaG005Bqg8h+jr7+e5NwmwIs72QnKN46pWd4M3Iqc4tl8BDbvfzxfOpfnA3GDV28Yid5miIeg6Qs43WONcsCYFcKsG61BghOCM8k7oDHSSwDuNLw5Q34LYznoUgNtVWeKXaNVSGMJAJE7zvAcVESGNRIlwp3SlAP0DHARwVlm64JYWLfmoKWVKnZOercccgM6FVdz3OrgqwP5YnbeLtTA4cI8vxqcjL034/Gp6PxB38y/PfVaDK04lTppqSFRg12bvGNq/J1JFy8dV+/znYy060R6BPyhvrCpADxEAjIsy5iSuKwV+G2i1UVjHiZYeg6GMhYJRAfUFagFbvkdcjrF7A+Y9omo4KsWvKsliuWvCX+A6enJsHucu8Nznh7d4ShgIoQOxuRmb5qR0tOiaxhaBVmO+Qe/pJnSwPh+gYlGSxkh8BbevBoqgXtK8C20+DRoP7ju6Njh31VJ1ndsRbWHjWMxmpSLTUtCCBQE974BIvdRqVjxQp82ZqyT/0BoYsPajPU2OixawE2K+B+tZLKi/jqOuTH+zBWK5wwYDQrHVwHVI6rgeDnpqjxquQ8+msrpjSP22iQpl7Xa6KyEBlQGIs2gJ6TMAedYMeqhy0fqD5UssLQBm/1/qgTpFsVEM6wxMb95NEfLUiawsb+SVPSyH3NXL2aCYAvakFhxyHHcOd5iImDQIgQgFPFJZ8WdVk3tor4GbsF2bfa7OdCgcesWKXuNpp964vzWfXt/PFPM/RtwREEhY3EBuQEmGsj3zg8gtA63CnUlgrgPpCRdRNkQKTPqG+fyqz8vdDNWj074nK1v7kdwO8LmdXi6xXlulYThEW+EmB0hHbajQZSA1HAaKC7MDhCxnnRHTOaf8Y5MZ+ZEW0It+xptee26eDyX7rv558Ox6MCCuBvX0tuGBh1Ig8yLyNGT//rdrbBSFNCxQtK9jrnZn29pnigVqvpuOo7w06YBA/1qbtavn3bmLB0AOrJxabsV1A33JIiX7Y98v5U4BSof09BECi3jQu5ciEfPJvu3zRCkkPQR9l5dUUmuUJg61Ef3S0Gddlb5sCGkKLq486cuv4p2lMj2M7dkoX7+ZGqKlev0a027LyWVkivoQivftutFwNgvBUm+ub4QkMwvWDlmakc2sTsVjjszhyClNWKtG7xTwC586vpxVWznitn37MKX5B5rjeA6Ru9iy1DLOGu7iwRvqt151rlWGvTbdeSrRGs1q8/RhhJTujz0HgpAsqLIUF7B6VPAt2qv5iIQEzO5ejDeHjqn4ymg+nofHzpvz+/Gp86e8VOAqWhT0dW4G8iyJVZfgxQeXI1phrWnHOZ9rJ9ezGc/HA++Tjcd6nrHIpJkWU+qAbCtLB0zt9fDic/IRkjTCVXmA1RBoDVYW/OY1k4r3h6GwLX7S2+EsR/xZzf4t9iBAet44rOVnWkvuh/SmTs1s9lXr1iOCCbFR2J2kR4Tu2JeVlNta9aNn5Kg8PMUUtfbkDnjcZ1oGw3waHR+iAT3hP12Z80SWpntsHfyzMO/g7JOvgrQQIJ4QCUUE87e4MGmB5rfXg1gXvbkvcKFZRq05CP3wlKji4ZnlcvJXuFuXr2osduZRx6jk5QYPMm3Hjm/07uy1OTtk5MWSF6WzClUko+UdHsXNgez5hl7fHBfMdRjRnWel7QHg5J+a3nN4ZW+1kBJjFzfuzZ89eeBfyeNZldC9bxYu0OZKQUD5ao78ybfXZirqNVPUkjWtcH2lih8/sn7c5ZiVByX23WRIljizkg9b29i8M+yD9+WEXAxIqr7E2yWMhAFMuByEMoFQMInhB5VlHfvnEe57V2zHueiSFdYg+42YQJIMpDuMAkDM9L56P03l/kUQSLQ2WWOrPBm//47978Y/75qHf07t2j06PJXTyCxTO1RhzSdJvxYfjL8OQKM5L/w2B0BuFhG2MuuIz+uJ8zgXHJI87qTaB9yvxnak9T3j81eWd9P7G1GYuSTMGg2zi5j0R4I1BpUNCnSX6zLLqZZHRg3uaYmtDKAcU98vjlansN9Aj8tFXRFrPUy1gECaiL4oW+23J3Rx+HQ3kbCI2p9JkvBZXyTn+AAMDotuiSzW7n5bGAzjoGFOzfjKDPdsKiuNM7NZ09LWqrCE8H91SgX+rwrifPnHo4e6Z5YfhzC8rlfEIDS/71377VzWRLv3gGm901Tf4uNKZtTK08bzuUqZArAsi80u1CWubmieH265Xa0sXDA9YtEkmVUPHwAELYvIxVPdFUaba9P5x8m7Iarw4gajOXpld8PgQT8IindZI+SqikJojKs3mXqrCjQ0/FJsOToT4Smw4/TEbTX22QrkRS8g+tEuqKF1xtDWm49ZzCBwUZTHaXVJKzE50sT8zHX2/Lr4XeFt/k7VdKsLbQ9Pq11kolflj29P1u8FGnrStHatC3qq6qgfns3bwIQ08H6FbE/TKkXWQ43KbJWKaBeFjOOqyD/ERLVq++oyu7ffTRosg/r/m61YvUkvo/bEYWOfhlfXOyfb3LiFyOHNMZnE2Gg9Nf/ZPzjx9H0ylY07M5VHca6kP6+Rr24+62RKx26BPGw2zxS2y+aFgV0WsLuO7bY8IORSkpTR4PlvE7Q29nAJzZKDJ/zmN0K6/VY8wy1U+fXvqZ337OXkGU5F/1T2ypQreVs0taSuJo403THD/m45H53hPrDxhYFfm9RVY+TSsjqe54mA33yviwW74FULAG+IeN5ctJ1njB50ZdyBxsfpoi0uxxpotKRM61yrB4XXnaSFwWIlYGWwDSa0HYu0T42PkfGFw2MA=="))
    assert hashlib.sha256(legacy).hexdigest() == "34eaf1dd0d6f99b394e6e39670383d9e1bb416c3694ecc47b599c17bf62cf9f8"
    old = types.ModuleType("w4_exact_legacy_pitch")
    exec(compile(legacy, "<immutable-v1-compatibility-golden>", "exec"), old.__dict__)
    monkeypatch.setattr(old, "get_repository", lambda: workspace["repo"])
    args = arguments(workspace)
    pending = old.prepare_internal_pitch(**args)
    first = old.run_internal_pitch(pending["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    row_before = dict(workspace["repo"].query_one("SELECT * FROM ai_task WHERE task_id=?", (pending["task_id"],)))
    read = pitch.read_internal_pitch(pending["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    replay = pitch.run_internal_pitch(pending["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert read == first == replay
    assert pitch.prepare_internal_pitch(**args) == first
    row_after = dict(workspace["repo"].query_one("SELECT * FROM ai_task WHERE task_id=?", (pending["task_id"],)))
    assert row_after == row_before
    assert read["work"]["document"]["contract"] == pitch.CONTRACT
    assert "input_manifest" not in read["work"]["document"]["payload"]
    assert count(workspace) == 2
    old.engine.cache_clear()


@pytest.mark.parametrize("change", ["missing", "empty", "hash", "duplicate", "bound", "state", "time", "url", "private", "role", "size", "extra"])
def test_manuscript_manifest_rejects_incomplete_forged_or_bound_claims_before_task(workspace, change):
    args = manuscript_arguments(workspace)
    manifest = args["input_manifest"]
    material = manifest["materials"][0]
    expected = "INTERNAL_PITCH_INPUT_MANIFEST_INVALID"
    if change == "missing":
        args["input_manifest"] = None
    elif change == "empty":
        manifest["materials"] = []
    elif change == "hash":
        material["text"] += "undeclared modification"
    elif change == "duplicate":
        manifest["materials"][1]["material_id"] = material["material_id"]
    elif change == "bound":
        material["source_ref"] = "DOC-invented"
        expected = "INTERNAL_PITCH_SOURCE_BINDING_REQUIRED"
    elif change == "state":
        manifest["state"] = "VERIFIED"
    elif change == "time":
        material["captured_at"] = "2026-10-03"
    elif change == "url":
        material["source_url"] = "https://user:private@example.com/source"
    elif change == "private":
        material["local_path"] = "D:/private.docx"
        expected = "WORK_PRIVATE_FIELD_FORBIDDEN"
    elif change == "role":
        material["role"] = []
    elif change == "size":
        material["size_bytes"] = True
    elif change == "extra":
        manifest["licensed"] = True
    with pytest.raises(ValueError, match=expected):
        pitch.prepare_internal_pitch(**args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0
    assert count(workspace) == 0


def test_manuscript_profile_never_coerces_empty_shots_or_unknown_profile(workspace):
    for profile, shot in ((pitch.MANUSCRIPT_PROFILE, {"shots": []}), ("unknown/v2", None)):
        args = manuscript_arguments(workspace)
        args.update(profile=profile, shot_plan=shot)
        with pytest.raises(ValueError, match="INTERNAL_PITCH_PROFILE_INVALID"):
            pitch.prepare_internal_pitch(**args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0


def test_manuscript_pending_cancel_and_committed_cancel_keep_history(workspace):
    args = manuscript_arguments(workspace)
    task = pitch.prepare_internal_pitch(**args)
    cancelled = pitch.cancel_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert cancelled["state"] == "cancelled" and cancelled["action_receipt"] is None
    with pytest.raises(ValueError, match="INTERNAL_PITCH_TASK_NOT_RUNNABLE"):
        run(workspace, task)
    made = run(workspace, pitch.prepare_internal_pitch(**manuscript_arguments(workspace, "manuscript-after-cancel-0002")))
    after = pitch.cancel_internal_pitch(made["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert after["action_receipt"] == made["action_receipt"] and after["cancellation"] == "ALREADY_COMMITTED"
    assert count(workspace) == 2


def test_manuscript_cannot_rewrite_legacy_contract_or_stale_parent(workspace):
    old = run(workspace, prepare(workspace))
    doc = old["work"]["document"]
    args = manuscript_arguments(workspace)
    args.update(work_id=doc["work_id"], previous_artifact_id=old["action_receipt"]["artifact_id"], expected_draft_digest=doc["draft_digest"])
    with pytest.raises(ValueError, match="WORK_CONTRACT_CHANGE_FORBIDDEN"):
        run(workspace, pitch.prepare_internal_pitch(**args))
    assert pitch.read_internal_pitch(old["task_id"], actor=workspace["actor"], scope=workspace["scope"])["work"]["document"] == doc
    assert count(workspace) == 2


def test_manuscript_changed_material_under_same_request_is_conflict(workspace):
    args = manuscript_arguments(workspace)
    pitch.prepare_internal_pitch(**args)
    args["input_manifest"]["materials"][0]["text"] += "\nNew operator input."
    args["input_manifest"]["materials"][0]["text_sha256"] = hashlib.sha256(args["input_manifest"]["materials"][0]["text"].encode()).hexdigest()
    with pytest.raises(ValueError, match="INTERNAL_PITCH_REQUEST_REUSE_CONFLICT"):
        pitch.prepare_internal_pitch(**args)
    assert count(workspace) == 0


# Bounded M review regressions; technical inputs are not source-use admission.
@pytest.mark.parametrize("display_name", ["客户原创提案.docx", "브랜드 제안서.DOCX", "Client Brief résumé.docx"])
def test_input_compatibility_unicode_upload_retains_actual_file_and_declarations(workspace, display_name):
    args = manuscript_arguments(workspace, "unicode-input-0001")
    path = workspace["root"] / display_name
    document = Document()
    document.add_paragraph("A complete local technical brief.")
    document.save(path)
    raw = path.read_bytes()
    material = args["input_manifest"]["materials"][0]
    text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(raw)).paragraphs)
    material.update(filename=display_name, size_bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest(),
                    text=text, text_sha256=hashlib.sha256(text.encode()).hexdigest(), source_url=None,
                    source_origin={"kind": "approved-file", "ref": "TEST-UPLOAD-DECLARATION-0001"})
    expected = copy.deepcopy(args["input_manifest"])
    first = run(workspace, pitch.prepare_internal_pitch(**args))
    stored = first["work"]["document"]["payload"]
    assert stored["input_manifest"] == expected
    assert stored["truth_state"]["input_materials"] == "OPERATOR_DECLARED_NOT_INGESTED"
    assert stored["truth_state"]["source_admission"] == "PENDING"
    assert stored["input_manifest"]["materials"][0]["source_ref"] is None
    assert path.name == display_name and path.read_bytes() == raw
    read = pitch.read_internal_pitch(first["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert read["work"]["document"]["payload"]["input_manifest"] == expected


@pytest.mark.parametrize("name", ["../提案.docx", "..\\제안.docx", "C:提案.docx", "提案/稿.docx",
    "提案\x00.docx", "提案\n.docx", "提案\u202e.docx", "", " .docx", ".docx",
    "CON.docx", "LPT1.docx", "brief .docx", "brief.docx ", "brief.txt"])
def test_input_compatibility_unsafe_basename_is_rejected(workspace, name):
    args = manuscript_arguments(workspace)
    args["input_manifest"]["materials"][0]["filename"] = name
    with pytest.raises(ValueError, match="INTERNAL_PITCH_INPUT_MANIFEST_INVALID"):
        pitch.prepare_internal_pitch(**args)
    assert count(workspace) == 0


@pytest.mark.parametrize("origin", [None, {"kind": "approved-file", "ref": ""},
    {"kind": "upload", "ref": "TEST-UPLOAD-0001"}, {"kind": "approved-file", "ref": "../private"},
    {"kind": "approved-file", "ref": "TEST-UPLOAD-0001", "approved": True}])
def test_input_compatibility_null_url_needs_existing_origin_shape(workspace, origin):
    args = manuscript_arguments(workspace)
    material = args["input_manifest"]["materials"][0]
    material["source_url"] = None
    if origin is not None:
        material["source_origin"] = origin
    with pytest.raises(ValueError, match="INTERNAL_PITCH_INPUT_MANIFEST_INVALID"):
        pitch.prepare_internal_pitch(**args)
    assert count(workspace) == 0


def test_input_compatibility_public_manifest_unchanged_and_origin_not_mixed(workspace):
    args = manuscript_arguments(workspace)
    expected = copy.deepcopy(args["input_manifest"])
    prepared = pitch.prepare_internal_pitch(**args)
    assert prepared["task_id"] and args["input_manifest"] == expected
    invalid = manuscript_arguments(workspace, "mixed-input-0002")
    invalid["input_manifest"]["materials"][0]["source_origin"] = {"kind": "approved-file", "ref": "TEST-UPLOAD-0001"}
    with pytest.raises(ValueError, match="INTERNAL_PITCH_INPUT_MANIFEST_INVALID"):
        pitch.prepare_internal_pitch(**invalid)


def audiovisual_arguments(ws, request="av-plan-0001"):
    values = arguments(ws, request)
    declared = {"kind": "USER_EXPLICIT", "locator": "current_user_brief"}
    unknown = {"kind": "UNKNOWN", "locator": None}
    values["profile"] = "audiovisual-planning/v5"
    values["shot_plan"] = {
        "schema": "lumenis.audiovisual-plan/v1", "brief_text": "保留手部三个关键操作，35秒竖屏完全静音。",
        "method_source": {"reference": "dot-av-v2", "sha256": "a" * 64},
        "outputs": [{"output_id": "silent-craft", "medium": "video", "purpose": "三个关键操作清晰可读",
                     "canvas_ratio": "9:16", "duration": {"min_seconds": 35, "max_seconds": 35,
                                                              "tolerance_seconds": 0},
                     "sound": {name: "forbidden" for name in ("original", "voice", "music")},
                     "languages": ["zh"], "constraints": {"must_preserve": ["hands", "tools", "causality"]},
                     "provenance": {name: dict(declared) for name in
                                    ("purpose", "canvas_ratio", "duration", "languages", "constraints", "original", "voice", "music")},
                     "plan": None}],
        "rule_assessments": [{"rule_ref": "dot-av-v2:R10", "output_ids": ["silent-craft"],
                              "required_conditions": [{"predicate": "same_action_chain_verified", "value": None,
                                                       "provenance": unknown}], "exclusions": []}],
    }
    return values


def test_audiovisual_profile_persists_reopens_revises_and_downloads_without_media_claim(workspace):
    ws = workspace
    args = audiovisual_arguments(ws)
    first = run(ws, pitch.prepare_internal_pitch(**args))
    document = first["work"]["document"]
    assert document["contract"] == pitch.AV_PLAN_CONTRACT
    plan = document["payload"]["audiovisual_plan"]
    assert plan["input"] == args["shot_plan"]
    assert plan["rule_decisions"][0]["decision"] == "NEEDS_REVIEW"
    assert plan["execution"] == "NOT_EXECUTED" and plan["rendition_binding"] is None
    assert set(plan["reviews"].values()) == {"NEEDS_REVIEW"}
    assert document["citations"] == []
    downloaded = pitch.download_internal_pitch(first["task_id"], actor=ws["actor"], scope=ws["scope"])
    assert hashlib.sha256(downloaded["content"]).hexdigest() == first["action_receipt"]["output"]["sha256"]
    assert Document(io.BytesIO(downloaded["content"])).paragraphs
    revised = audiovisual_arguments(ws, "av-plan-0002")
    revised["shot_plan"]["brief_text"] += " 允许按动作调整镜头顺序。"
    revised["shot_plan"]["outputs"][0]["constraints"]["preserve_order"] = False
    second = run(ws, pitch.prepare_internal_pitch(**revised, work_id=document["work_id"],
                previous_artifact_id=first["action_receipt"]["artifact_id"],
                expected_draft_digest=document["draft_digest"]))
    assert second["action_receipt"]["version"] == 2
    assert second["action_receipt"]["parent_sha256"] == first["action_receipt"]["sha256"]
    new_plan = second["work"]["document"]["payload"]["audiovisual_plan"]
    assert new_plan["brief_sha256"] != plan["brief_sha256"]
    reopened = pitch.read_internal_pitch(first["task_id"], actor=ws["actor"], scope=ws["scope"])
    assert reopened["work"]["document"]["payload"]["audiovisual_plan"] == plan
    assert run(ws, first)["action_receipt"] == first["action_receipt"]
    assert count(ws) == 4


def test_audiovisual_changed_brief_cannot_reuse_request_or_erase_unknowns(workspace):
    args = audiovisual_arguments(workspace)
    task = pitch.prepare_internal_pitch(**args)
    args["shot_plan"]["brief_text"] += " Changed."
    with pytest.raises(ValueError, match="INTERNAL_PITCH_REQUEST_REUSE_CONFLICT"):
        pitch.prepare_internal_pitch(**args)
    cancelled = pitch.cancel_internal_pitch(task["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert cancelled["state"] == "cancelled" and count(workspace) == 0


@pytest.mark.parametrize("reference", ["brief", "plan_evidence", "shot_evidence", "podcast_evidence"])
def test_audiovisual_unbound_source_references_cannot_create_task(workspace, reference):
    from lumenis.shared.content_creation.contracts import Shot, ShotPlan
    args = audiovisual_arguments(workspace)
    plan = ShotPlan((Shot(shot_purpose="Show the permitted action", duration_sec=35),)).to_dict()
    if reference == "brief":
        plan["brief_ref"] = "unbound-brief"
    elif reference == "plan_evidence":
        plan["evidence_refs"] = ["unbound-evidence"]
    elif reference == "shot_evidence":
        plan["shots"][0]["evidence_refs"] = ["unbound-evidence"]
    output = args["shot_plan"]["outputs"][0]
    if reference == "podcast_evidence":
        from podcast.intelligence.episode_runtime import PodcastPlan, PodcastSegment
        plan = PodcastPlan("Interview", "Preserve original speech", "zh", (
            PodcastSegment("segment-1", "UNRESOLVED_MEDIA", 0, 35, "原话", "原话", "Guest",
                           "Opening", ("unbound-evidence",)),
        )).to_dict()
        output.update(medium="podcast", canvas_ratio=None)
        output["provenance"]["canvas_ratio"] = {"kind": "UNKNOWN", "locator": None}
        output["sound"]["original"] = "required"
    output["plan"] = plan
    with pytest.raises(ValueError, match="INTERNAL_PITCH_SOURCE_BINDING_REQUIRED"):
        pitch.prepare_internal_pitch(**args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0
    assert count(workspace) == 0




def test_audiovisual_planned_video_and_current_podcast_survive_actual_saved_work(workspace):
    from lumenis.shared.content_creation.contracts import Shot, ShotPlan
    from podcast.intelligence.episode_runtime import PodcastPlan, PodcastSegment
    args = audiovisual_arguments(workspace)
    args["shot_plan"]["outputs"][0]["plan"] = ShotPlan((Shot(shot_purpose="Preserve the action", duration_sec=35),)).to_dict()
    audio = copy.deepcopy(args["shot_plan"]["outputs"][0])
    audio.update(output_id="spoken-version", medium="podcast", canvas_ratio=None)
    audio["sound"]["original"] = "required"
    audio["provenance"]["canvas_ratio"] = {"kind": "UNKNOWN", "locator": None}
    audio["plan"] = PodcastPlan("Spoken version", "Declared existing media", "zh", (
        PodcastSegment("speech", "UNRESOLVED", 0, 35, "原文", "修订文字", "Host", "Opening", (), 1),
    ), revision=3).to_dict()
    args["shot_plan"]["outputs"].append(audio)
    result = run(workspace, pitch.prepare_internal_pitch(**args))
    retained = result["work"]["document"]["payload"]["audiovisual_plan"]
    assert retained["input"] == args["shot_plan"]
    assert [item["estimated_duration_seconds"] for item in retained["outputs"]] == [35, 35]
    assert {item["status"] for item in retained["outputs"]} == {"PLANNED_NOT_EXECUTED"}
    assert retained["rendition_binding"] is None
    assert pitch.read_internal_pitch(result["task_id"], actor=workspace["actor"], scope=workspace["scope"])["work"]["document"]["payload"]["audiovisual_plan"] == retained


@pytest.mark.parametrize("mixed", [{"input_manifest": {}}, {"upstream_ref": {"kind": "saved_work"}}])
def test_audiovisual_profile_rejects_mixed_contract_inputs_before_task(workspace, mixed):
    args = audiovisual_arguments(workspace)
    args.update(mixed)
    with pytest.raises(ValueError, match="INTERNAL_PITCH_PROFILE_INVALID"):
        pitch.prepare_internal_pitch(**args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0


def test_audiovisual_cannot_rewrite_existing_v1_work_contract(workspace):
    first = run(workspace, prepare(workspace))
    parent = first["work"]["document"]
    args = audiovisual_arguments(workspace)
    args.update(work_id=parent["work_id"], previous_artifact_id=first["action_receipt"]["artifact_id"],
                expected_draft_digest=parent["draft_digest"])
    pending = pitch.prepare_internal_pitch(**args)
    with pytest.raises(ValueError, match="WORK_CONTRACT_CHANGE_FORBIDDEN"):
        run(workspace, pending)
    assert pitch.read_internal_pitch(first["task_id"], actor=workspace["actor"], scope=workspace["scope"])["work"]["document"] == parent
    assert count(workspace) == 2


def test_audiovisual_read_compares_exact_saved_snapshot(workspace, monkeypatch):
    committed = run(workspace, pitch.prepare_internal_pitch(**audiovisual_arguments(workspace)))
    original = work.reopen_work
    def substituted(*args, **kwargs):
        result = original(*args, **kwargs)
        result["document"]["payload"]["audiovisual_plan"]["input"]["brief_text"] += " wrong owner payload"
        return result
    monkeypatch.setattr(work, "reopen_work", substituted)
    with pytest.raises(ValueError, match="INTERNAL_PITCH_AV_PLAN_INTEGRITY_FAILED"):
        pitch.read_internal_pitch(committed["task_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert count(workspace) == 2


def test_audiovisual_actual_domain_http_prepare_run_read_and_export(workspace, client):
    args = audiovisual_arguments(workspace)
    args.pop("actor")
    reply = client.post(PREFIX, json=args)
    assert reply.status_code == 202
    task = reply.get_json()["result"]["task_id"]
    result = client.post(PREFIX + "/" + task + "/run", json={"scope": workspace["scope"]})
    assert result.status_code == 200
    assert result.get_json()["result"]["work"]["document"]["contract"] == pitch.AV_PLAN_CONTRACT
    read = client.get(PREFIX + "/" + task, query_string=query(workspace))
    assert read.status_code == 200
    snapshot = read.get_json()["result"]["work"]["document"]["payload"]["audiovisual_plan"]
    assert snapshot["execution"] == "NOT_EXECUTED"
    output = client.get(PREFIX + "/" + task + "/export", query_string=query(workspace))
    assert output.status_code == 200
    assert hashlib.sha256(output.data).hexdigest() == output.headers["X-Artifact-SHA256"]
    assert args["manuscript"]["sections"][0]["body"] in [paragraph.text for paragraph in Document(io.BytesIO(output.data)).paragraphs]


def test_audiovisual_shared_path_reuses_saved_work_owner(workspace):
    args = audiovisual_arguments(workspace)
    args["scope"] = {"scope_kind": "shared_workspace", "version": 1}
    pending = pitch.prepare_internal_pitch(**args)
    saved = pitch.run_internal_pitch(pending["task_id"], actor=args["actor"], scope=args["scope"])
    assert saved["scope"] == args["scope"]
    assert saved["work"]["document"]["payload"]["audiovisual_plan"]["input"] == args["shot_plan"]
    assert saved["work"]["document"]["context"]["scope_kind"] == "shared_workspace"


@pytest.mark.parametrize("language,body", [
    ("zh", "这是用户自己的文稿。保留未修改的段落。"),
    ("en", "This is the user's own manuscript. Keep unchanged paragraphs."),
    ("ko", "사용자가 직접 작성한 원고입니다. 수정하지 않은 문단은 유지합니다."),
])
def test_authored_v4_shared_workspace_save_revise_reopen_list_and_export(workspace, client, language, body):
    ws = workspace
    scope = {"scope_kind": "shared_workspace", "version": 1}
    args = http_input(ws)
    args.pop("shot_plan")
    args.update(profile=pitch.AUTHORING_PROFILE, scope=scope, request_id="authored-shared-0001")
    args["manuscript"]["title"] = "Independent manuscript"
    args["manuscript"]["language"] = language
    args["manuscript"]["sections"] = [{"heading": "Draft", "body": body},
                                        {"heading": "Retained", "body": "Keep this section."}]

    prepared = client.post(PREFIX, json=args)
    assert prepared.status_code == 202, prepared.get_json()
    task_id = prepared.get_json()["result"]["task_id"]
    completed = client.post(PREFIX + "/" + task_id + "/run", json={"scope": scope})
    assert completed.status_code == 200, completed.get_json()
    result = completed.get_json()["result"]
    first = result["work"]["document"]
    receipt = result["action_receipt"]
    assert first["contract"] == pitch.AUTHORING_CONTRACT
    assert first["payload"]["truth_state"]["source_admission"] == "USER_AUTHORED_CONTENT"
    assert first["payload"]["truth_state"]["external_fact_verification"] == "NOT_INFERRED_FROM_AUTHORSHIP"
    assert first["payload"]["truth_state"]["citation_admission"] == "NO_SIGNED_QUOTE"
    assert "input_manifest" not in first["payload"] and "shot_plan" not in first["payload"]
    assert first["citations"] == []

    listed = work.search_work(actor=ws["actor"], scope=scope, query="Independent manuscript")
    assert any(item["document"]["work_id"] == first["work_id"] for item in listed["items"])
    reopened = work.reopen_work(receipt["artifact_id"], actor=ws["actor"], scope=scope)
    assert reopened["document"] == first
    saved_docx = work.export_work(receipt["artifact_id"], actor=ws["actor"], scope=scope, format="docx")
    assert body in [paragraph.text for paragraph in Document(io.BytesIO(saved_docx["content"])).paragraphs]
    exported = client.get(PREFIX + "/" + task_id + "/export", query_string={
        "scope_kind": "shared_workspace", "version": "1"})
    assert exported.status_code == 200
    assert body in [paragraph.text for paragraph in Document(io.BytesIO(exported.data)).paragraphs]

    revised_args = copy.deepcopy(args)
    revised_args.update(request_id="authored-shared-0002", work_id=first["work_id"],
                        previous_artifact_id=receipt["artifact_id"],
                        expected_draft_digest=first["draft_digest"])
    revised_args["manuscript"]["sections"][0]["body"] += " Revised."
    revised_prepared = client.post(PREFIX, json=revised_args)
    assert revised_prepared.status_code == 202, revised_prepared.get_json()
    revised_id = revised_prepared.get_json()["result"]["task_id"]
    revised_response = client.post(PREFIX + "/" + revised_id + "/run", json={"scope": scope})
    assert revised_response.status_code == 200, revised_response.get_json()
    revised = revised_response.get_json()["result"]
    assert revised["action_receipt"]["version"] == 2
    assert revised["action_receipt"]["parent_artifact_id"] == receipt["artifact_id"]
    assert revised["action_receipt"]["parent_sha256"] == receipt["sha256"]
    assert revised["work"]["document"]["payload"]["manuscript"]["sections"][1] == args["manuscript"]["sections"][1]
    old = client.get(PREFIX + "/" + task_id, query_string={"scope_kind": "shared_workspace", "version": "1"})
    assert old.status_code == 200
    assert old.get_json()["result"]["action_receipt"] == receipt


@pytest.mark.parametrize("extra", ["shot_plan", "input_manifest", "upstream_ref"])
def test_authored_v4_rejects_mixed_source_or_media_inputs_before_task(workspace, extra):
    args = arguments(workspace)
    args.pop("shot_plan")
    args.update(profile=pitch.AUTHORING_PROFILE)
    args[extra] = {"forged": True}
    with pytest.raises(ValueError, match="INTERNAL_PITCH_PROFILE_INVALID"):
        pitch.prepare_internal_pitch(**args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0] == 0
    assert count(workspace) == 0
