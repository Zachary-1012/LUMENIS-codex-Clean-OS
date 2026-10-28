"""Actual local raster/vector/Office encoding on isolated original SavedWork."""
import base64
import copy
import hashlib
import io
import xml.etree.ElementTree as ET

import pytest
from PIL import Image

from tests.test_saved_document_workflow import workspace
from lumenis.shared.content_creation.audio_visuals import (
    validate_visual_document,visual_document_markdown,export_visual_document)
from lumenis.shared.content_creation.video_editing import VideoEditingError


def document():
    image = io.BytesIO();Image.new('RGB',(24,12),'blue').save(image,format='PNG');raw=image.getvalue()
    common = dict(x=0,y=0,width=100,height=70,rotation=0,opacity=1,visible=True,locked=False)
    return {'schema':'lumenis.visual-document/v1','canvas':{'width':320,'height':200,'background':'#ffffff'},
        'assets':[{'id':'photo','media_type':'image/png','sha256':hashlib.sha256(raw).hexdigest(),
                   'data_base64':base64.b64encode(raw).decode()}],
        'layers':[
            {**common,'id':'background','name':'Red block','kind':'rectangle','x':5,'y':5,
             'fill':'#ff0000','stroke':'#000000','stroke_width':0},
            {**common,'id':'title','name':'Editable title','kind':'text','x':125,'width':180,'height':90,
             'text':'OEVIRA\nVisual studio','font_size':20,'color':'#111111','align':'left','line_height':1.1},
            {**common,'id':'picture','name':'Original image','kind':'image','x':5,'y':105,
             'asset_id':'photo','fit':'contain'},
            {**common,'id':'ellipse','name':'Green circle','kind':'ellipse','x':200,'y':110,'width':60,'height':60,
             'fill':'#00ff00','stroke':'#111111','stroke_width':2},
        ]}


def save(ws,value,**extra):
    args = dict(title='Editable visual',markdown=visual_document_markdown(value),kind='report',
                payload={'visual_document':value},citations=[],request_id='visual-request-0001',
                actor=ws['actor'],scope=ws['scope'],work_id=None,previous_artifact_id=None)
    args.update(extra)
    resource=ws['work'].work_write_resource(**args)
    ws['gate'].grant_approval('data_export',resource=resource,granted_by='isolated visual encoding test')
    return ws['work'].save_work(**args)


def test_actual_raster_svg_and_editable_powerpoint():
    value=document();original=copy.deepcopy(value)
    png=export_visual_document(value,format='png')
    with Image.open(io.BytesIO(png)) as image:
        assert image.size==(320,200)
        assert image.getpixel((30,30))==(255,0,0,255)
        assert image.getpixel((30,135))==(0,0,255,255)
    svg=ET.fromstring(export_visual_document(value,format='svg'))
    assert len(svg.findall('{http://www.w3.org/2000/svg}g'))==4
    assert 'Visual studio' in ''.join(svg.itertext())
    assert export_visual_document(value,format='pdf').startswith(b'%PDF-')
    from pptx import Presentation
    deck=Presentation(io.BytesIO(export_visual_document(value,format='pptx')))
    assert len(deck.slides)==1 and len(deck.slides[0].shapes)==4
    text=deck.slides[0].shapes[1]
    assert text.text=='OEVIRA\nVisual studio'
    text.text='Revised in native object model'
    output=io.BytesIO();deck.save(output)
    assert Presentation(io.BytesIO(output.getvalue())).slides[0].shapes[1].text=='Revised in native object model'
    assert value==original


def test_save_revise_old_version_reopen_and_export(workspace):
    ws=workspace;value=document();first=save(ws,value);ref=first['artifact']['artifact_id']
    assert ws['work'].reopen_work(ref,actor=ws['actor'],scope=ws['scope'])['document']['payload']['visual_document']==value
    revision=dict(request_id='visual-revise-0002',expected_sha256=first['artifact']['sha256'],reason='Move and change title',
        edits=[{'pointer':'/payload/visual_document/layers/1/text','expected':'OEVIRA\nVisual studio','value':'User revision'},
               {'pointer':'/payload/visual_document/layers/0/x','expected':5,'value':18}])
    prepared=ws['work']._prepare_work_revision(ref,actor=ws['actor'],scope=ws['scope'],**revision)[1]
    second=save(ws,prepared['payload']['visual_document'],request_id='visual-revise-0002',work_id=first['document']['work_id'],
        previous_artifact_id=ref,expected_draft_digest=first['document']['draft_digest'])
    assert second['document']['version']==2
    cold=ws['work'].reopen_work(second['artifact']['artifact_id'],actor=ws['actor'],scope=ws['scope'])
    assert cold['document']['payload']['visual_document']['layers'][1]['text']=='User revision'
    assert ws['work'].reopen_work(ref,actor=ws['actor'],scope=ws['scope'])['document']['payload']['visual_document']==value
    for fmt in ['png','svg','pptx','pdf','json']:
        bound={'id':second['artifact']['artifact_id'],'sha256':second['artifact']['sha256'],'format':fmt}
        ws['gate'].grant_approval('data_export',resource='saved-work-export:v1:'+ws['work']._digest(bound),granted_by='isolated export')
        result=ws['work'].export_work(second['artifact']['artifact_id'],actor=ws['actor'],scope=ws['scope'],format=fmt)
        assert result['sha256']==hashlib.sha256(result['content']).hexdigest()
    assert save(ws,value)['replayed'] is True


def test_title_only_revision_keeps_legacy_visual_object_as_visual(workspace):
    ws = workspace
    value = document()
    first = save(ws, value, title='Legacy visual title')
    artifact_id = first['artifact']['artifact_id']
    revision = {
        'request_id': 'visual-title-only-0002',
        'expected_sha256': first['artifact']['sha256'],
        'reason': 'Update the work title; retain the visual original.',
        'edits': [{'pointer': '/title', 'expected': 'Legacy visual title', 'value': '寰意 · 工作稿'}],
    }
    _item, prepared = ws['work']._prepare_work_revision(
        artifact_id, actor=ws['actor'], scope=ws['scope'], **revision,
    )
    resource = ws['work'].work_write_resource(
        title=prepared['title'], markdown=prepared['markdown'], kind=prepared['kind'],
        payload=prepared['payload'], citations=prepared['citations'],
        request_id=revision['request_id'], work_id=prepared['work_id'],
        previous_artifact_id=artifact_id, expected_draft_digest=first['document']['draft_digest'],
        contract=prepared['contract'], actor=ws['actor'], scope=ws['scope'],
    )
    ws['gate'].grant_approval('data_export', resource=resource, granted_by='isolated visual revision test')
    second = ws['work'].revise_work(
        artifact_id, actor=ws['actor'], scope=ws['scope'], **revision,
    )
    assert second['document']['version'] == 2
    assert second['document']['title'] == '寰意 · 工作稿'
    assert second['document']['markdown'] == first['document']['markdown']
    assert second['document']['payload']['visual_document'] == value
    assert ws['work'].reopen_work(
        artifact_id, actor=ws['actor'], scope=ws['scope'],
    )['document'] == first['document']


def test_locked_object_and_original_raster_cannot_change_silently():
    old=document();old['layers'][2]['locked']=True;new=copy.deepcopy(old)
    new['layers'][2]['x']=40
    with pytest.raises(VideoEditingError,match='VISUAL_LAYER_LOCKED'):validate_visual_document(new,previous=old)
    new['layers'][2]['locked']=False
    validate_visual_document(new,previous=old)
    old['assets'][0]['data_base64']='PHN2ZyBvbmxvYWQ9ImFsZXJ0KDEpIi8+' # executable SVG is never a raster
    with pytest.raises(VideoEditingError,match='VISUAL_ASSET_INTEGRITY_FAILED'):validate_visual_document(old)


def test_hidden_layers_preserved_and_text_overflow_never_truncates():
    value=document();value['layers'][0]['visible']=False
    png=export_visual_document(value,format='png')
    with Image.open(io.BytesIO(png)) as image:assert image.getpixel((30,30))==(255,255,255,255)
    value['layers'][1]['height']=2
    with pytest.raises(VideoEditingError,match='VISUAL_TEXT_OVERFLOW_REQUIRES_EDIT'):export_visual_document(value,format='png')


def test_cancel_and_image_hash_failure_return_no_export():
    def cancelled():raise PermissionError('explicit cancellation')
    with pytest.raises(PermissionError,match='explicit cancellation'):
        export_visual_document(document(),format='png',check_cancel=cancelled)
    value=document();value['assets'][0]['sha256']='0'*64
    with pytest.raises(VideoEditingError,match='VISUAL_ASSET_INTEGRITY_FAILED'):export_visual_document(value,format='svg')


def test_http_visual_prepare_confirm_revision_and_binary_export_use_real_owner(workspace):
    import json
    from experience_runtime import experience_bp
    from lumenis.system.security import permission
    ws = workspace
    ws['app'].register_blueprint(experience_bp)
    api = ws['app'].test_client()
    headers = {'X-Lumenis-Workspace-Scope': json.dumps({'scope_kind': 'shared_workspace', 'version': 1})}
    value = document()
    body = {'title': 'Original editable visual', 'markdown': '',
            'kind': 'report', 'payload': {'visual_document': value}, 'citations': [],
            'request_id': 'http-visual-create-0001'}
    prepared = api.post('/api/experience/saved-work', headers=headers, json={**body, 'mode': 'prepare'})
    assert prepared.status_code == 200, prepared.get_json()
    draft = prepared.get_json()['result']
    assert draft['saved'] is False and draft['effect_performed'] is False
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == 0
    assert ws['gate'].get_valid_approval_token('data_export', draft['resource'], exact=True) is None
    confirm = {**body, 'mode': 'confirm', 'confirm': True, 'confirmed_resource': draft['resource']}
    assert api.post('/api/experience/saved-work', headers=headers, json=confirm).status_code == 403
    ws['gate'].grant_permission(permission.PermissionLevel.HUMAN_APPROVAL)
    assert api.post('/api/experience/saved-work', headers=headers,
                    json={**confirm, 'confirmed_resource': 'wrong'}).status_code == 403
    assert api.post('/api/experience/saved-work', headers={**headers, 'Origin': 'https://outside.example'},
                    json=confirm).status_code in (401, 403)
    saved = api.post('/api/experience/saved-work', headers=headers, json=confirm)
    assert saved.status_code == 200, saved.get_json()
    first = saved.get_json()['result']; first_id = first['artifact']['artifact_id']
    assert first['document']['version'] == 1
    assert first['document']['markdown'] == visual_document_markdown(value)
    assert api.post('/api/experience/saved-work', headers=headers, json=confirm).get_json()['result']['artifact']['artifact_id'] == first_id
    revision = {'request_id': 'http-visual-revise-0002', 'expected_sha256': first['artifact']['sha256'],
        'reason': 'Explicit user revision', 'edits': [{'pointer': '/payload/visual_document/layers/1/text',
            'expected': value['layers'][1]['text'], 'value': 'Revised by the author'}]}
    revision_url = '/api/experience/saved-work/' + first_id + '/revise'
    preview = api.post(revision_url, headers=headers, json={**revision, 'mode': 'prepare'})
    assert preview.status_code == 200, preview.get_json()
    second = api.post(revision_url, headers=headers, json={**revision, 'mode': 'confirm', 'confirm': True,
        'confirmed_resource': preview.get_json()['result']['resource']})
    assert second.status_code == 200, second.get_json()
    second = second.get_json()['result']; second_id = second['artifact']['artifact_id']
    assert second['document']['version'] == 2
    for identifier, expected in ((first_id, value['layers'][1]['text']), (second_id, 'Revised by the author')):
        cold = ws['app'].test_client().get('/api/experience/saved-work/' + identifier, headers=headers)
        assert cold.status_code == 200, cold.get_json()
        assert cold.get_json()['result']['document']['payload']['visual_document']['layers'][1]['text'] == expected
    export_url = '/api/experience/saved-work/' + second_id + '/export'
    preview = api.post(export_url, headers=headers, json={'format': 'svg', 'mode': 'prepare'})
    assert preview.status_code == 200, preview.get_json()
    assert preview.get_json()['result']['effect_performed'] is False
    exported = api.post(export_url, headers=headers, json={'format': 'svg', 'mode': 'confirm',
        'confirm': True, 'confirmed_resource': preview.get_json()['result']['resource']})
    assert exported.status_code == 200, exported.get_json()
    assert 'Revised by the author' in ''.join(ET.fromstring(exported.data).itertext())
    assert exported.headers['X-Content-SHA256'] == hashlib.sha256(exported.data).hexdigest()


def test_http_legacy_visual_text_revision_requires_explicit_continue_as_report(workspace):
    import json
    from experience_runtime import experience_bp
    from lumenis.shared.reporting import saved_work
    from lumenis.system.security import permission

    ws = workspace
    ws['app'].config.update(TESTING=True, LUMENIS_OPEN_WORKSPACE=True)
    ws['app'].register_blueprint(experience_bp)
    api = ws['app'].test_client()
    headers = {'X-Lumenis-Workspace-Scope': json.dumps({'scope_kind': 'shared_workspace', 'version': 1})}
    value = document()
    value['layers'][1]['text'] = '寰意\n从资料到作品'
    original_markdown = visual_document_markdown(value)
    new_markdown = ('寰意\n从资料到作品\n\n资料有来源，结论可修订。\n\n'
                    '成果可以保存、重新打开并继续编辑。')
    create = {'title': 'Editable visual document · user-authored, source rights not independently verified',
        'markdown': original_markdown, 'kind': 'report', 'payload': {'visual_document': value},
        'citations': [], 'request_id': 'legacy-visual-create-0001'}
    saved_response = api.post('/api/experience/saved-work', headers=headers, json=create)
    assert saved_response.status_code == 200, saved_response.get_json()
    first = saved_response.get_json()['result']
    first_id = first['artifact']['artifact_id']
    edit = {'request_id': 'legacy-visual-edit-0002', 'expected_sha256': first['artifact']['sha256'],
        'reason': '更新文案；保留上一版本。', 'edits': [
            {'pointer': '/title', 'expected': create['title'], 'value': '寰意 · 工作稿'},
            {'pointer': '/markdown', 'expected': original_markdown, 'value': new_markdown},
        ]}
    before_count = ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0]
    rejected = api.post(f'/api/experience/saved-work/{first_id}/revise', headers=headers, json=edit)
    assert rejected.status_code == 400, rejected.get_json()
    assert rejected.get_json()['code'] == 'WORK_VISUAL_DOCUMENT_CONTINUE_MODE_REQUIRED'
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == before_count

    invalid_flag = api.post(f'/api/experience/saved-work/{first_id}/revise', headers=headers,
        json={**edit, 'request_id': 'legacy-visual-invalid-mode-0003', 'continue_as_report': 1})
    assert invalid_flag.status_code == 400
    assert invalid_flag.get_json()['code'] == 'WORK_VISUAL_CONTINUE_MODE_INVALID'

    ws['gate'].grant_permission(permission.PermissionLevel.HUMAN_APPROVAL)
    preview = api.post(f'/api/experience/saved-work/{first_id}/revise', headers=headers,
        json={**edit, 'continue_as_report': True, 'mode': 'prepare'})
    assert preview.status_code == 200, preview.get_json()
    prepared = preview.get_json()['result']
    assert prepared['saved'] is False and prepared['effect_performed'] is False
    assert prepared['state'] == 'PREPARED'
    confirmed = api.post(f'/api/experience/saved-work/{first_id}/revise', headers=headers,
        json={**edit, 'continue_as_report': True, 'mode': 'confirm', 'confirm': True,
              'confirmed_resource': prepared['resource']})
    assert confirmed.status_code == 200, confirmed.get_json()
    second = confirmed.get_json()['result']
    second_id = second['artifact']['artifact_id']
    assert second['document']['kind'] == 'report'
    assert second['document']['title'] == '寰意 · 工作稿'
    assert second['document']['markdown'] == new_markdown
    assert 'visual_document' not in second['document']['payload']
    assert second['document']['parent_artifact_id'] == first_id
    assert second['document']['parent_sha256'] == first['artifact']['sha256']
    assert second['document']['version'] == 2

    old_read = api.get(f'/api/experience/saved-work/{first_id}', headers=headers)
    new_read = api.get(f'/api/experience/saved-work/{second_id}', headers=headers)
    assert old_read.status_code == new_read.status_code == 200
    old_document = old_read.get_json()['result']['document']
    new_document = new_read.get_json()['result']['document']
    assert old_document['markdown'] == original_markdown
    assert old_document['payload']['visual_document'] == value
    assert new_document['markdown'] == new_markdown

    with ws['app'].test_request_context('/', headers={'Host': 'localhost'},
                                        environ_overrides={'REMOTE_ADDR': '127.0.0.1'}):
        actor = permission.resolve_local_workspace_principal(module='reports', action='read')
    export = saved_work.prepare_work_export(second_id, actor=actor,
        scope={'scope_kind': 'shared_workspace', 'version': 1}, format='md')
    ws['gate'].grant_approval('data_export', resource=export['resource'],
                              granted_by='isolated visual export test')
    exported = api.get(f'/api/experience/saved-work/{second_id}/export?format=md', headers=headers)
    assert exported.status_code == 200, exported.get_json()
    export_text = exported.data.decode('utf-8')
    assert export_text.startswith(new_markdown)
    assert 'Editable visual document · user-authored' not in export_text

    # Explicit continuation converts a legacy visual report even when the
    # revision only changes its title.
    title_create = {**create, 'title': 'Legacy title-only visual',
                    'request_id': 'legacy-title-create-0001'}
    title_response = api.post('/api/experience/saved-work', headers=headers, json=title_create)
    assert title_response.status_code == 200, title_response.get_json()
    title_parent = title_response.get_json()['result']
    title_parent_id = title_parent['artifact']['artifact_id']
    title_edit = {'request_id': 'legacy-title-continue-0002',
        'expected_sha256': title_parent['artifact']['sha256'],
        'reason': '继续写作时更新标题。', 'edits': [{'pointer': '/title',
            'expected': 'Legacy title-only visual', 'value': '寰意 · 标题已更新'}]}
    title_revision_url = f'/api/experience/saved-work/{title_parent_id}/revise'
    title_preview = api.post(title_revision_url, headers=headers,
        json={**title_edit, 'continue_as_report': True, 'mode': 'prepare'})
    assert title_preview.status_code == 200, title_preview.get_json()
    title_confirm = api.post(title_revision_url, headers=headers,
        json={**title_edit, 'continue_as_report': True, 'mode': 'confirm', 'confirm': True,
              'confirmed_resource': title_preview.get_json()['result']['resource']})
    assert title_confirm.status_code == 200, title_confirm.get_json()
    title_child = title_confirm.get_json()['result']
    assert title_child['document']['title'] == '寰意 · 标题已更新'
    assert title_child['document']['markdown'] == original_markdown
    assert 'visual_document' not in title_child['document']['payload']
    assert title_child['document']['parent_artifact_id'] == title_parent_id
    assert title_child['document']['parent_sha256'] == title_parent['artifact']['sha256']
    assert api.get(f'/api/experience/saved-work/{title_parent_id}', headers=headers
        ).get_json()['result']['document']['payload']['visual_document'] == value

    title_count = ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0]
    changed_mode = api.post(title_revision_url, headers=headers, json=title_edit)
    assert changed_mode.status_code == 409, changed_mode.get_json()
    assert changed_mode.get_json()['code'] == 'WORK_REQUEST_REUSE_CONFLICT'
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == title_count


def test_http_continue_as_report_flag_is_rejected_for_normal_report(workspace):
    import json
    from experience_runtime import experience_bp

    ws = workspace
    ws['app'].config.update(TESTING=True, LUMENIS_OPEN_WORKSPACE=True)
    ws['app'].register_blueprint(experience_bp)
    api = ws['app'].test_client()
    headers = {'X-Lumenis-Workspace-Scope': json.dumps({'scope_kind': 'shared_workspace', 'version': 1})}
    body = {'title': 'Normal report', 'markdown': 'A normal editable report.', 'kind': 'report',
            'payload': {}, 'citations': [], 'request_id': 'normal-report-create-0001'}
    saved = api.post('/api/experience/saved-work', headers=headers, json=body)
    assert saved.status_code == 200, saved.get_json()
    first = saved.get_json()['result']
    rejected = api.post(f"/api/experience/saved-work/{first['artifact']['artifact_id']}/revise",
        headers=headers, json={'request_id': 'normal-report-invalid-mode-0002',
            'expected_sha256': first['artifact']['sha256'], 'reason': 'Invalid mode on a normal report.',
            'continue_as_report': True, 'edits': [{'pointer': '/markdown',
                'expected': body['markdown'], 'value': 'Edited normal report.'}]})
    assert rejected.status_code == 400
    assert rejected.get_json()['code'] == 'WORK_VISUAL_CONTINUE_NOT_APPLICABLE'
    assert api.get(f"/api/experience/saved-work/{first['artifact']['artifact_id']}",
        headers=headers).get_json()['result']['document']['markdown'] == body['markdown']


def test_continue_as_report_rejects_typed_saved_work_contract(workspace):
    import json
    from experience_runtime import experience_bp

    ws = workspace
    ws['app'].config.update(TESTING=True, LUMENIS_OPEN_WORKSPACE=True)
    ws['app'].register_blueprint(experience_bp)
    api = ws['app'].test_client()
    headers = {'X-Lumenis-Workspace-Scope': json.dumps({'scope_kind': 'shared_workspace', 'version': 1})}
    value = document()
    markdown = visual_document_markdown(value)
    created = api.post('/api/experience/saved-work', headers=headers, json={
        'title': 'Typed visual report', 'markdown': markdown, 'kind': 'report',
        'payload': {'visual_document': value}, 'citations': [],
        'request_id': 'typed-visual-create-0001',
        'contract': {'id': 'lumenis.domain-report/v1', 'version': '1', 'required_payload_paths': []},
    })
    assert created.status_code == 200, created.get_json()
    first = created.get_json()['result']
    rejected = api.post(f"/api/experience/saved-work/{first['artifact']['artifact_id']}/revise",
        headers=headers, json={'request_id': 'typed-visual-continue-0002',
            'expected_sha256': first['artifact']['sha256'], 'reason': 'Typed reports stay with their owner.',
            'continue_as_report': True, 'edits': [{'pointer': '/title',
                'expected': 'Typed visual report', 'value': 'Converted title'}]})
    assert rejected.status_code == 400, rejected.get_json()
    assert rejected.get_json()['code'] == 'WORK_REVISION_OWNER_REQUIRED'
    reopened = api.get(f"/api/experience/saved-work/{first['artifact']['artifact_id']}", headers=headers)
    assert reopened.status_code == 200
    assert reopened.get_json()['result']['document']['payload']['visual_document'] == value


def test_http_large_original_image_full_object_reorder_versions_and_repeat_recovery(workspace):
    import json
    from experience_runtime import experience_bp
    from lumenis_app_api import _internal_pitch_body
    from lumenis.system.security import permission
    ws = workspace
    ws['app'].register_blueprint(experience_bp)
    api = ws['app'].test_client()
    headers = {'X-Lumenis-Workspace-Scope': json.dumps({'scope_kind': 'shared_workspace', 'version': 1})}
    image = io.BytesIO()
    Image.new('RGB', (768, 512), 'white').save(image, format='PNG', compress_level=0)
    original = image.getvalue()
    assert len(original) > 1024 * 1024
    value = document()
    value['assets'][0].update(sha256=hashlib.sha256(original).hexdigest(), data_base64=base64.b64encode(original).decode())
    payload = {'visual_document': value, 'retained_user_notes': 'Preserve my other approved work fields'}
    body = {'title': 'Original image and editable layers', 'markdown': visual_document_markdown(value),
            'kind': 'report', 'payload': payload, 'citations': [], 'request_id': 'large-visual-http-0001'}
    assert len(json.dumps(body).encode()) > 1024 * 1024
    with ws['app'].test_request_context('/bounded-default', method='POST', json=body):
        with pytest.raises(ValueError, match='INTERNAL_PITCH_REQUEST_INVALID'):
            _internal_pitch_body()
    ws['gate'].grant_permission(permission.PermissionLevel.HUMAN_APPROVAL)
    url = '/api/experience/saved-work'
    def confirm_save(values):
        prepared = api.post(url, headers=headers, json={**values, 'mode': 'prepare'})
        assert prepared.status_code == 200, prepared.get_json()
        resource = prepared.get_json()['result']['resource']
        response = api.post(url, headers=headers, json={**values, 'mode': 'confirm', 'confirm': True,
                                                       'confirmed_resource': resource})
        assert response.status_code == 200, response.get_json()
        return response.get_json()['result']
    first = confirm_save(body)
    first_id = first['artifact']['artifact_id']
    revised = copy.deepcopy(first['document']['payload'])
    layers = revised['visual_document']['layers']
    revised['visual_document']['layers'] = [layers[2], layers[0], layers[3], layers[1]]
    revised['visual_document']['layers'][0]['x'] = 35
    revised['visual_document']['layers'][3]['text'] = 'Author reordered the complete object structure'
    second_body = {**body, 'title': 'Reordered editable work', 'payload': revised,
        'markdown': visual_document_markdown(revised['visual_document']), 'request_id': 'large-visual-http-0002',
        'work_id': first['document']['work_id'], 'previous_artifact_id': first_id,
        'expected_draft_digest': first['document']['draft_digest'], 'contract': first['document']['contract']}
    second = confirm_save(second_body)
    assert second['document']['version'] == 2
    assert confirm_save(second_body)['artifact']['artifact_id'] == second['artifact']['artifact_id']
    for selected, expected in ((first, payload), (second, revised)):
        cold = ws['app'].test_client().get(url + '/' + selected['artifact']['artifact_id'], headers=headers)
        assert cold.status_code == 200, cold.get_json()
        stored = cold.get_json()['result']['document']['payload']
        assert stored == expected
        assert stored['retained_user_notes'] == payload['retained_user_notes']
        assert base64.b64decode(stored['visual_document']['assets'][0]['data_base64']) == original
    # The generic patch API must keep its container replacement restriction.
    rejected = api.post(url + '/' + second['artifact']['artifact_id'] + '/revise', headers=headers, json={
        'request_id': 'forbidden-container-replace-0003', 'expected_sha256': second['artifact']['sha256'],
        'reason': 'This requires the complete-object version command', 'edits': [{
            'pointer': '/payload/visual_document/layers', 'expected': revised['visual_document']['layers'],
            'value': value['layers']}]})
    assert rejected.status_code == 400, rejected.get_json()
