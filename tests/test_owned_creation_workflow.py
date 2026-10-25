"""Typed Shared creation regressions; the model replies are explicit fixtures.

Real Importer, Task, signed Evidence, Office objects and version transactions run
on disposable data. This is not evidence of a real model or professional quality.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

import ai_task_engine
import content_generation as creation
from lumenis.shared.content_creation import internal_pitch as pitch
from lumenis.shared.reporting import saved_work as work, business_brief
from lumenis.system.ai_gateway import core as gateway, owned_service
from tests import test_saved_document_workflow as source_helpers

source_workspace = source_helpers.workspace


@pytest.fixture(params=['cth', 'shared'])
def workspace(source_workspace, monkeypatch, request):
    ws = source_workspace
    if request.param == 'shared':
        ws['scope'] = {'scope_kind':'shared_workspace','version':1}
    for module in (pitch, ai_task_engine, creation):
        monkeypatch.setattr(module, 'get_repository', lambda: ws['repo'])
    with ws['repo'].transaction() as con:
        con.executescript(ai_task_engine._TASK_ENGINE_SCHEMA)
    pitch.engine.cache_clear()
    yield ws
    pitch.engine.cache_clear()


def arguments(ws, request_id='owned-creation-0001', **changes):
    args = {'actor':ws['actor'],'scope':ws['scope'],'manuscript':None,'profile':pitch.OWNED_PROFILE,
        'request_id':request_id,'generation':{'brief':{'title':'A careful proposal','question':'Explain a useful proposal; do not invent an audience result.',
            'language':'en','constraints':['Keep unknown outcomes explicit.']},'max_output_tokens':4096}}
    args.update(changes)
    return args


def run(ws, task):
    return pitch.run_internal_pitch(task['task_id'], actor=ws['actor'], scope=ws['scope'])


def original_source(ws):
    path = ws['root']/'self-authored-reference.txt'
    path.write_text('Original permitted wording.\nThe response from customers remains unknown.', encoding='utf-8')
    source_scope = ({'domain_id':'cth','subject_id':None,'brand_focus_id':None,'global_scope':True}
                    if ws['scope'].get('scope_kind') == 'shared_workspace' else ws['scope'])
    options = {'actor':ws['actor'],'scope':source_scope,'provenance':{
        'source_url':None,'published_at':None,'captured_at':'2026-10-25T00:00:00Z','rights':'owner-approved',
        'title':'Self-authored isolated reference','intended_use':'research_reference',
        'source_origin':{'kind':'approved-file','ref':'owned-creation-isolated-input'},
        'content_origin':{'kind':'self-authored','claim_type':'synthetic_sample'},
        'rights_statement':'Self-authored isolated test input; no external facts claimed.'}}
    resource = ws['documents'].document_import_resource(str(path), **options)
    ws['gate'].grant_approval('data_import', resource=resource, granted_by='isolated-owner-exact-import')
    return ws['documents'].import_document(str(path), **options)


def analysis(request):
    material = json.loads(request.user_payload)['material']
    selected = material['selected_section_ids']
    output = {'sections':[{'section_id':key,'heading':f'Revised {number}',
        'body':f'New paragraph {number}.\n\nIts outcome remains unverified.',
        'reason':'Clarify the proposed action while retaining uncertainty.'} for number,key in enumerate(selected or ['new-one','new-two'], 1)],
        'source_quotes':[{'source_index':0,'start':0,'end':27}] if material['sources'] else [],
        'limitations':['Model response is a test fixture; review remains required.']}
    if not selected:
        output.update(title='A careful proposal', summary='A proposed approach, pending review.')
    return output


def response(request, *, actor, output=None):
    before = creation.validate_creation_source_use(request, actor=actor)
    after = creation.validate_creation_source_use(request, actor=actor)
    value = {key:hashlib.sha256(('fixture:'+key).encode()).hexdigest() for key in business_brief._OWNED_HASH_FIELDS}
    value.update({key:'fixture-'+key for key in business_brief._OWNED_TEXT_FIELDS})
    value.update(ok=True,status='COMPLETED',provider='owned',operation=request.operation,domain_id=request.domain_id,scope=request.scope,
        request_id=request.request_id,request_digest=owned_service.owned_service_request_digest(request),
        input_revision_sha256=request.input_revision_sha256,source_binding_sha256=request.source_binding_sha256,
        source_use_policy_version=request.source_use_policy_version,source_use_policy_sha256=request.source_use_policy_sha256,
        source_evidence_refs=list(request.source_evidence_refs),real_input_bound=True,response_schema_state='SUBSET_VALIDATED',
        source_audit_ids={'scope_before':before.scope_audit_id,'source_before':before.source_audit_id,
                          'scope_after':after.scope_audit_id,'source_after':after.source_audit_id},
        admission_audit_hashes={'fixture':hashlib.sha256(b'fixture-admission').hexdigest()},
        evaluation_state='NOT_RUN',quality_state='NOT_EVALUATED',promotion_status='NOT_PROMOTED',product_acceptance=False,
        context_window_tokens=65536,output_token_limit=32768,prompt_tokens=100,completion_tokens=400,total_tokens=500,
        analysis=output if output is not None else analysis(request))
    return value


def fake_gateway(ws, monkeypatch, transform=None):
    calls = []
    def execute(request, **kwargs):
        assert kwargs['is_cancelled']() is False
        calls.append(request)
        output = analysis(request)
        if transform:
            transform(request, output)
        return response(request, actor=ws['actor'], output=output)
    monkeypatch.setattr(gateway, 'run_owned_request', execute)
    monkeypatch.setattr(gateway, 'chat', lambda *a, **k: pytest.fail('generic chat is never a creation fallback'))
    return calls


def test_signed_source_save_reopen_export_and_committed_replay(workspace, monkeypatch):
    ws = workspace
    source = original_source(ws)
    calls = fake_gateway(ws, monkeypatch)
    args = arguments(ws)
    args['generation']['source_evidence_refs'] = [source['evidence_id']]
    task = pitch.prepare_internal_pitch(**args)
    saved = run(ws, task)
    assert saved['state'] == 'succeeded' and len(calls) == 1
    assert saved['work']['document']['contract'] == pitch.AUTHORING_CONTRACT
    payload = saved['work']['document']['payload']
    assert payload['truth_state']['owned_model_execution'] == 'COMPLETED'
    assert payload['truth_state']['professional_quality'] == 'NOT_EVALUATED'
    assert payload['creation']['source_bindings'][0]['evidence_id'] == source['evidence_id']
    assert [item['section_id'] for item in payload['creation']['changes']] == [item['section_id'] for item in saved['revision_sections']]
    assert saved['creation_evidence']['integrity_digest']
    artifact = saved['action_receipt']['artifact_id']
    reopened = work.reopen_work(artifact, actor=ws['actor'], scope=ws['scope'])
    assert reopened['revision_sections'] == saved['revision_sections']
    assert reopened['creation_evidence'] == saved['creation_evidence']
    output = pitch.download_internal_pitch(task['task_id'], actor=ws['actor'], scope=ws['scope'])
    assert 'New paragraph 1.' in [p.text for p in Document(io.BytesIO(output['content'])).paragraphs]
    assert run(ws, task)['action_receipt'] == saved['action_receipt']
    assert pitch.prepare_internal_pitch(**args)['action_receipt'] == saved['action_receipt']
    assert len(calls) == 1
    # Retained integration DTO for the frontend owner; contains synthetic data only.
    if ws['scope'].get('scope_kind') == 'shared_workspace':
        directory = Path('work/product-closure-20261011/workspace-20261024')
        directory.mkdir(parents=True, exist_ok=True)
        (directory/'owned-creation-hermetic-response-20261025.json').write_text(json.dumps(
            {'evidence_scope':'HERMETIC_GATEWAY_FIXTURE_NOT_MODEL_OR_USER_ACCEPTANCE','result':saved}, ensure_ascii=False, indent=2), encoding='utf-8')


def test_retained_gateway_reply_cannot_change_provenance_binding(workspace, monkeypatch):
    ws = workspace
    task = pitch.prepare_internal_pitch(**arguments(ws))
    def fake(request, **kwargs):
        result = response(request, actor=ws['actor'])
        result['input_revision_sha256'] = 'f'*64
        return result
    monkeypatch.setattr(gateway, 'run_owned_request', fake)
    with pytest.raises(ValueError, match='OWNED_RESULT_BINDING_INVALID'):
        run(ws, task)
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == 0


def authored_parent(ws):
    manuscript = {'title':'Preserve this title','summary':'Preserve the human summary.','language':'en','authorship':'user_supplied',
        'sections':[{'heading':'Original first','body':'The first human paragraph stays unchanged.'},
                    {'heading':'Original second','body':'Improve only this paragraph.'},
                    {'heading':'Original last','body':'The last human paragraph stays unchanged.'}]}
    task = pitch.prepare_internal_pitch(manuscript=manuscript, request_id='human-parent-0001', profile=pitch.AUTHORING_PROFILE,
                                       actor=ws['actor'], scope=ws['scope'])
    return run(ws, task)


def revision(ws, parent, request_id='selected-revision-0002'):
    doc = parent['work']['document']
    args = arguments(ws, request_id, work_id=doc['work_id'], previous_artifact_id=parent['action_receipt']['artifact_id'],
                     expected_draft_digest=doc['draft_digest'])
    args['generation']['selected_section_ids'] = [parent['revision_sections'][1]['section_id']]
    return args


def test_selected_revision_preserves_human_text_and_original_office_parts(workspace, monkeypatch):
    ws = workspace
    first = authored_parent(ws)
    before = pitch.download_internal_pitch(first['task_id'], actor=ws['actor'], scope=ws['scope'])['content']
    calls = fake_gateway(ws, monkeypatch)
    args = revision(ws, first)
    second = run(ws, pitch.prepare_internal_pitch(**args))
    original = first['work']['document']['payload']['manuscript']
    edited = second['work']['document']['payload']['manuscript']
    assert edited['title'] == original['title'] and edited['summary'] == original['summary']
    assert edited['sections'][0] == original['sections'][0] and edited['sections'][2] == original['sections'][2]
    assert edited['sections'][1]['body'].startswith('New paragraph 1.')
    assert second['action_receipt']['version'] == 2
    assert second['action_receipt']['parent_artifact_id'] == first['action_receipt']['artifact_id']
    assert second['work']['document']['payload']['creation']['office_edit_receipt']['modified_parts'] == ['word/document.xml']
    after = pitch.download_internal_pitch(second['task_id'], actor=ws['actor'], scope=ws['scope'])['content']
    with zipfile.ZipFile(io.BytesIO(before)) as prior, zipfile.ZipFile(io.BytesIO(after)) as current:
        assert set(prior.namelist()) == set(current.namelist())
        for name in prior.namelist():
            if name != 'word/document.xml':
                assert prior.read(name) == current.read(name)
    # Repeat preparation after the child commits must recover, not infer again.
    assert pitch.prepare_internal_pitch(**args)['action_receipt'] == second['action_receipt']
    assert pitch.read_internal_pitch(first['task_id'], actor=ws['actor'], scope=ws['scope'])['action_receipt'] == first['action_receipt']
    assert len(calls) == 1


@pytest.mark.parametrize('language', ['zh','en','ko','ja','fr','de','es','pt','ru','ar','hi','id'])
def test_source_free_brief_is_independent_and_exports_twelve_languages(workspace, monkeypatch, language):
    ws = workspace
    calls = fake_gateway(ws, monkeypatch)
    args = arguments(ws)
    args['generation']['brief']['language'] = language
    saved = run(ws, pitch.prepare_internal_pitch(**args))
    assert calls[0].source_evidence_refs[0].startswith('internal-pitch-brief:')
    assert saved['work']['document']['payload']['creation']['source_bindings'] == []
    assert saved['work']['document']['payload']['truth_state']['source_admission'] == 'USER_BRIEF_ONLY'
    raw = pitch.download_internal_pitch(saved['task_id'], actor=ws['actor'], scope=ws['scope'])['content']
    assert Document(io.BytesIO(raw)).core_properties.language == language


def test_stale_parent_rejected_before_model_call(workspace, monkeypatch):
    ws = workspace
    first = authored_parent(ws)
    calls = fake_gateway(ws, monkeypatch)
    stale = pitch.prepare_internal_pitch(**revision(ws, first, 'stale-revision-0003'))
    run(ws, pitch.prepare_internal_pitch(**revision(ws, first)))
    with pytest.raises(ValueError, match='WORK_REVISION_CONFLICT'):
        run(ws, stale)
    assert len(calls) == 1


@pytest.mark.parametrize('fault', ['duplicate_section','unknown_section','bad_quote','unselected_title'])
def test_invalid_generated_changes_never_commit(workspace, monkeypatch, fault):
    ws = workspace
    first = authored_parent(ws)
    def corrupt(request, output):
        if fault == 'duplicate_section':
            output['sections'].append(copy.deepcopy(output['sections'][0]))
        elif fault == 'unknown_section':
            output['sections'][0]['section_id'] = 'not-selected'
        elif fault == 'bad_quote':
            output['source_quotes'] = [{'source_index':0,'start':0,'end':1}]
        else:
            output['title'] = 'A title the user did not select'
    fake_gateway(ws, monkeypatch, corrupt)
    task = pitch.prepare_internal_pitch(**revision(ws, first))
    with pytest.raises(ValueError, match='CREATION_OUTPUT'):
        run(ws, task)
    assert pitch.read_internal_pitch(task['task_id'], actor=ws['actor'], scope=ws['scope'])['state'] == 'failed'
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == 2


def test_cancellation_during_generation_commits_no_work(workspace, monkeypatch):
    ws = workspace
    task = pitch.prepare_internal_pitch(**arguments(ws))
    def cancel(request, **kwargs):
        pitch.cancel_internal_pitch(task['task_id'], actor=ws['actor'], scope=ws['scope'])
        kwargs['is_cancelled']()
        pytest.fail('cancel must interrupt the original Task worker')
    monkeypatch.setattr(gateway, 'run_owned_request', cancel)
    with pytest.raises(ValueError, match='INTERNAL_PITCH_CANCELLED'):
        run(ws, task)
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == 0
    assert pitch.read_internal_pitch(task['task_id'], actor=ws['actor'], scope=ws['scope'])['state'] == 'cancelled'


def test_unavailable_is_not_manual_authorship_or_success(workspace, monkeypatch):
    ws = workspace
    task = pitch.prepare_internal_pitch(**arguments(ws))
    monkeypatch.setattr(gateway, 'run_owned_request', lambda *a, **k:{'ok':False,'code':'OWNED_SERVICE_EXECUTOR_NOT_BOUND'})
    with pytest.raises(RuntimeError, match='OWNED_SERVICE_EXECUTOR_NOT_BOUND'):
        run(ws, task)
    assert pitch.read_internal_pitch(task['task_id'], actor=ws['actor'], scope=ws['scope'])['state'] == 'failed'
    assert ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == 0


def test_validator_rejects_altered_request_and_signed_execution_tamper(workspace, monkeypatch):
    ws = workspace
    calls = fake_gateway(ws, monkeypatch)
    captured = []
    original = gateway.run_owned_request
    def check(request, **kwargs):
        with pytest.raises(PermissionError, match='CREATION_INPUT_OR_SOURCE_CHANGED'):
            creation.validate_creation_source_use(replace(request, input_revision_sha256='f'*64), actor=ws['actor'])
        resource = f'content-creation-ai:v1:{request.job_id}:{request.input_revision_sha256}'
        assert creation.validate_creation_ai_resource(resource, actor=ws['actor'], scope=ws['scope'])
        assert not creation.validate_creation_ai_resource(resource[:-1]+'x', actor=ws['actor'], scope=ws['scope'])
        captured.append(request)
        return original(request, **kwargs)
    monkeypatch.setattr(gateway, 'run_owned_request', check)
    saved = run(ws, pitch.prepare_internal_pitch(**arguments(ws)))
    with ws['repo'].transaction() as con:
        con.execute('UPDATE app_evidence_records SET outputs_json=? WHERE id=?', ('{}',saved['creation_evidence']['evidence_id']))
    with pytest.raises(ValueError, match='CREATION_EXECUTION_EVIDENCE_INVALID'):
        work.reopen_work(saved['action_receipt']['artifact_id'], actor=ws['actor'], scope=ws['scope'])
    assert len(calls) == len(captured) == 1


def test_private_uploaded_original_uses_existing_custody_version(workspace, monkeypatch):
    ws = workspace
    if ws['scope'].get('scope_kind') != 'shared_workspace':
        pytest.skip('Private uploaded originals use the exact independent Shared scope')
    from lumenis.importer import documents
    path = ws['root']/'private-original.docx'
    original = Document()
    original.add_heading('Private self-authored source', level=1)
    original.add_paragraph('Original permitted wording. Customer reaction remains unknown.')
    original.save(path)
    source = documents.save_private_document(str(path), actor=ws['actor'], scope=ws['scope'], request_id='private-source-0001')
    assert source['document']['payload']['source_material']['rights'] == 'UNKNOWN'
    calls = fake_gateway(ws, monkeypatch)
    args = arguments(ws)
    args['generation']['source_artifacts'] = [{'artifact_id':source['artifact']['artifact_id'],'sha256':source['artifact']['sha256']}]
    saved = run(ws, pitch.prepare_internal_pitch(**args))
    binding = saved['work']['document']['payload']['creation']['source_bindings'][0]
    assert binding['source_kind'] == 'private_saved_source' and binding['rights'] == 'UNKNOWN'
    assert binding['source_artifact_sha256'] == source['artifact']['sha256']
    assert binding['external_redistribution'] is False and binding['training_eligible'] is False
    assert calls[0].source_evidence_refs[1].startswith('saved_work:')
    assert saved['work']['document']['citations'] == []
    assert saved['work']['document']['payload']['creation']['private_source_quotes'][0]['quote']
    assert work.reopen_work(saved['action_receipt']['artifact_id'], actor=ws['actor'], scope=ws['scope'])['creation_evidence']
    Path('work/product-closure-20261011/workspace-20261024/owned-creation-private-hermetic-response-20261025.json').write_text(
        json.dumps({'evidence_scope':'HERMETIC_GATEWAY_FIXTURE_NOT_MODEL_OR_USER_ACCEPTANCE','source':{
            'artifact_id':source['artifact']['artifact_id'],'sha256':source['artifact']['sha256']},'result':saved},
            ensure_ascii=False,indent=2),encoding='utf-8')
    wrong = {**args, 'generation':copy.deepcopy(args['generation'])}
    wrong['request_id'] = 'changed-source-0002'
    wrong['generation']['source_artifacts'][0]['sha256'] = 'f'*64
    with pytest.raises(PermissionError, match='CREATION_SOURCE_ARTIFACT_CHANGED'):
        pitch.prepare_internal_pitch(**wrong)
    assert len(calls) == 1


def test_missing_execution_signer_rejects_before_inference(workspace, monkeypatch):
    ws = workspace
    calls = fake_gateway(ws, monkeypatch)
    task = pitch.prepare_internal_pitch(**arguments(ws))
    monkeypatch.delenv('LUMENIS_EVIDENCE_HMAC_KEY_RING')
    with pytest.raises(RuntimeError, match='CREATION_EXECUTION_SIGNING_UNAVAILABLE'):
        run(ws, task)
    assert not calls and ws['repo'].query_one('SELECT COUNT(*) FROM lumenis_shared_artifacts')[0] == 0


def test_manual_twelve_language_manuscript_stays_manual(workspace):
    ws = workspace
    manuscript = {'title':'قصة قصيرة','summary':'مخطوطة قدّمها المستخدم','language':'ar','authorship':'user_supplied',
                  'sections':[{'heading':'مقدمة','body':'هذه فقرة كتبها المستخدم.'}]}
    task = pitch.prepare_internal_pitch(manuscript=manuscript, request_id='arabic-authored-0001', profile=pitch.AUTHORING_PROFILE,
                                       actor=ws['actor'], scope=ws['scope'])
    saved = run(ws, task)
    assert saved['work']['document']['payload']['truth_state']['owned_model_execution'] == 'NOT_RUN'
    assert 'creation' not in saved['work']['document']['payload']


def test_actual_api_uses_existing_task_routes_and_unavailable_recovery(workspace, monkeypatch):
    from tests.test_internal_pitch_workflow import client as client_fixture
    ws = workspace
    client = client_fixture.__wrapped__(ws)
    args = arguments(ws)
    args.pop('actor')
    prefix = '/api/experience/creation/internal-pitches'
    pending = client.post(prefix, json=args)
    assert pending.status_code == 202
    task_id = pending.get_json()['result']['task_id']
    monkeypatch.setattr(gateway, 'run_owned_request', lambda *a, **k:{'ok':False,'code':'OWNED_SERVICE_EXECUTOR_NOT_BOUND'})
    unavailable = client.post(prefix+'/'+task_id+'/run', json={'scope':ws['scope']})
    assert unavailable.status_code == 503 and unavailable.get_json()['code'] == 'OWNED_SERVICE_EXECUTOR_NOT_BOUND'
    calls = fake_gateway(ws, monkeypatch)
    recovered = client.post(prefix+'/'+task_id+'/run', json={'scope':ws['scope']})
    assert recovered.status_code == 200 and recovered.get_json()['result']['state'] == 'succeeded'
    query = {'scope_kind':'shared_workspace','version':'1'} if ws['scope'].get('scope_kind') else {
        key:str(value).lower() if type(value) is bool else value or '' for key,value in ws['scope'].items()}
    reopened = client.get(prefix+'/'+task_id, query_string=query)
    assert reopened.status_code == 200
    assert reopened.get_json()['result']['creation_evidence'] == recovered.get_json()['result']['creation_evidence']
    assert len(calls) == 1


def test_unquoted_source_revocation_in_artifact_transaction_is_not_ignored(workspace, monkeypatch):
    ws = workspace
    source = original_source(ws)
    args = arguments(ws)
    args['generation']['source_evidence_refs'] = [source['evidence_id']]
    fake_gateway(ws, monkeypatch, lambda _request, output:output.update(source_quotes=[]))
    task = pitch.prepare_internal_pitch(**args)
    original_save = work.save_work
    def changing(*args, **kwargs):
        commit = kwargs['before_commit']
        def before_commit(con, artifact, document):
            con.execute('UPDATE app_evidence_records SET collection_metadata=? WHERE id=?', ('{}',source['evidence_id']))
            commit(con, artifact, document)
        return original_save(*args, **{**kwargs,'before_commit':before_commit})
    monkeypatch.setattr(work, 'save_work', changing)
    before = tuple(ws['repo'].query_one('SELECT COUNT(*) FROM '+table)[0] for table in ('lumenis_shared_artifacts','app_evidence_records'))
    with pytest.raises(ValueError, match='WORK_EVIDENCE_CONTEXT_MISMATCH'):
        run(ws, task)
    assert tuple(ws['repo'].query_one('SELECT COUNT(*) FROM '+table)[0] for table in ('lumenis_shared_artifacts','app_evidence_records')) == before
