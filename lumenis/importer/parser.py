"""
文件解析器
支持 PDF / Word(DOCX) / Excel(XLSX) / CSV / TXT
不依赖重型库，优先用标准库+可选依赖
"""
import os
import re
import hashlib
import json
from typing import List, Dict
from ..config import IMPORTER_CONFIG


class ParseResult:
    """解析结果"""
    def __init__(self):
        self.text: str = ""
        self.tables: List[List[List[str]]] = []  # 表格数据
        self.metadata: Dict = {}
        self.paragraphs: List[str] = []
        self.locators: List[Dict] = []
        self.source_tables: List[Dict] = []
        self.source_blocks: List[Dict] = []


def parse_file(file_path: str) -> ParseResult:
    """
    根据文件扩展名自动选择解析器
    返回 ParseResult
    """
    if not isinstance(file_path, str) or not file_path.strip() or "://" in file_path:
        raise ValueError("LOCAL_DOCUMENT_REQUIRED")
    if not os.path.isfile(file_path):
        raise FileNotFoundError(file_path)
    if os.path.getsize(file_path) > IMPORTER_CONFIG["max_file_size_mb"] * 1024 * 1024:
        raise ValueError("DOCUMENT_TOO_LARGE")
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        result = parse_pdf(file_path)
    elif ext == ".docx":
        result = parse_word(file_path)
    elif ext == ".doc":
        result = parse_legacy_word(file_path)
    elif ext == ".xlsx":
        result = parse_excel(file_path)
    elif ext == ".xls":
        result = parse_legacy_excel(file_path)
    elif ext == ".csv":
        result = parse_csv(file_path)
    elif ext in (".txt", ".json", ".xml", ".xbrl"):
        result = parse_txt(file_path)
    elif ext in (".pptx", ".html", ".htm"):
        result = parse_markitdown(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")
    digest = hashlib.sha256()
    with open(file_path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    result.metadata.update(source_name=os.path.basename(file_path),
                           source_fingerprint=digest.hexdigest(), network_mode="local_only")
    # Offsets refer to the exact retained parsed text, not guessed PDF pages or
    # source-byte positions. Keep repeated paragraphs independently locatable.
    if not result.locators:
        offset = 0
        for index, line in enumerate(result.text.splitlines(keepends=True), 1):
            text = line.rstrip("\r\n")
            if text.strip():
                result.locators.append({"unit": "parsed_line", "index": index,
                                        "start": offset, "end": offset + len(text)})
            offset += len(line)
    result.metadata["parsed_text_sha256"] = hashlib.sha256(result.text.encode("utf-8")).hexdigest()
    return result


def parse_markitdown(file_path: str) -> ParseResult:
    """Convert one bounded local document through optional MarkItDown.

    MarkItDown is an importer implementation, not a source, Evidence or
    Canonical owner.  It is deliberately lazy and local-only: a missing
    optional dependency is reported honestly and remote URI inputs are never
    handed to a third-party converter.
    """
    if not isinstance(file_path, str) or not file_path.strip():
        raise ValueError("文档路径为空")
    local_path = os.path.abspath(file_path)
    if not os.path.isfile(local_path):
        raise FileNotFoundError(local_path)
    if os.path.getsize(local_path) > IMPORTER_CONFIG["max_file_size_mb"] * 1024 * 1024:
        raise ValueError(f"文件超过{IMPORTER_CONFIG['max_file_size_mb']}MB")

    try:
        from markitdown import MarkItDown
    except ImportError as exc:
        raise ImportError(
            "PPTX/HTML 解析需要可选依赖 markitdown；请按 requirements.txt 安装后重试"
        ) from exc

    result = ParseResult()
    result.metadata.update({
        "format": os.path.splitext(local_path)[1].lstrip(".").lower(),
        "parser": "markitdown",
        "network_mode": "local_only",
        "source_name": os.path.basename(local_path),
    })
    digest = hashlib.sha256()
    with open(local_path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    result.metadata["source_fingerprint"] = digest.hexdigest()

    converter = MarkItDown(enable_plugins=False)
    convert_local = getattr(converter, "convert_local", None)
    converted = convert_local(local_path) if callable(convert_local) else converter.convert(local_path)
    markdown = getattr(converted, "markdown", None)
    if markdown is None:
        markdown = getattr(converted, "text_content", None)
    if not isinstance(markdown, str) or not markdown.strip():
        raise ValueError("MarkItDown 未返回可用文本")
    result.text = markdown
    result.paragraphs = [line.strip() for line in markdown.splitlines() if line.strip()]
    version = getattr(__import__("markitdown"), "__version__", None)
    if version is None:
        from importlib.metadata import version as distribution_version
        version = distribution_version("markitdown")
    result.metadata["parser_version"] = str(version)
    return result



def decode_local_text(raw: bytes):
    """Bounded, deterministic local decoding; never rewrite original bytes.

    A successful non-BOM fallback is an assumption, not proof of source encoding.
    Binary controls and invalid code units are rejected, not silently replaced.
    """
    import unicodedata
    if not isinstance(raw, bytes) or not 0 < len(raw) <= 8 * 1024 * 1024:
        raise ValueError("DOCUMENT_TEXT_SIZE_INVALID")
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encodings, certainty = ("utf-32",), "bom_declared"
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings, certainty = ("utf-16",), "bom_declared"
    elif raw.startswith(b"\xef\xbb\xbf"):
        encodings, certainty = ("utf-8-sig",), "bom_declared"
    else:
        encodings, certainty = ("utf-8", "gb18030"), "assumed_decoder"
    for encoding in encodings:
        try:
            text = raw.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
        if any(unicodedata.category(char) in {"Cc", "Cs"} and char not in "\t\n\r\f" for char in text):
            continue
        warnings = [] if certainty == "bom_declared" else ["Encoding inferred by deterministic decoder order; source encoding not independently verified."]
        return text, {"encoding": encoding, "encoding_state": certainty, "decoding_warnings": warnings,
                      "original_bytes_modified": False}
    raise ValueError("DOCUMENT_TEXT_ENCODING_OR_CONTROL_INVALID")

def parse_txt(file_path: str) -> ParseResult:
    """纯文本解析"""
    result = ParseResult()
    result.metadata["format"] = "txt"
    with open(file_path, "rb") as f:
        result.text, decoding = decode_local_text(f.read(8 * 1024 * 1024 + 1))
    result.metadata.update(decoding)
    result.paragraphs = [p.strip() for p in result.text.split("\n") if p.strip()]
    return result


def parse_csv(file_path: str) -> ParseResult:
    """Parse one bounded CSV without silently treating another dialect as one column.

    The existing entry has no delimiter declaration. Infer only a unique,
    rectangular multi-column interpretation among comma, semicolon and tab,
    using the actual CSV quote/newline grammar over the entire bounded input.
    A declaration-free single-column file remains supported. Ambiguous dialects,
    inconsistent columns and malformed quoting fail explicitly; none is padded,
    flattened into a successful column, or treated as an encoding decision.
    """
    import csv
    import io

    result = ParseResult()
    result.metadata["format"] = "csv"
    with open(file_path, "rb") as source:
        text, decoding = decode_local_text(source.read(8 * 1024 * 1024 + 1))
    result.metadata.update(decoding)
    candidates = []
    saw_multiple_columns = False
    saw_parse_error = False
    # Three linear, bounded scans keep only shape summaries, not three tables.
    # Blank records are retained but do not decide the field width.
    for delimiter in (",", ";", "\t"):
        widths = set()
        try:
            for row in csv.reader(io.StringIO(text, newline=""),
                                  delimiter=delimiter, strict=True):
                if row:
                    widths.add(len(row))
            multiple = any(width > 1 for width in widths)
            saw_multiple_columns = saw_multiple_columns or multiple
            if len(widths) == 1 and multiple:
                candidates.append(delimiter)
        except csv.Error:
            saw_parse_error = True
    if len(candidates) > 1:
        raise ValueError("CSV_DIALECT_AMBIGUOUS")
    if candidates:
        delimiter = candidates[0]
        dialect_state = "inferred_unique_rectangular"
    elif saw_multiple_columns:
        raise ValueError("CSV_COLUMNS_INCONSISTENT")
    elif saw_parse_error:
        raise ValueError("CSV_DIALECT_UNDETERMINED")
    else:
        delimiter = ","
        dialect_state = "single_column_or_empty"
    try:
        table = list(csv.reader(io.StringIO(text, newline=""),
                                delimiter=delimiter, strict=True))
    except csv.Error as exc:
        raise ValueError("CSV_PARSE_INVALID") from exc
    result.metadata.update(csv_delimiter=delimiter,
                           csv_dialect_state=dialect_state,
                           csv_dialect_candidates=[",", ";", "\t"],
                           csv_row_count=len(table),
                           csv_column_count=next((len(row) for row in table if row), 0),
                           csv_blank_record_count=sum(not row for row in table),
                           csv_missing_cell_count=sum(cell == "" for row in table for cell in row),
                           csv_dialect_warnings=[
                               "Delimiter inferred from bounded CSV structure; not a publisher declaration."
                           ])
    result.tables = [table]
    # Preserve the established comma-normalized text/locator contract. Original
    # bytes remain separately retained; tables preserve quotes/newlines/empties.
    result.text = "\n".join([",".join(row) for row in table])
    result.paragraphs = [",".join(row) for row in table if any(cell.strip() for cell in row)]
    return result


def parse_pdf(file_path: str) -> ParseResult:
    """
    PDF解析
    优先用 pdfplumber，回退到 PyPDF2，都没有则报错提示安装
    """
    result = ParseResult()
    result.metadata["format"] = "pdf"
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            result.metadata["pages"] = len(pdf.pages)
            result.metadata["parser_version"] = "typed-source-v2"
            result.metadata["document_block_projection"] = "pdf-page-text-block-v1"
            all_text = []
            page_start = 0
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                all_text.append(text)
                text_blocks = _pdf_text_block_locators(page, text)
                locator = {
                    "unit": "pdf_page", "index": i + 1, "page_index": i + 1,
                    "source_category": "pdf_page",
                    "start": page_start, "end": page_start + len(text),
                    "text_blocks": text_blocks,
                }
                result.locators.append(locator)
                result.source_blocks.append({
                    "source_category": "pdf_page",
                    "text": text,
                    "source_fields": [],
                    "locator": locator,
                })
                page_start += len(text) + 2
                # 提取表格
                tables = page.extract_tables()
                for t in tables:
                    result.tables.append(t)
            result.text = "\n\n".join(all_text)
            result.paragraphs = [p.strip() for p in result.text.split("\n") if p.strip()]
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            result.metadata["pages"] = len(reader.pages)
            result.metadata["parser_version"] = "typed-source-v2"
            result.metadata["document_block_projection"] = "pdf-page-text-block-v1"
            all_text = []
            page_start = 0
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                all_text.append(text)
                locator = {
                    "unit": "pdf_page", "index": i + 1, "page_index": i + 1,
                    "source_category": "pdf_page",
                    "start": page_start, "end": page_start + len(text),
                    "text_blocks": _pdf_text_block_locators(None, text),
                }
                result.locators.append(locator)
                result.source_blocks.append({
                    "source_category": "pdf_page",
                    "text": text,
                    "source_fields": [],
                    "locator": locator,
                })
                page_start += len(text) + 2
            result.text = "\n\n".join(all_text)
            result.paragraphs = [p.strip() for p in result.text.split("\n") if p.strip()]
        except ImportError:
            raise ImportError("PDF解析需要 pdfplumber 或 PyPDF2，请运行: pip install pdfplumber")
    return result


def _pdf_text_block_locators(page, page_text: str) -> list[dict]:
    """Return page-local line locators without changing retained page text."""
    line_records = None
    extract_lines = getattr(page, "extract_text_lines", None) if page is not None else None
    if callable(extract_lines):
        try:
            line_records = extract_lines(return_chars=False)
        except (TypeError, NotImplementedError):
            line_records = None

    if not isinstance(line_records, list):
        line_records = [{"text": line} for line in (page_text or "").splitlines()]

    locators = []
    search_from = 0
    for index, record in enumerate(line_records, 1):
        if not isinstance(record, dict):
            continue
        text = record.get("text")
        if not isinstance(text, str):
            continue
        start = (page_text or "").find(text, search_from) if text else -1
        if start < 0:
            start = None
            end = None
        else:
            end = start + len(text)
            search_from = end
        block = {
            "unit": "pdf_text_block",
            "offset_basis": "page_text",
            "index": index,
            "start": start,
            "end": end,
            "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        }
        coordinates = {}
        for key in ("x0", "top", "x1", "bottom"):
            value = record.get(key)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                coordinates[key] = value
        if coordinates:
            block["bbox"] = coordinates
        locators.append(block)
    return locators


def parse_word(file_path: str) -> ParseResult:
    """
    Word DOCX 解析
    需要 python-docx
    """
    result = ParseResult()
    result.metadata["format"] = "word"
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(file_path)
        result.metadata["parser_version"] = "typed-source-v2"
        result.metadata["document_block_projection"] = "docx-body-block-v1"
        all_text = []
        offset, paragraph_index, table_index = 0, 0, 0
        for element in doc.element.body:
            if element.tag.endswith("}p"):
                paragraph_index += 1
                text = Paragraph(element, doc).text
                block_index = len(result.source_blocks) + 1
                start = offset + (2 if all_text else 0)
                end = start + len(text)
                locator = {"unit": "word_paragraph", "source_category": "docx_paragraph",
                           "index": paragraph_index,
                           "block_index": block_index, "start": start, "end": end}
                result.locators.append(locator)
                result.source_blocks.append({
                    "source_category": "docx_paragraph",
                    "text": text,
                    "source_fields": [],
                    "locator": locator,
                })
                if all_text:
                    offset += 2
                all_text.append(text)
                result.paragraphs.append(text)
                offset += len(text)
            elif element.tag.endswith("}tbl"):
                table_index += 1
                rows = [[cell.text for cell in row.cells] for row in Table(element, doc).rows]
                result.tables.append(rows)
                for row_index, cells in enumerate(rows, 1):
                    text = " | ".join(cells)
                    if all_text:
                        offset += 2
                    block_index = len(result.source_blocks) + 1
                    row_locator = {
                        "unit": "word_table_row", "table": table_index,
                        "source_category": "docx_table_row",
                        "row": row_index, "block_index": block_index,
                        "start": offset, "end": offset + len(text),
                        "cell_count": len(cells),
                    }
                    result.locators.append(row_locator)
                    result.source_blocks.append({
                        "source_category": "docx_table_row",
                        "text": text,
                        "source_fields": [
                            {"column_index": column, "value": cell,
                             "cell_type": "string"}
                            for column, cell in enumerate(cells, 1)
                        ],
                        "locator": row_locator,
                    })
                    cell_offset = offset
                    for column, cell in enumerate(cells, 1):
                        result.locators.append({"unit": "word_table_cell", "table": table_index,
                                                "row": row_index, "column": column,
                                                "start": cell_offset, "end": cell_offset + len(cell)})
                        cell_offset += len(cell) + 3
                    all_text.append(text)
                    offset += len(text)
        result.text = "\n\n".join(all_text)
    except ImportError:
        raise ImportError("Word解析需要 python-docx，请运行: pip install python-docx")
    return result


def parse_excel(file_path: str) -> ParseResult:
    """
    Excel XLSX 解析
    需要 openpyxl（不以 pandas 推断替代物理单元格）
    """
    result = ParseResult()
    result.metadata["format"] = "excel"
    try:
        import openpyxl
        # Retain formula expressions as source statements. A missing cached
        # calculation is never silently converted to a blank or zero.
        wb = openpyxl.load_workbook(file_path, data_only=False, read_only=True, keep_links=False)
        try:
            result.metadata["sheets"] = wb.sheetnames
            result.metadata["parser_version"] = "typed-source-v2"
            result.metadata["formula_evaluation"] = "never"
            result.metadata["excel_row_projection"] = "typed-xlsx-v1"
            all_text = []
            offset = 0
            retained_cells = 0
            # An 8 MB sparse workbook can declare billions of physical cells.
            # Reject over-budget preservation explicitly; never silently trim it.
            cell_limit = 100_000
            extents = [_actual_excel_extent(wb, sheet, cell_limit) for sheet in wb]
            if sum(rows * columns for rows, columns in extents) > cell_limit:
                raise ValueError("DOCUMENT_EXCEL_CELL_LIMIT")
            for table_index, sheet_name in enumerate(wb.sheetnames):
                ws = wb[sheet_name]
                table = []
                source_rows = []
                max_row, max_column = extents[table_index]
                physical_rows = (ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column)
                                 if max_row else ())
                for row_index, row in enumerate(physical_rows, 1):
                    retained_cells += len(row)
                    if retained_cells > cell_limit:
                        raise ValueError("DOCUMENT_EXCEL_CELL_LIMIT")
                    typed_values = [_typed_excel_cell(cell) for cell in row]
                    row_data = [
                        _legacy_excel_cell_text(cell.value) for cell in row
                    ]
                    source_row = {
                        "row_index": row_index,
                        "values": typed_values,
                    }
                    source_content_payload = json.dumps(
                        {"values": typed_values}, ensure_ascii=False,
                        sort_keys=True, separators=(",", ":"),
                    )
                    source_row["source_content_sha256"] = hashlib.sha256(
                        source_content_payload.encode("utf-8"),
                    ).hexdigest()
                    source_row_payload = json.dumps(
                        {"sheet_name": sheet_name, "table_index": table_index, **source_row},
                        ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                    )
                    source_row["source_row_sha256"] = hashlib.sha256(
                        source_row_payload.encode("utf-8"),
                    ).hexdigest()
                    source_rows.append(source_row)
                    table.append(row_data)
                    text = " | ".join(row_data)
                    if all_text:
                        offset += 1
                    cell_offset = offset
                    for column, cell in enumerate(row_data, 1):
                        result.locators.append({"unit": "spreadsheet_cell", "sheet": sheet_name,
                                                "row": row_index, "column": column,
                                                "start": cell_offset, "end": cell_offset + len(cell)})
                        cell_offset += len(cell) + 3
                    all_text.append(text)
                    offset += len(text)
                result.tables.append(table)
                result.source_tables.append({
                    "sheet_name": sheet_name,
                    "table_index": table_index,
                    "sheet_state": ws.sheet_state,
                    "rows": source_rows,
                })
            result.text = "\n".join(all_text)
            result.paragraphs = all_text
        finally:
            wb.close()

    except ImportError as exc:
        raise ImportError("Excel typed source preservation requires openpyxl") from exc

    return result


def _actual_excel_extent(workbook, worksheet, cell_limit):
    """Read actual XML positions, never trust the advisory dimension attribute.

    Bound the rectangular preservation cost before openpyxl materializes gaps.
    A missing/understated/overstated dimension cannot hide cells or evade bounds.
    """
    from xml.etree.ElementTree import iterparse
    from openpyxl.utils.cell import coordinate_to_tuple
    max_row = max_column = previous_row = 0
    current_row = previous_column = 0
    with workbook._archive.open(worksheet._worksheet_path) as source:
        for event, element in iterparse(source, events=("start", "end")):
            tag = element.tag.rsplit("}", 1)[-1]
            if event == "start" and tag == "row":
                raw = element.get("r")
                if raw is None:
                    current_row = previous_row + 1
                elif not re.fullmatch(r"[1-9][0-9]*", raw):
                    raise ValueError("DOCUMENT_EXCEL_COORDINATE_INVALID")
                else:
                    current_row = int(raw)
                if current_row <= previous_row:
                    raise ValueError("DOCUMENT_EXCEL_ROW_ORDER_INVALID")
                previous_row = current_row
                previous_column = 0
                max_row = current_row
            elif event == "start" and tag == "c":
                coordinate = element.get("r")
                if coordinate is None:
                    row, column = current_row, previous_column + 1
                else:
                    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", coordinate):
                        raise ValueError("DOCUMENT_EXCEL_COORDINATE_INVALID")
                    row, column = coordinate_to_tuple(coordinate)
                if not current_row or row != current_row or column <= previous_column:
                    raise ValueError("DOCUMENT_EXCEL_CELL_ORDER_INVALID")
                previous_column = column
                max_column = max(max_column, column)
            if max_row * max(1, max_column) > cell_limit:
                raise ValueError("DOCUMENT_EXCEL_CELL_LIMIT")
            if event == "end":
                element.clear()
    return max_row, max(1, max_column) if max_row else 0


def _legacy_excel_cell_text(value) -> str:
    """Keep the historical string table/text projection stable."""
    return str(getattr(value, "text", value)) if value is not None else ""


def _typed_excel_cell(cell) -> dict:
    """Return a deterministic JSON-safe source value without evaluating formulas."""
    import math
    from datetime import date, datetime, time, timedelta

    value = cell.value
    cell_type = getattr(cell, "data_type", None)
    if value is None:
        kind, safe_value = "blank", None
    elif cell_type == "f":
        kind, safe_value = "formula", str(getattr(value, "text", value))
    elif cell_type == "e":
        kind, safe_value = "error", str(value)
    elif isinstance(value, bool):
        kind, safe_value = "boolean", value
    elif isinstance(value, datetime):
        kind, safe_value = "datetime", value.isoformat()
    elif isinstance(value, date):
        kind, safe_value = "date", value.isoformat()
    elif isinstance(value, time):
        kind, safe_value = "time", value.isoformat()
    elif isinstance(value, timedelta):
        kind, safe_value = "duration", str(value)
    elif isinstance(value, int):
        kind, safe_value = "integer", value
    elif isinstance(value, float):
        kind = "number"
        safe_value = value if math.isfinite(value) else str(value)
    elif isinstance(value, str):
        kind, safe_value = "string", value
    else:
        # openpyxl's supported cell value types are covered above. Preserve an
        # unexpected typed value as text instead of dropping it or coercing it
        # to a numeric zero.
        kind, safe_value = "other", str(value)

    return {
        "type": kind,
        "value": safe_value,
        "cell_type": str(cell_type or ""),
    }


def parse_legacy_word(file_path: str) -> ParseResult:
    """Convert binary Word with a private, inert local LibreOffice profile.

    Original bytes remain the canonical Source input. Converted paragraph/cell
    offsets refer to retained text, not original layout/pages. Runtime packaging
    and decoder isolation remain W6 responsibilities for production use.
    """
    import shutil
    import signal
    import subprocess
    import tempfile
    from pathlib import Path
    executable = shutil.which("soffice")
    if executable is None:
        raise ImportError("二进制 DOC 解析需要本地 LibreOffice；不能用 DOCX 解析器替代")
    executable = str(Path(executable).resolve())
    configuration = '''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting">
<prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>
<prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop>
</item>
<item oor:path="/org.openoffice.Office.Writer/Content/Update">
<prop oor:name="Link" oor:op="fuse"><value>2</value></prop>
<prop oor:name="Field" oor:op="fuse"><value>false</value></prop>
</item></oor:items>'''
    with tempfile.TemporaryDirectory(prefix="lumenis-doc-decode-") as directory:
        work = Path(directory)
        profile = work / "profile"
        (profile / "user").mkdir(parents=True)
        (profile / "user" / "registrymodifications.xcu").write_text(configuration, encoding="utf-8")
        # The decoder receives only its public runtime configuration, never
        # caller credentials, private environment variables or a user profile.
        environment = {"PATH": os.defpath, "LANG": "C.UTF-8", "TMPDIR": str(work)}
        if os.environ.get("ORT_DISABLE_TELEMETRY") == "1":
            environment["ORT_DISABLE_TELEMETRY"] = "1"
        version = subprocess.run([executable, "-env:UserInstallation=" + profile.as_uri(), "--version"],
                                 capture_output=True, text=True, env=environment, timeout=5, check=False)
        parser_version = version.stdout.strip()[:200] if version.returncode == 0 else "not_reported"
        with tempfile.TemporaryFile() as output, tempfile.TemporaryFile() as errors:
            process = subprocess.Popen([executable, "-env:UserInstallation=" + profile.as_uri(),
                                        "--headless", "--convert-to", "docx:Office Open XML Text",
                                        "--outdir", str(work), os.path.abspath(file_path)],
                                       stdout=output, stderr=errors, env=environment,
                                       start_new_session=(os.name == "posix"))
            try:
                returncode = process.wait(timeout=30)
            except subprocess.TimeoutExpired as exc:
                if os.name == "posix":
                    os.killpg(process.pid, signal.SIGKILL)
                else:
                    process.kill()
                process.wait()
                raise ValueError("LEGACY_DOC_CONVERSION_TIMEOUT") from exc
        converted = work / (Path(file_path).stem + ".docx")
        if returncode != 0 or not converted.is_file():
            raise ValueError("LEGACY_DOC_CONVERSION_FAILED")
        if not 0 < converted.stat().st_size <= 8 * 1024 * 1024:
            raise ValueError("LEGACY_DOC_CONVERSION_SIZE_INVALID")
        converted_sha = hashlib.sha256(converted.read_bytes()).hexdigest()
        result = parse_word(str(converted))
    for locator in result.locators:
        locator["unit"] = "converted_" + locator["unit"]
    result.metadata.update(format="doc", parser="libreoffice", parser_version=parser_version,
                           intermediate_docx_sha256=converted_sha,
                           intermediate_state="transient_conversion_not_committed_artifact",
                           conversion_mode="local_doc_to_docx", positions="retained_converted_text",
                           embedded_macros="disabled_by_ephemeral_profile",
                           external_links="updates_disabled_by_ephemeral_profile",
                           active_content="disabled_by_ephemeral_profile")
    return result


def parse_legacy_excel(file_path: str) -> ParseResult:
    """Read binary XLS locally; preserve original bytes and cached-value limits.

    xlrd does not expose formula expressions. Values are statements retained in
    the input workbook, never a claim that a formula was freshly calculated.
    Exact source cell coordinates remain available alongside the raw archive.
    """
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError("XLS解析需要可选依赖 xlrd；不能用 XLSX 解析器替代") from exc
    result = ParseResult()
    workbook = xlrd.open_workbook(file_path, on_demand=True)
    result.metadata.update(format="xls", parser="xlrd", parser_version=xlrd.__version__,
                           sheets=workbook.sheet_names(),
                           calculation_state="retained_cached_values_not_recalculated",
                           formula_expressions="not_exposed_by_xlrd")
    lines, offset = [], 0
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            table = []
            for row_index in range(sheet.nrows):
                cells = []
                for column in range(sheet.ncols):
                    cell = sheet.cell(row_index, column)
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = ""
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = "True" if cell.value else "False"
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(cell.value, "#UNKNOWN_ERROR")
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                        value = str(int(cell.value))
                    else:
                        value = str(cell.value)
                    cells.append(value)
                if not any(value.strip() for value in cells):
                    continue
                table.append(cells)
                text = " | ".join(cells)
                if lines:
                    offset += 1
                cell_offset = offset
                for column, value in enumerate(cells, 1):
                    if value:
                        result.locators.append({"unit": "spreadsheet_cell", "sheet": sheet_name,
                                                "row": row_index + 1, "column": column,
                                                "start": cell_offset, "end": cell_offset + len(value)})
                    cell_offset += len(value) + 3
                lines.append(text)
                offset += len(text)
            result.tables.append(table)
        result.text = "\n".join(lines)
        result.paragraphs = lines
    finally:
        workbook.release_resources()
    return result


def split_into_chunks(text: str, max_chunk_size: int = 500) -> List[str]:
    """将长文本切分成适合AI分析的块"""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks = []
    current = ""
    for p in paragraphs:
        if len(current) + len(p) > max_chunk_size and current:
            chunks.append(current.strip())
            current = p
        else:
            current += "\n" + p
    if current.strip():
        chunks.append(current.strip())
    return chunks
