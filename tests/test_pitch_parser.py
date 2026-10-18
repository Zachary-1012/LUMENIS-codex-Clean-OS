# -*- coding: utf-8 -*-
"""方案需求拆解分析引擎测试。"""
import json
import hashlib
import os
import pytest
from scripts.pitch_parser import (analyze_pitch, analyze_pitch_file, detect_category,
                                  detect_circle, detect_tonality, detect_channels,
                                  detect_goals, detect_budget, PITCH_TOC)


def test_analyze_pitch_full():
    text = ("我们想做一款高端小众香水品牌「岚序」，预算300万，目标人群是高净值贵妇和小众藏家，"
            "主要渠道小红书和私域，希望做新品上市和品牌建设，创意方向是东方节气与稀缺仪式感，"
            "可以比稿。")
    r = analyze_pitch(text, n_ideas=2)
    e = r["extracted"]
    assert e["category"] == "美妆个护"
    assert e["circle"] == "高奢小众"
    assert e["tonality"] == "luxury"
    assert "小红书" in e["channels"]
    assert set(e["goals"]) >= {"新品上市", "品牌建设"}
    assert e["budget"]["value"] == 300
    assert e["budget"]["unit"] == "万"
    assert len(r["ideas"]["ideas"]) == 2
    # 高奢调性 → 不追热点叙事
    assert "不参与流量狂欢" in r["ideas"]["ideas"][0]["timeliness"] or \
           "不追热点" in r["ideas"]["ideas"][0]["timeliness"]
    assert len(r["pitch_toc"]) == 8
    assert r["pitch_toc"][0]["layer"] == "数据层"


def test_mass_tonality_differs():
    r1 = analyze_pitch("大众美妆品牌，预算50万，目标年轻女性，抖音小红书，爆款破圈", n_ideas=1)
    r2 = analyze_pitch("高奢珠宝品牌，预算1000万，目标高净值贵妇，私域，稀缺限量", n_ideas=1)
    assert r1["extracted"]["tonality"] != r2["extracted"]["tonality"]
    assert r1["ideas"]["ideas"][0]["tonality"] == "大众"
    assert r2["ideas"]["ideas"][0]["tonality"] == "高奢"


def test_pitch_file_formats(tmp_path):
    # md
    md = tmp_path / "req.md"
    md.write_text("餐饮品牌「拾味」，预算200万，目标都市白领，线下+抖音，做开业活动引爆", encoding="utf-8")
    r = analyze_pitch_file(str(md), n_ideas=1)
    assert r["source_file"] == str(md)
    assert r["extracted"]["category"] == "餐饮"
    # json
    jf = tmp_path / "req.json"
    jf.write_text(json.dumps({"brand": "岚序", "type": "香水", "budget": "300万"}, ensure_ascii=False),
                  encoding="utf-8")
    rj = analyze_pitch_file(str(jf), n_ideas=1)
    assert "美妆" in rj["extracted"]["category"]
    assert rj["extracted"]["budget"] == {"value": 300.0, "unit": "万"}
    # csv
    cf = tmp_path / "req.csv"
    cf.write_text("字段,内容\n品牌,拾味\n品类,餐饮\n预算,200万\n", encoding="utf-8-sig")
    rc = analyze_pitch_file(str(cf), n_ideas=1)
    assert rc["extracted"]["category"] == "餐饮"
    assert rc["extracted"]["budget"] == {"value": 200.0, "unit": "万"}


def test_budget_requires_label_unit_and_unambiguous_amount():
    assert detect_budget("2026年15周年 预算200万") == {"value": 200.0, "unit": "万"}
    assert detect_budget("预算：300元") == {"value": 300.0, "unit": "元"}
    assert detect_budget("budget: 1百万") == {"value": 1.0, "unit": "百万"}
    assert detect_budget("投入约2亿") == {"value": 2.0, "unit": "亿"}
    assert detect_budget("Budget £300k") == {"value": 300.0, "unit": "£k"}
    assert detect_budget("budget of £2700") == {"value": 2700.0, "unit": "£"}
    assert detect_budget("budget is GBP2700") == {"value": 2700.0, "unit": "GBP"}
    assert detect_budget("预算 USD 12.5") == {"value": 12.5, "unit": "USD"}
    assert detect_budget("预算") is None

    for text in (
        "15周年，年份2026年",
        "预算待确认，目标覆盖200万人，发布时间2026-05-01",
        "预算300，投入未定",
        "预算200万-300万",
        "预算200万至300",
        "预算200万平方米",
        "预算200万 方案执行",
        "预算200万，投入300万",
        "面向300万受众，预算缺失",
        "预算增长20%",
    ):
        assert detect_budget(text) is None, text


def test_file_input_provenance_hashes_exact_source_and_extracted_text(tmp_path):
    source = tmp_path / "需求.txt"
    content = "品牌拾味，预算200万，广州开业活动".encode("utf-8")
    source.write_bytes(content)

    result = analyze_pitch_file(str(source), n_ideas=1)
    provenance = result["input_provenance"]
    assert provenance["kind"] == "file_import_input"
    assert provenance["filename"] == "需求.txt"
    assert provenance["format"] == "txt"
    assert provenance["source_bytes_sha256"] == hashlib.sha256(content).hexdigest()
    assert provenance["extracted_text"] == content.decode("utf-8")
    assert provenance["extracted_text_sha256"] == hashlib.sha256(content).hexdigest()


def test_docx_and_pptx_extract_real_document_text(tmp_path):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    import io

    docx_path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("餐饮拾味开业活动")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "预算"
    table.cell(0, 1).text = "200万"
    docx_bytes = io.BytesIO()
    document.save(docx_bytes)
    docx_path.write_bytes(docx_bytes.getvalue())
    docx_result = analyze_pitch_file(str(docx_path), n_ideas=1)
    assert docx_result["extracted"]["category"] == "餐饮"
    assert docx_result["extracted"]["budget"] == {"value": 200.0, "unit": "万"}
    assert "预算 | 200万" in docx_result["input_provenance"]["extracted_text"]

    pptx_path = tmp_path / "brief.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "拾味开业活动 预算200万"
    pptx_bytes = io.BytesIO()
    presentation.save(pptx_bytes)
    pptx_path.write_bytes(pptx_bytes.getvalue())
    pptx_result = analyze_pitch_file(str(pptx_path), n_ideas=1)
    assert pptx_result["extracted"]["budget"] == {"value": 200.0, "unit": "万"}
    assert pptx_result["input_provenance"]["format"] == "pptx"


def test_file_import_rejects_invalid_text_and_unknown_formats(tmp_path):
    invalid_utf8 = tmp_path / "bad.txt"
    invalid_utf8.write_bytes(b"brand=\xff\xfe")
    with pytest.raises(UnicodeDecodeError):
        analyze_pitch_file(str(invalid_utf8), n_ideas=1)

    binary_text = tmp_path / "binary.md"
    binary_text.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError, match="二进制"):
        analyze_pitch_file(str(binary_text), n_ideas=1)

    unsupported = tmp_path / "brief.pdf"
    unsupported.write_bytes(b"%PDF-1.7")
    with pytest.raises(ValueError, match="不支持"):
        analyze_pitch_file(str(unsupported), n_ideas=1)


def test_detectors():
    assert detect_category("我们要做新能源汽车") == "汽车"
    assert detect_category("文旅景区门票活动") == "文旅"
    assert detect_circle("目标人群是宝妈和亲子家庭") == "母婴家庭"
    assert detect_tonality("顶奢限量稀缺") == "luxury"
    assert detect_channels("小红书+抖音+线下") == ["小红书", "抖音", "线下"]
    assert set(detect_goals("要破圈上热搜")) & {"声量传播"}
    assert len(PITCH_TOC) == 8


def test_pitch_has_marketing_field():
    """方案拆解返回营销节点适配（当月节点+情感主题+建议调性）。"""
    from scripts.pitch_parser import analyze_pitch
    r = analyze_pitch("广州太古汇 商业地产 15周年 预算200万 高净值 小红书 核心诉求提升到访", "广州太古汇", 2, "15周年")
    assert "marketing" in r, "缺 marketing 字段"
    m = r["marketing"]
    assert m is None or ("nodes" in m and "themes" in m and "tonality_advice" in m), m

def test_new_topic_roles_covered():
    """美妆/母婴/户外/酒饮 语义角色可识别并产出意象。"""
    from scripts.idea_generator import _topic_role
    assert _topic_role("面霜") == "beauty"
    assert _topic_role("亲子乐园") == "family"
    assert _topic_role("露营") == "sport"
    assert _topic_role("威士忌") == "wine"


def test_explicit_brief_date_reaches_creative_calendar(tmp_path):
    for date, market, expected in [('2026-02-17', 'CN', 'reference_only_not_live_verified'),
                                   ('2027-02-17', 'CN', 'year_not_covered'),
                                   ('2026-02-17', 'HK', 'market_not_covered')]:
        value = analyze_pitch('社区汤店，预算3000，轮班工作者单人晚餐', target_date=date, market=market)
        assert value['marketing']['state'] == expected
        assert value['ideas']['meta']['calendar_reference']['date'] == date
        assert value['ideas']['meta']['calendar_reference']['state'] == expected
        if expected != 'reference_only_not_live_verified':
            assert value['ideas']['meta']['bound_hotspot'] == '（未绑定节点）'
    assert analyze_pitch('独立书店')['marketing']['state'] == 'date_required'
    path = tmp_path / 'brief.txt'
    path.write_text('独立书店500元阅读交流活动')
    result = analyze_pitch_file(str(path), target_date='2026-04-23')
    assert result['ideas']['meta']['bound_hotspot'] == '世界读书日'
