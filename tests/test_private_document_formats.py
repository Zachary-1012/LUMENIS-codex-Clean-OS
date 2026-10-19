"""Actual generated local files, retained originals and parser output, without AI."""
import hashlib
import os
import shutil
import subprocess
import pytest
from tests.test_saved_document_workflow import workspace


def _make_source(root, extension):
    marker = "Private format source 2026"
    path = root / ("format-source." + extension)
    if extension == "txt":
        path.write_text(marker, encoding="utf-8")
    elif extension == "csv":
        path.write_text("title,value\n" + marker + ",42\n", encoding="utf-8")
    elif extension in {"html", "htm"}:
        path.write_text("<html><body><h1>" + marker + "</h1><p>Original local source.</p></body></html>", encoding="utf-8")
    elif extension == "pdf":
        from reportlab.pdfgen import canvas
        document = canvas.Canvas(str(path))
        document.drawString(50, 750, marker)
        document.save()
    elif extension in {"docx", "doc"}:
        from docx import Document
        document = Document()
        document.add_paragraph(marker)
        docx = path.with_suffix(".docx")
        document.save(str(docx))
        if extension == "doc":
            executable = shutil.which("soffice")
            assert executable, "Local LibreOffice is required for this legacy DOC acceptance"
            profile = root / "synthetic-doc-profile"
            result = subprocess.run([executable, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", "doc:MS Word 97", "--outdir", str(root), str(docx)], capture_output=True, timeout=30, env={"PATH": os.defpath, "LANG": "C.UTF-8"})
            assert result.returncode == 0 and path.is_file(), result.stderr.decode(errors="replace")
    elif extension == "xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["title", "value"])
        workbook.active.append([marker, 42])
        workbook.create_sheet("Second").append(["second sheet retained", 9])
        workbook.save(path)
    elif extension == "xls":
        import xlwt
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("First")
        sheet.write(0, 0, "title")
        sheet.write(1, 0, marker)
        workbook.add_sheet("Second").write(0, 0, "second sheet retained")
        workbook.save(str(path))
    elif extension == "pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = marker
        document.save(path)
    return path, marker


@pytest.mark.parametrize("extension", ["txt", "csv", "html", "htm", "pdf", "docx", "doc", "xlsx", "xls", "pptx"])
def test_actual_private_source_formats(workspace, extension):
    ws = workspace
    path, marker = _make_source(ws["root"], extension)
    raw = path.read_bytes()
    scope = {"scope_kind": "shared_workspace", "version": 1}
    result = ws["documents"].save_private_document(str(path), actor=ws["actor"], scope=scope, request_id="actual-format-" + extension)
    document = result["document"]
    assert marker in document["payload"]["parsed_source"]["text"]
    assert document["payload"]["basic_analysis"]["ai_state"] == "NOT_RUN"
    assert document["payload"]["source_material"]["sha256"] == hashlib.sha256(raw).hexdigest()
    original = ws["work"].download_attachment(result["artifact"]["artifact_id"], name=document["attachments"][0]["name"], actor=ws["actor"], scope=scope)
    assert original["content"] == raw
    if extension in {"xls", "xlsx"}:
        assert len(document["payload"]["parsed_source"]["tables"]) == 2
    if extension == "doc":
        metadata = document["payload"]["parsed_source"]["metadata"]
        assert metadata["embedded_macros"] == "disabled_by_ephemeral_profile"
        assert metadata["active_content"] == "disabled_by_ephemeral_profile"
        assert metadata["external_links"] == "updates_disabled_by_ephemeral_profile"


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32", "gb18030"])
@pytest.mark.parametrize("extension", ["txt", "csv"])
def test_original_text_encodings_are_preserved(workspace, encoding, extension):
    ws = workspace
    text = "标题,值\n合法私有材料,42\n"
    path = ws["root"] / ("encoded." + extension)
    raw = text.encode(encoding)
    path.write_bytes(raw)
    scope = {"scope_kind": "shared_workspace", "version": 1}
    result = ws["documents"].save_private_document(str(path), actor=ws["actor"], scope=scope,
        request_id="encoding-" + encoding + "-" + extension)
    source = result["document"]["payload"]["parsed_source"]
    assert "合法私有材料" in source["text"]
    assert source["metadata"]["original_bytes_modified"] is False
    assert source["metadata"]["encoding_state"] == ("assumed_decoder" if encoding == "gb18030" else "bom_declared")
    assert result["document"]["payload"]["source_material"]["sha256"] == hashlib.sha256(raw).hexdigest()


@pytest.mark.parametrize("raw", [b"MZ\0binary", b"text\x01control", b"\xff\xfeb", b"\xff\xff\xff"])
def test_text_binary_and_malformed_bytes_are_rejected(raw):
    from lumenis.importer.parser import decode_local_text
    with pytest.raises(ValueError):
        decode_local_text(raw)


def test_owned_biff8_decoder_zero_boolean_and_coordinates(tmp_path):
    """Independent XLS decoder proof; does not validate Office conversion."""
    import xlwt
    from lumenis.importer.parser import parse_file

    path = tmp_path / "owned-biff8.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Owned sheet")
    sheet.write(0, 0, "Own source statement")
    sheet.write(0, 1, 0)
    sheet.write(0, 2, False)
    sheet.write(1, 1, "Exact second-row source")
    book.save(str(path))
    original = path.read_bytes()
    assert original.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    result = parse_file(str(path))
    assert result.tables[0][0] == ["Own source statement", "0", "False"]
    cell = next(item for item in result.locators if item["row"] == 2 and item["column"] == 2)
    assert cell["sheet"] == "Owned sheet"
    assert result.text[cell["start"]:cell["end"]] == "Exact second-row source"
    assert result.metadata["calculation_state"] == "retained_cached_values_not_recalculated"
    assert result.metadata["formula_expressions"] == "not_exposed_by_xlrd"
    assert result.metadata["source_fingerprint"] == hashlib.sha256(original).hexdigest()
    assert path.read_bytes() == original
