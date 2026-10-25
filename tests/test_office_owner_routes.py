import io
import json

import pytest
from docx import Document

from tests import test_saved_document_workflow
from tests.test_root_owner_auth_composition import _build_root_hooks

workspace = test_saved_document_workflow.workspace


@pytest.mark.parametrize('account_mode', [False, True])
def test_office_http_edit_reopen_attachment_conflict_and_scope(workspace, monkeypatch, account_mode):
    import experience_runtime

    app = workspace['app']
    _build_root_hooks(app)
    actor = {'current': None}
    if account_mode:
        import lumenis.auth.service
        from account.models import ACCOUNT_TABLES_SQL

        with workspace['repo'].transaction() as connection:
            connection.executescript(ACCOUNT_TABLES_SQL)
            connection.execute('INSERT INTO v4_users(id,username,password_hash,role,status) VALUES(?,?,?,?,?)',
                               (921, 'office-owner', 'isolated-only', 'owner', 'active'))
        app.config['LUMENIS_OPEN_WORKSPACE'] = False
        actor['current'] = {'id': 921}
        monkeypatch.setattr(lumenis.auth.service, 'get_current_user', lambda **kwargs: actor['current'])
    app.register_blueprint(experience_runtime.experience_bp)
    client = app.test_client()
    headers = {'X-Lumenis-Workspace-Scope': json.dumps(
        {'scope_kind': 'shared_workspace', 'version': 1}), 'Origin': 'http://localhost'}
    document = Document()
    document.add_paragraph('Original client wording.')
    document.add_paragraph('This paragraph is unchanged.')
    raw = io.BytesIO()
    document.save(raw)
    raw.seek(0)
    uploaded = client.post('/api/experience/materials/documents', headers=headers,
        data={'file': (raw, 'client.docx'), 'request_id': 'office-http-upload-01'})
    assert uploaded.status_code == 200, uploaded.get_json()
    saved = uploaded.get_json()['result']
    artifact_id = saved['artifact']['artifact_id']
    base = '/api/experience/saved-work/' + artifact_id
    page = client.get(base + '/office', headers=headers)
    assert page.status_code == 200, page.get_json()
    original = page.get_json()['result']
    item = next(value for value in original['objects'] if value['value'] == 'Original client wording.')
    body = {'request_id': 'office-http-edit-001',
        'expected_sha256': original['source_artifact_sha256'],
        'expected_package_sha256': original['package_sha256'],
        'reason': 'Correct one paragraph', 'commands': [{'object_id': item['object_id'],
            'operation': 'replace_text', 'expected_value': item['value'], 'value': 'Revised client wording.'}]}
    edited = client.post(base + '/office/edit', headers=headers, json=body)
    assert edited.status_code == 200, edited.get_json()
    result = edited.get_json()['result']
    child = '/api/experience/saved-work/' + result['artifact']['artifact_id']
    assert result['document']['version'] == 2
    assert result['document']['parent_artifact_id'] == artifact_id
    reread = client.get(child + '/office', headers=headers)
    assert reread.status_code == 200
    values = [entry['value'] for entry in reread.get_json()['result']['objects']]
    assert 'Revised client wording.' in values
    assert 'This paragraph is unchanged.' in values
    parent = client.get(base + '/office', headers=headers)
    assert any(entry['value'] == 'Original client wording.' for entry in parent.get_json()['result']['objects'])
    name = result['office_attachment']['name']
    download = client.get(child + '/attachments/' + name, headers=headers)
    assert download.status_code == 200
    reopened = Document(io.BytesIO(download.data))
    assert [p.text for p in reopened.paragraphs] == ['Revised client wording.', 'This paragraph is unchanged.']
    conflict = client.post(child + '/office/edit', headers=headers, json={**body, 'request_id': 'office-stale-001'})
    assert conflict.status_code == 409, conflict.get_json()
    for query in ('?limit=501', '?offset=1&offset=2', '?actor=other'):
        assert client.get(base + '/office' + query, headers=headers).status_code == 400
    denied = client.post(child + '/office/edit', headers={**headers, 'Origin': 'https://other.example'}, json=body)
    assert denied.status_code == 403
    if not account_mode:
        denied = client.get(child + '/office', headers=headers, environ_overrides={'REMOTE_ADDR': '203.0.113.1'})
        assert denied.status_code in {401, 403}
    app.config['LUMENIS_OPEN_WORKSPACE'] = False
    actor['current'] = None
    assert client.get(child + '/office', headers=headers).status_code in {401, 403}
