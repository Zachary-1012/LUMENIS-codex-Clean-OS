"""Exact app route/hook extraction, real Word owner, isolated Permission only."""
import ast
from concurrent.futures import ThreadPoolExecutor
import io
from pathlib import Path
import re
import threading

from docx import Document
from flask import Flask, jsonify, request, send_file
import pytest

from lumenis.system.security import permission as p
from scripts import pitch_export, pitch_parser

ANALYZE = '/api/creative/pitch-analyze'
EXPORT = '/api/creative/pitch-export'


@pytest.fixture
def route(monkeypatch):
    p.set_testing_bypass(False)
    monkeypatch.delenv('PERMISSION_GATE_TESTING', raising=False)
    gate = p.PermissionGate()
    monkeypatch.setattr(p, '_default_gate', gate)
    app = Flask('isolated_current_pitch_routes')
    names = {'_creative_pitch_request', '_creative_pitch_analysis',
             '_creative_pitch_check_export', 'creative_pitch_analyze',
             'creative_pitch_export', '_pg_resolve_operation_type',
             '_permission_gate_write_check'}
    constants = {'_PG_WRITE_METHODS', '_PG_BYPASS_PATHS'}
    tree = ast.parse(Path('app.py').read_text(encoding='utf-8-sig'))
    nodes = [node for node in tree.body if
             (isinstance(node, ast.FunctionDef) and node.name in names) or
             (isinstance(node, ast.Assign) and any(isinstance(target, ast.Name)
              and target.id in constants for target in node.targets))]
    assert {node.name for node in nodes if isinstance(node, ast.FunctionDef)} == names
    namespace = dict(app=app, request=request, jsonify=jsonify, send_file=send_file,
                     _pg=gate, OperationType=p.OperationType, re=re)
    exec(compile(ast.Module(body=nodes, type_ignores=[]), '<current-app-pitch>', 'exec'), namespace)
    yield app, gate, namespace
    p.set_testing_bypass(True)


def manuscript(language='en'):
    return {'title': 'A local reading proposal', 'summary': 'Use the supplied brief; unknown outcomes stay pending.',
            'language': language, 'authorship': 'user_supplied',
            'sections': [{'heading': f'Section {index}', 'body': f'Actual editable paragraph {index}.\n\nKeep the original words.'}
                         for index in range(1, 10)]}


@pytest.mark.parametrize('language,label', [('en', 'Review and usage approval pending.'),
                                          ('zh', '审阅及使用授权待确认'),
                                          ('ko', '검토 및 사용 승인 대기')])
def test_all_eleven_supplied_items_and_actual_prose_survive_download(route, language, label):
    app, _, _ = route
    draft = manuscript(language)
    response = app.test_client().post(EXPORT, json={'text': 'Library youth communication brief.', 'manuscript': draft})
    assert response.status_code == 200
    document = Document(io.BytesIO(response.data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    required = [draft['title'], draft['summary'], *[section['heading'] for section in draft['sections']]]
    assert len(required) == 11 and all(item in paragraphs for item in required)
    assert all(paragraph in paragraphs for section in draft['sections'] for paragraph in section['body'].split('\n\n'))
    assert any(label in paragraph for paragraph in paragraphs)
    assert document.core_properties.language == language
    assert document.core_properties.comments == 'authorship=user_supplied; approval=not_verified'
    assert response.headers['Cache-Control'] == 'no-store'


@pytest.mark.parametrize('invalid', [None, [], {}, {'title': 'Incomplete'}, {'title': 'T', 'summary': 'S', 'sections': []}])
def test_explicit_invalid_manuscript_never_falls_back_to_diagnostic(route, invalid):
    app, _, _ = route
    response = app.test_client().post(EXPORT, json={'text': 'Readable brief.', 'manuscript': invalid})
    assert response.status_code == 400 and response.get_json()['code'] == 'INVALID_PITCH_INPUT'


@pytest.mark.parametrize('body', [{'text': []}, {'text': 'B', 'brand': 42}, {'text': 'B', 'hotspot': False},
                                 {'text': 'B', 'n': True}, {'text': 'B', 'n': 0}, {'text': 'B', 'n': -1},
                                 {'text': 'B', 'n': '3'}, {'text': 'B', 'actor_id': 1},
                                 {'text': 'B', 'approval_token': 'untrusted'}])
def test_invalid_fields_are_rejected_before_analysis(route, monkeypatch, body):
    app, _, _ = route
    def must_not_analyze(*args):
        pytest.fail('invalid request reached parser')
    monkeypatch.setattr(pitch_parser, 'analyze_pitch', must_not_analyze)
    assert app.test_client().post(EXPORT, json=body).status_code == 400


def test_duplicate_json_and_nonfinite_values_and_body_size_are_rejected(route):
    client = route[0].test_client()
    for body in ('{"text":"a","text":"b"}', '{"text":"a","n":NaN}', '[]'):
        assert client.post(EXPORT, data=body, content_type='application/json').status_code == 400
    assert client.post(EXPORT, data=b' ' * (1024 * 1024 + 1), content_type='application/json').status_code == 400


@pytest.mark.parametrize('field,value', [('language', []), ('authorship', {}), ('approved', True)])
def test_malformed_manuscript_types_or_authority_are_input_errors(route, field, value):
    draft = manuscript()
    draft[field] = value
    response = route[0].test_client().post(EXPORT, json={'text': 'Brief', 'manuscript': draft})
    assert response.status_code == 400 and response.get_json()['code'] == 'INVALID_PITCH_INPUT'


def test_oversized_multipart_text_field_is_an_input_error(route):
    response = route[0].test_client().post(EXPORT, data={'text': 'x' * 600000}, content_type='multipart/form-data')
    assert response.status_code == 400


def test_legacy_request_stays_honest_diagnostic(route):
    response = route[0].test_client().post(EXPORT, json={'text': 'A youth library communication brief.', 'n': 10})
    assert response.status_code == 200
    assert any('不是完整提案成稿' in paragraph.text for paragraph in Document(io.BytesIO(response.data)).paragraphs)


def test_same_client_filename_has_separate_live_paths_and_cleanup(route, monkeypatch):
    app, _, _ = route
    rendezvous = threading.Barrier(2)
    paths = []
    paths_lock = threading.Lock()
    def parse(path, *args):
        path = Path(path)
        with paths_lock:
            paths.append(path)
        rendezvous.wait(timeout=5)
        assert path.read_bytes() in (b'first', b'second')
        return {'source_file': str(path), 'input_provenance': {'filename': path.name}}
    monkeypatch.setattr(pitch_parser, 'analyze_pitch_file', parse)
    def post(content):
        with app.test_client() as client:
            return client.post(ANALYZE, data={'file': (io.BytesIO(content), '../../same.txt')}).get_json()
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(post, (b'first', b'second')))
    assert len(paths) == 2 and paths[0] != paths[1]
    assert all(not path.parent.exists() for path in paths)
    assert all(result['data']['source_file'] == 'same.txt' for result in results)
    assert all(result['data']['input_provenance']['filename'] == 'same.txt' for result in results)


def test_parser_failure_cleans_upload_and_hides_server_path(route, monkeypatch):
    seen = []
    def fail(path, *args):
        seen.append(Path(path))
        raise RuntimeError('private server path ' + path)
    monkeypatch.setattr(pitch_parser, 'analyze_pitch_file', fail)
    response = route[0].test_client().post(ANALYZE, data={'file': (io.BytesIO(b'brief'), 'same.txt')})
    assert response.status_code == 500 and 'private server' not in response.get_data(as_text=True)
    assert seen and not seen[0].parent.exists()


def test_multipart_manuscript_and_original_byte_provenance(route):
    import hashlib
    import json
    client = route[0].test_client()
    draft = manuscript()
    response = client.post(EXPORT, data={'file': (io.BytesIO(b'Library brief.'), 'brief.txt'),
                                        'n': '3', 'manuscript': json.dumps(draft)})
    assert response.status_code == 200
    assert Document(io.BytesIO(response.data)).paragraphs[0].text == draft['title']
    analysis = client.post(ANALYZE, data={'file': (io.BytesIO(b'Library brief.'), 'brief.txt')}).get_json()['data']
    assert analysis['source_file'] == 'brief.txt'
    assert analysis['input_provenance']['source_bytes_sha256'] == hashlib.sha256(b'Library brief.').hexdigest()


@pytest.mark.parametrize('filename,content', [('brief.pdf', b'unsupported'), ('empty.txt', b''),
                                            ('bad.txt', b'\x00binary'), ('large.txt', b'x' * (16 * 1024 * 1024 + 1))],
                         ids=['unsupported', 'empty', 'binary', 'oversized'])
def test_unsupported_empty_binary_and_oversized_uploads_fail_explicitly(route, filename, content):
    response = route[0].test_client().post(ANALYZE, data={'file': (io.BytesIO(content), filename)})
    assert response.status_code == 400


def test_initial_permission_denial_has_no_parser_or_export_effect(route, monkeypatch):
    app, gate, _ = route
    gate.revoke_permission(p.PermissionLevel.VERIFY)
    def forbidden(*args, **kwargs):
        pytest.fail('denied request performed work')
    monkeypatch.setattr(pitch_parser, 'analyze_pitch', forbidden)
    monkeypatch.setattr(pitch_export, 'export_pitch_docx', forbidden)
    assert app.test_client().post(EXPORT, json={'text': 'Brief'}).status_code == 403


def test_revocation_after_file_publish_blocks_emission_and_cleans_output(route, monkeypatch):
    app, gate, _ = route
    original_link = pitch_export.os.link
    destinations = []
    def link_then_revoke(source, destination):
        original_link(source, destination)
        destinations.append(Path(destination))
        gate.revoke_permission(p.PermissionLevel.VERIFY)
    monkeypatch.setattr(pitch_export.os, 'link', link_then_revoke)
    response = app.test_client().post(EXPORT, json={'text': 'Brief', 'manuscript': manuscript()})
    assert response.status_code == 403 and response.get_json()['code'] == 'PITCH_EXPORT_DENIED'
    assert destinations and all(not path.parent.exists() for path in destinations)


def test_export_failure_cleans_owned_directory_and_hides_internal_error(route, monkeypatch):
    seen = []
    def fail(analysis, path, **kwargs):
        seen.append(Path(path))
        Path(path).write_bytes(b'partial')
        raise RuntimeError('private disk location ' + path)
    monkeypatch.setattr(pitch_export, 'export_pitch_docx', fail)
    response = route[0].test_client().post(EXPORT, json={'text': 'Brief'})
    assert response.status_code == 500 and 'private disk' not in response.get_data(as_text=True)
    assert seen and not seen[0].parent.exists()


def test_testing_bypass_cannot_authorize_export_and_missing_gate_is_closed(route):
    app, _, namespace = route
    p.set_testing_bypass(True)
    assert app.test_client().post(EXPORT, json={'text': 'Brief'}).status_code == 503
    p.set_testing_bypass(False)
    namespace['_pg'] = None
    assert app.test_client().post(EXPORT, json={'text': 'Brief'}).status_code == 503
