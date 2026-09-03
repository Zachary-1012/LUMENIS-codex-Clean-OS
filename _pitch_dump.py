# -*- coding: utf-8 -*-
"""解析8个太古汇真实方案文件 → 文本dump，供精读学习。"""
import os, sys, io, glob

OUT = r"D:\Taikoo-monitor-new\_pitch_dump"
os.makedirs(OUT, exist_ok=True)

FILES = [
    r"C:\Users\28049\Downloads\提案Final【25年年度review】-2025年广州太古汇小红书运营项目_MSG251128.pptx",
    r"C:\Users\28049\Downloads\For：Zachary.pdf",
    r"C:\Users\28049\Downloads\v2-索尼音乐打卡活动Review-2025年广州太古汇小红书运营项目-250428.pptx",
    r"C:\Users\28049\Downloads\【Study-NIO蔚来】25年广州太古汇对话时刻播客项目-板块2-0113.xlsx",
    r"C:\Users\28049\Downloads\2024年广州太古汇播客专项-1104.pptx",
    r"C:\Users\28049\Downloads\播客先声.pdf",
    r"C:\Users\28049\Downloads\日谈公园介绍2024版.pdf",
    r"C:\Users\28049\Downloads\uth又识mediakit2025.pdf",
]


def dump_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for si, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t.strip():
                    parts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    parts.append(" | ".join(cells))
            if shape.shape_type == 6:  # group
                try:
                    for sub in shape.shapes:
                        if sub.has_text_frame:
                            t = "\n".join(p.text for p in sub.text_frame.paragraphs if p.text.strip())
                            if t.strip():
                                parts.append(t)
                except Exception:
                    pass
        if parts:
            out.append("=== Slide %d ===\n%s" % (si, "\n".join(parts)))
    return "\n\n".join(out)


def dump_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None and str(c).strip()]
            if vals:
                rows.append(" | ".join(vals))
        if rows:
            out.append("=== Sheet: %s ===\n%s" % (ws.title, "\n".join(rows)))
    return "\n\n".join(out)


def dump_pdf(path):
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for pi, page in enumerate(pdf.pages, 1):
            t = page.extract_text() or ""
            if t.strip():
                out.append("=== Page %d ===\n%s" % (pi, t))
    return "\n\n".join(out)


for p in FILES:
    name = os.path.splitext(os.path.basename(p))[0]
    ext = os.path.splitext(p)[1].lower()
    try:
        if ext == ".pptx":
            text = dump_pptx(p)
        elif ext == ".xlsx":
            text = dump_xlsx(p)
        elif ext == ".pdf":
            text = dump_pdf(p)
        else:
            text = open(p, encoding="utf-8", errors="ignore").read()
        out_path = os.path.join(OUT, name + ".txt")
        io.open(out_path, "w", encoding="utf-8").write(text)
        print("OK  %s  (%d chars)" % (name, len(text)))
    except Exception as e:
        print("ERR %s  %s" % (name, e))
