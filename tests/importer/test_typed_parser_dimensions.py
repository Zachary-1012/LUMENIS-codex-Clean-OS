from zipfile import ZipFile, ZIP_DEFLATED
import re
from openpyxl import Workbook
from lumenis.importer.parser import parse_file

def test_understated_dimension_must_not_silently_drop_real_cells(tmp_path):
    original = tmp_path / 'original.xlsx'
    modified = tmp_path / 'understated.xlsx'
    wb = Workbook(); ws = wb.active
    ws['A1'] = 'header'; ws['B2'] = 'MUST_RETAIN'
    wb.save(original)
    with ZipFile(original) as zin, ZipFile(modified, 'w', ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension ref="[^"]+"/>', b'<dimension ref="A1:A1"/>', data)
                assert b'MUST_RETAIN' in data
            zout.writestr(item, data)
    try:
        parsed = parse_file(str(modified))
    except ValueError:
        return
    values = [c['value'] for sheet in parsed.source_tables for row in sheet['rows'] for c in row['values']]
    assert 'MUST_RETAIN' in values, ('silent source cell loss', values)

def test_missing_dimension_sparse_extent_must_obey_cell_budget(tmp_path):
    original = tmp_path / 'original_sparse.xlsx'
    modified = tmp_path / 'dimensionless_sparse.xlsx'
    wb = Workbook(); ws = wb.active
    ws['A100001'] = 'last'
    wb.save(original)
    with ZipFile(original) as zin, ZipFile(modified, 'w', ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension ref="[^"]+"/>', b'', data)
            zout.writestr(item, data)
    try:
        parsed = parse_file(str(modified))
    except ValueError as exc:
        assert 'LIMIT' in str(exc)
        return
    assert False, ('over-budget physical extent accepted', len(parsed.source_tables[0]['rows']))


def test_overstated_dimension_does_not_expand_or_erase_actual_cells(tmp_path):
    original=tmp_path/'real.xlsx'; modified=tmp_path/'lying.xlsx'
    wb=Workbook();wb.active['B2']='KEEP';wb.save(original)
    with ZipFile(original) as zin, ZipFile(modified,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw=zin.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                raw=re.sub(rb'<dimension ref="[^"]+"/>',b'<dimension ref="A1:XFD1048576"/>',raw)
            zout.writestr(item,raw)
    parsed=parse_file(str(modified))
    assert len(parsed.source_tables[0]['rows'])==2
    assert parsed.source_tables[0]['rows'][1]['values'][1]['value']=='KEEP'


def test_blank_row_without_cells_still_has_bounded_physical_extent(tmp_path):
    import pytest
    original=tmp_path/'real.xlsx';modified=tmp_path/'blankfar.xlsx'
    wb=Workbook();wb.active['A1']='KEEP';wb.save(original)
    with ZipFile(original) as zin, ZipFile(modified,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw=zin.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                raw=re.sub(rb'<sheetData>.*?</sheetData>',b'<sheetData><row r="100001"></row></sheetData>',raw)
                raw=re.sub(rb'<dimension ref="[^"]+"/>',b'',raw)
            zout.writestr(item,raw)
    with pytest.raises(ValueError,match='DOCUMENT_EXCEL_CELL_LIMIT'):
        parse_file(str(modified))
