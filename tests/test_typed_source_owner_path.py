"""Typed originals use the existing actual Evidence/Canonical/SavedWork owners."""
import copy
import hashlib
import json
import pytest
from tests.test_saved_document_workflow import workspace, _save
from lumenis.importer.parser import parse_file
from lumenis.importer.analyzer import extract_source_rows, source_rows_fingerprint
from lumenis.system.security import permission


def _file(ws, kind):
    path = ws['root'] / ('typed.' + kind)
    if kind == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        first = book.active
        first.title = 'One'
        first.append(['title', 'content', 'number', 'flag', 'missing', 'formula'])
        first.append(['same', 'original first', 0, False, None, '=1+1'])
        second = book.create_sheet('Two')
        second.append(['title', 'content', 'number', 'flag'])
        second.append(['same', 'original second', 1, True])
        book.save(path)
    elif kind == 'docx':
        from docx import Document
        doc = Document(); doc.add_paragraph('original first');doc.add_paragraph('')
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='original second'
        doc.add_paragraph('original final');doc.save(path)
    else:
        from reportlab.pdfgen.canvas import Canvas
        doc=Canvas(str(path));doc.drawString(30,750,'original first');doc.showPage()
        doc.drawString(30,750,'original second');doc.save()
    return path


def _import(ws, path, parsed, rows, **extra):
    args={'actor':ws['actor'],'scope':ws['scope'],'provenance':{
        'source_url':'https://www.swireproperties.com/typed', 'published_at':None,
        'captured_at':'2026-10-18T01:00:00Z','rights':'public-reference-only','title':'Typed fixture'}}
    resource=ws['documents'].document_import_resource(str(path),**args)
    ws['gate'].grant_approval(permission.OperationType.DATA_IMPORT.value,resource=resource,granted_by='isolated-test')
    selection={'selected_row_indices':list(range(len(rows))),'preview_row_count':len(rows),
               'expected_preview_rows_sha256':source_rows_fingerprint(parsed,rows,count=len(rows))}
    selection.update(extra)
    return ws['documents'].import_document(str(path),**args,**selection),args


@pytest.mark.parametrize('kind',['xlsx','docx','pdf'])
def test_typed_actual_import_quote_save_reopen_export(workspace,kind):
    ws=workspace;path=_file(ws,kind);parsed=parse_file(str(path));rows=extract_source_rows(parsed)
    assert rows
    source,args=_import(ws,path,parsed,rows)
    assert source['canonical_rows_inserted']==len(rows)
    seen=[]
    for item,row in zip(source['canonical_row_records'],rows):
        stored=ws['repo'].query_one('SELECT * FROM canonical_data WHERE canonical_id=?',(item['canonical_id'],))
        provenance=json.loads(stored['provenance']);record=provenance['source_record']
        assert record==row.source_record
        assert provenance['source_fields']==row.source_fields
        assert hashlib.sha256(provenance['source_row_payload'].encode()).hexdigest()==row.source_row_sha256
        assert stored['sentiment']=='unknown' and stored['verification_status']=='review'
        seen.append(item['canonical_id'])
    assert len(set(seen))==len(rows)
    if kind=='xlsx':
        assert [row.source_locator['sheet_name'] for row in rows]==['One','One','Two','Two']
        values=rows[1].source_record['values']
        assert values[2]['value']==0 and values[3]['value'] is False and values[4]['value'] is None
        assert values[5]['type']=='formula' and values[5]['value']=='=1+1'
    start=parsed.text.index('original second');end=start+len('original second')
    located=ws['documents'].locate_quote(source['evidence_id'],actor=ws['actor'],scope=ws['scope'],start=start,end=end)
    assert located['citation']['quote']=='original second'
    saved=_save(ws,citations=[located['citation']]);artifact=saved['artifact']['artifact_id']
    assert ws['work'].reopen_work(artifact,actor=ws['actor'],scope=ws['scope'])['document']==saved['document']
    exported=ws['work'].export_work(artifact,actor=ws['actor'],scope=ws['scope'],format='json')
    assert json.loads(exported['content'])['citations'][0]['quote']=='original second'
    original=ws['documents'].reopen_imported_source(source['evidence_id'],actor=ws['actor'],scope=ws['scope'])
    assert original['content']==path.read_bytes()
    replay=ws['documents'].import_document(str(path),**args,selected_row_indices=list(range(len(rows))),
        preview_row_count=len(rows),expected_preview_rows_sha256=source_rows_fingerprint(parsed,rows,count=len(rows)))
    assert replay['replayed'] and replay['canonical_rows_inserted']==0
    ws['app'].config['LUMENIS_OPEN_WORKSPACE']=False
    with pytest.raises((ValueError,PermissionError)):
        ws['work'].reopen_work(artifact,actor=ws['actor'],scope=ws['scope'])


def test_cross_sheet_order_and_preview_version_are_bound(workspace):
    ws=workspace;path=_file(ws,'xlsx');parsed=parse_file(str(path));rows=extract_source_rows(parsed)
    reordered=copy.deepcopy(parsed);reordered.source_tables.reverse()
    with pytest.raises(ValueError,match='ORDER'):
        extract_source_rows(reordered)
    wrong=source_rows_fingerprint(parsed,list(reversed(rows)),count=len(rows))
    with pytest.raises(ValueError,match='SELECTION_CHANGED'):
        _import(ws,path,parsed,rows,expected_preview_rows_sha256=wrong)
    assert ws['repo'].query_one('SELECT COUNT(*) FROM canonical_data')[0]==0


def test_typed_private_shared_save_retains_all_original_structures(workspace):
    ws=workspace;path=_file(ws,'xlsx');parsed=parse_file(str(path))
    result=ws['documents'].save_private_document(str(path),actor=ws['actor'],
        scope={'scope_kind':'shared_workspace','version':1},request_id='typed-private-001')
    reopened=ws['work'].reopen_work(result['artifact']['artifact_id'],actor=ws['actor'],scope={'scope_kind':'shared_workspace','version':1})
    retained=reopened['document']['payload']['parsed_source']
    assert retained['source_tables']==parsed.source_tables
    assert retained['source_blocks']==[]
    assert reopened['document']['payload']['source_material']['rights']=='UNKNOWN'


def test_source_row_payload_tamper_rolls_back_real_import(workspace,monkeypatch):
    from lumenis.importer import analyzer
    ws=workspace;path=_file(ws,'xlsx');parsed=parse_file(str(path));rows=extract_source_rows(parsed)
    original=analyzer.extract_source_rows
    def changed(parsed):
        result=original(parsed)
        result[1].raw_chunk += 'tampered'
        return result
    monkeypatch.setattr(analyzer,'extract_source_rows',changed)
    with pytest.raises(ValueError,match='ROW_PROJECTION_INVALID'):
        _import(ws,path,parsed,rows)
    assert ws['repo'].query_one('SELECT COUNT(*) FROM canonical_data')[0]==0
    assert ws['repo'].query_one('SELECT COUNT(*) FROM app_evidence_records')[0]==0


def test_source_blocks_and_preview_share_the_same_complete_order(workspace):
    from lumenis.importer.analyzer import analyze_parsed
    for kind in ['xlsx','docx','pdf']:
        parsed=parse_file(str(_file(workspace,kind)))
        preview=analyze_parsed(parsed)
        exact=extract_source_rows(parsed)
        assert len(preview)==len(exact)
        assert [n.source_row_sha256 for n in preview]==[n.source_row_sha256 for n in exact]
        assert all(n.analysis_status=='NOT_RUN' for n in preview)


from tests.test_independent_growth_real_app import real_app, headers


@pytest.mark.parametrize('kind',['xlsx','docx','pdf'])
def test_typed_shared_real_session_upload_reopen_and_foreign_denial(real_app,kind):
    import io
    ws=real_app;path=_file(ws,kind)
    response=ws['api'].post('/api/experience/materials/documents',headers=headers(201),
        data={'file':(io.BytesIO(path.read_bytes()),path.name),'request_id':'typed-real-'+kind},
        content_type='multipart/form-data')
    assert response.status_code==200,response.get_json()
    result=response.get_json()['result'];ref=result['artifact']['artifact_id']
    own=ws['api'].get('/api/experience/saved-work/'+ref,headers=headers(201))
    assert own.status_code==200,own.get_json()
    retained=result['document']['payload']['parsed_source']
    assert retained['source_tables'] if kind=='xlsx' else retained['source_blocks']
    denied=ws['api'].get('/api/experience/saved-work/'+ref,headers=headers(202))
    assert denied.status_code in (400,403,404),denied.get_json()
    assert 'original second' not in denied.get_data(as_text=True)
    assert 'source_tables' not in denied.get_data(as_text=True)
    with ws['repo'].transaction() as con:
        con.execute("DELETE FROM v4_sessions WHERE user_id=201")
    assert ws['api'].get('/api/experience/saved-work/'+ref,headers=headers(201)).status_code in (401,403)


def test_preview_exposes_only_server_retained_row_identity(workspace):
    from lumenis.importer import service
    path=_file(workspace,'xlsx');parsed=parse_file(str(path));rows=extract_source_rows(parsed)
    visible=[service._note_to_dict(row) for row in rows]
    assert visible[1]['source_record']==rows[1].source_record
    assert visible[1]['source_fields'][3]['value']['value'] is False
    assert visible[1]['source_locator']['sheet_name']=='One'
    assert visible[3]['source_locator']['sheet_name']=='Two'
    assert visible[1]['source_row_sha256']!=visible[3]['source_row_sha256']
