"""
文件解析器
支持 PDF / Word(DOCX) / Excel(XLSX) / CSV / TXT
不依赖重型库，优先用标准库+可选依赖
"""
import os
import re
from typing import List, Dict
from ..config import IMPORTER_CONFIG


class ParseResult:
    """解析结果"""
    def __init__(self):
        self.text: str = ""
        self.tables: List[List[List[str]]] = []  # 表格数据
        self.metadata: Dict = {}
        self.paragraphs: List[str] = []


def parse_file(file_path: str) -> ParseResult:
    """
    根据文件扩展名自动选择解析器
    返回 ParseResult
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return parse_word(file_path)
    elif ext in (".xlsx", ".xls"):
        return parse_excel(file_path)
    elif ext == ".csv":
        return parse_csv(file_path)
    elif ext == ".txt":
        return parse_txt(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def parse_txt(file_path: str) -> ParseResult:
    """纯文本解析"""
    result = ParseResult()
    result.metadata["format"] = "txt"
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        result.text = f.read()
    result.paragraphs = [p.strip() for p in result.text.split("\n") if p.strip()]
    return result


def parse_csv(file_path: str) -> ParseResult:
    """CSV解析"""
    import csv
    result = ParseResult()
    result.metadata["format"] = "csv"
    table = []
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            table.append(row)
    result.tables = [table]
    result.text = "\n".join([",".join(row) for row in table])
    result.paragraphs = [",".join(row) for row in table if any(cell.strip() for cell in row)]
    return result


def parse_pdf(file_path: str) -> ParseResult:
    """
    PDF解析
    优先用 pdfplumber，回退到 PyPDF2，都没有则报错提示安装
    """
    result = ParseResult()
    result.metadata["format"] = "pdf"
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            result.metadata["pages"] = len(pdf.pages)
            all_text = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                all_text.append(text)
                # 提取表格
                tables = page.extract_tables()
                for t in tables:
                    result.tables.append(t)
            result.text = "\n\n".join(all_text)
            result.paragraphs = [p.strip() for p in result.text.split("\n") if p.strip()]
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            result.metadata["pages"] = len(reader.pages)
            all_text = []
            for page in reader.pages:
                all_text.append(page.extract_text() or "")
            result.text = "\n\n".join(all_text)
            result.paragraphs = [p.strip() for p in result.text.split("\n") if p.strip()]
        except ImportError:
            raise ImportError("PDF解析需要 pdfplumber 或 PyPDF2，请运行: pip install pdfplumber")
    return result


def parse_word(file_path: str) -> ParseResult:
    """
    Word DOCX 解析
    需要 python-docx
    """
    result = ParseResult()
    result.metadata["format"] = "word"
    try:
        from docx import Document
        doc = Document(file_path)
        all_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                all_text.append(para.text.strip())
                result.paragraphs.append(para.text.strip())
        # 提取表格
        for table in doc.tables:
            t_data = []
            for row in table.rows:
                t_data.append([cell.text.strip() for cell in row.cells])
            result.tables.append(t_data)
        result.text = "\n\n".join(all_text)
    except ImportError:
        raise ImportError("Word解析需要 python-docx，请运行: pip install python-docx")
    return result


def parse_excel(file_path: str) -> ParseResult:
    """
    Excel XLSX 解析
    需要 openpyxl 或 pandas
    """
    result = ParseResult()
    result.metadata["format"] = "excel"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        result.metadata["sheets"] = wb.sheetnames
        all_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            table = []
            for row in ws.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else "" for cell in row]
                if any(cell.strip() for cell in row_data):
                    table.append(row_data)
                    all_text.append(" | ".join(row_data))
            result.tables.append(table)
        result.text = "\n".join(all_text)
        result.paragraphs = all_text
    except ImportError:
        try:
            import pandas as pd
            xls = pd.ExcelFile(file_path)
            result.metadata["sheets"] = xls.sheet_names
            all_text = []
            for sheet in xls.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet)
                result.tables.append(df.values.tolist())
                all_text.append(df.to_string())
            result.text = "\n\n".join(all_text)
            result.paragraphs = all_text
        except ImportError:
            raise ImportError("Excel解析需要 openpyxl 或 pandas，请运行: pip install openpyxl")
    return result


def split_into_chunks(text: str, max_chunk_size: int = 500) -> List[str]:
    """将长文本切分成适合AI分析的块"""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks = []
    current = ""
    for p in paragraphs:
        if len(current) + len(p) > max_chunk_size and current:
            chunks.append(current.strip())
            current = p
        else:
            current += "\n" + p
    if current.strip():
        chunks.append(current.strip())
    return chunks
