import copy
import hashlib
import io
import json
from pathlib import Path

import pytest
from docx import Document

from tests.test_internal_pitch_workflow import workspace, arguments
from tests.test_saved_document_workflow import _source, _save
from tests.test_source_brief_job_routes import client, GLOBAL, URL
from tests.test_source_brief_job import job_db
from lumenis.shared.content_creation import internal_pitch as pitch
from lumenis.shared.content_creation.contracts import ContentContractError, resolve_creation_upstream


@pytest.fixture
def upstream(workspace,monkeypatch):
    ws=workspace
    import canonical_data
    from lumenis.shared.evidence import store
    from lumenis.importer import documents
    from lumenis.shared.reporting import saved_work
    monkeypatch.setenv("LUMENIS_EVIDENCE_HMAC_KEY_RING",json.dumps({"test":"isolated-creation-test-key"}))
    monkeypatch.setenv("LUMENIS_EVIDENCE_HMAC_ACTIVE_KEY_ID","test")
    for module in (store,documents):
        monkeypatch.setattr(module,"get_repository",lambda:ws["repo"])
    store.init_evidence_store(str(ws["repo"].db_path))
    with ws["repo"].transaction() as con:
        store.ensure_evidence_schema(con)
    canonical_data.ensure_canonical_table(str(ws["repo"].db_path))
    ws["documents"]=documents;ws["work"]=saved_work
    _,_,source=_source(ws)
    cite=documents.locate_quote(source["evidence_id"],actor=ws["actor"],scope=ws["scope"],start=0,end=20)["citation"]
    payload={"macro_user_work":{"question":"Explain the selected source","brief_id":"actual-saved-research-brief"},
             "selected_materials":[{"evidence_id":source["evidence_id"],"start":0,"end":20,
                                    "selection_reason":"This exact range supports the chosen question"}],
             "resource_constraints":{"currency":"GBP","maximum":2700,"duration_seconds":68},
             "unknowns":["Audience usefulness and actual owned model review remain open"]}
    research=_save(ws,request_id="upstream-material-01",payload=payload,citations=[cite])
    ref={"kind":"saved_work","artifact_id":research["artifact"]["artifact_id"],
         "expected_sha256":research["artifact"]["sha256"],"expected_version":research["document"]["version"]}
    return ws,research,ref


def values(ws,ref,request):
    manuscript={"title":"Original source-bound proposal","summary":"A complete proposal with retained research context.",
       "language":"en","authorship":"assisted","sections":[
          {"heading":"Task and source","body":"The chosen source and budget come from the saved research version."},
          {"heading":"Original approach","body":"Retain the original scope and a useful editable result."},
          {"heading":"Conditions","body":"The work remains a proposal; no external action or outcome has occurred."}]}
    return {"manuscript":manuscript,"request_id":request,"actor":ws["actor"],"scope":ws["scope"],
            "profile":"source-bound/v3","upstream_ref":ref}


def run(ws,params):
    task=pitch.prepare_internal_pitch(**params)
    return pitch.run_internal_pitch(task["task_id"],actor=ws["actor"],scope=ws["scope"])


def test_source_to_real_word_revision_return_and_reopen(upstream):
    ws,research,ref=upstream
    first=run(ws,values(ws,ref,"source-production-01"))
    doc=first["work"]["document"];binding=doc["payload"]["upstream_binding"]
    assert binding["document"]["payload"]==research["document"]["payload"]
    assert binding["document"]["citations"]==research["document"]["citations"]
    assert first["research_return_ref"]["artifact_id"]==research["artifact"]["artifact_id"]
    assert doc["parent_artifact_id"] is None
    download=pitch.download_internal_pitch(first["task_id"],actor=ws["actor"],scope=ws["scope"])
    assert hashlib.sha256(download["content"]).hexdigest()==first["action_receipt"]["output"]["sha256"]
    assert "The work remains a proposal; no external action or outcome has occurred." in [
        p.text for p in Document(io.BytesIO(download["content"])).paragraphs]
    revised=copy.deepcopy(research["document"]["payload"])
    revised["resource_constraints"]["maximum"]=1200
    second_research=_save(ws,request_id="upstream-material-02",payload=revised,citations=research["document"]["citations"],
                         work_id=research["document"]["work_id"],previous_artifact_id=research["artifact"]["artifact_id"])
    updated={"kind":"saved_work","artifact_id":second_research["artifact"]["artifact_id"],
             "expected_sha256":second_research["artifact"]["sha256"],"expected_version":2}
    params=values(ws,updated,"source-production-02")
    params["manuscript"]["sections"][1]["body"]="The reduced budget changes the original production scope."
    params.update(work_id=doc["work_id"],previous_artifact_id=first["action_receipt"]["artifact_id"],
                  expected_draft_digest=doc["draft_digest"])
    second=run(ws,params)
    assert second["work"]["document"]["version"]==2
    assert second["work"]["document"]["parent_artifact_id"]==first["action_receipt"]["artifact_id"]
    assert second["research_return_ref"]["artifact_id"]==second_research["artifact"]["artifact_id"]
    assert second["work"]["document"]["payload"]["upstream_binding"]["document"]["payload"]==revised
    assert pitch.read_internal_pitch(first["task_id"],actor=ws["actor"],scope=ws["scope"])["research_return_ref"][
        "artifact_id"]==research["artifact"]["artifact_id"]
    assert pitch.cancel_internal_pitch(second["task_id"],actor=ws["actor"],scope=ws["scope"])["cancellation"]=="ALREADY_COMMITTED"


@pytest.mark.parametrize("change",[
    {"expected_sha256":"0"*64},{"expected_version":2},{"expected_version":True},
    {"trusted_payload":{"grant":True}},{"kind":"client_context"},
])
def test_stale_or_forged_reference_creates_no_task(upstream,change):
    ws,research,ref=upstream
    before=ws["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0]
    ref={**ref,**change}
    with pytest.raises((ContentContractError,ValueError,PermissionError)):
        pitch.prepare_internal_pitch(**values(ws,ref,"bad-reference"))
    assert ws["repo"].query_one("SELECT COUNT(*) FROM ai_task")[0]==before


def test_same_transaction_source_withdrawal_prevents_new_work(upstream,monkeypatch):
    ws,research,ref=upstream
    prepared=pitch.prepare_internal_pitch(**values(ws,ref,"withdraw-before-commit"))
    original=ws["work"].save_work
    def save_with_withdrawal(*args,**kwargs):
        actual=kwargs["before_commit"]
        def before_commit(con,artifact,document):
            con.execute("UPDATE lumenis_subject_registry SET status='inactive' WHERE subject_id=?",("SUBJ-PITCH",))
            actual(con,artifact,document)
        kwargs["before_commit"]=before_commit
        return original(*args,**kwargs)
    monkeypatch.setattr(ws["work"],"save_work",save_with_withdrawal)
    before=ws["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]
    with pytest.raises(ValueError,match="WORK_PERMISSION_DENIED"):
        pitch.run_internal_pitch(prepared["task_id"],actor=ws["actor"],scope=ws["scope"])
    assert ws["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]==before
    assert ws["repo"].query_one("SELECT status FROM lumenis_subject_registry WHERE subject_id='SUBJ-PITCH'")[0]=="configured"
    assert pitch.read_internal_pitch(prepared["task_id"],actor=ws["actor"],scope=ws["scope"])["state"]=="failed"


def test_prepared_cancel_preserves_upstream_and_no_word(upstream):
    ws,research,ref=upstream
    task=pitch.prepare_internal_pitch(**values(ws,ref,"source-cancel-01"))
    before=ws["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]
    cancelled=pitch.cancel_internal_pitch(task["task_id"],actor=ws["actor"],scope=ws["scope"])
    assert cancelled["state"]=="cancelled" and cancelled["work"] is None
    assert ws["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]==before
    assert ws["work"].read_committed_work(research["artifact"]["artifact_id"],actor=ws["actor"],scope=ws["scope"])[
        "document"]["payload"]==research["document"]["payload"]


def test_read_revoked_source_never_returns_saved_production(upstream):
    ws,research,ref=upstream
    result=run(ws,values(ws,ref,"source-read-revoked"))
    with ws["repo"].transaction() as con:
        con.execute("UPDATE lumenis_subject_registry SET status='inactive' WHERE subject_id='SUBJ-PITCH'")
    with pytest.raises(ValueError,match="WORK_PERMISSION_DENIED"):
        pitch.read_internal_pitch(result["task_id"],actor=ws["actor"],scope=ws["scope"])


def test_source_brief_normal_reader_to_complete_word_and_return(client,monkeypatch):
    api,app,db,entry,gate=client
    headers={**GLOBAL,"X-Lumenis-Context-Mode":"FOCUS","X-Subject-ID":entry["subject_id"]}
    response=api.post(URL+"/generate",headers=headers,json={})
    assert response.status_code==200
    report=response.get_json()["result"]
    from lumenis.system.security import permission
    from lumenis.shared.database import get_repository
    from lumenis.shared.reporting.source_brief_job import _media_scope
    import ai_task_engine
    import shared_artifacts
    repository=get_repository()
    monkeypatch.setattr(pitch,"get_repository",lambda:repository)
    monkeypatch.setattr(ai_task_engine,"get_repository",lambda:repository)
    with repository.transaction() as con:
        con.executescript(ai_task_engine._TASK_ENGINE_SCHEMA)
    shared_artifacts.ensure_artifact_table()
    pitch.engine.cache_clear()
    try:
        with app.test_request_context("/",headers={**headers,"Host":"localhost"},
                                     environ_overrides={"REMOTE_ADDR":"127.0.0.1"}):
            actor=permission.resolve_local_workspace_principal(module="reports",action="write")
            scope=_media_scope(report["source_brief"]["scope"])
            ref={"kind":"source_brief_job","job_id":report["job_id"],
                 "expected_artifact_hash":report["artifact_outcome"]["artifact_hash"]}
            ws={"actor":actor,"scope":scope}
            finished=run(ws,values(ws,ref,"source-brief-production-01"))
            inherited=finished["work"]["document"]["payload"]["upstream_binding"]
            assert inherited["document"]["source_brief"]==report["source_brief"]
            assert inherited["document"]["requested_source_ids"]==report["requested_source_ids"]
            assert inherited["source_role"]=="source_scope_proof_report_not_shared_artifact_or_new_observation"
            assert finished["research_return_ref"]["job_id"]==report["job_id"]
            output=pitch.download_internal_pitch(finished["task_id"],actor=actor,scope=scope)
            assert hashlib.sha256(output["content"]).hexdigest()==finished["action_receipt"]["output"]["sha256"]
    finally:
        pitch.engine.cache_clear()
