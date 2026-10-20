"""Editable prose export, revision safety and cancellation at publication."""

import hashlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.pitch_export import _remove_title_borders, export_pitch_docx


def manuscript(body="先确认需求中的预算与日期。\n\n保留未知信息。\nReview the source."):
    return {
        "title": "图书馆青年传播提案",
        "summary": "邀请年轻人在图书馆完成自己正在做的事。",
        "authorship": "assisted",
        "sections": [{"heading": "策略与执行", "body": body}],
    }


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def minimal_audiovisual_snapshot():
    from lumenis.shared.content_creation.contracts import AudiovisualPlan

    proposed = {"kind": "PLANNER_PROPOSAL", "locator": "synthetic-test-input"}
    unknown = {"kind": "UNKNOWN", "locator": None}
    plan = AudiovisualPlan.from_dict({
        "schema": "lumenis.audiovisual-plan/v1",
        "brief_text": "Synthetic style regression input",
        "method_source": {"reference": "synthetic-method", "sha256": "0" * 64},
        "outputs": [{
            "output_id": "synthetic-video", "medium": "video", "purpose": "Style regression",
            "canvas_ratio": "16:9",
            "duration": {"min_seconds": None, "max_seconds": None, "tolerance_seconds": None},
            "sound": {"original": "optional", "voice": "optional", "music": "optional"},
            "languages": ["zh"], "constraints": {}, "plan": None,
            "provenance": {"purpose": proposed, "canvas_ratio": proposed, "duration": unknown,
                           "languages": proposed, "constraints": proposed, "original": proposed,
                           "voice": proposed, "music": proposed},
        }],
        "rule_assessments": [],
    })
    return plan.snapshot()


def test_actual_manuscript_survives_word_reopen(tmp_path):
    draft = manuscript()
    path = tmp_path / "proposal-v1.docx"
    result = export_pitch_docx({"manuscript": draft}, path)
    reopened = Document(result)
    paragraphs = [p.text for p in reopened.paragraphs]
    assert paragraphs[0] == draft["title"]
    assert paragraphs[1] == draft["summary"]
    assert paragraphs[-3:] == [draft["sections"][0]["heading"], *draft["sections"][0]["body"].split("\n\n")]
    assert reopened.paragraphs[0].style.name == "Title"
    assert reopened.styles["Normal"].paragraph_format.space_after.pt == 6
    assert "authorship=assisted" in reopened.core_properties.comments
    assert "approval=not_verified" in reopened.core_properties.comments
    assert not list(tmp_path.glob(".pitch-stage-*"))


def test_normal_and_av_manuscripts_keep_six_point_body_spacing_and_plain_title(tmp_path):
    for name, audiovisual_plan in (("normal", None), ("av", minimal_audiovisual_snapshot())):
        path = tmp_path / f"{name}.docx"
        export_pitch_docx({}, path, manuscript=manuscript(), audiovisual_plan=audiovisual_plan)
        reopened = Document(path)
        assert reopened.styles["Normal"].paragraph_format.space_after.pt == 6
        assert not list(reopened.styles["Title"].element.iter(qn("w:pBdr")))
        assert not list(reopened.paragraphs[0]._p.iter(qn("w:pBdr")))


def test_title_border_cleanup_removes_style_and_direct_paragraph_borders():
    doc = Document()
    title = doc.add_paragraph("Synthetic title", style="Title")
    for element in (doc.styles["Title"].element, title._p):
        properties = element.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        borders.append(bottom)
        properties.append(borders)

    _remove_title_borders(doc)

    assert not list(doc.styles["Title"].element.iter(qn("w:pBdr")))
    assert not list(title._p.iter(qn("w:pBdr")))


def test_revision_preserves_original_and_reopens_new_prose(tmp_path):
    first = tmp_path / "proposal-v1.docx"
    second = tmp_path / "proposal-v2.docx"
    export_pitch_docx({}, first, manuscript=manuscript())
    original = digest(first)
    revised = manuscript("改为先试行一场活动，审阅结果后决定下一轮。")
    export_pitch_docx({}, second, manuscript=revised)
    assert digest(first) == original
    assert Document(second).paragraphs[-1].text == revised["sections"][0]["body"]
    assert digest(second) != original
    with pytest.raises(FileExistsError):
        export_pitch_docx({}, first, manuscript=revised)
    assert digest(first) == original


def test_cancellation_after_staging_has_no_published_output(tmp_path):
    destination = tmp_path / "cancelled.docx"
    checks = 0

    def cancelled_at_publish():
        nonlocal checks
        checks += 1
        if checks == 3:
            raise RuntimeError("task cancelled")

    with pytest.raises(RuntimeError, match="task cancelled"):
        export_pitch_docx({}, destination, manuscript=manuscript(), check_continue=cancelled_at_publish)
    assert checks == 3
    assert not destination.exists()
    assert not list(tmp_path.glob(".pitch-stage-*"))


def test_failed_publish_does_not_leave_partial_result(tmp_path, monkeypatch):
    destination = tmp_path / "failed.docx"

    def unavailable_link(*args):
        raise OSError("filesystem cannot publish atomically")

    monkeypatch.setattr("scripts.pitch_export.os.link", unavailable_link)
    with pytest.raises(OSError, match="cannot publish atomically"):
        export_pitch_docx({}, destination, manuscript=manuscript())
    assert not destination.exists()
    assert not list(tmp_path.glob(".pitch-stage-*"))


def test_destination_created_during_export_is_never_replaced(tmp_path, monkeypatch):
    destination = tmp_path / "same-name.docx"
    concurrent = b"another writer's revision"

    def collision(*args):
        destination.write_bytes(concurrent)
        raise FileExistsError("destination won by another writer")

    monkeypatch.setattr("scripts.pitch_export.os.link", collision)
    with pytest.raises(FileExistsError):
        export_pitch_docx({}, destination, manuscript=manuscript())
    assert destination.read_bytes() == concurrent
    assert not list(tmp_path.glob(".pitch-stage-*"))


def test_invalid_ooxml_body_leaves_no_partial_output(tmp_path):
    with pytest.raises(ValueError):
        export_pitch_docx({}, tmp_path / "invalid.docx", manuscript=manuscript("bad\x00text"))
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("draft", [
    {"title": "x", "summary": "x", "sections": []},
    {"title": "x", "summary": "x", "sections": [{"heading": "x", "body": " "}]},
    {**manuscript(), "authorship": "client_approved"},
])
def test_incomplete_or_false_approval_manuscript_is_rejected(tmp_path, draft):
    with pytest.raises(ValueError):
        export_pitch_docx({}, tmp_path / "invalid.docx", manuscript=draft)
    assert not list(tmp_path.iterdir())


def test_diagnostic_export_is_explicit_about_its_limits(tmp_path):
    path = export_pitch_docx({"extracted": {"brand": "待确认", "budget": None}}, tmp_path / "analysis.docx")
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "不是完整提案成稿" in text
    assert "历史对标尚未经本次核验" in text


@pytest.mark.parametrize("language,label", [("en", "Client approval pending"),
                                           ("ko", "고객 승인 대기")])
def test_declared_document_language_keeps_approval_state_understandable(tmp_path, language, label):
    draft = {**manuscript(), "language": language}
    reopened = Document(export_pitch_docx({}, tmp_path / f"{language}.docx", manuscript=draft))
    assert reopened.core_properties.language == language
    assert label in reopened.paragraphs[2].text
