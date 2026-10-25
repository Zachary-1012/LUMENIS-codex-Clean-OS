"""Offline, source-bound audiograms and quote cards for the existing media owner.

The caller owns project/Artifact persistence, permission and source-use approval.
An explicitly supplied excerpt and literal copy are not ASR, semantic quote
verification, story selection, platform-specific packaging or publication.
"""
from __future__ import annotations

from array import array
import hashlib
import base64
import io
import json
import math
import re
from pathlib import Path
import sys
import tempfile
import unicodedata
import wave
from copy import deepcopy
from typing import Any, Callable

from . import rendering
from .video_editing import (
    MEDIA_INPUT_DEMUXERS, MediaAsset, VideoEditingError, _ffprobe_path,
    fingerprint_local_media,
)


def validate_visual_document(value, *, previous=None):
    """Editable user material in the original SavedWork, never external Evidence.

    Embedded raster originals remain byte-identical in JSON. Exports decode and
    re-encode them locally; SVG/script/URL inputs are not an image source.
    """
    from PIL import Image
    if (type(value) is not dict or set(value) != {'schema','canvas','assets','layers'}
            or value['schema'] != 'lumenis.visual-document/v1'):
        raise VideoEditingError('VISUAL_DOCUMENT_INVALID')
    canvas = value['canvas']
    if (type(canvas) is not dict or set(canvas) != {'width','height','background'}
            or any(type(canvas[k]) is not int or not 16 <= canvas[k] <= 8192 for k in ('width','height'))
            or canvas['width'] * canvas['height'] > 32_000_000):
        raise VideoEditingError('VISUAL_CANVAS_INVALID')
    color = lambda item: isinstance(item,str) and re.fullmatch(r'#[0-9a-fA-F]{6}(?:[0-9a-fA-F]{2})?',item)
    if not color(canvas['background']):
        raise VideoEditingError('VISUAL_COLOR_INVALID')
    assets = value['assets']
    if type(assets) is not list or len(assets) > 64:
        raise VideoEditingError('VISUAL_ASSETS_INVALID')
    asset_ids = set(); total = 0
    for asset in assets:
        if (type(asset) is not dict or set(asset) != {'id','media_type','sha256','data_base64'}
                or not isinstance(asset['id'],str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,64}',asset['id'])
                or asset['id'] in asset_ids or not isinstance(asset['media_type'],str)
                or asset['media_type'] not in {'image/png','image/jpeg','image/webp'}
                or not isinstance(asset['data_base64'],str) or len(asset['data_base64']) > 8*1024*1024):
            raise VideoEditingError('VISUAL_ASSET_INVALID')
        try:
            raw = base64.b64decode(asset['data_base64'],validate=True); total += len(raw)
            if not raw or total > 6*1024*1024 or hashlib.sha256(raw).hexdigest() != asset['sha256']:
                raise ValueError('original bytes')
            with Image.open(io.BytesIO(raw)) as image:
                if (image.format != {'image/png':'PNG','image/jpeg':'JPEG','image/webp':'WEBP'}[asset['media_type']]
                        or image.width*image.height > 32_000_000 or getattr(image,'n_frames',1) != 1):
                    raise ValueError('static bounded raster required')
                image.verify()
        except (ValueError,TypeError,OSError,Image.DecompressionBombError) as exc:
            raise VideoEditingError('VISUAL_ASSET_INTEGRITY_FAILED') from exc
        asset_ids.add(asset['id'])
    if type(value['layers']) is not list or not 1 <= len(value['layers']) <= 256:
        raise VideoEditingError('VISUAL_LAYERS_INVALID')
    common = {'id','kind','name','x','y','width','height','rotation','opacity','visible','locked'}
    specifics = {'text':{'text','font_size','color','align','line_height'},
                 'rectangle':{'fill','stroke','stroke_width'},'ellipse':{'fill','stroke','stroke_width'},
                 'image':{'asset_id','fit'}}
    ids = set()
    for layer in value['layers']:
        if (type(layer) is not dict or not isinstance(layer.get('kind'),str) or layer['kind'] not in specifics
                or set(layer) != common|specifics[layer['kind']]
                or not isinstance(layer['id'],str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,64}',layer['id']) or layer['id'] in ids
                or not isinstance(layer['name'],str) or len(layer['name']) > 160
                or type(layer['visible']) is not bool or type(layer['locked']) is not bool):
            raise VideoEditingError('VISUAL_LAYER_INVALID')
        for key,minimum,maximum in [('x',-16384,16384),('y',-16384,16384),('width',1,16384),
                                    ('height',1,16384),('rotation',-360,360),('opacity',0,1)]:
            number = layer[key]
            if type(number) not in (int,float) or not math.isfinite(number) or not minimum <= number <= maximum:
                raise VideoEditingError('VISUAL_GEOMETRY_INVALID')
        # A single off-canvas tile must not evade the canvas memory budget.
        if layer['width']*layer['height'] > 32_000_000:
            raise VideoEditingError('VISUAL_LAYER_RESOURCE_BUDGET')
        if layer['kind'] == 'text':
            _copy_text(layer['text'],64000)
            if (type(layer['font_size']) not in (int,float) or not 1 <= layer['font_size'] <= 1024
                    or type(layer['line_height']) not in (int,float) or not 1 <= layer['line_height'] <= 3
                    or not isinstance(layer['align'],str) or layer['align'] not in {'left','center','right'} or not color(layer['color'])):
                raise VideoEditingError('VISUAL_TYPOGRAPHY_INVALID')
        elif layer['kind'] == 'image':
            if (not isinstance(layer['asset_id'],str) or layer['asset_id'] not in asset_ids
                    or not isinstance(layer['fit'],str) or layer['fit'] not in {'contain','cover','stretch'}):
                raise VideoEditingError('VISUAL_IMAGE_BINDING_INVALID')
        elif (not color(layer['fill']) or not color(layer['stroke'])
              or type(layer['stroke_width']) not in (int,float) or not 0 <= layer['stroke_width'] <= 256):
            raise VideoEditingError('VISUAL_SHAPE_STYLE_INVALID')
        ids.add(layer['id'])
    if previous is not None:
        validate_visual_document(previous)
        current = {item['id']:item for item in value['layers']}
        current_assets = {item['id']:item['sha256'] for item in assets}
        old_assets = {item['id']:item['sha256'] for item in previous['assets']}
        for index,old in enumerate(previous['layers']):
            new = current.get(old['id'])
            if old['locked'] and (new is None or new['locked'] and (
                    new != old or value['layers'].index(new) != index
                    or old['kind'] == 'image' and current_assets.get(old['asset_id']) != old_assets[old['asset_id']])):
                raise VideoEditingError('VISUAL_LAYER_LOCKED')
    return deepcopy(value)


def visual_document_markdown(value):
    value = validate_visual_document(value)
    return ('Editable visual document · user-authored, source rights not independently verified\n\n'
            + '\n\n'.join(layer['text'] for layer in value['layers'] if layer['kind'] == 'text'))


def export_visual_document(value, *, format, check_cancel=None):
    """Local, literal PNG/SVG/PDF/PPTX; editable JSON remains the master."""
    from PIL import Image, ImageDraw, ImageFont, ImageOps
    from xml.etree import ElementTree as ET
    value = validate_visual_document(value)
    if format not in {'png','svg','pdf','pptx'}:
        raise VideoEditingError('VISUAL_EXPORT_FORMAT_INVALID')
    width,height = value['canvas']['width'],value['canvas']['height']
    if not rendering.CAPTION_FONT.is_file():
        raise VideoEditingError('VISUAL_FONT_UNAVAILABLE')
    font_hash = hashlib.sha256(rendering.CAPTION_FONT.read_bytes()).hexdigest()
    canvas = Image.new('RGBA',(width,height),value['canvas']['background'])
    assets = {item['id']:item for item in value['assets']}
    namespace = 'http://www.w3.org/2000/svg'
    ET.register_namespace('',namespace)
    tag = lambda name:'{'+namespace+'}'+name
    svg = ET.Element(tag('svg'),{'width':str(width),'height':str(height),'viewBox':f'0 0 {width} {height}',
        'version':'1.1'})
    ET.SubElement(svg,tag('title')).text = 'Editable user-authored visual; rights and claims require review'
    def rgba(value):
        from PIL.ImageColor import getcolor
        return getcolor(value,'RGBA')
    def svg_color(value):
        r,g,b,a = rgba(value)
        return f'#{r:02x}{g:02x}{b:02x}',str(a/255)
    bg,alpha = svg_color(value['canvas']['background'])
    ET.SubElement(svg,tag('rect'),{'width':str(width),'height':str(height),'fill':bg,'fill-opacity':alpha})
    pptx_layers = []
    for layer in value['layers']:
        _check(check_cancel)
        w,h = max(1,round(layer['width'])),max(1,round(layer['height']))
        tile = Image.new('RGBA',(w,h));draw = ImageDraw.Draw(tile)
        group = ET.SubElement(svg,tag('g'),{'id':layer['id'],'data-name':layer['name'],
            'data-locked':str(layer['locked']).lower(),'opacity':str(layer['opacity']),
            'display':'inline' if layer['visible'] else 'none',
            'transform':f"translate({layer['x']} {layer['y']}) rotate({layer['rotation']} {layer['width']/2} {layer['height']/2})"})
        lines = None; canonical_image = None
        if layer['kind'] == 'text':
            rendering._verify_caption_glyphs(layer['text'])
            font = ImageFont.truetype(str(rendering.CAPTION_FONT),max(1,round(layer['font_size'])))
            lines = _wrap_copy(draw,layer['text'],font,w).split('\n')
            ascent,descent = font.getmetrics(); step = (ascent+descent)*layer['line_height']
            if len(lines)*step > h:
                raise VideoEditingError('VISUAL_TEXT_OVERFLOW_REQUIRES_EDIT')
            for index,line in enumerate(lines):
                bounds = draw.textbbox((0,0),line,font=font)
                text_width = bounds[2]-bounds[0]
                if text_width > w:
                    raise VideoEditingError('VISUAL_TEXT_OVERFLOW_REQUIRES_EDIT')
                x = 0 if layer['align'] == 'left' else (w-text_width)/2 if layer['align'] == 'center' else w-text_width
                draw.text((x,index*step-bounds[1]),line,font=font,fill=layer['color'])
                fill,opacity = svg_color(layer['color'])
                ET.SubElement(group,tag('text'),{'x':str(x),'y':str(index*step+ascent),
                    'font-size':str(layer['font_size']),'font-family':font.getname()[0],
                    'fill':fill,'fill-opacity':opacity,'xml:space':'preserve'}).text = line
        elif layer['kind'] in {'rectangle','ellipse'}:
            shape = draw.rectangle if layer['kind'] == 'rectangle' else draw.ellipse
            stroke = round(layer['stroke_width'])
            shape((0,0,w-1,h-1),fill=layer['fill'],outline=layer['stroke'] if stroke else None,width=max(1,stroke))
            fill,fill_alpha = svg_color(layer['fill']);stroke_color,stroke_alpha = svg_color(layer['stroke'])
            geometry = ({'x':'0','y':'0','width':str(layer['width']),'height':str(layer['height'])}
                        if layer['kind'] == 'rectangle' else
                        {'cx':str(layer['width']/2),'cy':str(layer['height']/2),'rx':str(layer['width']/2),'ry':str(layer['height']/2)})
            ET.SubElement(group,tag('rect' if layer['kind'] == 'rectangle' else 'ellipse'),
                {**geometry,'fill':fill,'fill-opacity':fill_alpha,'stroke':stroke_color,
                 'stroke-opacity':stroke_alpha,'stroke-width':str(layer['stroke_width'])})
        else:
            raw = base64.b64decode(assets[layer['asset_id']]['data_base64'],validate=True)
            with Image.open(io.BytesIO(raw)) as original:
                original = ImageOps.exif_transpose(original).convert('RGBA')
                image = (ImageOps.fit(original,(w,h),method=Image.Resampling.LANCZOS) if layer['fit'] == 'cover' else
                         ImageOps.contain(original,(w,h),method=Image.Resampling.LANCZOS) if layer['fit'] == 'contain' else
                         original.resize((w,h),Image.Resampling.LANCZOS))
                tile.alpha_composite(image,((w-image.width)//2,(h-image.height)//2))
            buffer = io.BytesIO();tile.save(buffer,format='PNG');canonical_image = buffer.getvalue()
            ET.SubElement(group,tag('image'),{'width':str(layer['width']),'height':str(layer['height']),
                'href':'data:image/png;base64,'+base64.b64encode(canonical_image).decode('ascii')})
        pptx_layers.append((layer,lines,canonical_image))
        if not layer['visible']:
            continue
        if layer['opacity'] != 1:
            tile.putalpha(tile.getchannel('A').point(lambda number:round(number*layer['opacity'])))
        rotated = tile.rotate(-layer['rotation'],resample=Image.Resampling.BICUBIC,expand=True)
        canvas.alpha_composite(rotated,(round(layer['x']+(w-rotated.width)/2),round(layer['y']+(h-rotated.height)/2)))
    _check(check_cancel)
    if hashlib.sha256(rendering.CAPTION_FONT.read_bytes()).hexdigest() != font_hash:
        raise VideoEditingError('VISUAL_FONT_CHANGED')
    if format == 'svg':
        return ET.tostring(svg,encoding='utf-8',xml_declaration=True)
    output = io.BytesIO()
    if format == 'pptx':
        from pptx import Presentation
        from pptx.util import Emu,Pt
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN,MSO_AUTO_SIZE
        from pptx.dml.color import RGBColor
        from pptx.oxml.xmlchemy import OxmlElement
        presentation = Presentation();presentation.slide_width = Emu(width*9525);presentation.slide_height = Emu(height*9525)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb = RGBColor.from_string(bg[1:])
        def set_color(color_format,color_value,opacity):
            red,green,blue,alpha = rgba(color_value);color_format.rgb = RGBColor(red,green,blue)
            element = color_format._xFill.srgbClr
            alpha_element = OxmlElement('a:alpha');alpha_element.set('val',str(round(alpha/255*opacity*100000)))
            element.append(alpha_element)
        for layer,lines,image in pptx_layers:
            left,top,w,h = [Emu(round(layer[key]*9525)) for key in ('x','y','width','height')]
            if layer['kind'] == 'image':
                shape = slide.shapes.add_picture(io.BytesIO(image),left,top,width=w,height=h)
                alpha = OxmlElement('a:alphaModFix');alpha.set('amt',str(round(layer['opacity']*100000)))
                shape._element.blipFill.blip.append(alpha)
            elif layer['kind'] == 'text':
                shape = slide.shapes.add_textbox(left,top,w,h);text_frame = shape.text_frame
                text_frame.auto_size = MSO_AUTO_SIZE.NONE;text_frame.word_wrap = False
                text_frame.margin_left = text_frame.margin_right = text_frame.margin_top = text_frame.margin_bottom = 0
                for index,line in enumerate(lines):
                    paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
                    paragraph.text = line;paragraph.font.name = font.getname()[0]
                    paragraph.font.size = Pt(layer['font_size']*.75)
                    paragraph.line_spacing = layer['line_height'];paragraph.space_after = Pt(0)
                    paragraph.alignment = {'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[layer['align']]
                    set_color(paragraph.font.color,layer['color'],layer['opacity'])
            else:
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if layer['kind'] == 'rectangle' else MSO_SHAPE.OVAL,left,top,w,h)
                shape.fill.solid();set_color(shape.fill.fore_color,layer['fill'],layer['opacity'])
                if layer['stroke_width']:
                    shape.line.width = Pt(layer['stroke_width']*.75);set_color(shape.line.color,layer['stroke'],layer['opacity'])
                else:
                    shape.line.fill.background()
            shape.name = layer['name'];shape.rotation = layer['rotation']%360
            shape._element.nvSpPr.cNvPr.set('hidden','1' if not layer['visible'] else '0') if layer['kind'] != 'image' else shape._element.nvPicPr.cNvPr.set('hidden','1' if not layer['visible'] else '0')
        presentation.save(output)
    elif format == 'png':
        canvas.save(output,format='PNG')
    else:
        flattened = Image.new('RGB',canvas.size,'white');flattened.paste(canvas,mask=canvas.getchannel('A'))
        flattened.save(output,format='PDF',resolution=96,title='User-authored visual document')
    return output.getvalue()


MAX_EXCERPT_SECONDS = 180.0
MIN_EXCERPT_SECONDS = .2
MAX_OUTPUT_BYTES = 128 * 1024 * 1024
SAMPLE_RATE = 48000
FRAME_RATE = 30
WAVEFORM_BINS = 256


def _check(callback: Callable[[], None] | None) -> None:
    if callback:
        callback()


def _copy_text(value: Any, limit: int) -> str:
    # Rasterize literal text with Pillow; never insert copy in a filter, shell,
    # filename, subtitle markup or encoder metadata. Reject bidi/control tricks.
    if (not isinstance(value, str) or not value.strip() or len(value) > limit
            or any(unicodedata.category(char).startswith("C") and char != "\n" for char in value)):
        raise VideoEditingError("AUDIO_VISUAL_COPY_INVALID")
    return value


def _number(value: float) -> str:
    return f"{value:.9f}"


def _timecode(seconds: float) -> str:
    milliseconds = round(seconds * 1000)
    hours, remaining = divmod(milliseconds, 3600000)
    minutes, remaining = divmod(remaining, 60000)
    secs, millis = divmod(remaining, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"


def _source_unchanged(asset: MediaAsset, check_cancel) -> None:
    _check(check_cancel)
    current = fingerprint_local_media(asset.path)
    if current.to_dict() != asset.to_dict():
        raise VideoEditingError("AUDIO_VISUAL_SOURCE_CHANGED")
    _check(check_cancel)


def _decode_excerpt(asset: MediaAsset, *, ffmpeg: str, work: Path, start: float,
                    duration: float, check_cancel) -> tuple[Path, dict[str, Any]]:
    output = work / "excerpt.wav"
    rendering._run([
        ffmpeg, "-v", "error", "-xerror", "-nostdin", "-n", "-threads", "1",
        "-err_detect", "explode", "-ss", _number(start), "-protocol_whitelist", "file",
        "-format_whitelist", MEDIA_INPUT_DEMUXERS, "-i", asset.path, "-map", "0:a:0",
        "-vn", "-sn", "-dn", "-t", _number(duration), "-map_metadata", "-1",
        "-ac", "2", "-ar", str(SAMPLE_RATE), "-c:a", "pcm_s16le", "-threads:a", "1",
        "-f", "wav", str(output),
    ], cwd=work, timeout=120, check_continue=check_cancel)
    if output.stat().st_size > int(MAX_EXCERPT_SECONDS * SAMPLE_RATE * 4) + 4096:
        raise VideoEditingError("AUDIO_VISUAL_DECODE_LIMIT_EXCEEDED")
    try:
        with wave.open(str(output), "rb") as pcm:
            count = pcm.getnframes()
            if (pcm.getnchannels() != 2 or pcm.getframerate() != SAMPLE_RATE
                    or pcm.getsampwidth() != 2 or pcm.getcomptype() != "NONE"
                    or abs(count - round(duration * SAMPLE_RATE)) > 2):
                # A short/missing track must not become a silently padded clip.
                raise ValueError("excerpt does not cover the requested interval")
            raw = pcm.readframes(count)
        if len(raw) != count * 4 or count == 0:
            raise ValueError("truncated PCM")
    except (OSError, EOFError, wave.Error, ValueError) as exc:
        raise VideoEditingError("AUDIO_VISUAL_AUDIO_COVERAGE_INVALID") from exc
    samples = array("h")
    samples.frombytes(raw)
    if sys.byteorder != "little":
        samples.byteswap()
    peaks = []
    # The visual uses the channel extrema, avoiding cancellation of opposite-
    # phase stereo. Each bin covers every decoded sample in its time interval.
    for index in range(WAVEFORM_BINS):
        _check(check_cancel)
        first, last = index * count // WAVEFORM_BINS, (index + 1) * count // WAVEFORM_BINS
        chunk = samples[first * 2:last * 2]
        peaks.append({"min": min(chunk) / 32768, "max": max(chunk) / 32768,
                      "rms": math.sqrt(sum(value * value for value in chunk) / len(chunk)) / 32768})
    return output, {
        "method": "all_decoded_pcm_samples_channel_extrema",
        "stage": "excerpt_pcm_before_aac_encoding", "sample_rate": SAMPLE_RATE,
        "channels": 2, "sample_format": "s16le", "decoded_samples_per_channel": count,
        "actual_duration_sec": count / SAMPLE_RATE,
        "pcm_sha256": hashlib.sha256(raw).hexdigest(), "bins": peaks,
        "normalization": "none", "channel_conversion": "ffmpeg_stereo",
    }


def _wrap_copy(draw, text: str, font, width: int) -> str:
    lines = []
    for paragraph in text.split("\n"):
        clusters: list[str] = []
        for char in paragraph:
            if clusters and unicodedata.category(char).startswith("M"):
                clusters[-1] += char
            else:
                clusters.append(char)
        line: list[str] = []
        for cluster in clusters:
            candidate = "".join(line) + cluster
            bounds = draw.textbbox((0, 0), candidate, font=font)
            if line and bounds[2] - bounds[0] > width:
                # Keep CJK closing punctuation with a preceding character and
                # opening punctuation with the following character. No copy is
                # dropped; explicit user newlines remain unchanged.
                move = int(len(line) > 1 and (
                    cluster[0] in "，。；：！？、,.!?:;"
                    or unicodedata.category(cluster[0]) in {"Pe", "Pf"}
                    or unicodedata.category(line[-1][0]) in {"Ps", "Pi"}))
                if move:
                    lines.append("".join(line[:-1]))
                    line = [line[-1], cluster]
                else:
                    lines.append("".join(line))
                    line = [cluster]
            else:
                line.append(cluster)
        lines.append("".join(line))
    return "\n".join(lines)


def _draw_copy(draw, text: str, box: tuple[int, int, int, int], *, maximum: int,
               color: tuple[int, int, int], check_cancel) -> dict[str, Any]:
    from PIL import ImageFont
    x, y, right, bottom = box
    for size in range(maximum, 17, -1):
        _check(check_cancel)
        font = ImageFont.truetype(str(rendering.CAPTION_FONT), size, index=0)
        wrapped = _wrap_copy(draw, text, font, right - x)
        spacing = max(4, size // 4)
        bounds = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=spacing)
        width, height = bounds[2] - bounds[0], bounds[3] - bounds[1]
        if width <= right - x and height <= bottom - y:
            draw.multiline_text((x - bounds[0], y - bounds[1]), wrapped, font=font,
                                spacing=spacing, fill=color)
            return {"font_size": size, "lines": wrapped.count("\n") + 1,
                    "bounds": [x, y, x + width, y + height], "safe_box": list(box)}
    raise VideoEditingError("AUDIO_VISUAL_COPY_LAYOUT_OVERFLOW")


def _card(*, title: str, quote: str, stamp: str, width: int, height: int,
          waveform: dict[str, Any] | None, check_cancel, typography_scale: int = 1) -> tuple[bytes, dict[str, Any]]:
    from PIL import Image, ImageDraw
    canvas = Image.new("RGB", (width, height), (15, 23, 38))
    draw = ImageDraw.Draw(canvas)
    margin = max(36, min(width, height) // 15)
    right = width - margin
    draw.rounded_rectangle((margin, margin, margin + 72 * typography_scale, margin + 8 * typography_scale), radius=4 * typography_scale,
                           fill=(71, 214, 192))
    layout = {
        "title": _draw_copy(draw, title, (margin, margin + 35, right, int(height * .27)),
                            maximum=46 * typography_scale, color=(232, 239, 245), check_cancel=check_cancel),
        "quote": _draw_copy(draw, quote, (margin, int(height * .31), right,
                                           int(height * (.66 if waveform else .79))),
                            maximum=(44 if waveform else 52) * typography_scale, color=(255, 255, 255),
                            check_cancel=check_cancel),
    }
    if waveform:
        top, bottom = int(height * .72), int(height * .86)
        centre, amplitude = (top + bottom) / 2, (bottom - top) / 2
        draw.line((margin, centre, right, centre), fill=(64, 79, 98), width=1)
        for index, peak in enumerate(waveform["bins"]):
            x = margin + (index + .5) * (right - margin) / len(waveform["bins"])
            draw.line((x, centre - peak["max"] * amplitude,
                       x, centre - peak["min"] * amplitude),
                      fill=(71, 214, 192), width=max(1, (right - margin) // WAVEFORM_BINS - 1))
        layout["waveform"] = {"bounds": [margin, top, right, bottom], "scale": "fixed_full_scale_pcm"}
    layout["timecode"] = _draw_copy(
        draw, stamp, (margin, int(height * .91), right, height - margin // 2),
        maximum=24 * typography_scale, color=(174, 192, 209), check_cancel=check_cancel,
    )
    _check(check_cancel)
    buffer = io.BytesIO()
    canvas.save(buffer, format="PNG", optimize=False)
    data = buffer.getvalue()
    # Verify the actual encoded artifact instead of returning canvas assumptions.
    with Image.open(io.BytesIO(data)) as image:
        image.load()
        if image.format != "PNG" or image.size != (width, height) or image.mode != "RGB":
            raise VideoEditingError("AUDIO_VISUAL_IMAGE_VERIFICATION_FAILED")
    return data, {"width": width, "height": height, "mode": "RGB", "layout": layout,
                  "decode_all_pixels": "passed"}


def _verify_video(path: Path, *, ffmpeg: str, ffprobe: str, width: int, height: int,
                  duration: float, check_cancel) -> dict[str, Any]:
    work = path.parent
    probe_file = work / "probe.json"
    rendering._run([
        ffprobe, "-v", "error", "-protocol_whitelist", "file", "-format_whitelist", "mov",
        "-count_frames", "-show_entries",
        "stream=codec_type,codec_name,width,height,pix_fmt,sample_aspect_ratio,r_frame_rate,"
        "start_time,duration,nb_read_frames,sample_rate,channels:format=duration,format_name",
        "-of", "json", "-o", str(probe_file), str(path),
    ], cwd=work, timeout=120, check_continue=check_cancel)
    try:
        probe = json.loads(probe_file.read_text(encoding="utf-8"))
        videos = [stream for stream in probe["streams"] if stream["codec_type"] == "video"]
        audios = [stream for stream in probe["streams"] if stream["codec_type"] == "audio"]
        if len(probe["streams"]) != 2 or len(videos) != 1 or len(audios) != 1:
            raise ValueError("unexpected streams")
        video, audio = videos[0], audios[0]
        expected_frames = math.ceil(round(duration * FRAME_RATE, 7))
        actual_duration = float(probe["format"]["duration"])
        if ("mp4" not in probe["format"]["format_name"].split(",")
                or video["codec_name"] != "h264" or audio["codec_name"] != "aac"
                or (int(video["width"]), int(video["height"])) != (width, height)
                or video["pix_fmt"] != "yuv420p" or video["sample_aspect_ratio"] != "1:1"
                or video["r_frame_rate"] != f"{FRAME_RATE}/1"
                or int(video["nb_read_frames"]) != expected_frames
                or not math.isfinite(float(video["duration"]))
                or abs(float(video["duration"]) - expected_frames / FRAME_RATE) > 2 / SAMPLE_RATE
                or int(audio["sample_rate"]) != SAMPLE_RATE or int(audio["channels"]) != 2
                or int(audio["nb_read_frames"]) < 1
                or not math.isfinite(actual_duration) or abs(actual_duration - duration) > 2 / FRAME_RATE
                or any(not math.isfinite(float(stream["start_time"]))
                       or abs(float(stream["start_time"])) > 1 / FRAME_RATE for stream in (video, audio))
                or not math.isfinite(float(audio["duration"]))
                or abs(float(audio["duration"]) - duration) > 2 / SAMPLE_RATE):
            raise ValueError("output metadata mismatch")
    except (OSError, ValueError, KeyError, TypeError, IndexError, OverflowError) as exc:
        raise VideoEditingError("AUDIO_VISUAL_VIDEO_VERIFICATION_FAILED") from exc
    progress = work / "decode.progress"
    rendering._run([
        ffmpeg, "-v", "error", "-xerror", "-nostdin", "-threads", "1",
        "-err_detect", "explode", "-protocol_whitelist", "file", "-format_whitelist", "mov",
        "-i", str(path), "-map", "0:v:0", "-map", "0:a:0", "-progress", str(progress),
        "-nostats", "-f", "null", "-",
    ], cwd=work, timeout=120, check_continue=check_cancel)
    values = dict(line.split("=", 1) for line in progress.read_text().splitlines() if "=" in line)
    if values.get("progress") != "end" or values.get("frame") != str(expected_frames):
        raise VideoEditingError("AUDIO_VISUAL_FULL_DECODE_INCOMPLETE")
    return {"width": int(video["width"]), "height": int(video["height"]),
            "duration_sec": actual_duration, "audio_duration_sec": float(audio["duration"]),
            "video_duration_sec": float(video["duration"]),
            "stream_start_seconds": {"video": float(video["start_time"]), "audio": float(audio["start_time"])},
            "video_frames": int(video["nb_read_frames"]), "frame_rate": video["r_frame_rate"],
            "sample_aspect_ratio": video["sample_aspect_ratio"], "pixel_format": video["pix_fmt"],
            "video_codec": video["codec_name"], "audio_codec": audio["codec_name"],
            "audio_sample_rate": int(audio["sample_rate"]), "audio_channels": int(audio["channels"]),
            "audio_decoded_frames": int(audio["nb_read_frames"]),
            "decode_all_frames": "passed", "strict_decode": True}


def build_audio_visual_files(asset: MediaAsset, *, start_sec: float, end_sec: float,
                             quote: str, title: str, format_ratio: str = "9:16",
                             check_cancel: Callable[[], None] | None = None) -> dict[str, Any]:
    """Build a static waveform image, quote card and playable H.264/AAC excerpt.

    Only a single-audio-stream asset is accepted by this bounded interface. The
    supplied quote is user copy, never a claim that those words were spoken.
    Returns only fixed, safe basenames and bytes; temporary inputs are removed
    on success, failure and cancellation. No publication or external calls.
    """
    from shared_artifacts import _ffmpeg_path, validate_local_asset
    _check(check_cancel)
    if not isinstance(asset, MediaAsset):
        raise VideoEditingError("AUDIO_VISUAL_FINGERPRINTED_ASSET_REQUIRED")
    asset.validate()
    if asset.audio_stream_count == 0:
        raise VideoEditingError("AUDIO_VISUAL_AUDIO_REQUIRED")
    if asset.audio_stream_count != 1:
        raise VideoEditingError("AUDIO_VISUAL_SINGLE_AUDIO_STREAM_REQUIRED")
    if (any(type(value) not in (int, float) for value in (start_sec, end_sec))
            or not 0 <= start_sec < end_sec <= asset.duration_sec
            or not all(math.isfinite(value) for value in (start_sec, end_sec))
            or not MIN_EXCERPT_SECONDS - 1e-9 <= end_sec - start_sec <= MAX_EXCERPT_SECONDS + 1e-9):
        raise VideoEditingError("AUDIO_VISUAL_RANGE_INVALID")
    if not isinstance(format_ratio, str) or format_ratio not in rendering.FRAME_SIZES:
        raise VideoEditingError("AUDIO_VISUAL_FORMAT_RATIO_INVALID")
    title, quote = _copy_text(title, 80), _copy_text(quote, 400)
    validate_local_asset(asset.path)
    _source_unchanged(asset, check_cancel)
    ffmpeg, ffprobe = _ffmpeg_path(), _ffprobe_path()
    if not ffmpeg or not ffprobe or not rendering.CAPTION_FONT.is_file():
        raise VideoEditingError("AUDIO_VISUAL_TOOLCHAIN_UNAVAILABLE")
    stamp = f"{_timecode(start_sec)} - {_timecode(end_sec)}"
    rendering._verify_caption_glyphs(title + quote + stamp)
    font_sha256 = hashlib.sha256(rendering.CAPTION_FONT.read_bytes()).hexdigest()
    width, height = rendering.FRAME_SIZES[format_ratio]
    duration = end_sec - start_sec
    files: dict[str, bytes] = {}
    outputs: dict[str, Any] = {}
    with tempfile.TemporaryDirectory(prefix="lumenis-audio-visual-") as temporary:
        work = Path(temporary)
        audio, waveform = _decode_excerpt(asset, ffmpeg=ffmpeg, work=work, start=start_sec,
                                          duration=duration, check_cancel=check_cancel)
        for name, included_waveform in (("audio-visual.png", waveform), ("quote-card.png", None)):
            files[name], outputs[name] = _card(title=title, quote=quote, stamp=stamp,
                                              width=width, height=height, waveform=included_waveform,
                                              check_cancel=check_cancel)
            outputs[name]["mime_type"] = "image/png"
        (work / "audio-visual.png").write_bytes(files["audio-visual.png"])
        video = work / "audio-visual.mp4"
        rendering._run([
            ffmpeg, "-v", "error", "-xerror", "-nostdin", "-n", "-threads", "1",
            "-filter_threads", "1", "-loop", "1", "-framerate", str(FRAME_RATE),
            "-protocol_whitelist", "file", "-f", "image2", "-i", "audio-visual.png",
            "-protocol_whitelist", "file", "-format_whitelist", "wav", "-i", str(audio),
            "-map", "0:v:0", "-map", "1:a:0", "-map_metadata", "-1", "-map_chapters", "-1",
            "-frames:v", str(math.ceil(round(duration * FRAME_RATE, 7))),
            "-vf", "setsar=1", "-c:v", "libx264", "-preset", "veryfast",
            "-tune", "stillimage", "-crf", "20", "-pix_fmt", "yuv420p", "-threads:v", "1",
            "-c:a", "aac", "-b:a", "192k", "-ar", str(SAMPLE_RATE), "-ac", "2", "-threads:a", "1",
            "-fflags", "+bitexact", "-flags:v", "+bitexact", "-flags:a", "+bitexact",
            "-movflags", "+faststart", "-movie_timescale", str(SAMPLE_RATE),
            "-video_track_timescale", str(SAMPLE_RATE), "-fs", str(MAX_OUTPUT_BYTES), str(video),
        ], cwd=work, timeout=180, check_continue=check_cancel)
        if not 0 < video.stat().st_size < MAX_OUTPUT_BYTES:
            raise VideoEditingError("AUDIO_VISUAL_OUTPUT_LIMIT_EXCEEDED")
        outputs[video.name] = _verify_video(video, ffmpeg=ffmpeg, ffprobe=ffprobe, width=width,
                                          height=height, duration=duration, check_cancel=check_cancel)
        outputs[video.name]["mime_type"] = "video/mp4"
        files[video.name] = video.read_bytes()
    _source_unchanged(asset, check_cancel)
    for name, data in files.items():
        outputs[name].update({"filename": name, "sha256": hashlib.sha256(data).hexdigest(),
                              "size_bytes": len(data), "source_sha256": asset.sha256,
                              "source_start_sec": start_sec, "source_end_sec": end_sec})
    report = {
        "schema": "lumenis.audio-visual-excerpt/v1", "state": "rendered_verified",
        "source_asset_id": asset.asset_id, "source_sha256": asset.sha256,
        "source_duration_sec": asset.duration_sec, "source_rights_status": asset.rights_status,
        "source_start_sec": start_sec, "source_end_sec": end_sec,
        "requested_duration_sec": duration, "audio_stream_index": 0,
        "audio_stream_selection": "only_available_audio_stream",
        "format_ratio": format_ratio, "width": width, "height": height,
        "copy": {"title": title, "quote": quote},
        "copy_sha256": hashlib.sha256(json.dumps({"title": title, "quote": quote},
                                                  ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
        "copy_check": "USER_SUPPLIED_LITERAL_NOT_VERIFIED_AGAINST_AUDIO",
        "font": {"filename": rendering.CAPTION_FONT.name, "sha256": font_sha256,
                 "face_index": 0, "all_requested_glyphs_verified": True},
        "waveform": waveform, "outputs": outputs, "visual_mode": "static_waveform_and_literal_copy",
        "ai_generated": False, "asr_performed": False, "semantic_story_selection": "not_performed",
        "platform_specific_repackaging": "not_performed", "published": False,
        "limitations": ["User must verify quote wording and excerpt context.",
                        "No automatic 2-3 minute narrative or story-node selection.",
                        "Caller must authorize source use and register the existing Artifact."],
    }
    files["audio-visual.json"] = json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False).encode()
    _check(check_cancel)
    return {"files": files, "report": report}
