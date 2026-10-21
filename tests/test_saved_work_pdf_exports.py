"""Isolated PDF regressions; generated text is not real-user acceptance."""
import io
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pdfplumber
import pytest
from flask import Flask

from lumenis.shared.reporting.word_pdf_export import ReportExporter
from tests.test_saved_document_workflow import workspace, _save, _source
from tests.test_shared_work_owner_routes import route_workspace, WORKS, GLOBAL


TEXT = "中文：保留原始来源与旧版本。\nEnglish: Keep the prior version.\n한국어: 출처 수정 후 이전 버전을 보존합니다."


def pdf_text(raw):
    assert raw.startswith(b"%PDF-")
    assert b"/FontFile2" in raw and b"STSong-Light" not in raw
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def compact(text):
    return re.sub(r"\s+", "", text)


def snapshot(marker="VERSION-1"):
    return {"title": "材料 / Research / 자료", "markdown": TEXT + "\n<b>literal markup</b> & " + marker,
            "kind": "research", "payload": {"parent": "SOURCE-VERSION-0"}, "citations": []}


def test_saved_work_pdf_preserves_snapshot_and_source_lineage(workspace):
    ws = workspace
    _, _, source = _source(ws)
    citation = ws["documents"].locate_quote(source["evidence_id"], actor=ws["actor"],
                                          scope=ws["scope"], start=0, end=20)["citation"]
    saved = _save(ws, markdown=snapshot()["markdown"], payload=snapshot()["payload"], citations=[citation])
    exported = ws["work"].export_work(saved["artifact"]["artifact_id"], actor=ws["actor"],
                                      scope=ws["scope"], format="pdf")
    text = pdf_text(exported["content"])
    assert compact(TEXT) in compact(text)
    assert "<b>literal markup</b> & VERSION-1" in text
    assert "SOURCE-VERSION-0" in text and saved["document"]["work_id"] in text
    assert exported["filename"] == saved["artifact"]["artifact_id"] + ".pdf"
    assert citation["quote"] in text and "public-reference-only" in text


def test_pdf_reports_missing_font_assets(monkeypatch, tmp_path):
    import lumenis.shared.database as database
    monkeypatch.setattr(database, "resolve_path", lambda path: str(tmp_path / "missing.ttf"))
    with pytest.raises(RuntimeError, match="^WORK_EXPORT_FONT_UNAVAILABLE$"):
        ReportExporter.export_saved_document(snapshot(), sources=[], format="pdf")


def test_pdf_rejects_unrenderable_characters():
    document = snapshot()
    document["markdown"] += "\U00010348"
    with pytest.raises(RuntimeError, match="^WORK_EXPORT_GLYPH_UNAVAILABLE$"):
        ReportExporter.export_saved_document(document, sources=[], format="pdf")


def test_concurrent_pdf_exports_keep_each_snapshot_separate():
    markers = ["OWNED-VERSION-" + str(index) for index in range(3)]
    with ThreadPoolExecutor(max_workers=3) as pool:
        outputs = list(pool.map(lambda marker: ReportExporter.export_saved_document(
            snapshot(marker), sources=[], format="pdf"), markers))
    for marker, raw in zip(markers, outputs):
        text = pdf_text(raw)
        assert marker in text and compact(TEXT) in compact(text)
        assert all(other not in text for other in markers if other != marker)


def disable_word_conversion(monkeypatch):
    def unavailable(*args):
        raise OSError("Word converter unavailable in isolated test")
    monkeypatch.setitem(sys.modules, "docx2pdf", SimpleNamespace(convert=unavailable))


def test_reportlab_fallback_preserves_languages_and_unknown_values(monkeypatch, tmp_path):
    disable_word_conversion(monkeypatch)
    exporter = ReportExporter(str(tmp_path))
    report = {"brand_name": "材料 자료", "period_label": "Research", "executive_summary": TEXT,
              "key_metrics": [{"name": "missing-value", "value": None}, {"name": "genuine-zero", "value": 0}],
              "risk_events": [{"title": "unknown-score"}, {"title": "zero-score", "risk_score": 0}]}
    path = exporter.export_pdf(report, "fallback.pdf")
    assert path.endswith(".pdf")
    from pathlib import Path
    text = pdf_text(Path(path).read_bytes())
    assert compact(TEXT) in compact(text)
    assert "missing-value—" in compact(text) and "genuine-zero0" in compact(text)
    assert "unknown-score-风险指数—" in compact(text) and "zero-score-风险指数0" in compact(text)


def test_pdf_failure_never_returns_the_word_file(monkeypatch, tmp_path):
    import lumenis.shared.database as database
    disable_word_conversion(monkeypatch)
    monkeypatch.setattr(database, "resolve_path", lambda path: str(tmp_path / "missing.ttf"))
    exporter = ReportExporter(str(tmp_path))
    assert exporter.export_pdf({"brand_name": "Owned"}, "unavailable.pdf") == ""
    assert (tmp_path / "unavailable.docx").is_file()


def test_pdf_does_not_accept_unverified_word_converter_output(monkeypatch, tmp_path):
    from pathlib import Path
    called = []
    def unverified(word, pdf):
        called.append(word)
        Path(pdf).write_bytes(b"%PDF-unverified-fonts")
    monkeypatch.setitem(sys.modules, "docx2pdf", SimpleNamespace(convert=unverified))
    result = ReportExporter(str(tmp_path)).export_pdf({"brand_name": "Owned", "executive_summary": TEXT}, "verified.pdf")
    assert compact(TEXT) in compact(pdf_text(Path(result).read_bytes()))
    assert called == []


@pytest.mark.parametrize("code", ["WORK_EXPORT_FONT_UNAVAILABLE", "WORK_EXPORT_GLYPH_UNAVAILABLE"])
def test_saved_work_export_api_preserves_font_failure_code(route_workspace, monkeypatch, code):
    ws = route_workspace
    saved = _save(ws, markdown=TEXT + ("\U00010348" if code.endswith("GLYPH_UNAVAILABLE") else ""))
    if code.endswith("FONT_UNAVAILABLE"):
        import lumenis.shared.database as database
        monkeypatch.setattr(database, "resolve_path", lambda path: str(ws["root"] / "missing-font.ttf"))
    response = ws["api"].get(WORKS + "/" + saved["artifact"]["artifact_id"] + "/export",
                             headers=GLOBAL, query_string={"format": "pdf"})
    assert response.status_code == 503
    assert response.json["ok"] is False and response.json["code"] == code
    assert response.json["publish_status"] == "not_performed"


def test_reportlab_fallback_preserves_complete_report(monkeypatch, tmp_path):
    disable_word_conversion(monkeypatch)
    report = {
        "report_id": "REPORT-SOURCE-VERSION-7", "brand_name": "Owned", "executive_summary": TEXT,
        "risk_events": [{"title": "risk-detail", "alert_message": "alert-detail", "source_count": 0}],
        "sentiment_analysis": {"overall_label": "sentiment-detail", "distribution": {"positive": 0}},
        "propagation_analysis": {"engagement_level": "propagation-detail", "total_engagement": 0,
                                 "top_notes": [{"title": "note-detail", "author": "source-author"}]},
        "key_insights": [{"type": "insight-type", "insight": "insight-detail", "metric": "metric-detail"}],
        "action_items": [{"action": "action-detail", "owner": "action-owner", "deadline": "2026-10-08"}],
        "data_quality": {"warnings": ["quality-detail"], "total_collected": 0},
        "cth_specific": {"reference_scope": True,
                         "tone_evaluation": {"overall": "tone-detail", "warm_ratio": 0},
                         "active_activity_nodes": [{"name": "activity-detail", "month": "October"}]},
    }
    from pathlib import Path
    text = pdf_text(Path(ReportExporter(str(tmp_path)).export_pdf(report, "complete.pdf")).read_bytes())
    for value in ("REPORT-SOURCE-VERSION-7", "alert-detail", "sentiment-detail", "propagation-detail",
                  "note-detail", "source-author", "insight-detail", "metric-detail", "action-detail",
                  "action-owner", "quality-detail", "tone-detail", "activity-detail"):
        assert value in text
    normalized = compact(text)
    for value in ("正面0条", "负面—条", "总互动量：0", "提及数：—", "来源数：0",
                  "质量评分：—/100", "采集完整：—", "温暖0%", "精致—%"):
        assert value in normalized
    from docx import Document
    word = Document(tmp_path / "complete.docx")
    word_text = compact("\n".join(p.text for p in word.paragraphs))
    for value in ("正面0条", "负面—条", "总互动量：0", "提及数：—", "来源数：0",
                  "质量评分：—/100", "采集完整：—", "温暖0%", "精致—%"):
        assert value in word_text


@pytest.mark.parametrize("missing", ["word", "pdf"])
def test_reporting_api_does_not_claim_an_incomplete_export(monkeypatch, missing):
    from lumenis.shared import reporting
    from lumenis.shared.reporting import routes
    from lumenis.shared.reporting.api import reporting_bp
    monkeypatch.setenv("LUMENIS_REPORTING", "1")
    actor = {"id": 17}
    scope = {"domain_id": "cth", "product_id": "CTH", "subject_id": "SUBJ-REPORT",
             "brand_focus_id": None, "global_scope": False}
    monkeypatch.setattr(routes, "_authorized_report_scope",
                        lambda **_kwargs: (actor, scope, [], {"subject_id": "SUBJ-REPORT"}))
    paths = {"word_path": "word.docx", "pdf_path": "report.pdf"}
    paths[missing + "_path"] = ""
    monkeypatch.setattr(reporting, "ReportExporter", lambda: SimpleNamespace(generate_and_export=lambda **kwargs: paths))
    app = Flask(__name__)
    app.register_blueprint(reporting_bp)
    result = app.test_client().post("/api/reporting/export", json={})
    assert result.status_code == 503
    assert result.json["ok"] is False and result.json["code"] == "REPORT_EXPORT_INCOMPLETE"
    assert result.json["unavailable_formats"] == [missing]


def test_report_exporter_checks_data_export_and_revalidates_source_use(monkeypatch, tmp_path):
    from lumenis.shared.reporting import export_api, saved_work

    actor = {"id": 17, "role": "owner"}
    scope = {"domain_id": "cth", "product_id": "CTH", "subject_id": "SUBJ-REPORT",
             "brand_focus_id": None, "global_scope": False}
    selected = {"subject_id": "SUBJ-REPORT", "name_zh": "隔离测试主体"}
    notes = [{"source_id": "SRC-1"}]
    events = []
    built = []
    effects = []
    source_checks = []

    def build_report(kind, start, end, **kwargs):
        built.append((kind, start, end, kwargs))
        return kind, "2026-10-01", "2026-10-02", {"report_id": "R-1"}, 1, notes, events

    monkeypatch.setattr(export_api, "_build_report", build_report)
    monkeypatch.setattr(export_api, "_verify_report_source_use",
                        lambda *args, **kwargs: source_checks.append((args, kwargs)))
    monkeypatch.setattr(saved_work, "_check_effect",
                        lambda *args, **kwargs: effects.append((args, kwargs)))

    exporter = ReportExporter(str(tmp_path))

    def write_word(report, filename):
        path = tmp_path / filename
        path.write_bytes(b"scoped report")
        return str(path)

    monkeypatch.setattr(exporter, "export_word", write_word)
    result = exporter.generate_and_export(
        brand_mode="cth", period="daily", export_formats=["word"], actor=actor,
        scope=scope, scope_entry=selected,
    )

    assert result["word_path"].endswith(".docx")
    assert built == [("daily", None, None, {
        "actor": actor, "scope": {**scope, "global_scope": False}, "scope_entry": selected,
    })]
    assert len(effects) == 2
    assert all(call[0][0] is actor for call in effects)
    assert len(source_checks) == 1
    assert source_checks[0][0] == (actor, {**scope, "global_scope": False}, notes, events)


def test_report_exporter_removes_partial_pdf_and_preserves_word_output(monkeypatch, tmp_path):
    from pathlib import Path

    from lumenis.shared.reporting import export_api, saved_work

    actor = {"id": 17, "role": "owner"}
    scope = {"domain_id": "cth", "product_id": "CTH", "subject_id": "SUBJ-REPORT",
             "brand_focus_id": None, "global_scope": False}
    monkeypatch.setattr(export_api, "_build_report",
                        lambda *args, **kwargs: ("daily", "2026-10-01", "2026-10-01",
                                                 {"report_id": "R-2"}, 0, [], []))
    monkeypatch.setattr(export_api, "_verify_report_source_use", lambda *args, **kwargs: None)
    monkeypatch.setattr(saved_work, "_check_effect", lambda *args, **kwargs: None)
    exporter = ReportExporter(str(tmp_path))

    def write_word(_report, filename):
        path = tmp_path / filename
        path.write_bytes(b"complete word report")
        return str(path)

    def partial_pdf(_report, filename):
        (tmp_path / filename).write_bytes(b"%PDF-partial")
        (tmp_path / Path(filename).with_suffix(".docx")).write_bytes(b"partial intermediate word")
        return ""

    monkeypatch.setattr(exporter, "export_word", write_word)
    monkeypatch.setattr(exporter, "export_pdf", partial_pdf)

    result = exporter.generate_and_export(
        export_formats=["word", "pdf"], actor=actor, scope=scope,
    )

    assert Path(result["word_path"]).read_bytes() == b"complete word report"
    assert result["pdf_path"] == ""
    assert {path.name for path in tmp_path.iterdir()} == {Path(result["word_path"]).name}


def test_report_exporter_removes_word_file_after_partial_write_error(monkeypatch, tmp_path):
    from lumenis.shared.reporting import export_api, saved_work

    actor = {"id": 17, "role": "owner"}
    scope = {"domain_id": "cth", "product_id": "CTH", "subject_id": "SUBJ-REPORT",
             "brand_focus_id": None, "global_scope": False}
    monkeypatch.setattr(export_api, "_build_report",
                        lambda *args, **kwargs: ("daily", "2026-10-01", "2026-10-01",
                                                 {"report_id": "R-3"}, 0, [], []))
    monkeypatch.setattr(export_api, "_verify_report_source_use", lambda *args, **kwargs: None)
    monkeypatch.setattr(saved_work, "_check_effect", lambda *args, **kwargs: None)
    exporter = ReportExporter(str(tmp_path))

    def partial_word(_report, filename):
        (tmp_path / filename).write_bytes(b"partial word")
        raise OSError("simulated interrupted file write")

    monkeypatch.setattr(exporter, "export_word", partial_word)

    with pytest.raises(OSError, match="simulated interrupted file write"):
        exporter.generate_and_export(export_formats=["word"], actor=actor, scope=scope)

    assert list(tmp_path.iterdir()) == []


def test_special_report_pdf_preserves_multilingual_literal_text():
    from lumenis.shared.reporting.routes import report_bp
    app = Flask(__name__)
    app.register_blueprint(report_bp)
    result = app.test_client().post("/api/report/special/export", json={
        "format": "pdf", "report": {"title": "材料 / Research / 자료", "event_overview": TEXT + "\n<b>literal</b> &"}})
    assert result.status_code == 200
    text = pdf_text(result.data)
    assert compact(TEXT) in compact(text) and "<b>literal</b> &" in text


@pytest.mark.parametrize("reference_scope", [None, False])
def test_nonreference_report_does_not_export_reference_only_material(tmp_path, reference_scope):
    from pathlib import Path
    from docx import Document
    cth = {"tone_evaluation": {"overall": "REFERENCE-TONE-ONLY"},
           "active_activity_nodes": [{"name": "REFERENCE-ACTIVITY-ONLY"}]}
    if reference_scope is not None:
        cth["reference_scope"] = reference_scope
    report = {"brand_name": "Isolated ordinary subject", "executive_summary": TEXT,
              "cth_specific": cth}
    path = ReportExporter(str(tmp_path)).export_pdf(report, "ordinary.pdf")
    assert path
    text = pdf_text(Path(path).read_bytes())
    word = "\n".join(p.text for p in Document(tmp_path / "ordinary.docx").paragraphs)
    for rendered in (text, word):
        assert compact(TEXT) in compact(rendered)
        assert "REFERENCE-TONE-ONLY" not in rendered
        assert "REFERENCE-ACTIVITY-ONLY" not in rendered
        assert "参考主体专项分析" not in rendered


def test_report_pdf_preserves_long_multilingual_summary_across_pages(tmp_path):
    from pathlib import Path
    lines = [f"段落{index:03d} 中文边界 English 한국어 END{index:03d}" for index in range(180)]
    report = {"brand_name": "Isolated pagination", "executive_summary": "\n".join(lines),
              "cth_specific": {"reference_scope": True,
                               "tone_evaluation": {"overall": "FINAL-TONE", "warm_ratio": 0},
                               "active_activity_nodes": [{"name": "FINAL-ACTIVITY"}]}}
    path = ReportExporter(str(tmp_path)).export_pdf(report, "multipage.pdf")
    assert path
    raw = Path(path).read_bytes()
    text = pdf_text(raw)
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        assert len(pdf.pages) >= 3
    normalized = compact(text)
    for line in lines:
        assert normalized.count(compact(line)) == 1
    assert "FINAL-TONE" in text and "FINAL-ACTIVITY" in text
    assert normalized.index(compact(lines[0])) < normalized.index(compact(lines[-1]))
    assert normalized.index(compact(lines[-1])) < normalized.index("FINAL-TONE")
