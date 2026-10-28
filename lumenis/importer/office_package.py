"""Object edits inside an original Office package; untouched parts stay byte exact."""
from __future__ import annotations

import copy
import hashlib
import io
import posixpath
import re
import zipfile
from bisect import bisect_left
from decimal import Decimal, InvalidOperation
from urllib.parse import unquote, urlsplit

from lxml import etree

from lumenis.shared.reporting import saved_work

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
_P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
_S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_R = '{http://schemas.openxmlformats.org/package/2006/relationships}'
_DOCUMENT_REL_ID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
_XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'
_OBJECT_OPERATIONS = {'text_run': 'replace_text', 'cell_text_run': 'replace_text',
    'cell_number': 'set_number', 'cell_boolean': 'set_boolean', 'cell_formula': 'replace_formula',
    'word_run_format': 'set_run_format', 'drawing_run_format': 'set_run_format',
    'word_paragraph_format': 'set_paragraph_format', 'shape_geometry': 'set_geometry',
    'cell_style': 'set_cell_style'}
_MAX_OBJECTS = 100000
_SIGNATURE_RELATIONSHIP_TYPES = {
    'http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/' + kind
    for kind in ('origin', 'signature', 'certificate')
}
_SIGNATURE_CONTENT_TYPES = {
    'application/vnd.openxmlformats-package.digital-signature-origin',
    'application/vnd.openxmlformats-package.digital-signature-xmlsignature+xml',
    'application/vnd.openxmlformats-package.digital-signature-certificate',
}


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _xml(raw):
    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
        raise ValueError('OFFICE_XML_ENTITY_FORBIDDEN')
    try:
        return etree.parse(io.BytesIO(raw), etree.XMLParser(
            resolve_entities=False, load_dtd=False, no_network=True,
            remove_blank_text=False, strip_cdata=False, huge_tree=False))
    except etree.XMLSyntaxError as exc:
        raise ValueError('OFFICE_XML_INVALID') from exc


def _serialized(tree):
    return etree.tostring(tree, encoding='UTF-8', xml_declaration=True,
                          standalone=tree.docinfo.standalone)


def _cell_position(reference):
    match = re.fullmatch(r'\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6})', reference)
    if match is None:
        raise ValueError('OFFICE_FORMULA_RANGE_INVALID')
    column = 0
    for letter in match[1]:
        column = column * 26 + ord(letter) - ord('A') + 1
    return int(match[2]), column


def _formula_result_cells(root):
    cells = list(root.iter(_S + 'c'))
    results = {cell for cell in cells if cell.find(_S + 'f') is not None}
    events = []
    for cell in results:
        reference = cell.find(_S + 'f').get('ref')
        if reference is None:
            continue
        ends = reference.split(':')
        if not 1 <= len(ends) <= 2:
            raise ValueError('OFFICE_FORMULA_RANGE_INVALID')
        first, last = _cell_position(ends[0]), _cell_position(ends[-1])
        if first[0] > last[0] or first[1] > last[1]:
            raise ValueError('OFFICE_FORMULA_RANGE_INVALID')
        events.extend(((first[0], first[1], last[1] + 1, 1),
                       (last[0] + 1, first[1], last[1] + 1, -1)))
    if not events:
        return results
    # Sweep existing cells, including array/data-table outputs, without expanding a grid.
    positions = sorted(((*_cell_position(cell.get('r', '')), cell) for cell in cells), key=lambda item: item[:2])
    columns = sorted({column for _row, column, _cell in positions})
    counts, event_index = [0] * (len(columns) + 1), 0
    events.sort()
    for row, column, cell in positions:
        while event_index < len(events) and events[event_index][0] <= row:
            _row, start, end, delta = events[event_index]
            for boundary, change in ((start, delta), (end, -delta)):
                index = bisect_left(columns, boundary) + 1
                while index < len(counts):
                    counts[index] += change
                    index += index & -index
            event_index += 1
        index, covered = bisect_left(columns, column) + 1, 0
        while index:
            covered += counts[index]
            index -= index & -index
        if covered:
            results.add(cell)
    return results


class OfficePackage:
    def __init__(self, raw, *, filename, media_type):
        self.spec = saved_work._attachment_spec(filename, media_type)
        if (self.spec is None or self.spec['extension'] not in {'.docx', '.pptx', '.xlsx'}
                or not saved_work._valid_ooxml_attachment(raw, self.spec)):
            raise ValueError('OFFICE_PACKAGE_INVALID')
        self.raw = raw
        self.sha256 = _sha(raw)
        self.filename = filename
        self.media_type = media_type
        with zipfile.ZipFile(io.BytesIO(raw)) as source:
            self.infos = source.infolist()
            self.comment = source.comment
            self.parts = {info.filename: source.read(info) for info in self.infos}
        normalized = [unquote(name).casefold() for name in self.parts]
        if len(set(normalized)) != len(normalized):
            raise ValueError('OFFICE_PART_IDENTITY_AMBIGUOUS')
        self.trees = {}
        self.external = []
        self.relationships = {}
        self._verify_relationships()
        self.part_labels = {}
        self.paragraphs = {}
        self._read_part_labels()
        self.signatures = self._signature_parts()
        self.active = [name for name in self.parts if 'vbaproject' in name.casefold()
                       or '/activex/' in name.casefold()]
        self.objects = []
        self.worksheet_parts = []
        self.unsupported = []
        self.bindings = {}
        self._project()

    def _tree(self, part):
        if part not in self.trees:
            self.trees[part] = _xml(self.parts[part])
        return self.trees[part]

    def _signature_parts(self):
        found = {name for name in self.parts if name.startswith('_xmlsignatures/')}
        for relations in self.relationships.values():
            found.update(relation['target'] for relation in relations.values()
                         if relation['type'] in _SIGNATURE_RELATIONSHIP_TYPES)
        defaults = {}
        for node in self._tree('[Content_Types].xml').getroot():
            if node.get('ContentType') not in _SIGNATURE_CONTENT_TYPES:
                continue
            if etree.QName(node).localname == 'Override':
                part = unquote(node.get('PartName', '')).lstrip('/')
                if part in self.parts:
                    found.add(part)
            elif etree.QName(node).localname == 'Default':
                defaults[node.get('Extension')] = True
        found.update(name for name in self.parts if name.rsplit('.', 1)[-1] in defaults)
        return sorted(found)

    def _verify_relationships(self):
        for name in self.parts:
            if not name.endswith('.rels'):
                continue
            if name == '_rels/.rels':
                source, directory = '', ''
            else:
                parent, basename = posixpath.split(name)
                if not parent.endswith('/_rels') or not basename.endswith('.rels'):
                    raise ValueError('OFFICE_RELATIONSHIP_SOURCE_INVALID')
                directory = parent[:-len('/_rels')]
                source = posixpath.join(directory, basename[:-5])
                if source not in self.parts:
                    raise ValueError('OFFICE_RELATIONSHIP_SOURCE_MISSING')
            root = self._tree(name).getroot()
            if root.tag != _R + 'Relationships':
                raise ValueError('OFFICE_RELATIONSHIP_INVALID')
            identities = set()
            for relation in root:
                identity, target = relation.get('Id'), relation.get('Target')
                if (relation.tag != _R + 'Relationship' or not identity or identity in identities
                        or not target or not relation.get('Type')
                        or relation.get('TargetMode') not in {None, 'Internal', 'External'}):
                    raise ValueError('OFFICE_RELATIONSHIP_INVALID')
                identities.add(identity)
                if relation.get('TargetMode') == 'External':
                    self.external.append({'part': name, 'id': identity, 'state': 'PRESERVED_NOT_OPENED'})
                    continue
                uri = urlsplit(target)
                path = unquote(uri.path)
                if uri.scheme or uri.netloc or uri.query or '\\' in path or '\x00' in path:
                    raise ValueError('OFFICE_RELATIONSHIP_TARGET_INVALID')
                resolved = posixpath.normpath(path.lstrip('/') if path.startswith('/') else
                                              posixpath.join(directory, path)) if path else source
                if resolved.startswith('../') or resolved not in self.parts:
                    raise ValueError('OFFICE_RELATIONSHIP_TARGET_MISSING')
                self.relationships.setdefault(source, {})[identity] = {
                    'target': resolved, 'type': relation.get('Type')}

    def _read_part_labels(self):
        if self.spec['extension'] == '.xlsx':
            workbook = 'xl/workbook.xml'
            for index, sheet in enumerate(self._tree(workbook).iter(_S + 'sheet'), 1):
                relation = self.relationships.get(workbook, {}).get(sheet.get(_DOCUMENT_REL_ID))
                if relation and relation['type'].endswith('/worksheet') and sheet.get('name'):
                    part = relation['target']
                    if part in self.part_labels:
                        raise ValueError('OFFICE_WORKSHEET_IDENTITY_AMBIGUOUS')
                    self.part_labels[part] = {'sheet_name': sheet.get('name'), 'sheet_index': index}
        elif self.spec['extension'] == '.pptx':
            presentation = 'ppt/presentation.xml'
            for number, slide in enumerate(self._tree(presentation).iter(_P + 'sldId'), 1):
                relation = self.relationships.get(presentation, {}).get(slide.get(_DOCUMENT_REL_ID))
                if relation and relation['type'].endswith('/slide'):
                    part = relation['target']
                    if part in self.part_labels:
                        raise ValueError('OFFICE_SLIDE_IDENTITY_AMBIGUOUS')
                    self.part_labels[part] = {'slide_number': number, 'notes': False}
                    for child in self.relationships.get(part, {}).values():
                        if child['type'].endswith('/notesSlide'):
                            if child['target'] in self.part_labels:
                                raise ValueError('OFFICE_NOTES_IDENTITY_AMBIGUOUS')
                            self.part_labels[child['target']] = {'slide_number': number, 'notes': True}

    def _add(self, part, kind, value, *, node=None, cell=None, text_index=None, shared=None, editable=True,
             readonly_reason='GROUPED_FORMULA_RESULT'):
        path = self._tree(part).getpath(cell if cell is not None else node)
        locator = {'part': part, 'node_path': path}
        locator.update(self.part_labels.get(part, {}))
        if kind == 'text_run':
            paragraph_tag = _W + 'p' if self.spec['extension'] == '.docx' else _A + 'p'
            if part not in self.paragraphs:
                self.paragraphs[part] = {paragraph: index for index, paragraph in
                    enumerate(self._tree(part).iter(paragraph_tag), 1)}
            paragraph = next((parent for parent in node.iterancestors() if parent.tag == paragraph_tag), None)
            if paragraph is not None:
                locator['paragraph_index'] = self.paragraphs[part][paragraph]
            if self.spec['extension'] == '.pptx':
                shape = next((parent for parent in node.iterancestors()
                    if parent.tag in {_P + 'sp', _P + 'graphicFrame'}), None)
                properties = next(shape.iter(_P + 'cNvPr'), None) if shape is not None else None
                if properties is not None and properties.get('name'):
                    locator['shape_name'] = properties.get('name')
        if cell is not None:
            locator['cell'] = cell.get('r')
        if text_index is not None:
            locator['text_index'] = text_index
        identity = 'office-' + _sha((part + '\0' + path + '\0' + kind + '\0' + str(text_index)).encode())[:32]
        if identity in self.bindings or len(self.bindings) >= _MAX_OBJECTS:
            raise ValueError('OFFICE_OBJECT_LIMIT_OR_IDENTITY_INVALID')
        item = {'object_id': identity, 'kind': kind, 'value': value, 'locator': locator,
                'part_sha256': _sha(self.parts[part]), 'editable': editable,
                'supported_operations': [_OBJECT_OPERATIONS[kind]]}
        if editable:
            self.objects.append(item)
        else:
            self.unsupported.append({**item, 'part': part, 'cell': cell.get('r') if cell is not None else None,
                'state': 'PRESERVED_NOT_EDITABLE', 'reason': readonly_reason})
        self.bindings[identity] = {'public': item, 'node': node, 'cell': cell,
                                   'shared': shared, 'text_index': text_index, 'part': part}

    @staticmethod
    def _word_property(parent, tag, attribute='val'):
        item = parent.find(_W + tag) if parent is not None else None
        return item.get(_W + attribute) if item is not None else None

    def _run_format(self, run, word):
        properties = run.find((_W if word else _A) + 'rPr')
        if word:
            value = {key: self._word_property(properties, tag) for key, tag in
                     (('bold', 'b'), ('italic', 'i'), ('underline', 'u'), ('size_half_points', 'sz'), ('color', 'color'))}
            for key in ('bold', 'italic'):
                node = properties.find(_W + ('b' if key == 'bold' else 'i')) if properties is not None else None
                value[key] = None if node is None else value[key] not in {'0', 'false', 'off'}
            if value['size_half_points'] is not None:
                value['size_half_points'] = int(value['size_half_points'])
            return value
        value = {key: properties.get(attribute) if properties is not None else None
                 for key, attribute in (('bold', 'b'), ('italic', 'i'), ('underline', 'u'), ('size_hundredth_points', 'sz'))}
        for key in ('bold', 'italic'):
            value[key] = None if value[key] is None else value[key] not in {'0', 'false', 'off'}
        if value['size_hundredth_points'] is not None:
            value['size_hundredth_points'] = int(value['size_hundredth_points'])
        rgb = properties.find(_A + 'solidFill/' + _A + 'srgbClr') if properties is not None else None
        value['color'] = rgb.get('val') if rgb is not None else None
        return value

    def _project_word_format(self, part):
        for paragraph in self._tree(part).iter(_W + 'p'):
            properties = paragraph.find(_W + 'pPr')
            value = {'alignment': self._word_property(properties, 'jc')}
            for key, tag in (('keep_with_next', 'keepNext'), ('page_break_before', 'pageBreakBefore')):
                node = properties.find(_W + tag) if properties is not None else None
                value[key] = None if node is None else node.get(_W + 'val') not in {'0', 'false', 'off'}
            for key, attribute in (('space_before_twips', 'before'), ('space_after_twips', 'after')):
                raw = self._word_property(properties, 'spacing', attribute)
                value[key] = int(raw) if raw is not None else None
            self._add(part, 'word_paragraph_format', value, node=paragraph)
        for run in self._tree(part).iter(_W + 'r'):
            if run.find(_W + 't') is not None:
                self._add(part, 'word_run_format', self._run_format(run, True), node=run)

    @staticmethod
    def _geometry(node):
        offset, extent = node.find(_A + 'off'), node.find(_A + 'ext')
        if offset is None or extent is None:
            return None
        return {'x_emu': int(offset.get('x')), 'y_emu': int(offset.get('y')),
                'width_emu': int(extent.get('cx')), 'height_emu': int(extent.get('cy')),
                'rotation_60000deg': int(node.get('rot', '0')),
                'flip_horizontal': node.get('flipH', '0') in {'1', 'true'},
                'flip_vertical': node.get('flipV', '0') in {'1', 'true'}}

    def _project(self):
        ct = self._tree('[Content_Types].xml').getroot()
        types = {node.get('PartName', '').lstrip('/'): node.get('ContentType', '')
                 for node in ct if etree.QName(node).localname == 'Override'}
        shared_parts = [name for name, content_type in types.items() if content_type.endswith('spreadsheetml.sharedStrings+xml')]
        shared = list(self._tree(shared_parts[0]).getroot()) if len(shared_parts) == 1 else []
        for part, content_type in types.items():
            if content_type in _SIGNATURE_CONTENT_TYPES:
                part = unquote(part)
            if part not in self.parts:
                raise ValueError('OFFICE_CONTENT_TYPE_PART_MISSING')
            if self.spec['extension'] == '.docx' and (
                    content_type.endswith('wordprocessingml.document.main+xml') or
                    any(content_type.endswith('wordprocessingml.' + suffix + '+xml') for suffix in
                        ('header', 'footer', 'footnotes', 'endnotes', 'comments'))):
                for node in self._tree(part).iter(_W + 't'):
                    self._add(part, 'text_run', node.text or '', node=node)
                self._project_word_format(part)
            elif self.spec['extension'] == '.pptx' and any(content_type.endswith('presentationml.' + suffix + '+xml')
                    for suffix in ('slide', 'notesSlide')):
                for node in self._tree(part).iter(_A + 't'):
                    self._add(part, 'text_run', node.text or '', node=node)
                for run in self._tree(part).iter(_A + 'r'):
                    self._add(part, 'drawing_run_format', self._run_format(run, False), node=run)
                for geometry in self._tree(part).getroot().iter(_A + 'xfrm', _P + 'xfrm'):
                    value = self._geometry(geometry)
                    if value is not None:
                        concrete = value['width_emu'] > 0 and value['height_emu'] > 0
                        self._add(part, 'shape_geometry', value, node=geometry, editable=concrete,
                                  readonly_reason='INHERITED_OR_ZERO_GEOMETRY')
            elif self.spec['extension'] == '.xlsx' and content_type.endswith('spreadsheetml.worksheet+xml'):
                self.worksheet_parts.append(part)
                root = self._tree(part).getroot()
                formula_results = _formula_result_cells(root)
                for cell in root.iter(_S + 'c'):
                    self._add(part, 'cell_style', {'style_index': int(cell.get('s', '0'))}, node=cell)
                for cell in root.iter(_S + 'c'):
                    if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', cell.get('r', '')):
                        raise ValueError('OFFICE_CELL_REFERENCE_INVALID')
                    formula, value = cell.find(_S + 'f'), cell.find(_S + 'v')
                    if formula is not None:
                        if formula.get('t') in {None, 'normal'} and 'ref' not in formula.attrib and 'si' not in formula.attrib:
                            self._add(part, 'cell_formula', formula.text or '', node=formula, cell=cell)
                        else:
                            self.unsupported.append({'part': part, 'cell': cell.get('r'),
                                'kind': 'grouped_formula', 'editable': False, 'state': 'PRESERVED_NOT_EDITABLE'})
                        continue
                    if cell.get('t') in {'s', 'inlineStr'}:
                        si = None
                        if cell.get('t') == 's':
                            if value is None or not re.fullmatch(r'\d+', value.text or '') or int(value.text) >= len(shared):
                                raise ValueError('OFFICE_SHARED_STRING_INVALID')
                            si = shared[int(value.text)]
                            text_nodes = list(si.iter(_S + 't'))
                        else:
                            text_nodes = list(cell.iter(_S + 't'))
                        for index, node in enumerate(text_nodes):
                            self._add(part, 'cell_text_run', node.text or '', node=node, cell=cell,
                                      text_index=index, shared=si, editable=cell not in formula_results)
                    elif value is not None and cell.get('t') in {None, 'n', 'b'}:
                        self._add(part, 'cell_boolean' if cell.get('t') == 'b' else 'cell_number',
                                  value.text or '', node=value, cell=cell, editable=cell not in formula_results)

    def _remove_calculation_chain(self, modified):
        content_types = self._tree('[Content_Types].xml').getroot()
        overrides = {node.get('PartName', '').lstrip('/'): node.get('ContentType', '')
                     for node in content_types if etree.QName(node).localname == 'Override'}
        defaults = {node.get('Extension'): node.get('ContentType', '')
                    for node in content_types if etree.QName(node).localname == 'Default'}
        chains = {name for name in self.parts if overrides.get(name,
                  defaults.get(name.rsplit('.', 1)[-1], '')).endswith('spreadsheetml.calcChain+xml')}
        rel_part = 'xl/_rels/workbook.xml.rels'
        relations = self._tree(rel_part).getroot()
        chain_relations = [node for node in relations if node.get('Type', '').endswith('/calcChain')]
        targets = {self.relationships.get('xl/workbook.xml', {}).get(node.get('Id'), {}).get('target')
                   for node in chain_relations}
        if len(chains) > 1 or len(chain_relations) > 1 or chains != targets:
            raise ValueError('OFFICE_CALCULATION_CHAIN_INVALID')
        if not chains:
            return set()
        chain = next(iter(chains))
        if chain not in self.parts or self._tree(chain).getroot().tag != _S + 'calcChain':
            raise ValueError('OFFICE_CALCULATION_CHAIN_INVALID')
        for source, mapping in self.relationships.items():
            if any(relation['target'] == chain and (source != 'xl/workbook.xml' or not relation['type'].endswith('/calcChain'))
                   for relation in mapping.values()):
                raise ValueError('OFFICE_CALCULATION_CHAIN_INVALID')
        directory, basename = posixpath.split(chain)
        own_relations = posixpath.join(directory, '_rels', basename + '.rels')
        removed = {chain}
        if own_relations in self.parts:
            if len(self._tree(own_relations).getroot()):
                raise ValueError('OFFICE_CALCULATION_CHAIN_INVALID')
            removed.add(own_relations)
        for node in list(content_types):
            if node.get('PartName', '').lstrip('/') in removed:
                content_types.remove(node)
                modified.add('[Content_Types].xml')
        for node in chain_relations:
            relations.remove(node)
        modified.add(rel_part)
        return removed

    def _invalidate_xlsx_calculation(self, modified):
        required, invalidated = False, 0
        removed = self._remove_calculation_chain(modified)
        for part in self.worksheet_parts:
            cells = _formula_result_cells(self._tree(part).getroot())
            required = required or bool(cells)
            for cell in cells:
                cached = cell.find(_S + 'v')
                if cached is not None:
                    cell.remove(cached)
                    invalidated += 1
                    modified.add(part)
        if required:
            workbook = self._tree('xl/workbook.xml').getroot()
            calc = workbook.find(_S + 'calcPr')
            if calc is None:
                calc = etree.Element(_S + 'calcPr')
                following = {_S + name for name in ('oleSize', 'customWorkbookViews', 'pivotCaches',
                    'smartTagPr', 'smartTagTypes', 'webPublishing', 'fileRecoveryPr', 'webPublishObjects', 'extLst')}
                position = next((index for index, child in enumerate(workbook) if child.tag in following), len(workbook))
                workbook.insert(position, calc)
            for key, value in {'calcMode': 'auto', 'calcOnSave': '1', 'calcCompleted': '0',
                               'fullCalcOnLoad': '1', 'forceFullCalc': '1'}.items():
                calc.set(key, value)
            modified.add('xl/workbook.xml')
        return removed, {'formula_evaluation': 'NATIVE_RECALCULATION_REQUIRED' if required else 'NOT_REQUIRED',
                         'formula_cache_invalidation': 'WORKBOOK_WIDE_NO_DEPENDENCY_GRAPH' if required else 'NOT_APPLICABLE',
                         'formula_cache_cells_invalidated': invalidated,
                         'calculation_chain': 'REMOVED_FOR_NATIVE_REBUILD' if removed else 'NOT_PRESENT'}

    def projection(self):
        formulas = [cell for part in self.worksheet_parts for cell in self._tree(part).iter(_S + 'c')
                    if cell.find(_S + 'f') is not None]
        calc = self._tree('xl/workbook.xml').getroot().find(_S + 'calcPr') if self.worksheet_parts else None
        requires_calculation = bool(formulas) and (any(cell.find(_S + 'v') is None or not cell.find(_S + 'v').text for cell in formulas)
            or calc is not None and (calc.get('fullCalcOnLoad') in {'1', 'true'} or calc.get('forceFullCalc') in {'1', 'true'}))
        return {'schema': 'lumenis.office-objects/v1', 'package_sha256': self.sha256,
                'format': self.spec['extension'][1:], 'objects': self.objects,
                'external_relationships': self.external, 'signature_parts': self.signatures,
                'active_content_parts': self.active,
                'style_options': self._cell_style_options() if self.worksheet_parts else [],
                'property_editing': {'mode': 'direct_properties', 'inherited_values': None,
                    'resolved_theme_appearance': 'NOT_VERIFIED',
                    'unrecognized_unchanged_values': 'PRESERVED',
                    'new_unknown_properties': 'REJECTED'},
                'unsupported_objects': self.unsupported,
                'preserved_uneditable_parts': sorted(name for name in self.parts
                    if name not in {item['locator']['part'] for item in self.objects}),
                'formula_evaluation': 'NATIVE_RECALCULATION_REQUIRED' if requires_calculation else 'NOT_RUN',
                'native_application_acceptance': 'NOT_VERIFIED',
                'part_count': len(self.parts), 'object_count': len(self.objects)}

    def _cell_style_options(self):
        part = next((item['target'] for item in self.relationships.get('xl/workbook.xml', {}).values()
                     if item['type'].endswith('/styles')), None)
        if part is None:
            return []
        root = self._tree(part).getroot()
        formats, fonts, fills = (root.find(_S + tag) for tag in ('cellXfs', 'fonts', 'fills'))
        if formats is None:
            return []
        names = {node.get('xfId'): node.get('name') for node in root.iter(_S + 'cellStyle')}
        custom = {node.get('numFmtId'): node.get('formatCode') for node in root.iter(_S + 'numFmt')}
        from openpyxl.styles.numbers import BUILTIN_FORMATS
        options = []
        for index, style in enumerate(formats):
            font_id, fill_id = int(style.get('fontId', '0')), int(style.get('fillId', '0'))
            font = fonts[font_id] if fonts is not None and font_id < len(fonts) else None
            fill = fills[fill_id] if fills is not None and fill_id < len(fills) else None
            def font_value(tag, attribute='val'):
                child = font.find(_S + tag) if font is not None else None
                return child.get(attribute) if child is not None else None
            color = font_value('color', 'rgb')
            background = fill.find(_S + 'patternFill/' + _S + 'fgColor') if fill is not None else None
            size = font_value('sz')
            format_id = int(style.get('numFmtId', '0'))
            preview = {'font_name': font_value('name'), 'font_size_points': float(size) if size else None,
                'bold': font is not None and font.find(_S + 'b') is not None,
                'italic': font is not None and font.find(_S + 'i') is not None,
                'font_color': '#' + color[-6:] if color and re.fullmatch('[0-9a-fA-F]{8}', color) else None,
                'fill_color': '#' + background.get('rgb')[-6:] if background is not None and re.fullmatch(
                    '[0-9a-fA-F]{8}', background.get('rgb', '')) else None,
                'number_format': custom.get(str(format_id), BUILTIN_FORMATS.get(format_id)),
                'theme_colors': 'PRESERVED_NOT_RESOLVED'}
            style_name = names.get(style.get('xfId'))
            details = [preview['font_name'], (f"{preview['font_size_points']:g} pt" if preview['font_size_points'] else None),
                       'Bold' if preview['bold'] else None, 'Italic' if preview['italic'] else None,
                       preview['number_format'] if format_id else None]
            options.append({'style_index': index, 'name': ' · '.join(str(value) for value in
                            [style_name, *details] if value) or f'Style {index + 1}',
                            'preview': preview, 'editable': True, 'source': 'existing_workbook_style'})
        return options

    def _validate_object_properties(self, binding, value):
        kind, before = binding['public']['kind'], binding['public']['value']
        if type(value) is not dict or set(value) != set(before):
            raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
        if kind in {'word_run_format', 'drawing_run_format'}:
            for key in ('bold', 'italic'):
                if value[key] is not None and type(value[key]) is not bool:
                    raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            size_key = 'size_half_points' if kind == 'word_run_format' else 'size_hundredth_points'
            if value[size_key] is not None and (type(value[size_key]) is not int or not 2 <= value[size_key] <= 200000):
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            if value['color'] is not None and (type(value['color']) is not str or re.fullmatch(r'[0-9a-fA-F]{6}|auto', value['color']) is None):
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            underlines = ({'none','single','words','double','thick','dotted','dottedHeavy','dash','dashedHeavy',
                'dashLong','dashLongHeavy','dotDash','dashDotHeavy','dotDotDash','dashDotDotHeavy','wave','wavyHeavy','wavyDouble'}
                if kind == 'word_run_format' else {'none','sng','words','dbl','heavy','dotted','dottedHeavy','dash',
                'dashHeavy','dashLong','dashLongHeavy','dotDash','dotDashHeavy','dotDotDash','dotDotDashHeavy','wavy','wavyHeavy','wavyDbl'})
            if value['underline'] != before['underline'] and value['underline'] is not None and value['underline'] not in underlines:
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            if (kind == 'drawing_run_format' and value['color'] == 'auto' and value['color'] != before['color']):
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
        elif kind == 'word_paragraph_format':
            if value['alignment'] is not None and value['alignment'] not in {'left','right','center','both','distribute','start','end','thaiDistribute','numTab','mediumKashida','highKashida','lowKashida'}:
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            for key in ('keep_with_next', 'page_break_before'):
                if value[key] is not None and type(value[key]) is not bool:
                    raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            for key in ('space_before_twips', 'space_after_twips'):
                if value[key] is not None and (type(value[key]) is not int or not 0 <= value[key] <= 31680):
                    raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
        elif kind == 'shape_geometry':
            for key in ('x_emu','y_emu','width_emu','height_emu','rotation_60000deg'):
                if type(value[key]) is not int or not -(1 << 31) < value[key] < (1 << 31):
                    raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            if value['width_emu'] <= 0 or value['height_emu'] <= 0 or not 0 <= value['rotation_60000deg'] < 21600000:
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
            if any(type(value[key]) is not bool for key in ('flip_horizontal','flip_vertical')):
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
        elif kind == 'cell_style':
            style_part = next((item['target'] for item in self.relationships.get('xl/workbook.xml', {}).values()
                               if item['type'].endswith('/styles')), None)
            styles = self._tree(style_part).getroot().find(_S + 'cellXfs') if style_part else None
            if type(value['style_index']) is not int or styles is None or not 0 <= value['style_index'] < len(styles):
                raise ValueError('OFFICE_OBJECT_PROPERTIES_INVALID')
        return copy.deepcopy(value)

    @staticmethod
    def _insert_property(parent, node, names, namespace):
        ranks = {namespace + name: index for index, name in enumerate(names)}
        rank = ranks.get(node.tag, len(ranks))
        position = next((index for index, child in enumerate(parent)
                         if ranks.get(child.tag, -1) > rank), len(parent))
        parent.insert(position, node)

    def _apply_object_properties(self, binding, value):
        kind, node, before = binding['public']['kind'], binding['node'], binding['public']['value']
        changed = {key for key in value if value[key] != before[key]}
        if kind == 'shape_geometry':
            for tag, fields in (('off', {'x':'x_emu','y':'y_emu'}), ('ext', {'cx':'width_emu','cy':'height_emu'})):
                child = node.find(_A + tag)
                for attr, key in fields.items():
                    if key in changed:
                        child.set(attr, str(value[key]))
            for key, attr in (('rotation_60000deg','rot'),('flip_horizontal','flipH'),('flip_vertical','flipV')):
                if key in changed:
                    node.set(attr, str(int(value[key])))
            return
        if kind == 'cell_style':
            node.set('s', str(value['style_index']))
            return
        word = kind.startswith('word_')
        namespace = _W if word else _A
        properties_tag = 'pPr' if kind == 'word_paragraph_format' else 'rPr'
        properties = node.find(namespace + properties_tag)
        if properties is None:
            properties = etree.Element(namespace + properties_tag)
            node.insert(0, properties)
        if not word:
            for key, attr in (('bold','b'),('italic','i'),('underline','u'),('size_hundredth_points','sz')):
                if key in changed:
                    if value[key] is None:
                        properties.attrib.pop(attr, None)
                    else:
                        properties.set(attr, str(int(value[key])) if type(value[key]) is bool else str(value[key]))
            if 'color' in changed:
                for child in list(properties):
                    if etree.QName(child).localname in {'noFill','solidFill','gradFill','blipFill','pattFill','grpFill'}:
                        properties.remove(child)
                if value['color'] is not None:
                    fill = etree.Element(_A + 'solidFill')
                    etree.SubElement(fill, _A + 'srgbClr', val=value['color'])
                    self._insert_property(properties, fill, ('ln','noFill','solidFill','gradFill','blipFill','pattFill','grpFill',
                        'effectLst','effectDag','highlight','uLnTx','uLn','uFillTx','uFill','latin','ea','cs','sym','hlinkClick','hlinkMouseOver','rtl','extLst'), _A)
            return
        paragraph = kind == 'word_paragraph_format'
        names = (('pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers',
            'pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN',
            'bidi','adjustRightInd','snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc',
            'textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange') if paragraph else
            ('rStyle','rFonts','b','bCs','i','iCs','caps','smallCaps','strike','dStrike','outline','shadow','emboss','imprint','noProof',
             'snapToGrid','vanish','webHidden','color','spacing','w','kern','position','sz','szCs','highlight','u','effect','bdr','shd',
             'fitText','vertAlign','rtl','cs','em','lang','eastAsianLayout','specVanish','oMath','rPrChange'))
        mapping = ({'alignment':('jc','val'),'keep_with_next':('keepNext','val'),'page_break_before':('pageBreakBefore','val'),
            'space_before_twips':('spacing','before'),'space_after_twips':('spacing','after')} if paragraph else
            {'bold':('b','val'),'italic':('i','val'),'underline':('u','val'),'size_half_points':('sz','val'),'color':('color','val')})
        for key in changed:
            tag, attribute = mapping[key]
            child = properties.find(_W + tag)
            if child is None and value[key] is not None:
                child = etree.Element(_W + tag)
                self._insert_property(properties, child, names, _W)
            if child is not None:
                if value[key] is None:
                    if tag == 'spacing':
                        child.attrib.pop(_W + attribute, None)
                    else:
                        properties.remove(child)
                else:
                    child.set(_W + attribute, str(int(value[key])) if type(value[key]) is bool else str(value[key]))

    def edit(self, commands, *, expected_package_sha256, acknowledge_signature_invalidation=False):
        if expected_package_sha256 != self.sha256:
            raise ValueError('OFFICE_PACKAGE_REVISION_CONFLICT')
        if self.active:
            raise ValueError('OFFICE_ACTIVE_CONTENT_EDIT_UNSUPPORTED')
        if self.signatures and acknowledge_signature_invalidation is not True:
            raise ValueError('OFFICE_SIGNATURE_INVALIDATION_CONFIRMATION_REQUIRED')
        if type(commands) is not list or not 1 <= len(commands) <= 200:
            raise ValueError('OFFICE_EDIT_COMMANDS_INVALID')
        seen, checked = set(), []
        for command in commands:
            if type(command) is not dict or set(command) != {'object_id', 'operation', 'expected_value', 'value'}:
                raise ValueError('OFFICE_EDIT_COMMAND_INVALID')
            identity = command['object_id']
            if type(identity) is not str or identity in seen or identity not in self.bindings:
                raise ValueError('OFFICE_EDIT_OBJECT_INVALID')
            seen.add(identity)
            binding = self.bindings[identity]
            if binding['public']['editable'] is not True:
                raise ValueError('OFFICE_EDIT_OPERATION_UNSUPPORTED')
            if command['expected_value'] != binding['public']['value']:
                raise ValueError('OFFICE_OBJECT_REVISION_CONFLICT')
            kind, operation, value = binding['public']['kind'], command['operation'], command['value']
            if operation != _OBJECT_OPERATIONS[kind]:
                raise ValueError('OFFICE_EDIT_OPERATION_UNSUPPORTED')
            if kind in {'word_run_format','drawing_run_format','word_paragraph_format','shape_geometry','cell_style'}:
                checked.append((command, binding, self._validate_object_properties(binding, value)))
                continue
            if kind == 'cell_boolean':
                if type(value) is not bool:
                    raise ValueError('OFFICE_CELL_VALUE_INVALID')
                value = '1' if value else '0'
            if type(value) is not str or len(value) > 32000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
                raise ValueError('OFFICE_EDIT_TEXT_INVALID')
            if kind == 'cell_number':
                try:
                    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', value) or not Decimal(value).is_finite():
                        raise ValueError('OFFICE_CELL_VALUE_INVALID')
                except InvalidOperation as exc:
                    raise ValueError('OFFICE_CELL_VALUE_INVALID') from exc
            if kind == 'cell_formula':
                value = value.removeprefix('=')
                if (not value.strip() or '[' in value or ']' in value or '://' in value or '|' in value
                        or re.search(r'(?i)(?:WEBSERVICE|HYPERLINK|RTD|DDE|CALL|REGISTER|EXEC|IMAGE|IMPORTXML)\s*\(', value)):
                    raise ValueError('OFFICE_EXTERNAL_FORMULA_FORBIDDEN')
            checked.append((command, binding, value))
        modified, cloned, applied = set(), {}, []
        for command, binding, value in checked:
            part, node, cell = binding['part'], binding['node'], binding['cell']
            structured = binding['public']['kind'] in {'word_run_format','drawing_run_format','word_paragraph_format','shape_geometry','cell_style'}
            if structured:
                self._apply_object_properties(binding, value)
            if binding['shared'] is not None:
                cell_key = id(cell)
                if cell_key not in cloned:
                    inline = copy.deepcopy(binding['shared'])
                    inline.tag = _S + 'is'
                    for old in list(cell):
                        if old.tag in {_S + 'v', _S + 'is'}:
                            cell.remove(old)
                    cell.set('t', 'inlineStr')
                    extension = cell.find(_S + 'extLst')
                    cell.insert(list(cell).index(extension) if extension is not None else len(cell), inline)
                    cloned[cell_key] = inline
                node = list(cloned[cell_key].iter(_S + 't'))[binding['text_index']]
            if not structured:
                node.text = value
            if binding['public']['kind'] in {'text_run', 'cell_text_run'} and (value[:1].isspace() or value[-1:].isspace()):
                node.set(_XML_SPACE, 'preserve')
            if binding['public']['kind'] == 'cell_formula':
                cell.attrib.pop('t', None)
            modified.add(part)
            applied.append({'object_id': command['object_id'], 'operation': command['operation'],
                            'before': command['expected_value'], 'after': value,
                            'locator': binding['public']['locator']})
        removed, calculation = set(), {'formula_evaluation': 'NOT_RUN'}
        if self.spec['extension'] == '.xlsx' and any(binding['public']['kind'] != 'cell_style' for _command, binding, _value in checked):
            removed, calculation = self._invalidate_xlsx_calculation(modified)
        return self._edited_output(modified, applied, removed=removed, calculation=calculation)

    def _edited_output(self, modified, applied, *, removed=None, calculation=None):
        removed = removed or set()
        calculation = calculation or {}
        replacements = {part: _serialized(self.trees[part]) for part in modified}
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as target:
            target.comment = self.comment
            for info in self.infos:
                if info.filename not in removed:
                    target.writestr(copy.copy(info), replacements.get(info.filename, self.parts[info.filename]))
        raw = output.getvalue()
        result = OfficePackage(raw, filename=self.filename, media_type=self.media_type)
        unchanged = [name for name in self.parts if name not in modified and name not in removed]
        if any(self.parts[name] != result.parts[name] for name in unchanged) or set(result.parts) != set(self.parts) - removed:
            raise ValueError('OFFICE_UNEDITED_PART_CHANGED')
        receipt = {'schema': 'lumenis.office-edit/v1', 'parent_package_sha256': self.sha256,
                   'package_sha256': result.sha256, 'modified_parts': sorted(modified), 'removed_parts': sorted(removed), 'edits': applied,
                   'preserved_parts': {name: _sha(self.parts[name]) for name in unchanged},
                   'relationship_parts_unchanged': all(name in unchanged for name in self.parts if name.endswith('.rels')),
                   'signature_state': 'INVALIDATED_ORIGINAL_SIGNATURE_PARTS_RETAINED' if self.signatures else 'NOT_PRESENT',
                   'external_relationships': self.external, 'active_content_execution': 'NOT_RUN',
                   **calculation,
                   'native_application_acceptance': 'NOT_VERIFIED'}
        return {'content': raw, 'receipt': receipt, 'projection': result.projection()}

    @staticmethod
    def _paragraph_text(paragraph):
        return ''.join(node.text or '' if node.tag == _W + 't' else '\t' if node.tag == _W + 'tab' else '\n'
                       for node in paragraph.iter() if node.tag in {_W+'t', _W+'tab', _W+'br', _W+'cr'})

    @staticmethod
    def _replace_paragraph_text(paragraph, text):
        """Keep run styles, hyperlinks, bookmarks and nontext objects in place."""
        from difflib import SequenceMatcher
        runs = list(paragraph.iter(_W + 'r'))
        if not runs:
            runs = [etree.SubElement(paragraph, _W + 'r')]
        pieces, owners = [], []
        text_tags = {_W+'t', _W+'tab', _W+'br', _W+'cr'}
        for index, run in enumerate(runs):
            nodes = [node for node in run if node.tag in text_tags]
            value = ''.join(node.text or '' if node.tag == _W+'t' else '\t' if node.tag == _W+'tab' else '\n' for node in nodes)
            pieces.append(value)
            owners.extend([index] * len(value))
        before = ''.join(pieces)
        revised = [''] * len(runs)
        for kind, start, end, after_start, after_end in SequenceMatcher(None, before, text).get_opcodes():
            if kind == 'equal':
                for position in range(start, end):
                    revised[owners[position]] += before[position]
            elif kind in {'insert', 'replace'}:
                owner = owners[min(start, len(owners)-1)] if owners else 0
                revised[owner] += text[after_start:after_end]
        for run, value in zip(runs, revised):
            positions = [index for index,node in enumerate(run) if node.tag in text_tags]
            position = positions[0] if positions else len(run)
            for node in list(run):
                if node.tag in text_tags:
                    run.remove(node)
            for piece in re.split('(\n|\t)', value):
                if not piece:
                    continue
                node = etree.Element(_W + ('br' if piece == '\n' else 'tab' if piece == '\t' else 't'))
                if piece not in {'\n','\t'}:
                    node.text = piece
                    if piece[:1].isspace() or piece[-1:].isspace():
                        node.set(_XML_SPACE, 'preserve')
                run.insert(position, node)
                position += 1

    def revise_word_manuscript(self, manuscript, replacements, *, expected_package_sha256, authorship_notice=None):
        """Revise selected original manuscript paragraphs, preserving the OPC package."""
        if self.spec['extension'] != '.docx' or self.sha256 != expected_package_sha256:
            raise ValueError('OFFICE_PACKAGE_REVISION_CONFLICT')
        if self.active or self.signatures:
            raise ValueError('OFFICE_ACTIVE_OR_SIGNED_REVISION_REQUIRES_EXPLICIT_OWNER')
        if (type(replacements) is not dict or not replacements or any(type(index) is not int or
                not 0 <= index < len(manuscript['sections']) for index in replacements)):
            raise ValueError('OFFICE_MANUSCRIPT_SELECTION_INVALID')
        part = next((item['target'] for item in self.relationships.get('', {}).values()
                     if item['type'].endswith('/officeDocument')), None)
        if part is None:
            raise ValueError('OFFICE_MANUSCRIPT_ANCHOR_INVALID')
        body = self._tree(part).find(_W+'body')
        paragraphs = [node for node in body if node.tag == _W+'p'] if body is not None else []
        if (len(paragraphs) < 3 or self._paragraph_text(paragraphs[0]) != manuscript['title']
                or self._paragraph_text(paragraphs[1]) != manuscript['summary']):
            raise ValueError('OFFICE_MANUSCRIPT_ANCHOR_CHANGED')
        offset, anchors = 3, []
        for section in manuscript['sections']:
            expected = [section['heading'], *section['body'].split('\n\n')]
            nodes = paragraphs[offset:offset+len(expected)]
            if len(nodes) != len(expected) or [self._paragraph_text(node) for node in nodes] != expected:
                raise ValueError('OFFICE_MANUSCRIPT_ANCHOR_CHANGED')
            anchors.append(nodes)
            offset += len(expected)
        applied = []
        if authorship_notice is not None:
            if type(authorship_notice) is not str or not authorship_notice.strip() or len(authorship_notice) > 1000:
                raise ValueError('OFFICE_AUTHORSHIP_NOTICE_INVALID')
            before_notice = self._paragraph_text(paragraphs[2])
            self._replace_paragraph_text(paragraphs[2], authorship_notice)
            applied.append({'operation':'update_authorship_notice','before':before_notice,'after':authorship_notice})
        for index, replacement in sorted(replacements.items()):
            nodes = anchors[index]
            values = [replacement['heading'], *replacement['body'].split('\n\n')]
            if any(type(value) is not str or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value) for value in values):
                raise ValueError('OFFICE_EDIT_TEXT_INVALID')
            # Existing paragraph objects remain even if the new prose has fewer
            # paragraphs. Empty objects retain comments, bookmarks and drawings.
            for number, node in enumerate(nodes):
                self._replace_paragraph_text(node, values[number] if number < len(values) else '')
            cursor = nodes[-1]
            for value in values[len(nodes):]:
                new = etree.Element(_W+'p', nsmap=cursor.nsmap)
                properties = cursor.find(_W+'pPr')
                if properties is not None:
                    new.append(copy.deepcopy(properties))
                run = etree.SubElement(new, _W+'r')
                formatting = next(cursor.iter(_W+'rPr'), None)
                if formatting is not None:
                    run.append(copy.deepcopy(formatting))
                self._replace_paragraph_text(new, value)
                cursor.addnext(new)
                cursor = new
            applied.append({'operation':'revise_manuscript_section', 'section_index':index,
                            'before':manuscript['sections'][index], 'after':replacement})
        return self._edited_output({part}, applied)
