import io
import json
import uuid
import pytest
from tests.test_independent_growth_real_app import real_app, headers
from tests.test_saved_document_workflow import workspace
from tests.test_growth_research_shared import BASE


def upload(ws):
    from openpyxl import Workbook
    wb=Workbook();one=wb.active;one.title='posts'
    columns=['title','content','platform','post_id','likes','sentiment','publish_date']
    one.append(columns)
    one.append(['first','verbatim first','test','P1',0,'positive','2026-10-18'])
    one.append(['second','verbatim second','test','P2',None,None,None])
    one.append(['third','verbatim third','test','P3','=2+2','negative',None])
    two=wb.create_sheet('comments')
    two.append(['title','content','platform','comment_id','parent_post_id','likes'])
    two.append(['reply','verbatim comment','test','C1','P1',5])
    stream=io.BytesIO();wb.save(stream)
    response=ws['api'].post('/api/experience/materials/documents',headers=headers(201),
        data={'file':(io.BytesIO(stream.getvalue()),'typed.xlsx'),'request_id':'typed-report-'+uuid.uuid4().hex})
    assert response.status_code==200,response.get_json()
    return response.get_json()['result']


def generate(ws,source):
    response=ws['api'].post(BASE+'/generate',headers=headers(201),json={
        'source_ids':[source['artifact']['artifact_id']],'business_question':'What explicit data did these sheets provide?'})
    assert response.status_code==200,response.get_json()
    return response.get_json()['result']


def test_typed_report_real_session_all_sheets_metrics_reopen_and_revoke(real_app):
    ws=real_app;source=upload(ws);report=generate(ws,source)
    analysis=report['typed_source_analysis']
    assert analysis['row_coverage']['total_rows']==4
    assert analysis['row_coverage']['preserved_header_rows']==2
    likes=next(m for m in analysis['metrics'] if m['metric']=='likes')['explicit_fields']
    by_sheet={field['table_locator']['sheet_name']:field for field in likes}
    assert by_sheet['posts']['known_count']==1 and by_sheet['posts']['zero_count']==1
    assert by_sheet['posts']['missing_count']==1 and by_sheet['posts']['formula_count']==1
    assert by_sheet['posts']['sum_of_known_imported_values']=='0'
    assert by_sheet['comments']['sum_of_known_imported_values']=='5'
    assert analysis['comment_parent_links']['matched_exactly']==1
    assert analysis['source_reported_sentiment']['fields'][0]['labels']=={'negative':1,'positive':1}
    assert 'sum of imported known values=5' in report['source_brief']['markdown']
    reopened=ws['api'].get(BASE+'/jobs/'+report['job_id'],headers=headers(201))
    assert reopened.status_code==200,reopened.get_json()
    assert reopened.get_json()['result']['typed_source_analysis']==analysis
    assert ws['api'].get(BASE+'/jobs/'+report['job_id'],headers=headers(202)).status_code in (403,404)
    with ws['repo'].transaction() as con:
        con.execute("DELETE FROM lumenis_shared_artifacts WHERE artifact_id=?",(source['artifact']['artifact_id'],))
    assert ws['api'].get(BASE+'/jobs/'+report['job_id'],headers=headers(201)).status_code != 200


def test_private_payload_typed_values_cannot_replace_original_attachment(real_app):
    ws=real_app;source=upload(ws)
    # Revision is legitimate user editing, but typed report re-parses the retained
    # bytes rather than interpreting user-authored parsed_source as original data.
    payload=source['document']['payload']
    payload['parsed_source']['source_tables'][0]['rows'][1]['values'][4]['value']=999999
    response=ws['api'].post('/api/experience/saved-work',headers=headers(201),json={
        'title':source['document']['title'],'markdown':source['document']['markdown'],
        'kind':'research','payload':payload,'citations':[],'request_id':'typed-edited-'+uuid.uuid4().hex,
        'work_id':source['document']['work_id'],'previous_artifact_id':source['artifact']['artifact_id'],
        'expected_draft_digest':source['document']['draft_digest']})
    assert response.status_code==200,response.get_json()
    revised=response.get_json()['result']
    report=generate(ws,revised)
    likes=next(m for m in report['typed_source_analysis']['metrics'] if m['metric']=='likes')['explicit_fields']
    assert next(f for f in likes if f['table_locator']['sheet_name']=='posts')['sum_of_known_imported_values']=='0'


def test_rehashed_forged_metric_snapshot_is_recomputed_and_rejected(real_app):
    from lumenis.shared.reporting import source_brief_job as owner
    ws=real_app;source=upload(ws);report=generate(ws,source)
    with ws['repo'].transaction() as con:
        row=con.execute('SELECT * FROM daily_reports WHERE report_id=?',(report['job_id'],)).fetchone()
        row=dict(row);snapshot=json.loads(row['source_snapshot'])
        metric=next(m for m in snapshot['typed_source_analysis']['metrics'] if m['metric']=='likes')
        metric['explicit_fields'][0]['sum_of_known_imported_values']='999999'
        encoded=owner._json(snapshot)
        digest=owner._hash_artifact(report_id=row['report_id'],report_type=row['report_type'],
            report_date=row['report_date'],total_notes=row['total_notes'],summary=row['summary'],source_snapshot=encoded)
        con.execute('UPDATE daily_reports SET source_snapshot=?,artifact_hash=? WHERE report_id=?',(encoded,digest,row['report_id']))
    response=ws['api'].get(BASE+'/jobs/'+report['job_id'],headers=headers(201))
    assert response.status_code!=200
    assert '999999' not in response.get_data(as_text=True)


def test_source_deleted_at_report_commit_cannot_create_typed_report(real_app,monkeypatch):
    from lumenis.shared.reporting import source_brief_job as owner
    ws=real_app;source=upload(ws);persist=owner._persist_report
    def remove_then_persist(*args,**kwargs):
        with ws['repo'].transaction() as con:
            con.execute('DELETE FROM lumenis_shared_artifacts WHERE artifact_id=?',(source['artifact']['artifact_id'],))
        return persist(*args,**kwargs)
    monkeypatch.setattr(owner,'_persist_report',remove_then_persist)
    response=ws['api'].post(BASE+'/generate',headers=headers(201),json={
        'source_ids':[source['artifact']['artifact_id']],'business_question':'Count the explicit rows'})
    assert response.status_code!=200
    assert ws['repo'].query_one('SELECT COUNT(*) FROM daily_reports')[0]==0


def test_explicit_stats_keep_boolean_unknown_and_sheet_identity_separate():
    from lumenis.shared.reporting.business_brief import analyze_shared_typed_source_records
    def row(index,sheet,likes):
        return {'index':index,'source_locator':{'sheet_name':sheet,'table_index':index,'row_index':2},
                'source_row_sha256':str(index)*64,'source_record':{},'content':'source',
                'source_fields':[{'column_index':1,'header':{'value':'likes'},'value':{'type':'boolean' if type(likes) is bool else 'integer','value':likes}}]}
    source={'evidence_ref':'saved_work:ART-'+'a'*16+':'+'b'*64,'source_row_count':2,
            'source_rows':[row(0,'posts',False),row(1,'comments',0)],'source_fingerprint':'c'*64,
            'content_sha256':'d'*64,'source_rows_sha256':'e'*64,'rights':'UNKNOWN','intended_use':'private_custody_and_analysis'}
    result=analyze_shared_typed_source_records(question='Count explicit fields',sources=[source])
    fields=next(m for m in result['metrics'] if m['metric']=='likes')['explicit_fields']
    by_sheet={f['table_locator']['sheet_name']:f for f in fields}
    assert by_sheet['posts']['invalid_count']==1 and by_sheet['posts']['sum_of_known_imported_values'] is None
    assert by_sheet['comments']['known_count']==1 and by_sheet['comments']['zero_count']==1


def test_metrics_only_header_after_preamble_has_explicit_locator_binding(real_app):
    from openpyxl import Workbook
    ws=real_app;book=Workbook();sheet=book.active;sheet.title='summary'
    sheet.append(['Publisher supplied totals'])
    sheet.append(['likes','views'])
    sheet.append([0,12])
    stream=io.BytesIO();book.save(stream)
    response=ws['api'].post('/api/experience/materials/documents',headers=headers(201),
        data={'file':(io.BytesIO(stream.getvalue()),'summary.xlsx'),'request_id':'metrics-only-'+uuid.uuid4().hex})
    assert response.status_code==200,response.get_json()
    report=generate(ws,response.get_json()['result']);analysis=report['typed_source_analysis']
    assert analysis['header_detection']['bindings'][0]['source_locator']['row_index']==2
    fields=next(m for m in analysis['metrics'] if m['metric']=='likes')['explicit_fields']
    assert fields[0]['known_count']==1 and fields[0]['sum_of_known_imported_values']=='0'
    assert analysis['row_coverage']['total_rows']==2  # preamble retained, header separated
    assert analysis['row_coverage']['preserved_header_rows']==1


# Independent regression cases, retained with their original assertions.
import pytest
from decimal import localcontext

def source(values, header='likes'):
    rows=[{'index':i,'source_locator':{'sheet_name':'posts','table_index':0,'row_index':i+2},'source_row_sha256':str(i%10)*64,'source_record':{},'content':'retained','source_fields':[{'column_index':1,'header':{'value':header},'value':{'type':'string','value':v}}]} for i,v in enumerate(values)]
    return {'evidence_ref':'saved_work:ART-'+'a'*16+':'+'b'*64,'source_row_count':len(rows),'source_rows':rows,'source_fingerprint':'c'*64,'content_sha256':'d'*64,'source_rows_sha256':'e'*64,'rights':'UNKNOWN','intended_use':'private_custody_and_analysis'}

def analyze(values, header='likes'):
    from lumenis.shared.reporting.business_brief import analyze_shared_typed_source_records
    return analyze_shared_typed_source_records(question='Count explicit retained fields',sources=[source(values,header)])

def total(values):
    return next(x for x in analyze(values)['metrics'] if x['metric']=='likes')['explicit_fields'][0]['sum_of_known_imported_values']

def test_exact_large_numeric_sum():
    assert total(['123456789012345678901234567890','1'])=='123456789012345678901234567891'

def test_calculation_does_not_depend_on_process_decimal_context():
    with localcontext() as c:
        c.prec=4
        actual=total(['123456','1'])
    assert actual=='123457'

def test_date_only_does_not_invent_midnight_precision():
    result=analyze(['2026-10-18','2026-10-18T12:00:00'],'publish_date')
    field=result['observation_time']['by_explicit_field'][0]
    assert field['range_status']=='INCOMPARABLE_PRECISION_OR_TIMEZONE'

@pytest.mark.parametrize('value',['0001-01-01T00:00:00+14:00','9999-12-31T23:59:59-12:00'])
def test_unrepresentable_utc_date_is_unknown_not_crash(value):
    result=analyze([value],'publish_date')
    assert result['observation_time']['by_explicit_field'][0]['unparseable_count']==1

from tests.test_independent_growth_real_app import real_app, headers, counts
from tests.test_saved_document_workflow import workspace
from tests.test_growth_research_shared import BASE
import io, uuid

def upload_csv(ws, text):
    r=ws['api'].post('/api/experience/materials/documents',headers=headers(201),data={'file':(io.BytesIO(text.encode()),'counts.csv'),'request_id':uuid.uuid4().hex})
    assert r.status_code==200,r.get_json()
    return r.get_json()['result']

def report_csv(ws,text):
    s=upload_csv(ws,text)
    r=ws['api'].post(BASE+'/generate',headers=headers(201),json={'source_ids':[s['artifact']['artifact_id']],'business_question':'Review the complete explicit counts'})
    assert r.status_code==200,r.get_json()
    return s,r.get_json()['result']

def test_real_bytes_csv_exact_sum(real_app):
    s,r=report_csv(real_app,'title,likes\nfirst,123456789012345678901234567890\nsecond,1\n')
    f=next(x for x in r['typed_source_analysis']['metrics'] if x['metric']=='likes')['explicit_fields'][0]
    assert f['sum_of_known_imported_values']=='123456789012345678901234567891'



def test_exact_sum_isolates_ambient_exponent_limits_and_traps():
    from decimal import Overflow, Rounded, Inexact
    with localcontext() as context:
        context.prec = 2
        context.Emax = 2
        context.Emin = -2
        context.traps[Overflow] = True
        context.traps[Rounded] = True
        context.traps[Inexact] = True
        assert total(['123456789012345678901234567890', '0.00000001']) == '123456789012345678901234567890.00000001'
