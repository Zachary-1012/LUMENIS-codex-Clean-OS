# -*- coding: utf-8 -*-
"""方案拆解结果 → Word 文档导出（可交付比稿/提案材料）。

用法（CLI）：
  python scripts/pitch_export.py <analysis.json> <out.docx>
  python scripts/pitch_export.py --text "方案需求文本" --brand 广州太古汇 --hotspot 15周年 --out out.docx

依赖：python-docx（ZBC .venv 已装 1.2.0）。个人业务资料只存本地，本模块不联网、不上传。
"""
import json
import os
import sys
import tempfile
import uuid

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _HAS_DOCX = True
except Exception:
    _HAS_DOCX = False

from scripts.pitch_parser import analyze_pitch

_ACCENT = RGBColor(0x2F, 0x3B, 0x66) if _HAS_DOCX else None
_SUB = RGBColor(0x6B, 0x72, 0x80) if _HAS_DOCX else None


def _h(doc, text, size=14, color=_ACCENT, space=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(2)
    return p


def _p(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(3)
    return p


def _kv_table(doc, rows):
    if not rows:
        return
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v) if v not in (None, "") else "—"


def _save_new_document(doc, out_path, check_continue):
    """Stage and validate locally; publish a new file without replacing a revision."""
    from lumenis.shared.reporting.artifact import assert_artifact_or_raise

    destination = os.path.abspath(os.fspath(out_path)) if out_path else os.path.join(
        tempfile.gettempdir(), "pitch_%s.docx" % uuid.uuid4().hex)
    if not destination.lower().endswith(".docx"):
        raise ValueError("输出必须为 .docx 文件")
    if os.path.lexists(destination):
        raise FileExistsError("输出已存在，请为修订选择新文件")
    if check_continue:
        check_continue()
    fd, staged = tempfile.mkstemp(prefix=".pitch-stage-", suffix=".docx",
                                  dir=os.path.dirname(destination))
    try:
        with os.fdopen(fd, "wb") as stream:
            doc.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        # The shared verifier proves bytes/hash only; reopen also verifies OOXML.
        assert_artifact_or_raise(staged)
        Document(staged)
        if check_continue:
            check_continue()
        os.link(staged, destination)
        return destination
    finally:
        os.unlink(staged)


def _remove_title_borders(doc):
    """Remove template borders from both the Title style and its paragraphs."""
    from docx.oxml.ns import qn

    for element in [doc.styles["Title"].element,
                    *[paragraph._p for paragraph in doc.paragraphs if paragraph.style.name == "Title"]]:
        for border in list(element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)


def _manuscript_document(manuscript):
    """Export supplied final prose. This function neither generates nor approves it."""
    if not isinstance(manuscript, dict):
        raise ValueError("成稿必须包含标题、说明和正文段落")
    title, summary = manuscript.get("title"), manuscript.get("summary")
    sections = manuscript.get("sections")
    if (not isinstance(title, str) or not title.strip()
            or not isinstance(summary, str) or not summary.strip()
            or not isinstance(sections, list) or not sections):
        raise ValueError("成稿标题、说明和正文不能为空")
    for section in sections:
        if (not isinstance(section, dict)
                or not isinstance(section.get("heading"), str)
                or not section["heading"].strip()
                or not isinstance(section.get("body"), str)
                or not section["body"].strip()):
            raise ValueError("每节成稿必须有标题和实际正文")
    authorship = manuscript.get("authorship", "assisted")
    labels = {
        "zh": {"assisted": "辅助创作草稿 审阅与批准待确认",
               "user_supplied": "用户提供的成稿 审阅及使用授权待确认"},
        "en": {"assisted": "Draft prepared with assistance. Review and approval pending.",
               "user_supplied": "Supplied manuscript. Review and usage approval pending."},
        "ko": {"assisted": "도움을 받아 작성한 초안 검토 및 승인 대기",
               "user_supplied": "제공된 원고 검토 및 사용 승인 대기"},
    }
    # These are document language disclosures, not claims of model execution.
    for code, assisted, supplied in (
        ('ja','支援を受けて作成した草稿。レビューと承認は未完了です。','ユーザー提供の原稿。レビューと利用許可は確認待ちです。'),
        ('fr','Brouillon rédigé avec assistance. Révision et approbation en attente.','Manuscrit fourni par l’utilisateur. Révision et autorisation d’utilisation à confirmer.'),
        ('de','Mit Unterstützung erstellter Entwurf. Prüfung und Freigabe stehen aus.','Vom Nutzer bereitgestelltes Manuskript. Prüfung und Nutzungsfreigabe stehen aus.'),
        ('es','Borrador elaborado con asistencia. Revisión y aprobación pendientes.','Manuscrito proporcionado por el usuario. Revisión y autorización de uso pendientes.'),
        ('pt','Rascunho elaborado com assistência. Revisão e aprovação pendentes.','Manuscrito fornecido pelo usuário. Revisão e autorização de uso pendentes.'),
        ('ru','Черновик подготовлен с помощью. Проверка и одобрение ещё не выполнены.','Рукопись предоставлена пользователем. Проверка и разрешение на использование ожидаются.'),
        ('ar','مسودة أُعدّت بمساعدة. المراجعة والموافقة قيد الانتظار.','مخطوطة قدّمها المستخدم. المراجعة وإذن الاستخدام قيد الانتظار.'),
        ('hi','सहायता से तैयार मसौदा। समीक्षा और अनुमोदन लंबित हैं।','उपयोगकर्ता द्वारा दी गई पांडुलिपि। समीक्षा और उपयोग की अनुमति लंबित हैं।'),
        ('id','Draf disusun dengan bantuan. Peninjauan dan persetujuan masih menunggu.','Naskah disediakan pengguna. Peninjauan dan izin penggunaan masih menunggu.'),
    ):
        labels[code] = {'assisted':assisted, 'user_supplied':supplied}
    language = manuscript.get("language", "zh")
    if language not in labels:
        raise ValueError("成稿语言不受支持")
    if authorship not in labels[language]:
        raise ValueError("成稿来源必须为 assisted 或 user_supplied")
    doc = Document()
    for name in ("Title", "Heading 1", "Normal"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.add_paragraph(title, style="Title")
    doc.add_paragraph(summary)
    doc.add_paragraph(labels[language][authorship])
    for section in sections:
        doc.add_heading(section["heading"], level=1)
        # Preserve blank paragraphs, line breaks and all authored words.
        for paragraph in section["body"].split("\n\n"):
            doc.add_paragraph(paragraph)
    doc.core_properties.title = title
    doc.core_properties.language = language
    doc.core_properties.comments = "authorship=%s; approval=not_verified" % authorship
    _remove_title_borders(doc)
    return doc


def _append_audiovisual_plan(doc, snapshot, language, check_continue=None):
    """Render the complete saved planning snapshot, without inferring media facts."""
    from lumenis.shared.content_creation.contracts import AudiovisualPlan
    if not isinstance(snapshot, dict) or "input" not in snapshot:
        raise ValueError("AV_PLAN_SNAPSHOT_INVALID")
    expected = AudiovisualPlan.from_dict(snapshot["input"])
    if not expected.matches_snapshot(snapshot):
        raise ValueError("AV_PLAN_SNAPSHOT_INVALID")
    text = {
        "zh": ("完整视听计划", "以下为保存的创作意图与计划，尚未制作媒体。来源声明未重新核验；研究方法引用不等于素材许可。视觉、声音、事实和权益均待审查。稿件与计划的一致性仍待审阅。",
               "未知／未提供", "未提供项目", "空文本", "是", "否"),
        "en": ("Complete audiovisual plan", "Saved authoring intent and planning only. Media has not been produced. Provenance declarations are not revalidated; method references are not material licences. Visual, audio, factual and rights reviews remain pending. Manuscript/plan consistency still needs review.",
               "Unknown / not provided", "No items supplied", "Empty text", "Yes", "No"),
        "ko": ("전체 시청각 계획", "저장된 창작 의도와 계획이며 미디어는 아직 제작되지 않았습니다. 출처 선언은 재검증되지 않았고 연구 방법 참조는 소재 사용 허가가 아닙니다. 영상, 음성, 사실 및 권리와 원고·계획의 일치 여부는 검토가 필요합니다.",
               "알 수 없음 / 미제공", "제공된 항목 없음", "빈 텍스트", "예", "아니요"),
    }[language]
    labels = {
        "input": ("创作意图", "Authoring intent"), "schema": ("计划格式版本", "Plan format"),
        "brief_text": ("完整原始需求", "Complete original brief"),
        "method_source": ("研究方法出处（非素材许可）", "Research method reference (not a material licence)"),
        "reference": ("参考出处", "Reference"), "sha256": ("声明的摘要", "Declared SHA-256"),
        "outputs": ("交付计划", "Planned outputs"), "output_id": ("输出标识", "Output identifier"),
        "medium": ("媒体类型", "Medium"), "purpose": ("用途", "Purpose"),
        "canvas_ratio": ("画幅比例", "Canvas ratio"), "duration": ("时长要求", "Duration requirements"),
        "min_seconds": ("最短秒数", "Minimum seconds"), "max_seconds": ("最长秒数", "Maximum seconds"),
        "tolerance_seconds": ("容差秒数", "Tolerance seconds"), "sound": ("声音策略", "Sound policy"),
        "original": ("原声", "Original sound"), "voice": ("配音", "Voice"), "music": ("音乐", "Music"),
        "languages": ("语言", "Languages"), "constraints": ("约束及媒体验证状态", "Constraints and media validation"),
        "provenance": ("各项来源声明", "Provenance declarations"), "kind": ("声明类型", "Declaration kind"),
        "locator": ("声明位置", "Declaration locator"), "plan": ("镜头或播客分段计划", "Shot or podcast segment plan"),
        "rule_assessments": ("规则条件与排除项", "Rule conditions and exclusions"),
        "rule_ref": ("规则出处", "Rule reference"), "output_ids": ("适用输出", "Applicable outputs"),
        "required_conditions": ("必要条件", "Required conditions"), "exclusions": ("排除条件", "Exclusions"),
        "predicate": ("待判断事项", "Condition"), "value": ("声明结果", "Declared value"),
        "rule_decisions": ("规则评估（非执行许可）", "Rule assessment (not execution permission)"),
        "decision": ("评估状态", "Assessment status"), "status": ("计划状态", "Plan status"),
        "input_sha256": ("保存计划摘要", "Saved plan SHA-256"), "brief_sha256": ("原始需求摘要", "Brief SHA-256"),
        "estimated_duration_seconds": ("声明的预计时长（秒）", "Declared estimated duration (seconds)"),
        "execution": ("媒体执行状态", "Media execution status"), "predicate_evidence": ("条件核验状态", "Condition verification status"),
        "reviews": ("待完成的审查", "Pending reviews"), "visual": ("视觉", "Visual"), "audio": ("声音", "Audio"),
        "facts": ("事实", "Facts"), "rights": ("权益", "Rights"), "rendition_binding": ("实际成品关联", "Actual rendition binding"),
        "shots": ("镜头", "Shots"), "segments": ("播客分段", "Podcast segments"),
        "brief_ref": ("绑定的需求引用", "Bound brief reference"), "evidence_refs": ("绑定的证据引用", "Bound evidence references"),
        "narrative_beat": ("叙事节拍", "Narrative beat"), "emotional_beat": ("情绪节拍", "Emotional beat"),
        "shot_purpose": ("镜头用途", "Shot purpose"), "visual_motif": ("视觉母题", "Visual motif"),
        "shot_size": ("景别", "Shot size"), "camera_motion": ("镜头运动", "Camera movement"),
        "subject_action": ("主体动作", "Subject action"), "continuity": ("连贯性", "Continuity"),
        "audio_role": ("声音作用", "Audio role"), "original_sound": ("原声内容", "Original sound content"),
        "voice_over": ("配音文字", "Voice-over text"), "text_subtitles": ("字幕", "Subtitles"),
        "transition_type": ("转场方式", "Transition type"), "transition_reason": ("转场原因", "Transition reason"),
        "duration_sec": ("镜头时长（秒）", "Shot duration (seconds)"), "hold_sec": ("停留时长（秒）", "Hold duration (seconds)"),
        "pacing": ("节奏", "Pacing"), "title": ("标题", "Title"), "description": ("说明", "Description"),
        "language": ("语言", "Language"), "revision": ("修订号", "Revision"), "version": ("格式版本", "Format version"),
        "segment_id": ("分段标识", "Segment identifier"), "chapter_title": ("章节标题", "Chapter title"),
        "media_asset_id": ("声明的媒体素材标识", "Declared media asset identifier"),
        "audio_stream_index": ("声明的音轨索引", "Declared audio stream index"),
        "source_in": ("源片段入点（秒）", "Source in point (seconds)"), "source_out": ("源片段出点（秒）", "Source out point (seconds)"),
        "original_text": ("原始文字", "Original text"), "corrected_text": ("修订文字", "Revised text"),
        "speaker": ("说话人", "Speaker"),
    }
    _remove_title_borders(doc)
    doc.add_page_break()
    doc.add_heading(text[0], level=1)
    doc.add_paragraph(text[1])

    def field_label(key):
        pair = labels.get(key)
        return pair[0 if language == "zh" else 1] if pair else str(key).replace("_", " ")

    def render(label, value, depth=0):
        if check_continue:
            check_continue()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(min(depth, 4) * 0.35)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = isinstance(value, (dict, list)) and bool(value)
        paragraph.add_run(label + ": ").bold = True
        if isinstance(value, dict) and value:
            priority = ("output_id", "medium", "purpose", "brief_text", "title", "description", "canvas_ratio",
                        "duration", "sound", "languages", "constraints", "plan", "provenance", "method_source",
                        "outputs", "rule_assessments", "schema")
            keys = [key for key in priority if key in value] + [key for key in value if key not in priority]
            for key in keys:
                render(field_label(key), value[key], depth + 1)
        elif isinstance(value, list) and value:
            for number, child in enumerate(value, 1):
                render(str(number), child, depth + 1)
        else:
            if value is None:
                display = text[2]
            elif value == [] or value == {}:
                display = text[3]
            elif type(value) is bool:
                display = text[5] if value else text[6]
            elif value == "":
                display = text[4]
            else:
                display = str(value)
            paragraph.add_run(display)

    # Render every field, including full original nested constraints and media
    # plan, then the derived snapshot statuses and hashes. Do not summarize away
    # unknowns, user text, source declarations or exclusions.
    for key, value in snapshot.items():
        render(field_label(key), value)


def export_pitch_docx(analysis, out_path=None, *, manuscript=None, check_continue=None, audiovisual_plan=None):
    """导出需求分析或实际提供的成稿；新版本须新路径，不覆盖旧件。"""
    if not _HAS_DOCX:
        raise RuntimeError("python-docx 不可用，无法导出 Word")
    if not isinstance(analysis, dict):
        raise ValueError("需求分析必须为对象")
    if check_continue:
        check_continue()
    if manuscript is None:
        manuscript = analysis.get("manuscript")
    if manuscript is not None:
        doc = _manuscript_document(manuscript)
        if audiovisual_plan is not None:
            _append_audiovisual_plan(doc, audiovisual_plan, manuscript.get("language", "zh"), check_continue)
        return _save_new_document(doc, out_path, check_continue)
    if audiovisual_plan is not None:
        raise ValueError("AV_PLAN_MANUSCRIPT_REQUIRED")
    ex = (analysis.get("extracted") or {})
    doc = Document()
    # 封面标题
    t = doc.add_heading("", level=0)
    r = t.add_run("%s · %s 方案拆解" % (ex.get("brand") or "品牌", ex.get("category") or "品类"))
    r.font.size = Pt(20)
    r.font.color.rgb = _ACCENT
    _p(doc, "本材料为需求分析和创意候选，不是完整提案成稿。历史对标尚未经本次核验。")
    _p(doc, "调性 %s ｜ 圈层 %s ｜ 预算 %s ｜ 来源：Zachary Creative 方案拆解引擎"
       % (ex.get("tonality") or "—", ex.get("circle") or "—",
          ("%s%s" % (ex.get("budget", {}).get("value", ""), ex.get("budget", {}).get("unit", ""))) if ex.get("budget") else "—"),
       size=9, color=_SUB)

    # 一、需求要素
    _h(doc, "一、需求要素提取")
    _kv_table(doc, [
        ("品牌", ex.get("brand")), ("品类", ex.get("category")),
        ("预算", ("%s%s" % (ex.get("budget", {}).get("value", ""), ex.get("budget", {}).get("unit", ""))) if ex.get("budget") else "—"),
        ("圈层", ex.get("circle")), ("调性", ex.get("tonality")),
        ("渠道", "、".join(ex.get("channels") or []) or "—"),
        ("目标", "、".join(ex.get("goals") or []) or "—"),
        ("创意线索", "；".join(ex.get("creative_leads") or []) or "—"),
    ])

    # 二、对标基准
    bm = analysis.get("benchmark") or {}
    _h(doc, "二、对标基准")
    _kv_table(doc, [("对标对象", bm.get("benchmark")), ("打法", bm.get("playbook"))])

    # 三、创意方向
    ideas = (analysis.get("ideas") or {}).get("ideas") or []
    if ideas:
        _h(doc, "三、创意方向")
        for i, idea in enumerate(ideas, 1):
            _p(doc, "方向%d ｜ %s（%s）" % (i, idea.get("mechanism", ""), idea.get("tension_type", "")), bold=True)
            _p(doc, idea.get("big_idea", ""), size=12, bold=True)
            _kv_table(doc, [
                ("角度", idea.get("angle")), ("张力", idea.get("tension_type")),
                ("预算", idea.get("budget")), ("时效", idea.get("timeliness")),
                ("触点", "、".join(idea.get("touchpoints") or [])),
                ("信号/依据", idea.get("why")), ("来源", idea.get("provenance")),
            ])
    else:
        _h(doc, "三、创意方向", color=_SUB)
        _p(doc, "（未提取到创意方向）")

    # 四、方案脉络（PITCH_TOC 为 list[dict]）
    toc = analysis.get("pitch_toc") or []
    if toc:
        _h(doc, "四、方案脉络")
        for item in toc:
            if isinstance(item, dict):
                _p(doc, "• %s（%s）：%s" % (item.get("sec", ""), item.get("layer", ""), item.get("desc", "")), size=10)
            else:
                _p(doc, "• %s" % item, size=10)

    return _save_new_document(doc, out_path, check_continue)


def main():
    import argparse
    ap = argparse.ArgumentParser(description="方案拆解结果导出 Word")
    ap.add_argument("analysis_json", nargs="?", default="", help="analyze_pitch 的 JSON 文件")
    ap.add_argument("out", nargs="?", default="", help="输出 docx 路径")
    ap.add_argument("--text", default="", help="或直接传方案需求文本")
    ap.add_argument("--brand", default="")
    ap.add_argument("--hotspot", default="")
    ap.add_argument("--n", type=int, default=3)
    a = ap.parse_args()
    if a.analysis_json:
        analysis = json.load(open(a.analysis_json, encoding="utf-8"))
    else:
        analysis = analyze_pitch(a.text or "（空方案）", a.brand, a.n, a.hotspot)
    out = export_pitch_docx(analysis, a.out or None)
    print("导出完成: %s" % out)


if __name__ == "__main__":
    main()
