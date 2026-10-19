"""G104 real existing owner chain, synthetic authored text (not quality acceptance)."""
import copy
import io
import pytest
from docx import Document
from scripts.campaign_plan import generate_plan, PLAN_SECTIONS
from scripts.perspective_adapter import adapt_plan
from tests.test_internal_pitch_workflow import workspace, manuscript_arguments, run, pitch

@pytest.mark.parametrize('role', ['client', 'agency', 'vendor'])
def test_ten_section_author_edit_save_revise_download_reopen(workspace, role):
    ws = workspace
    plan = generate_plan('测试品牌', goal='说明待验证方案', budget='尚未提供', audience='待Brief确认')
    # The existing generator supplies a scaffold, never a finished original draft.
    assert 'X%' in plan['sections']['媒介与资源']
    assert 'X周' in plan['sections']['执行节奏']
    for i, heading in enumerate(plan['sections']):
        plan['sections'][heading] = f'用户原创测试段落{i}：{heading}。预算与期程尚未确认，不作收益承诺。'
    view = adapt_plan(plan, role)
    args = manuscript_arguments(ws, 'campaign-first-' + role)
    args['manuscript'] = {'title': '用户原创十段方案', 'summary': '仅供内部审校，未知字段待补。',
        'sections': [{'heading': k, 'body': v} for k,v in view['sections'].items()],
        'language': 'zh', 'authorship': 'user_supplied'}
    first = run(ws, pitch.prepare_internal_pitch(**args))
    assert first['state'] == 'succeeded'
    original = copy.deepcopy(first['work']['document'])
    output = pitch.download_internal_pitch(first['task_id'], actor=ws['actor'], scope=ws['scope'])
    paras = [p.text for p in Document(io.BytesIO(output['content'])).paragraphs]
    assert all(s['body'] in paras for s in args['manuscript']['sections'])
    assert all(any(heading in p for p in paras) for _,heading in PLAN_SECTIONS)
    assert not any('X%' in p or 'X周' in p for p in paras)
    revision = {**args, 'manuscript': copy.deepcopy(args['manuscript'])}
    revision['request_id'] = 'campaign-revise-' + role
    revision['manuscript']['sections'][0]['body'] += '用户修订：先补实际Brief，再判断方案。'
    revision.update(work_id=original['work_id'], previous_artifact_id=first['action_receipt']['artifact_id'], expected_draft_digest=original['draft_digest'])
    second = run(ws, pitch.prepare_internal_pitch(**revision))
    assert second['action_receipt']['version'] == 2
    assert second['work']['document']['payload']['manuscript'] == revision['manuscript']
    assert pitch.read_internal_pitch(first['task_id'], actor=ws['actor'], scope=ws['scope'])['work']['document'] == original
    updated = pitch.download_internal_pitch(second['task_id'], actor=ws['actor'], scope=ws['scope'])
    updated_paras = [p.text for p in Document(io.BytesIO(updated['content'])).paragraphs]
    assert all(s['body'] in updated_paras for s in revision['manuscript']['sections'])
    assert second['action_receipt']['external_effect'] == 'NOT_PERFORMED'
    assert second['action_receipt']['business_outcome'] == 'NOT_OBSERVED'


def test_scaffold_placeholders_are_preserved_as_user_text_not_silently_filled(workspace):
    # Existing owner truthfully preserves manuscript. It does not generate original
    # finished copy or infer missing budgets. This documents the remaining G104 gap.
    plan = generate_plan('测试品牌')
    args = manuscript_arguments(workspace, 'scaffold-not-finished')
    args['manuscript']['sections'] = [{'heading': k, 'body': v} for k,v in plan['sections'].items()]
    result = run(workspace, pitch.prepare_internal_pitch(**args))
    output = pitch.download_internal_pitch(result['task_id'], actor=workspace['actor'], scope=workspace['scope'])
    text = '\n'.join(p.text for p in Document(io.BytesIO(output['content'])).paragraphs)
    assert 'X%' in text and 'X周' in text
    assert result['work']['document']['payload']['truth_state']['owned_model_execution'] == 'NOT_RUN'
