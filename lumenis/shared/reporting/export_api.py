"""Scope-authorized Shared Reporting previews, saved reports and exports."""

import csv
import hashlib
import io
import json
import logging
import os
import re
import sqlite3
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

from flask import Blueprint, jsonify, request, send_file

import database as db
from lumenis.shared.database import get_repository, main_db_path

# Shared Reporting owns export behavior; repo root remains the data/runtime anchor.
PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

export_bp = Blueprint('export', __name__)

# 数据库/导出路径保持原生产位置，不因模块迁移改变数据归属。
DB_PATH = main_db_path()
EXPORT_DIR = str(PROJECT_ROOT / 'data' / 'exports')
os.makedirs(EXPORT_DIR, exist_ok=True)

LOGGER = logging.getLogger(__name__)


def get_db():
    """获取数据库连接"""
    conn = get_repository().connect(timeout=10)
    conn.row_factory = sqlite3.Row
    return conn


def _i(v, default=0):
    """安全转整数，None/空/异常都回落到 default"""
    try:
        return int(v) if v not in (None, '') else default
    except (TypeError, ValueError):
        return default

SENT_CN = {'positive': '正面', 'negative': '负面', 'neutral': '中性'}
RISK_CN = {'high': '高风险', 'medium': '中风险', 'low': '低风险'}

def note_type_cn(t):
    return '视频' if t == 'video' else '图文'

def engagement(n):
    """互动量 = 点赞+评论+收藏+分享"""
    return _i(n.get('likes')) + _i(n.get('comments')) + _i(n.get('collects')) + _i(n.get('shares'))


def get_date_range(period, date_str=None):
    """获取日期范围"""
    if date_str:
        try:
            base_date = datetime.strptime(date_str, '%Y-%m-%d')
        except:
            base_date = datetime.now()
    else:
        base_date = datetime.now()

    if period == 'daily':
        start = base_date.replace(hour=0, minute=0, second=0)
        end = base_date.replace(hour=23, minute=59, second=59)
    elif period == 'weekly':
        # 自然周：周一到周日
        start = (base_date - timedelta(days=base_date.weekday())).replace(hour=0, minute=0, second=0)
        end = (start + timedelta(days=6)).replace(hour=23, minute=59, second=59)
    elif period == 'monthly':
        # 自然月：1号到月底
        start = base_date.replace(day=1, hour=0, minute=0, second=0)
        if base_date.month == 12:
            end = base_date.replace(year=base_date.year + 1, month=1, day=1) - timedelta(seconds=1)
        else:
            end = base_date.replace(month=base_date.month + 1, day=1) - timedelta(seconds=1)
    elif period == 'yearly':
        # 自然年：1.1到12.31
        start = base_date.replace(month=1, day=1, hour=0, minute=0, second=0)
        end = base_date.replace(month=12, day=31, hour=23, minute=59, second=59)
    else:
        start = base_date.replace(hour=0, minute=0, second=0)
        end = base_date.replace(hour=23, minute=59, second=59)

    return start, end


def _authorized_legacy_report_context(action='read'):
    """Compatibility calls must resolve the same authenticated report scope as HTTP routes."""
    try:
        return _request_report_scope(action)
    except RuntimeError as exc:
        raise PermissionError('REPORT_SCOPE_REQUIRED') from exc


def _legacy_report_bounds(period, date_str=None, start_date=None, end_date=None):
    normalized = 'yearly' if period == 'annual' else period
    if start_date or end_date:
        kind, start, end = _report_period(normalized, start_date, end_date)
        return kind, start, end
    start, end = get_date_range(normalized, date_str)
    return _report_period(normalized, start.strftime('%Y-%m-%d'), end.strftime('%Y-%m-%d'))


def filter_notes_by_period(period, date_str=None, require_negative_confirmed=False,
                           start_date=None, end_date=None):
    """Read the current authorized scope; never apply a reference Subject implicitly."""
    actor, scope, _entries, _selected = _authorized_legacy_report_context('read')
    kind, start, end = _legacy_report_bounds(period, date_str, start_date, end_date)
    rows = db.get_scoped_event_source_rows(
        scope, source_types=('note',), date_start=start, date_end=end, limit=5000)
    if rows.get('truncated') or rows.get('rejected_count'):
        raise ValueError('REPORT_SOURCE_COVERAGE_INCOMPLETE')
    notes = rows.get('sources') or []
    _verify_report_source_use(actor, db._normalize_event_scope(scope, allow_global=True), notes, [])
    if require_negative_confirmed and kind != 'daily':
        notes = [row for row in notes if row.get('sentiment') != 'negative'
                 or row.get('negative_confirmed') == 1]
    return notes, datetime.strptime(start, '%Y-%m-%d'), datetime.strptime(end, '%Y-%m-%d')


def generate_report_data(period, date_str=None, start_date=None, end_date=None):
    """Compatibility facade over the current scope-authorized Shared report path."""
    actor, scope, _entries, selected = _authorized_legacy_report_context('read')
    kind, start, end = _legacy_report_bounds(period, date_str, start_date, end_date)
    _kind, _start, _end, report, _source_count, _notes, _events = _build_report(
        kind, start, end, actor=actor, scope=scope, scope_entry=selected)
    return report


def bind_brand_matcher(_matcher):
    """Retained import compatibility; current reports use their authorized scope registry."""


def export_report_to_csv(report_type='daily', date_str=None, start_date=None, end_date=None):
    """Compatibility facade over the current permission- and SourceUse-checked export."""
    actor, scope, _entries, selected = _authorized_legacy_report_context('read')
    kind, start, end = _legacy_report_bounds(report_type, date_str, start_date, end_date)
    result = export_artifacts(kind, start, end, ('csv',), actor=actor, scope=scope,
                              scope_entry=selected)
    artifact = result['files']['csv']
    return artifact['content'].decode('utf-8'), artifact['filename']


def _report_period(kind, start=None, end=None):
    if kind == 'annual':
        kind = 'yearly'
    if kind not in {'daily', 'weekly', 'monthly', 'yearly'}:
        raise ValueError('REPORT_PERIOD_INVALID')
    date_str = end or start
    resolved_start, resolved_end = get_date_range(kind, date_str)
    if start:
        resolved_start = datetime.strptime(start, '%Y-%m-%d')
    if end:
        resolved_end = datetime.strptime(end, '%Y-%m-%d')
    if resolved_start > resolved_end:
        raise ValueError('REPORT_DATE_RANGE_INVALID')
    return kind, resolved_start.strftime('%Y-%m-%d'), resolved_end.strftime('%Y-%m-%d')


def _source_use_scope(scope, row):
    """A GLOBAL report remains an envelope; every source keeps its exact focus."""
    domain = 'cth' if scope.get('domain_id') == 'cth' else 'zbmp'
    focus_key = 'subject_id' if domain == 'cth' else 'brand_focus_id'
    other_key = 'brand_focus_id' if domain == 'cth' else 'subject_id'
    focus = row.get(focus_key)
    if not isinstance(focus, str) or not focus.strip() or row.get(other_key) not in (None, ''):
        raise ValueError('REPORT_SOURCE_SCOPE_UNVERIFIED')
    if not scope.get('global_scope') and focus != scope.get(focus_key):
        raise ValueError('REPORT_SOURCE_SCOPE_MISMATCH')
    return {'domain_id': domain, 'subject_id': focus if domain == 'cth' else None,
            'brand_focus_id': focus if domain == 'zbmp' else None, 'global_scope': False}


def _verify_report_source_use(actor, scope, notes, events, *, include_bindings=False):
    """Recheck each included source's original bytes and declared use rights."""
    from lumenis.shared.database import get_repository
    from lumenis.shared.reporting.business_brief import _source_use_material
    from lumenis.shared.reporting.operational_brief import _content_evidence_ref

    unique = {}
    for row in list(notes or []):
        if isinstance(row, dict):
            source_type = row.get('source_type') or 'note'
            source_id = row.get('source_id') or row.get('note_id')
            source_scope = _source_use_scope(scope, row)
            unique[(source_type, source_id, source_scope['domain_id'], source_scope['subject_id'],
                    source_scope['brand_focus_id'], str(row.get('scope_evidence_ref') or ''))] = row
    for event in list(events or []):
        for row in event.get('sources', []) if isinstance(event, dict) else []:
            if isinstance(row, dict):
                source_type = row.get('source_type')
                source_id = row.get('source_id')
                source_scope = _source_use_scope(scope, row)
                unique[(source_type, source_id, source_scope['domain_id'], source_scope['subject_id'],
                        source_scope['brand_focus_id'], str(row.get('scope_evidence_ref') or ''))] = row
    sources_by_scope = {}
    source_bindings = []
    with get_repository().connection(read_only=True, wal=False) as con:
        for (_source_type, source_id, _domain, _subject_id, _brand_focus_id, _proof_ref), original_row in unique.items():
            if not isinstance(source_id, str) or not source_id:
                raise ValueError('REPORT_SOURCE_ID_UNAVAILABLE')
            row = dict(original_row)
            source_scope = _source_use_scope(scope, row)
            proof_ref = str(row.get('scope_evidence_ref') or '')
            focus_key = 'subject_id' if source_scope['domain_id'] == 'cth' else 'brand_focus_id'
            other_key = 'brand_focus_id' if focus_key == 'subject_id' else 'subject_id'
            candidates = con.execute(
                'SELECT raw_ref,captured_at FROM canonical_data '
                'WHERE product_id=? AND domain_id=? AND source_id=? AND source_type=? '
                'AND scope_status=\'VERIFIED\' AND scope_evidence_ref=? AND '
                f'{focus_key}=? AND COALESCE({other_key},\'\')=\'\'',
                ('CTH' if source_scope['domain_id'] == 'cth' else 'ZBMP', source_scope['domain_id'],
                 source_id, str(row.get('source_type') or _source_type or ''), proof_ref,
                 source_scope[focus_key]),
            ).fetchall()
            if len(candidates) != 1:
                raise ValueError('REPORT_SOURCE_CANONICAL_CHANGED')
            row['raw_ref'] = candidates[0][0]
            row['captured_at'] = row.get('captured_at') or candidates[0][1]
            body = row.get('content')
            if not isinstance(body, str) or not body:
                raise ValueError('REPORT_SOURCE_CONTENT_UNAVAILABLE')
            content_ref = _content_evidence_ref(con, row, proof_ref)
            if not content_ref:
                raise PermissionError('REPORT_SOURCE_USE_UNVERIFIED')
            if include_bindings:
                evidence_match = re.fullmatch(r'app_evidence_records:([1-9][0-9]*)', content_ref)
                if evidence_match is None:
                    raise ValueError('REPORT_SOURCE_EVIDENCE_BINDING_INVALID')
                source_bindings.append({
                    'source_type': str(row.get('source_type') or _source_type or ''),
                    'source_id': source_id,
                    'scope': source_scope,
                    'scope_evidence_ref': proof_ref,
                    'content_evidence_id': int(evidence_match.group(1)),
                    'content_evidence_ref': content_ref,
                })
            published = row.get('publish_datetime') or row.get('publish_date') or row.get('published_at')
            published = published.strip() if isinstance(published, str) and published.strip() else None
            captured = row.get('captured_at')
            captured = captured.strip() if isinstance(captured, str) and captured.strip() else None
            platform = row.get('source_platform') or row.get('source_adapter') or _source_type
            item = {
                'source_id': source_id, 'content_evidence_ref': content_ref, 'text': body,
                'body_sha256': hashlib.sha256(body.encode('utf-8')).hexdigest(),
                'text_selection': None, 'publish_datetime': published,
                'captured_at': captured, 'source_platform': platform,
            }
            key = (source_scope['domain_id'], source_scope['subject_id'], source_scope['brand_focus_id'])
            sources_by_scope.setdefault(key, (source_scope, []))[1].append(item)
        for source_scope, sources in sources_by_scope.values():
            for offset in range(0, len(sources), 8):
                _source_use_material({'scope': source_scope, 'sources': sources[offset:offset + 8]},
                                     actor, connection=con)
    source_count = len({(key[0], key[1]) for key in unique})
    if include_bindings:
        return {'source_count': source_count, 'bindings': source_bindings}
    return source_count


def _verify_saved_report_source_use(actor, scope, item):
    """Revalidate each immutable report source binding before releasing a file."""
    from lumenis.shared.database import get_repository
    from lumenis.shared.reporting.business_brief import _source_use_material
    from lumenis.shared.reporting.operational_brief import _content_evidence_ref
    from lumenis.shared.reporting.saved_work import _verified_evidence, authorize

    document = item.get('document') if isinstance(item, dict) else None
    payload = document.get('payload') if isinstance(document, dict) else None
    report = payload.get('report') if isinstance(payload, dict) else None
    bindings = report.get('source_use_bindings') if isinstance(report, dict) else None
    citations = document.get('citations') if isinstance(document, dict) else None
    if not isinstance(bindings, list) or not isinstance(citations, list):
        raise ValueError('REPORT_SOURCE_BINDING_UNAVAILABLE')
    if not bindings:
        if report.get('source_details') or citations:
            raise ValueError('REPORT_SOURCE_BINDING_UNAVAILABLE')
        return 0

    normalized = []
    seen = set()
    for binding in bindings:
        if (not isinstance(binding, dict)
                or set(binding) != {'source_type', 'source_id', 'scope', 'scope_evidence_ref',
                                    'content_evidence_id', 'content_evidence_ref'}
                or not isinstance(binding.get('source_type'), str)
                or not isinstance(binding.get('source_id'), str)
                or not isinstance(binding.get('scope'), dict)
                or not isinstance(binding.get('scope_evidence_ref'), str)
                or type(binding.get('content_evidence_id')) is not int
                or binding['content_evidence_id'] <= 0):
            raise ValueError('REPORT_SOURCE_BINDING_INVALID')
        source_scope = _source_use_scope(scope, binding['scope'])
        if source_scope != binding['scope']:
            raise ValueError('REPORT_SOURCE_SCOPE_MISMATCH')
        evidence_ref = f"app_evidence_records:{binding['content_evidence_id']}"
        if binding['content_evidence_ref'] != evidence_ref:
            raise ValueError('REPORT_SOURCE_EVIDENCE_BINDING_INVALID')
        key = (binding['source_type'], binding['source_id'], source_scope['domain_id'],
               source_scope['subject_id'], source_scope['brand_focus_id'],
               binding['scope_evidence_ref'], evidence_ref)
        if key in seen:
            raise ValueError('REPORT_SOURCE_BINDING_DUPLICATE')
        seen.add(key)
        normalized.append((binding, source_scope))

    citations_by_evidence = {}
    for citation in citations:
        if isinstance(citation, dict) and type(citation.get('evidence_id')) is int:
            citations_by_evidence.setdefault(citation['evidence_id'], []).append(citation)
    references = {entry[0]['content_evidence_id'] for entry in normalized}
    if not set(citations_by_evidence).issubset(references):
        raise ValueError('REPORT_CITATION_SOURCE_BINDING_MISMATCH')

    detail_bindings = set()
    details = list(report.get('source_details') or [])
    for event in report.get('event_details') or []:
        if isinstance(event, dict):
            details.extend(event.get('sources') or [])
    for detail in details:
        if not isinstance(detail, dict):
            raise ValueError('REPORT_SOURCE_DETAIL_INVALID')
        source_type = str(detail.get('source_type') or 'note')
        source_id = detail.get('source_id') or detail.get('note_id')
        source_scope = _source_use_scope(scope, detail)
        detail_bindings.add((source_type, source_id, source_scope['domain_id'],
                             source_scope['subject_id'], source_scope['brand_focus_id'],
                             str(detail.get('scope_evidence_ref') or '')))
    binding_details = {(binding['source_type'], binding['source_id'], source_scope['domain_id'],
                        source_scope['subject_id'], source_scope['brand_focus_id'],
                        binding['scope_evidence_ref'])
                       for binding, source_scope in normalized}
    if detail_bindings != binding_details:
        raise ValueError('REPORT_SOURCE_DETAIL_BINDING_MISMATCH')

    by_scope = {}
    for binding, source_scope in normalized:
        by_scope.setdefault(json.dumps(source_scope, sort_keys=True), (source_scope, []))[1].append(binding)
    verified_count = 0
    with get_repository().connection(read_only=True, wal=False) as con:
        for source_scope, scoped_bindings in by_scope.values():
            for offset in range(0, len(scoped_bindings), 500):
                chunk = scoped_bindings[offset:offset + 500]
                mentions = [{'source_type': binding['source_type'], 'source_id': binding['source_id']}
                            for binding in chunk]
                _normalized_refs, raw_rows = db._verified_event_source_refs(con, source_scope, mentions)
                for binding in chunk:
                    source_type, source_id = binding['source_type'], binding['source_id']
                    table = db._EVENT_SOURCE_TABLES.get(source_type)
                    if not table:
                        raise ValueError('REPORT_SOURCE_TYPE_INVALID')
                    row = dict(raw_rows[(table, source_type, source_id)])
                    proof_ref = str(row.get('scope_evidence_ref') or '')
                    if proof_ref != binding['scope_evidence_ref']:
                        raise ValueError('REPORT_SOURCE_CANONICAL_CHANGED')
                    focus_key = 'subject_id' if source_scope['domain_id'] == 'cth' else 'brand_focus_id'
                    other_key = 'brand_focus_id' if focus_key == 'subject_id' else 'subject_id'
                    candidates = con.execute(
                        'SELECT raw_ref,captured_at FROM canonical_data '
                        'WHERE product_id=? AND domain_id=? AND source_id=? AND source_type=? '
                        'AND scope_status=\'VERIFIED\' AND scope_evidence_ref=? AND '
                        f'{focus_key}=? AND COALESCE({other_key},\'\')=\'\'',
                        ('CTH' if source_scope['domain_id'] == 'cth' else 'ZBMP', source_scope['domain_id'],
                         source_id, source_type, proof_ref, source_scope[focus_key]),
                    ).fetchall()
                    if len(candidates) != 1:
                        raise ValueError('REPORT_SOURCE_CANONICAL_CHANGED')
                    row['raw_ref'] = candidates[0][0]
                    row['captured_at'] = row.get('captured_at') or candidates[0][1]
                    content_ref = _content_evidence_ref(con, row, proof_ref)
                    if content_ref != binding['content_evidence_ref']:
                        raise ValueError('REPORT_SOURCE_EVIDENCE_BINDING_CHANGED')
                    evidence_id = binding['content_evidence_id']
                    context = authorize(actor, source_scope, action='read', connection=con)
                    evidence = _verified_evidence(con, evidence_id, context, actor)
                    body = evidence.get('raw_content')
                    if not isinstance(body, str) or not body:
                        raise ValueError('REPORT_SOURCE_CONTENT_UNAVAILABLE')
                    published = row.get('publish_datetime') or row.get('publish_date') or row.get('published_at')
                    published = published.strip() if isinstance(published, str) and published.strip() else None
                    captured = row.get('captured_at')
                    captured = captured.strip() if isinstance(captured, str) and captured.strip() else None
                    platform = row.get('source_platform') or row.get('source_adapter') or source_type
                    source = {
                        'source_id': source_id, 'content_evidence_ref': content_ref,
                        'text': body, 'body_sha256': hashlib.sha256(body.encode('utf-8')).hexdigest(),
                        'text_selection': None, 'publish_datetime': published,
                        'captured_at': captured, 'source_platform': platform,
                    }
                    _source_use_material({'scope': source_scope, 'sources': [source]}, actor, connection=con)
                    verified_count += 1
    return verified_count


def _build_report(kind, start, end, *, actor, scope, scope_entry=None):
    import database as business_db
    from lumenis.shared.reporting import ReportGenerator, daily

    kind, start, end = _report_period(kind, start, end)
    if scope.get('global_scope') is True:
        if kind != 'daily' or start != end:
            raise ValueError('GLOBAL_REPORT_PERIOD_UNAVAILABLE')
        digest = daily.generate_daily_digest(end, scope=scope)
        note_read = business_db.get_scoped_event_source_rows(
            scope, source_types=('note', 'web_mention'), date_start=start, date_end=end, limit=5000)
        event_read = business_db.get_scoped_events(
            500, scope=scope, date_start=start, date_end=end, include_sources=True)
        if (note_read.get('truncated') or note_read.get('rejected_count')
                or event_read.get('truncated') or event_read.get('rejected_count')):
            raise ValueError('REPORT_SOURCE_COVERAGE_INCOMPLETE')
        notes, events = note_read['sources'], event_read['events']
        domain = 'cth' if scope.get('domain_id') == 'cth' else 'zbmp'
        report = {
            'report_id': f"{kind.upper()}-{datetime.now().strftime('%Y%m%d-%H%M%S')}",
            'brand_mode': domain, 'brand_name': 'CTH 全域' if domain == 'cth' else 'BMP 全域',
            'scope': dict(scope), 'period': kind, 'period_label': '日报',
            'start_date': start, 'end_date': end, 'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'executive_summary': digest.get('executive_summary') or digest.get('ai_summary') or '',
            'key_metrics': [
                {'name': '发现记录', 'value': digest.get('total_found'), 'unit': '条'},
                {'name': '有效记录', 'value': digest.get('effective_count'), 'unit': '条'},
                {'name': '正面记录', 'value': digest.get('positive_count'), 'unit': '条'},
                {'name': '中性记录', 'value': digest.get('neutral_count'), 'unit': '条'},
                {'name': '风险记录', 'value': digest.get('risk_count'), 'unit': '条'},
                {'name': '风险指数', 'value': digest.get('risk_index'), 'unit': '/100'},
                {'name': '品牌健康度', 'value': digest.get('brand_health_score'), 'unit': '/100'},
            ],
            'risk_events': digest.get('top_events') or [],
            'sentiment_analysis': {'overall_label': '由已核验范围统计', 'distribution': {
                'positive': digest.get('positive_count'), 'neutral': digest.get('neutral_count'),
                'negative': digest.get('negative_count')}},
            'key_insights': digest.get('key_insights') or [],
            'action_items': digest.get('action_items') or [],
            'data_quality': {'total_collected': digest.get('total_found'), 'quality_score': None,
                             'collection_complete': None, 'warnings': ['GLOBAL 日报只使用当前范围内已核验来源。']},
            'appendix': {'note_count': len(notes), 'event_count': len(events),
                         'note_ids': [n.get('source_id') or n.get('note_id') for n in notes[:50]],
                         'event_ids': [e.get('id') for e in events]},
            'cth_specific': {},
        }
    else:
        domain = 'cth' if scope.get('domain_id') == 'cth' else 'zbmp'
        generator = ReportGenerator(brand_mode=domain)
        notes = generator._get_notes_from_db(start, end, scope)
        events = generator._get_events_from_db(start, end, scope)
        report = generator.generate(kind, start_date=start, end_date=end, notes=notes, events=events,
                                    scope=scope, scope_entry=scope_entry)
    source_use = _verify_report_source_use(actor, scope, notes, events, include_bindings=True)
    verified_sources = source_use['source_count']
    _attach_report_details(report, notes, events)
    # Bind the saved report to the exact Canonical scope proof and original
    # content Evidence used at generation. Exports revalidate these bindings.
    report['source_use_bindings'] = source_use['bindings']
    return kind, start, end, report, verified_sources, notes, events


def _attach_report_details(report, notes, events):
    """Retain verified source/event detail rows in saved and downloadable reports."""
    source_fields = (
        'source_type', 'source_id', 'note_id', 'platform', 'source_platform', 'url',
        'product_id', 'domain_id', 'subject_id', 'brand_focus_id', 'scope_status', 'scope_evidence_ref',
        'title', 'publish_date', 'publish_datetime', 'publish_at', 'collected_at',
        'note_type', 'views', 'impressions', 'cover_ctr', 'likes', 'comments',
        'collects', 'shares', 'follower_gain', 'author', 'author_nickname',
        'sentiment', 'risk_level', 'relevance_status',
    )
    event_fields = (
        'id', 'event_key', 'title', 'risk_level', 'mention_count', 'event_date',
        'first_seen', 'last_seen', 'summary', 'risk_judgment', 'recommended_action', 'status',
    )

    def project(row, fields):
        return {key: row.get(key) for key in fields if key in row}

    if not isinstance(report, dict):
        raise TypeError('REPORT_RESULT_INVALID')
    from canonical_metrics import engagement_count
    report['source_details'] = []
    for row in notes:
        if not isinstance(row, dict):
            continue
        item = project(row, source_fields)
        values = [engagement_count(row.get(key)) for key in ('likes', 'comments', 'collects', 'shares')]
        item['engagement'] = sum(values) if all(value is not None for value in values) else None
        report['source_details'].append(item)
    report['event_details'] = []
    for row in events:
        if not isinstance(row, dict):
            continue
        item = project(row, event_fields)
        item['sources'] = [project(source, ('source_type', 'source_id', 'note_id', 'title', 'url',
                                             'product_id', 'domain_id', 'subject_id', 'brand_focus_id',
                                             'scope_status', 'scope_evidence_ref'))
                           for source in row.get('sources', []) if isinstance(source, dict)]
        report['event_details'].append(item)
    return report


def _spreadsheet_text(value):
    """Keep source-derived text literal in spreadsheet-compatible exports."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, sort_keys=True)
    leading = text.lstrip()
    if (leading.startswith(('-', '+')) and re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', leading)):
        return text
    if leading.startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + text
    return text


def _report_cells(report):
    """Flatten the full report without dropping source, risk, or appendix fields."""
    def walk(path, item):
        if isinstance(item, dict):
            if not item:
                yield path, '{}'
            for key, nested in item.items():
                yield from walk(f'{path}.{key}' if path else str(key), nested)
        elif isinstance(item, (list, tuple)):
            if not item:
                yield path, '[]'
            for index, nested in enumerate(item):
                yield from walk(f'{path}[{index}]', nested)
        else:
            yield path, item
    yield from walk('', report if isinstance(report, dict) else {'report': report})


def _report_section_name(key):
    labels = {
        'executive_summary': '执行摘要', 'executive_insight': '管理层洞察',
        'key_metrics': '关键指标', 'risk_events': '风险事件',
        'sentiment_analysis': '情感分析', 'propagation_analysis': '传播分析',
        'key_insights': '关键洞察', 'action_items': '行动建议',
        'data_quality': '数据质量', 'appendix': '附录', 'sources': '来源',
        'source_details': '来源明细', 'event_details': '事件明细',
        'cth_specific': 'CTH专题', 'zbmp_specific': 'BMP专题',
    }
    return labels.get(str(key), '报告信息')


def _report_table_rows(value):
    """Render an existing section as readable rows without discarding nested fields."""
    if isinstance(value, dict):
        return [(str(key), nested) for key, nested in value.items()]
    if isinstance(value, (list, tuple)):
        rows = []
        for index, item in enumerate(value, 1):
            if isinstance(item, dict):
                rows.extend((f'{index}.{key}', nested) for key, nested in item.items())
            else:
                rows.append((str(index), item))
        return rows
    return [('内容', value)]


def _write_excel_row(sheet, values):
    row_index = sheet.max_row + 1
    for column_index, value in enumerate(values, 1):
        cell = sheet.cell(row=row_index, column=column_index)
        if value is None:
            cell.value = '—'
            cell.data_type = 's'
        elif isinstance(value, (str, dict, list, tuple)):
            cell.value = _spreadsheet_text(value)
            cell.data_type = 's'
        else:
            cell.value = value


def _render_report(report, fmt):
    if fmt == 'docx':
        from lumenis.shared.reporting.word_pdf_export import ReportExporter
        with tempfile.TemporaryDirectory(prefix='lumenis-report-export-') as temp_dir:
            exporter = ReportExporter(output_dir=temp_dir)
            path = exporter.export_word(report, 'report.docx')
            if not path or not os.path.isfile(path):
                raise RuntimeError('REPORT_EXPORT_DEPENDENCY_UNAVAILABLE')
            with open(path, 'rb') as stream:
                return stream.read(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if fmt == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(['板块', '字段', '内容'])
        source_report = report if isinstance(report, dict) else {'report': report}
        for key, section in source_report.items():
            section_name = _report_section_name(key)
            for path, value in _report_cells({key: section}):
                field = path.removeprefix(f'{key}.')
                if value is None or isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                writer.writerow([_spreadsheet_text(section_name), _spreadsheet_text(field), _spreadsheet_text(value)])
        return output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8-sig'
    if fmt == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        full = book.active
        full.title = '完整报告'
        _write_excel_row(full, ('板块', '字段', '内容'))
        for path, value in _report_cells(report):
            if value is None or isinstance(value, (dict, list, tuple)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            text = value
            if isinstance(text, str) and len(text) > 30000:
                for offset in range(0, len(text), 30000):
                    root = path.split('.', 1)[0]
                    _write_excel_row(full, (_report_section_name(root), f'{path}[chunk:{offset // 30000 + 1}]', text[offset:offset + 30000]))
            else:
                root = path.split('.', 1)[0]
                _write_excel_row(full, (_report_section_name(root), path, text))
        source_report = report if isinstance(report, dict) else {'report': report}
        section_titles = {
            'key_metrics': '指标', 'risk_events': '风险事件',
            'key_insights': '关键洞察', 'action_items': '行动建议',
            'sources': '来源', 'source_details': '来源明细',
            'event_details': '事件明细', 'appendix': '附录',
        }
        for key, title in section_titles.items():
            if key not in source_report:
                continue
            sheet = book.create_sheet(title[:31])
            _write_excel_row(sheet, ('字段', '内容'))
            for label, value in _report_table_rows(source_report[key]):
                if isinstance(value, str) and len(value) > 30000:
                    for offset in range(0, len(value), 30000):
                        _write_excel_row(sheet, (f'{label}[chunk:{offset // 30000 + 1}]', value[offset:offset + 30000]))
                else:
                    _write_excel_row(sheet, (label, value))
        output = io.BytesIO()
        book.save(output)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    raise ValueError('REPORT_EXPORT_FORMAT_INVALID')


def export_artifacts(kind='daily', start=None, end=None, formats=('docx',), *, actor, scope,
                     scope_entry=None):
    """Build transient downloads from the authenticated report scope; nothing is saved here."""
    allowed = {'docx', 'csv', 'xlsx'}
    requested = tuple(dict.fromkeys(str(x).lower() for x in formats))
    if not requested or any(fmt not in allowed for fmt in requested):
        raise ValueError('REPORT_EXPORT_FORMAT_INVALID')
    import database as business_db
    exact_scope = business_db._normalize_event_scope(scope, allow_global=True)
    from lumenis.shared.reporting.saved_work import _check_effect
    from lumenis.system.security.permission import OperationType
    kind, start, end = _report_period(kind, start, end)
    request_binding = {'scope': exact_scope, 'period': kind, 'start': start, 'end': end,
                       'formats': requested}
    resource = 'report-export-request:v1:' + hashlib.sha256(
        json.dumps(request_binding, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    ).hexdigest()
    _check_effect(actor, exact_scope, OperationType.DATA_EXPORT, resource)
    kind, start, end, report, source_count, notes, events = _build_report(
        kind, start, end, actor=actor, scope=exact_scope, scope_entry=scope_entry)
    outputs = {}
    for fmt in requested:
        raw, media_type = _render_report(report, fmt)
        binding = {'report_id': report.get('report_id'), 'scope': exact_scope, 'period': kind,
                   'start': start, 'end': end, 'sha256': hashlib.sha256(raw).hexdigest(), 'format': fmt}
        resource = 'report-export:v1:' + hashlib.sha256(
            json.dumps(binding, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
        ).hexdigest()
        _check_effect(actor, exact_scope, OperationType.DATA_EXPORT, resource)
        # Recheck source bytes, rights, scope, and export permission after conversion.
        _verify_report_source_use(actor, exact_scope, notes, events)
        _check_effect(actor, exact_scope, OperationType.DATA_EXPORT, resource)
        filename = f'lumenis_{kind}_{start}_{end}.{fmt}'
        outputs[fmt] = {'content': raw, 'filename': filename, 'media_type': media_type,
                        'sha256': hashlib.sha256(raw).hexdigest(), 'saved': False}
    return {'files': outputs,
            'payload': {'report_type': kind, 'period_start': start, 'period_end': end,
                        'report': report, 'source_count': source_count,
                        'saved': False, 'persistence_state': 'DOWNLOAD_ONLY_NOT_SAVED'}}


# ============================================
# API接口
# ============================================

def _request_report_scope(action='read'):
    from lumenis.shared.reporting.routes import _authorized_report_scope
    return _authorized_report_scope(action=action)


def _request_range(args):
    kind = args.get('type') or args.get('period') or 'daily'
    start = args.get('start') or args.get('start_date')
    end = args.get('end') or args.get('end_date')
    if not start and not end:
        resolved_start, resolved_end = get_date_range(kind, args.get('date'))
        start, end = resolved_start.strftime('%Y-%m-%d'), resolved_end.strftime('%Y-%m-%d')
    return kind, start, end


def _build_request_export(args, formats):
    kind, start, end = _request_range(args)
    actor, scope, _entries, selected = _request_report_scope('read')
    return export_artifacts(kind, start, end, formats, actor=actor, scope=scope, scope_entry=selected)


def _download_result(result, fmt):
    artifact = result['files'][fmt]
    response = send_file(io.BytesIO(artifact['content']), as_attachment=True,
                         download_name=artifact['filename'], mimetype=artifact['media_type'], max_age=0)
    response.headers['X-Lumenis-Saved'] = 'false'
    response.headers['X-Lumenis-Persistence-State'] = 'DOWNLOAD_ONLY_NOT_SAVED'
    response.headers['X-Lumenis-Content-SHA256'] = artifact['sha256']
    return response


def _report_error(exc):
    raw_code = str(exc)
    if re.fullmatch(r'(?:REPORT|EVENT|WORK|SAVED_ARTIFACT|ARTIFACT|SOURCE|SCOPE|PERMISSION|DATA_EXPORT|CANONICAL)_[A-Z0-9_]{2,80}', raw_code):
        code = raw_code
    else:
        code = 'REPORT_EXPORT_FAILED'
        LOGGER.exception('Shared Reporting export request failed')
    status = 403 if isinstance(exc, PermissionError) else 409 if isinstance(exc, ValueError) else 500
    return jsonify({'ok': False, 'status': 'BLOCKED' if status == 403 else 'ERROR',
                    'code': code, 'error': code}), status


def _report_citations(notes, events):
    """Bind a bounded citation sample to real source-content Evidence."""
    from lumenis.shared.reporting.operational_brief import _content_evidence_ref
    citations = []
    seen = set()
    rows = list(notes or [])
    rows.extend(source for event in (events or []) if isinstance(event, dict)
                for source in (event.get('sources') or []) if isinstance(source, dict))
    with get_repository().connection(read_only=True, wal=False) as con:
        for row in rows:
            source_id = row.get('source_id') or row.get('note_id')
            body = row.get('content')
            if (not isinstance(source_id, str) or not source_id or not isinstance(body, str) or not body
                    or not row.get('scope_evidence_ref')):
                continue
            evidence_ref = _content_evidence_ref(con, row, str(row['scope_evidence_ref']))
            match = re.fullmatch(r'app_evidence_records:([1-9][0-9]*)', evidence_ref or '')
            if match is None:
                continue
            evidence_id = int(match.group(1))
            if evidence_id in seen:
                continue
            seen.add(evidence_id)
            quote = body[:min(160, len(body))]
            citations.append({'evidence_id': evidence_id, 'start': 0, 'end': len(quote), 'quote': quote})
            if len(citations) == 100:
                break
    return citations


def _saved_report_markdown(report):
    title = str(report.get('brand_name') or '业务报告').strip()
    period = f"{report.get('start_date') or '—'} 至 {report.get('end_date') or '—'}"
    executive = str(report.get('executive_summary') or '当前没有可用的执行摘要。')
    return (f"# {title}报告\n\n周期：{period}\n\n## 执行摘要\n\n{executive}\n\n"
            "## 完整结构化内容\n\n```json\n"
            + json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2)
            + "\n```\n")


@export_bp.route('/api/export/preview')
def export_preview():
    try:
        kind, start, end = _request_range(request.args)
        actor, scope, _entries, selected = _request_report_scope('read')
        _kind, _start, _end, report, _source_count, _notes, _events = _build_report(
            kind, start, end, actor=actor, scope=scope, scope_entry=selected)
        return jsonify({'ok': True, 'report': report,
                        'saved': False, 'persistence_state': 'PREVIEW_NOT_SAVED'})
    except Exception as exc:
        return _report_error(exc)


@export_bp.route('/api/export/excel')
def export_excel():
    try:
        return _download_result(_build_request_export(request.args, ('xlsx',)), 'xlsx')
    except Exception as exc:
        return _report_error(exc)


@export_bp.route('/api/export/word')
def export_word():
    try:
        return _download_result(_build_request_export(request.args, ('docx',)), 'docx')
    except Exception as exc:
        return _report_error(exc)


@export_bp.route('/api/export/list')
def export_list():
    try:
        from lumenis.shared.reporting.saved_work import search_work
        actor, scope, _entries, _selected = _request_report_scope('read')
        result = search_work(actor=actor, scope=scope, kind='report', limit=50)
        return jsonify({'ok': True, 'items': result['items'], 'has_more': result['has_more'],
                        'verification_complete': result['verification_complete'], 'saved': True})
    except Exception as exc:
        return _report_error(exc)


@export_bp.route('/api/export/report')
def export_unified():
    fmt = (request.args.get('format') or 'xlsx').lower()
    fmt = {'excel': 'xlsx', 'word': 'docx'}.get(fmt, fmt)
    if fmt not in {'docx', 'csv', 'xlsx'}:
        return jsonify({'ok': False, 'code': 'REPORT_EXPORT_FORMAT_INVALID'}), 400
    try:
        return _download_result(_build_request_export(request.args, (fmt,)), fmt)
    except Exception as exc:
        return _report_error(exc)


def _saved_report_attachment(item, fmt):
    """Read the exact rendition already committed as this SavedWork attachment."""
    from shared_artifacts import artifact_file_path

    document = item.get('document') if isinstance(item, dict) else None
    attachments = document.get('attachments') if isinstance(document, dict) else None
    expected_name = f'report.{fmt}'
    binding = next((row for row in attachments or []
                    if isinstance(row, dict) and row.get('name') == expected_name), None)
    if binding is None:
        raise ValueError('SAVED_REPORT_ATTACHMENT_UNAVAILABLE')
    expected_media = {
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    }.get(fmt)
    if binding.get('media_type') != expected_media:
        raise ValueError('SAVED_REPORT_ATTACHMENT_INTEGRITY_FAILED')
    path = artifact_file_path(binding.get('artifact_id'))
    if path is None:
        raise ValueError('SAVED_REPORT_ATTACHMENT_UNAVAILABLE')
    raw = path.read_bytes()
    if (len(raw) != binding.get('size_bytes')
            or hashlib.sha256(raw).hexdigest() != binding.get('sha256')):
        raise ValueError('SAVED_REPORT_ATTACHMENT_INTEGRITY_FAILED')
    return raw, binding.get('media_type')


@export_bp.route('/api/export/generate', methods=['POST', 'GET'])
def export_generate():
    fmt = (request.values.get('format') or 'xlsx').lower()
    fmt = {'excel': 'xlsx', 'word': 'docx'}.get(fmt, fmt)
    if fmt not in {'docx', 'xlsx'}:
        return jsonify({'ok': False, 'code': 'REPORT_EXPORT_FORMAT_INVALID'}), 400
    try:
        kind, start, end = _request_range(request.values)
        actor, scope, _entries, selected = _request_report_scope('write')
        kind, start, end, report, source_count, notes, events = _build_report(
            kind, start, end, actor=actor, scope=scope, scope_entry=selected)
        citations = _report_citations(notes, events)
        rendered, media_type = _render_report(report, fmt)
        attachment_name = f'report.{fmt}'
        from lumenis.shared.reporting.saved_work import (
            reopen_work,
            save_work,
            search_work,
        )
        request_id = 'report-' + __import__('uuid').uuid4().hex
        saved_scope = {key: scope.get(key) for key in
                       ('domain_id', 'subject_id', 'brand_focus_id', 'global_scope')}
        saved = save_work(
            title=f"{report.get('brand_name') or '业务'} {report.get('period_label') or kind}报告 {start}",
            markdown=_saved_report_markdown(report), kind='report',
            payload={'report': report, 'requested_format': fmt,
                     'report_source_count': source_count,
                     'citation_coverage': {'cited': len(citations), 'limit': 100,
                                           'state': 'SAMPLE_LIMIT_REACHED' if len(citations) == 100
                                           else 'SAMPLE_WITHIN_LIMIT'}},
            citations=citations, request_id=request_id, actor=actor, scope=saved_scope,
            attachments=[{'name': attachment_name, 'content': rendered, 'media_type': media_type}],
        )
        artifact_id = saved['artifact']['artifact_id']
        # The returned readback proves the generated report is committed and can
        # be reopened through SavedWork before exposing a download link.
        reopened = reopen_work(artifact_id, actor=actor, scope=saved_scope)
        if (reopened['artifact']['sha256'] != saved['artifact']['sha256']
                or reopened['document']['payload'].get('report') != report):
            raise ValueError('REPORT_SAVED_READBACK_MISMATCH')
        _verify_saved_report_source_use(actor, saved_scope, reopened)
        listed = search_work(actor=actor, scope=saved_scope, kind='report', query=report.get('brand_name') or '', limit=100)
        if not any(item['artifact']['artifact_id'] == artifact_id for item in listed['items']):
            raise ValueError('REPORT_SAVED_LIST_READBACK_MISMATCH')
        _verify_report_source_use(actor, db._normalize_event_scope(scope, allow_global=True), notes, events)
        reopened = reopen_work(artifact_id, actor=actor, scope=saved_scope)
        _verify_saved_report_source_use(actor, saved_scope, reopened)
        saved_render, saved_media_type = _saved_report_attachment(reopened, fmt)
        if saved_render != rendered or saved_media_type != media_type:
            raise ValueError('REPORT_SAVED_ATTACHMENT_READBACK_MISMATCH')
        from urllib.parse import quote
        download_url = f"/api/export/file/{quote(artifact_id)}.{fmt}"
        return jsonify({'ok': True, 'name': f'{artifact_id}.{fmt}', 'download_url': download_url,
                        'size': len(rendered), 'artifact_hash': hashlib.sha256(rendered).hexdigest(),
                        'artifact_id': artifact_id, 'work_id': saved['document']['work_id'],
                        'version': saved['document']['version'], 'sha256': saved['artifact']['sha256'],
                        'format': fmt, 'source_count': source_count, 'saved': True,
                        'persistence_state': 'SAVEDWORK_COMMITTED_AND_REOPENED'})
    except Exception as exc:
        return _report_error(exc)


@export_bp.route('/api/export/file/<path:filename>')
def export_file(filename):
    match = re.fullmatch(r'(ART-[0-9a-f]{16})\.(json|md|txt|docx|pdf|xlsx|csv)', filename)
    if match is None:
        return jsonify({'ok': False, 'code': 'SAVED_ARTIFACT_ID_REQUIRED'}), 400
    try:
        from lumenis.shared.reporting.saved_work import (
            _check_effect,
            export_work,
            reopen_work,
        )
        actor, scope, _entries, _selected = _request_report_scope('read')
        fmt = match.group(2)
        saved_scope = {key: scope.get(key) for key in
                       ('domain_id', 'subject_id', 'brand_focus_id', 'global_scope')}
        item = reopen_work(match.group(1), actor=actor, scope=saved_scope)
        is_report = item['document'].get('kind') == 'report'
        report = item['document'].get('payload', {}).get('report')
        if is_report:
            if not isinstance(report, dict):
                raise ValueError('SAVED_REPORT_FORMAT_UNAVAILABLE')
            _verify_saved_report_source_use(actor, saved_scope, item)
        if fmt in {'docx', 'xlsx', 'csv'}:
            if not is_report or not isinstance(report, dict):
                raise ValueError('SAVED_REPORT_FORMAT_UNAVAILABLE')
            resource = 'saved-report-rendition:v1:' + hashlib.sha256(
                f"{match.group(1)}:{item['artifact']['sha256']}:{fmt}".encode()).hexdigest()
            from lumenis.system.security.permission import OperationType
            _check_effect(actor, saved_scope, OperationType.DATA_EXPORT, resource)
            if fmt in {'docx', 'xlsx'}:
                raw, media_type = _saved_report_attachment(item, fmt)
            else:
                raw, media_type = _render_report(report, fmt)
            filename_out = f'{match.group(1)}.{fmt}'
            result = {'content': raw, 'filename': filename_out, 'media_type': media_type,
                      'sha256': hashlib.sha256(raw).hexdigest()}
        else:
            result = export_work(match.group(1), actor=actor, scope=saved_scope, format=fmt)
        reopened = reopen_work(match.group(1), actor=actor, scope=saved_scope)
        if is_report:
            _verify_saved_report_source_use(actor, saved_scope, reopened)
        response = send_file(io.BytesIO(result['content']), as_attachment=True,
                             download_name=result['filename'], mimetype=result['media_type'], max_age=0)
        response.headers['X-Lumenis-Saved'] = 'true'
        response.headers['X-Lumenis-Persistence-State'] = 'SAVED_REPORT_RENDITION' if fmt in {'docx', 'xlsx', 'csv'} else 'SAVEDWORK_ARTIFACT'
        response.headers['X-Lumenis-Content-SHA256'] = result['sha256']
        return response
    except Exception as exc:
        return _report_error(exc)


print("[Export] 导出报告模板模块已加载", flush=True)
