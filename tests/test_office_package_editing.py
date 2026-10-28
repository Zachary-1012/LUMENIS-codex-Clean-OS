"""Original Office bytes and isolated SavedWork regressions; no native-app claim."""
import hashlib
import io
import zipfile

import pytest
from lxml import etree

from lumenis.importer.office_package import OfficePackage
from lumenis.shared.reporting import saved_work
from tests import test_saved_document_workflow

office_workspace = test_saved_document_workflow.workspace

SCOPE = {'scope_kind': 'shared_workspace', 'version': 1}


def _package(raw, extension):
    return OfficePackage(raw, filename='source.' + extension,
                         media_type=saved_work._OFFICE_FORMATS[extension]['media_type'])


def _command(item, value, operation='replace_text'):
    return {'object_id': item['object_id'], 'operation': operation,
            'expected_value': item['value'], 'value': value}


def _parts(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _rewrite(raw, changes):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, changes.get(info.filename, source.read(info)))
        for name, content in changes.items():
            if name not in source.namelist():
                target.writestr(name, content)
    return output.getvalue()


def _docx():
    from docx import Document
    document = Document()
    document.add_heading('Original editable manuscript', 0)
    paragraph = document.add_paragraph()
    paragraph.add_run('Original target text').bold = True
    paragraph.add_run(' Adjacent italic text').italic = True
    document.add_paragraph('Numbered content stays', style='List Number')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Table heading'
    table.cell(1, 0).text = 'Table value stays'
    document.sections[0].header.paragraphs[0].text = 'Original header stays'
    document.sections[0].footer.paragraphs[0].text = 'Original footer stays'
    output = io.BytesIO()
    document.save(output)
    return _rewrite(output.getvalue(), {'customXml/retained-extension.xml': b'<privateExtension keep="true">opaque original</privateExtension>'})


def _pptx():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    presentation = Presentation()
    for index in range(80):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f'Original slide {index + 1}'
        slide.placeholders[1].text = f'Content on slide {index + 1} stays'
        slide.notes_slide.notes_text_frame.text = f'Speaker note {index + 1} stays'
    data = CategoryChartData()
    data.categories = ['A', 'B']
    data.add_series('Original chart', (12, 18))
    presentation.slides[3].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(2), Inches(4), Inches(3), data)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _xlsx():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Original data'
    sheet.append(['Original rich text', 12, '=SUM(B1:B2)'])
    sheet.append(['Same shared text', 18])
    sheet['A1'].font = Font(bold=True)
    sheet['D1'] = __import__('datetime').date(2026, 10, 24)
    sheet.merge_cells('A4:B4')
    sheet['A4'] = 'Merged region stays'
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
    sheet.add_chart(chart, 'F2')
    workbook.create_sheet('Second').append(['Other sheet stays'])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_docx_run_patch_preserves_adjacent_styles_numbering_tables_headers_and_unknown_parts():
    from docx import Document
    original = _docx()
    package = _package(original, 'docx')
    target = next(item for item in package.objects if item['value'] == 'Original target text')
    result = package.edit([_command(target, ' Revised professional writing ')], expected_package_sha256=package.sha256)
    assert result['receipt']['modified_parts'] == ['word/document.xml']
    before, after = _parts(original), _parts(result['content'])
    assert set(before) == set(after)
    assert all(before[name] == after[name] for name in before if name != 'word/document.xml')
    document = Document(io.BytesIO(result['content']))
    assert document.paragraphs[1].runs[0].text == ' Revised professional writing '
    assert document.paragraphs[1].runs[0].bold is True
    assert document.paragraphs[1].runs[1].italic is True
    assert document.tables[0].cell(1, 0).text == 'Table value stays'
    assert document.sections[0].header.paragraphs[0].text == 'Original header stays'
    assert result['receipt']['relationship_parts_unchanged']


def test_word_character_and_paragraph_format_edit_preserves_original_objects():
    from docx import Document
    package = _package(_docx(), 'docx')
    text = next(item for item in package.objects if item['value'] == 'Original target text')
    run = next(item for item in package.objects if item['kind'] == 'word_run_format'
               and item['locator']['node_path'] == text['locator']['node_path'].rsplit('/', 1)[0])
    paragraph_path = run['locator']['node_path'].rsplit('/', 1)[0]
    paragraph = next(item for item in package.objects if item['kind'] == 'word_paragraph_format'
                     and item['locator']['node_path'] == paragraph_path)
    result = package.edit([
        _command(run, {**run['value'], 'bold': False, 'italic': True, 'size_half_points': 36, 'color': '234A38'}, 'set_run_format'),
        _command(paragraph, {**paragraph['value'], 'alignment': 'center', 'space_after_twips': 240}, 'set_paragraph_format')],
        expected_package_sha256=package.sha256)
    doc = Document(io.BytesIO(result['content']))
    assert doc.paragraphs[1].runs[0].text == 'Original target text'
    assert doc.paragraphs[1].runs[0].bold is False
    assert doc.paragraphs[1].runs[0].italic is True
    assert doc.paragraphs[1].runs[0].font.size.pt == 18
    assert doc.paragraphs[1].paragraph_format.space_after.pt == 12
    assert doc.paragraphs[1].runs[1].italic is True
    assert result['receipt']['modified_parts'] == ['word/document.xml']
    assert _parts(result['content'])['word/numbering.xml'] == package.parts['word/numbering.xml']


def test_powerpoint_geometry_roundtrip_preserves_notes_masters_and_chart():
    from pptx import Presentation
    package = _package(_pptx(), 'pptx')
    shape = next(item for item in package.objects if item['kind'] == 'shape_geometry')
    value = {**shape['value'], 'x_emu': shape['value']['x_emu'] + 914400,
             'rotation_60000deg': 900000, 'flip_horizontal': True}
    result = package.edit([_command(shape, value, 'set_geometry')], expected_package_sha256=package.sha256)
    reopened = _package(result['content'], 'pptx')
    assert next(item for item in reopened.objects if item['object_id'] == shape['object_id'])['value'] == value
    assert len(Presentation(io.BytesIO(result['content'])).slides) == 80
    assert result['receipt']['modified_parts'] == [shape['locator']['part']]
    assert all(package.parts[name] == reopened.parts[name] for name in package.parts
               if name != shape['locator']['part'])


def test_eighty_slide_two_object_edit_preserves_other_slides_masters_notes_charts_and_embedded_workbook():
    from pptx import Presentation
    original = _pptx()
    package = _package(original, 'pptx')
    targets = [next(item for item in package.objects if item['value'] == f'Original slide {index}') for index in (1, 80)]
    result = package.edit([_command(targets[0], 'Edited opening'), _command(targets[1], 'Edited conclusion')],
                          expected_package_sha256=package.sha256)
    changed = {'ppt/slides/slide1.xml', 'ppt/slides/slide80.xml'}
    assert set(result['receipt']['modified_parts']) == changed
    before, after = _parts(original), _parts(result['content'])
    assert all(before[name] == after[name] for name in before if name not in changed)
    assert any('/charts/' in name for name in before) and any('/embeddings/' in name for name in before)
    document = Presentation(io.BytesIO(result['content']))
    assert len(document.slides) == 80
    assert document.slides[0].shapes.title.text == 'Edited opening'
    assert document.slides[79].shapes.title.text == 'Edited conclusion'
    assert document.slides[3].notes_slide.notes_text_frame.text == 'Speaker note 4 stays'


def test_business_locators_follow_workbook_and_slide_relationship_order_without_changing_object_identity():
    package = _package(_xlsx(), 'xlsx')
    first = next(item for item in package.objects if item['value'] == 'Original rich text')
    second = next(item for item in package.objects if item['value'] == 'Other sheet stays')
    assert first['locator']['sheet_name'] == 'Original data' and first['locator']['sheet_index'] == 1
    assert second['locator']['sheet_name'] == 'Second' and second['locator']['sheet_index'] == 2
    assert first['locator']['cell'] == second['locator']['cell'] == 'A1'
    raw = _pptx()
    parts = _parts(raw)
    root = etree.fromstring(parts['ppt/presentation.xml'])
    slides = root.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst')
    last = slides[-1]
    slides.remove(last)
    slides.insert(0, last)
    reordered = _package(_rewrite(raw, {'ppt/presentation.xml': etree.tostring(root)}), 'pptx')
    title = next(item for item in reordered.objects if item['value'] == 'Original slide 80')
    note = next(item for item in reordered.objects if item['value'] == 'Speaker note 80 stays')
    assert title['locator']['slide_number'] == note['locator']['slide_number'] == 1
    assert title['locator']['notes'] is False and note['locator']['notes'] is True
    assert title['locator']['paragraph_index'] == 1 and title['locator']['shape_name']
    old = next(item for item in _package(raw, 'pptx').objects if item['value'] == 'Original slide 80')
    assert title['object_id'] == old['object_id'] and title['part_sha256'] == old['part_sha256']
    word = next(item for item in _package(_docx(), 'docx').objects if item['value'] == 'Original target text')
    assert word['locator']['paragraph_index'] == 2


def test_xlsx_edits_preserve_charts_dates_merges_styles_and_other_sheet_and_invalidate_formula_cache():
    from openpyxl import load_workbook
    original = _xlsx()
    package = _package(original, 'xlsx')
    text = next(item for item in package.objects if item['value'] == 'Original rich text')
    number = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    formula = next(item for item in package.objects if item['kind'] == 'cell_formula')
    result = package.edit([_command(text, 'Revised text'), _command(number, '21.5', 'set_number'),
        _command(formula, '=SUM(B1:B2)*2', 'replace_formula')], expected_package_sha256=package.sha256)
    before, after = _parts(original), _parts(result['content'])
    changed = {'xl/worksheets/sheet1.xml', 'xl/workbook.xml'}
    assert set(result['receipt']['modified_parts']) == changed
    assert all(before[name] == after[name] for name in before if name not in changed)
    book = load_workbook(io.BytesIO(result['content']), data_only=False)
    assert book.active['A1'].value == 'Revised text' and book.active['A1'].font.bold
    assert book.active['B1'].value == 21.5 and book.active['C1'].value == '=SUM(B1:B2)*2'
    assert book.active['D1'].value.date().isoformat() == '2026-10-24'
    assert 'A4:B4' in book.active.merged_cells and len(book.active._charts) == 1
    assert book['Second']['A1'].value == 'Other sheet stays'
    assert result['receipt']['formula_evaluation'] == 'NATIVE_RECALCULATION_REQUIRED'
    assert load_workbook(io.BytesIO(result['content']), data_only=True).active['C1'].value is None


def _cached_xlsx(*, cross_sheet=False, chain=False, calc_properties=True):
    raw = _xlsx()
    parts = _parts(raw)
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    sheet = etree.fromstring(parts['xl/worksheets/sheet1.xml'])
    row = sheet.find(f'{ns}sheetData/{ns}row')
    boolean = etree.SubElement(row, ns + 'c', r='E1', t='b')
    etree.SubElement(boolean, ns + 'v').text = '1'
    formula = sheet.find(f'.//{ns}c[@r="C1"]')
    formula.find(ns + 'f').text = 'IF(E1,SUM(B1:B2)+LEN(A1),0)'
    formula.find(ns + 'v').text = '48'
    workbook = etree.fromstring(parts['xl/workbook.xml'])
    calc = workbook.find(ns + 'calcPr')
    if calc_properties:
        calc.attrib.update({'calcMode': 'manual', 'fullCalcOnLoad': '0', 'forceFullCalc': '0',
                            'calcOnSave': '0', 'calcCompleted': '1', 'iterate': '1', 'iterateCount': '7'})
    else:
        workbook.remove(calc)
        etree.SubElement(workbook, ns + 'extLst')
    changes = {'xl/workbook.xml': etree.tostring(workbook), 'xl/worksheets/sheet1.xml': etree.tostring(sheet)}
    if cross_sheet:
        second = etree.fromstring(parts['xl/worksheets/sheet2.xml'])
        row = second.find(f'{ns}sheetData/{ns}row')
        for reference, text, cached, attributes in [
            ('B1', "'Original data'!C1*2", '96', {}),
            ('C1', 'INDIRECT("\'Original data\'!B1")', '12', {}),
            ('D1', "'Original data'!B1*2", '24', {'t': 'shared', 'si': '0', 'ref': 'D1:E1'}),
            ('E1', None, '36', {'t': 'shared', 'si': '0'}),
            ('F1', "'Original data'!B1*{1,2;3,4}", '12', {'t': 'array', 'ref': 'F1:G2'}),
        ]:
            cell = etree.SubElement(row, ns + 'c', r=reference)
            etree.SubElement(cell, ns + 'f', **attributes).text = text
            etree.SubElement(cell, ns + 'v').text = cached
        follower = etree.SubElement(row, ns + 'c', r='G1')
        etree.SubElement(follower, ns + 'v').text = '24'
        second_row = etree.SubElement(second.find(ns + 'sheetData'), ns + 'row', r='2')
        for reference, cached in [('F2', '36'), ('G2', '48'), ('H2', '99')]:
            follower = etree.SubElement(second_row, ns + 'c', r=reference)
            etree.SubElement(follower, ns + 'v').text = cached
        changes['xl/worksheets/sheet2.xml'] = etree.tostring(second)
    if chain:
        ct = etree.fromstring(parts['[Content_Types].xml'])
        etree.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
            PartName='/xl/custom-calculation.xml',
            ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml')
        rels = etree.fromstring(parts['xl/_rels/workbook.xml.rels'])
        etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
            Id='calculation-chain', Target='custom-calculation.xml',
            Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain')
        changes.update({'[Content_Types].xml': etree.tostring(ct), 'xl/_rels/workbook.xml.rels': etree.tostring(rels),
                        'xl/custom-calculation.xml': f'<calcChain xmlns="{ns[1:-1]}"><c r="C1" i="1"/></calcChain>'.encode()})
    return _rewrite(raw, changes)


@pytest.mark.parametrize('kind,cell,value,operation', [
    ('cell_number', 'B1', '21.5', 'set_number'),
    ('cell_text_run', 'A1', 'Changed input writing', 'replace_text'),
    ('cell_boolean', 'E1', False, 'set_boolean'),
])
def test_any_xlsx_input_edit_invalidates_cached_formula_and_overrides_manual_calculation(kind, cell, value, operation):
    from openpyxl import load_workbook
    original = _cached_xlsx()
    assert load_workbook(io.BytesIO(original), data_only=True).active['C1'].value == 48
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == kind and item['locator']['cell'] == cell)
    result = package.edit([_command(target, value, operation)], expected_package_sha256=package.sha256)
    before, after = _parts(original), _parts(result['content'])
    assert set(result['receipt']['modified_parts']) == {'xl/workbook.xml', 'xl/worksheets/sheet1.xml'}
    assert all(before[name] == after[name] for name in before if name not in result['receipt']['modified_parts'])
    formula = load_workbook(io.BytesIO(result['content']), data_only=False).active['C1']
    assert formula.value == '=IF(E1,SUM(B1:B2)+LEN(A1),0)'
    assert load_workbook(io.BytesIO(result['content']), data_only=True).active['C1'].value is None
    calc = etree.fromstring(after['xl/workbook.xml']).find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr')
    assert calc.get('calcMode') == 'auto' and calc.get('calcCompleted') == '0'
    assert all(calc.get(key) == '1' for key in ('calcOnSave', 'fullCalcOnLoad', 'forceFullCalc'))
    assert calc.get('iterate') == '1' and calc.get('iterateCount') == '7'
    assert result['receipt']['formula_evaluation'] == result['projection']['formula_evaluation'] == 'NATIVE_RECALCULATION_REQUIRED'
    assert result['receipt']['native_application_acceptance'] == 'NOT_VERIFIED'


def test_cross_sheet_indirect_shared_and_array_results_are_invalidated_without_changing_formulas():
    original = _cached_xlsx(cross_sheet=True)
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    result = package.edit([_command(target, '21', 'set_number')], expected_package_sha256=package.sha256)
    before, after = _parts(original), _parts(result['content'])
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    old_sheet, new_sheet = (etree.fromstring(parts['xl/worksheets/sheet2.xml']) for parts in (before, after))
    assert [etree.tostring(node) for node in old_sheet.iter(ns + 'f')] == [etree.tostring(node) for node in new_sheet.iter(ns + 'f')]
    assert all(new_sheet.find(f'.//{ns}c[@r="{reference}"]/{ns}v') is None for reference in ('B1', 'C1', 'D1', 'E1', 'F1', 'G1', 'F2', 'G2'))
    assert new_sheet.find(f'.//{ns}c[@r="H2"]/{ns}v').text == '99'
    assert etree.tostring(old_sheet.find(f'.//{ns}c[@r="A1"]')) == etree.tostring(new_sheet.find(f'.//{ns}c[@r="A1"]'))
    assert after['xl/charts/chart1.xml'] == before['xl/charts/chart1.xml']
    assert result['receipt']['formula_cache_invalidation'] == 'WORKBOOK_WIDE_NO_DEPENDENCY_GRAPH'


@pytest.mark.parametrize('formula_type', ['array', 'dataTable'])
def test_grouped_formula_result_cells_are_read_only_and_rejected_before_any_edit(formula_type):
    raw = _cached_xlsx(cross_sheet=True)
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    part = 'xl/worksheets/sheet2.xml'
    if formula_type == 'dataTable':
        sheet = etree.fromstring(_parts(raw)[part])
        formula = sheet.find(f'.//{ns}c[@r="F1"]/{ns}f')
        formula.set('t', 'dataTable')
        formula.set('dt2D', '1')
        formula.set('r1', 'B1')
        formula.set('r2', 'B2')
        formula.text = None
        raw = _rewrite(raw, {part: etree.tostring(sheet)})
    package = _package(raw, 'xlsx')
    projection = package.projection()
    outputs = [item for item in projection['unsupported_objects']
               if item.get('reason') == 'GROUPED_FORMULA_RESULT' and item['part'] == part]
    assert {item['cell'] for item in outputs} == {'G1', 'F2', 'G2'}
    assert all(item['editable'] is False and item['state'] == 'PRESERVED_NOT_EDITABLE' for item in outputs)
    assert not any(item['locator']['part'] == part and item['locator'].get('cell') in {'F1', 'G1', 'F2', 'G2'}
                   for item in projection['objects'])
    normal = next(item for item in projection['objects']
                  if item['locator']['part'] == part and item['locator'].get('cell') == 'H2')
    for output in outputs:
        with pytest.raises(ValueError, match='OFFICE_EDIT_OPERATION_UNSUPPORTED'):
            package.edit([_command(normal, '123', 'set_number'), _command(output, '999', 'set_number')],
                         expected_package_sha256=package.sha256)
    assert package.raw == raw and package.parts == _parts(raw)
    edited = package.edit([_command(normal, '123', 'set_number')], expected_package_sha256=package.sha256)
    reopened = _package(edited['content'], 'xlsx')
    assert any(item['locator'].get('cell') == 'H2' and item['value'] == '123' for item in reopened.objects)
    before, after = _parts(raw), _parts(edited['content'])
    assert [etree.tostring(node) for node in etree.fromstring(before[part]).iter(ns + 'f')] == [
        etree.tostring(node) for node in etree.fromstring(after[part]).iter(ns + 'f')]
    assert after['xl/charts/chart1.xml'] == before['xl/charts/chart1.xml']
    assert edited['projection']['formula_evaluation'] == 'NATIVE_RECALCULATION_REQUIRED'


def test_saved_array_result_rejection_preserves_parent_and_normal_edit_download_reopens(office_workspace):
    ws = office_workspace
    raw = _cached_xlsx(cross_sheet=True)
    path = ws['root'] / 'private-array-results.xlsx'
    path.write_bytes(raw)
    first = ws['documents'].save_private_document(path, actor=ws['actor'], scope=SCOPE,
                                                 request_id='office-array-original-001')
    artifact_id = first['artifact']['artifact_id']
    projected = ws['documents'].read_office_document(artifact_id, actor=ws['actor'], scope=SCOPE, limit=500)
    output = next(item for item in projected['unsupported_objects'] if item.get('cell') == 'G1')
    normal = next(item for item in projected['objects'] if item['locator'].get('cell') == 'H2')
    args = {'actor': ws['actor'], 'scope': SCOPE, 'request_id': 'office-array-edit-001',
            'expected_sha256': first['artifact']['sha256'], 'expected_package_sha256': projected['package_sha256'],
            'reason': 'Correct a manually entered value while preserving calculated outputs.'}
    with pytest.raises(ValueError, match='OFFICE_EDIT_OPERATION_UNSUPPORTED'):
        ws['documents'].edit_office_document(artifact_id, **args, commands=[_command(output, '999', 'set_number')])
    assert ws['documents'].read_office_document(artifact_id, actor=ws['actor'], scope=SCOPE)['current_version'] is True
    original = ws['work'].download_attachment(artifact_id, name=first['document']['attachments'][0]['name'],
                                               actor=ws['actor'], scope=SCOPE)
    assert original['content'] == raw
    second = ws['documents'].edit_office_document(artifact_id, **args, commands=[_command(normal, '123', 'set_number')])
    assert second['document']['version'] == 2 and second['document']['parent_artifact_id'] == artifact_id
    assert second['document']['work_id'] == first['document']['work_id']
    downloaded = ws['work'].download_attachment(second['artifact']['artifact_id'],
        name=second['office_attachment']['name'], actor=ws['actor'], scope=SCOPE)
    assert any(item['locator'].get('cell') == 'H2' and item['value'] == '123'
               for item in _package(downloaded['content'], 'xlsx').objects)
    reopened = ws['documents'].read_office_document(second['artifact']['artifact_id'], actor=ws['actor'], scope=SCOPE)
    assert reopened['package_sha256'] == hashlib.sha256(downloaded['content']).hexdigest()
    assert any(item['locator'].get('cell') == 'H2' and item['value'] == '123' for item in reopened['objects'])
    assert ws['work'].download_attachment(artifact_id, name=original['filename'], actor=ws['actor'], scope=SCOPE)['content'] == raw


@pytest.mark.parametrize('kind', ['cell_number', 'cell_formula'])
def test_xlsx_calc_chain_is_removed_with_only_its_relationship_and_content_type(kind):
    original = _cached_xlsx(chain=True)
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == kind)
    operation, value = ('set_number', '22') if kind == 'cell_number' else ('replace_formula', 'B1+B2')
    result = package.edit([_command(target, value, operation)], expected_package_sha256=package.sha256)
    before, after = _parts(original), _parts(result['content'])
    removed = {'xl/custom-calculation.xml'}
    modified = {'xl/workbook.xml', 'xl/worksheets/sheet1.xml', '[Content_Types].xml', 'xl/_rels/workbook.xml.rels'}
    assert set(result['receipt']['removed_parts']) == removed
    assert set(result['receipt']['modified_parts']) == modified
    assert set(after) == set(before) - removed
    for name in before.keys() - modified - removed:
        assert before[name] == after[name]
    for name, excluded in [('xl/_rels/workbook.xml.rels', 'Type'), ('[Content_Types].xml', 'ContentType')]:
        old = etree.fromstring(before[name])
        new = etree.fromstring(after[name])
        expected = [dict(node.attrib) for node in old if not node.get(excluded, '').endswith(('/calcChain', 'spreadsheetml.calcChain+xml'))]
        assert [dict(node.attrib) for node in new] == expected
    assert result['receipt']['relationship_parts_unchanged'] is False
    assert _package(result['content'], 'xlsx').projection()['formula_evaluation'] == 'NATIVE_RECALCULATION_REQUIRED'


def test_missing_xlsx_calculation_properties_are_inserted_before_extensions():
    package = _package(_cached_xlsx(calc_properties=False), 'xlsx')
    target = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    result = package.edit([_command(target, '22', 'set_number')], expected_package_sha256=package.sha256)
    root = etree.fromstring(_parts(result['content'])['xl/workbook.xml'])
    assert [etree.QName(child).localname for child in root][-2:] == ['calcPr', 'extLst']


def test_xlsx_without_formulas_preserves_workbook_and_does_not_claim_required_recalculation():
    original = _xlsx()
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    sheet = etree.fromstring(_parts(original)['xl/worksheets/sheet1.xml'])
    constant = sheet.find(f'.//{ns}c[@r="C1"]')
    constant.remove(constant.find(ns + 'f'))
    constant.find(ns + 'v').text = '30'
    original = _rewrite(original, {'xl/worksheets/sheet1.xml': etree.tostring(sheet)})
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    result = package.edit([_command(target, '22', 'set_number')], expected_package_sha256=package.sha256)
    assert result['receipt']['modified_parts'] == ['xl/worksheets/sheet1.xml']
    assert result['receipt']['formula_evaluation'] == 'NOT_REQUIRED'
    assert result['receipt']['formula_cache_cells_invalidated'] == 0
    assert _parts(original)['xl/workbook.xml'] == _parts(result['content'])['xl/workbook.xml']


def test_calc_chain_with_default_content_type_is_removed_without_rewriting_content_types():
    original = _cached_xlsx()
    parts = _parts(original)
    ct = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Default',
        Extension='calculation', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml')
    rels = etree.fromstring(parts['xl/_rels/workbook.xml.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='calculation-chain', Target='custom.calculation',
        Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain')
    raw_types = etree.tostring(ct)
    original = _rewrite(original, {'[Content_Types].xml': raw_types,
        'xl/_rels/workbook.xml.rels': etree.tostring(rels),
        'xl/custom.calculation': b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="C1" i="1"/></calcChain>'})
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    result = package.edit([_command(target, '22', 'set_number')], expected_package_sha256=package.sha256)
    assert result['receipt']['removed_parts'] == ['xl/custom.calculation']
    assert _parts(result['content'])['[Content_Types].xml'] == raw_types
    assert '[Content_Types].xml' not in result['receipt']['modified_parts']


@pytest.mark.parametrize('mutation', ['wrong_root', 'other_incoming_relation', 'external_chain_relation'])
def test_invalid_calculation_chain_never_removes_unrelated_parts_or_follows_relationships(mutation):
    original = _cached_xlsx(chain=True)
    parts = _parts(original)
    if mutation == 'wrong_root':
        original = _rewrite(original, {'xl/custom-calculation.xml': b'<unrelated-original>keep</unrelated-original>'})
    else:
        rels = etree.fromstring(parts['xl/_rels/workbook.xml.rels'])
        if mutation == 'other_incoming_relation':
            etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                Id='unrelated-original', Target='custom-calculation.xml', Type='urn:original:retained')
        else:
            chain = next(node for node in rels if node.get('Id') == 'calculation-chain')
            chain.set('TargetMode', 'External')
            chain.set('Target', 'https://example.invalid/never-open')
        original = _rewrite(original, {'xl/_rels/workbook.xml.rels': etree.tostring(rels)})
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    with pytest.raises(ValueError, match='CALCULATION_CHAIN_INVALID'):
        package.edit([_command(target, '22', 'set_number')], expected_package_sha256=package.sha256)
    assert package.raw == original


def test_xlsx_external_hyperlink_relationship_is_retained_without_being_opened():
    original = _cached_xlsx(chain=True)
    parts = _parts(original)
    name = 'xl/worksheets/_rels/sheet1.xml.rels'
    rels = etree.fromstring(parts[name])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='retained-hyperlink', TargetMode='External', Target='https://example.invalid/never-open',
        Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink')
    raw_relationships = etree.tostring(rels)
    original = _rewrite(original, {name: raw_relationships})
    package = _package(original, 'xlsx')
    target = next(item for item in package.objects if item['kind'] == 'cell_number' and item['locator']['cell'] == 'B1')
    result = package.edit([_command(target, '22', 'set_number')], expected_package_sha256=package.sha256)
    assert _parts(result['content'])[name] == raw_relationships
    assert result['receipt']['external_relationships'] == [{'part': name, 'id': 'retained-hyperlink', 'state': 'PRESERVED_NOT_OPENED'}]


def test_shared_string_cell_edit_clones_only_its_rich_runs():
    original = _xlsx()
    parts = _parts(original)
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    sheet = etree.fromstring(parts['xl/worksheets/sheet1.xml'])
    for reference in ('A1', 'A2'):
        cell = sheet.find(f'.//{{{ns}}}c[@r="{reference}"]')
        for child in list(cell):
            cell.remove(child)
        cell.set('t', 's')
        etree.SubElement(cell, '{' + ns + '}v').text = '0'
    ct = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
        PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
    rels = etree.fromstring(parts['xl/_rels/workbook.xml.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='shared-string-relation', Target='sharedStrings.xml', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings')
    shared = f'<sst xmlns="{ns}" count="2" uniqueCount="1"><si><r><rPr><b/></rPr><t>Shared</t></r><r><t xml:space="preserve"> tail</t></r></si></sst>'.encode()
    raw = _rewrite(original, {'xl/worksheets/sheet1.xml': etree.tostring(sheet),
        '[Content_Types].xml': etree.tostring(ct), 'xl/_rels/workbook.xml.rels': etree.tostring(rels), 'xl/sharedStrings.xml': shared})
    package = _package(raw, 'xlsx')
    target = next(item for item in package.objects if item['locator'].get('cell') == 'A1' and item['value'] == 'Shared')
    edited = package.edit([_command(target, 'Specific')], expected_package_sha256=package.sha256)
    assert edited['receipt']['modified_parts'] == ['xl/workbook.xml', 'xl/worksheets/sheet1.xml']
    assert _parts(edited['content'])['xl/sharedStrings.xml'] == shared
    projected = _package(edited['content'], 'xlsx').objects
    assert any(item['value'] == 'Shared' and item['locator'].get('cell') == 'A2' for item in projected)
    assert any(item['value'] == 'Specific' and item['locator'].get('cell') == 'A1' for item in projected)


def test_unsupported_commands_stale_hashes_and_external_formulas_never_emit_a_package():
    original = _xlsx()
    package = _package(original, 'xlsx')
    formula = next(item for item in package.objects if item['kind'] == 'cell_formula')
    with pytest.raises(ValueError, match='PACKAGE_REVISION_CONFLICT'):
        package.edit([_command(formula, 'SUM(B1:B2)', 'replace_formula')], expected_package_sha256='0' * 64)
    with pytest.raises(ValueError, match='EXTERNAL_FORMULA_FORBIDDEN'):
        package.edit([_command(formula, 'WEBSERVICE("https://example.com")', 'replace_formula')], expected_package_sha256=package.sha256)
    with pytest.raises(ValueError, match='OBJECT_REVISION_CONFLICT'):
        package.edit([{**_command(formula, '1+2', 'replace_formula'), 'expected_value': 'old'}], expected_package_sha256=package.sha256)
    assert package.raw == original


def test_external_relationship_is_preserved_without_following_it_and_signature_requires_explicit_acknowledgement():
    original = _docx()
    parts = _parts(original)
    relationships = etree.fromstring(parts['word/_rels/document.xml.rels'])
    etree.SubElement(relationships, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='never-follow', Target='https://example.invalid/private', TargetMode='External',
        Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink')
    raw = _rewrite(original, {'word/_rels/document.xml.rels': etree.tostring(relationships),
        '_xmlsignatures/synthetic-signature.xml': b'<signature-fixture>not a real signature</signature-fixture>'})
    package = _package(raw, 'docx')
    target = next(item for item in package.objects if item['value'] == 'Original target text')
    assert package.projection()['external_relationships'][0]['state'] == 'PRESERVED_NOT_OPENED'
    with pytest.raises(ValueError, match='SIGNATURE_INVALIDATION_CONFIRMATION_REQUIRED'):
        package.edit([_command(target, 'Updated')], expected_package_sha256=package.sha256)
    result = package.edit([_command(target, 'Updated')], expected_package_sha256=package.sha256,
        acknowledge_signature_invalidation=True)
    assert result['receipt']['signature_state'] == 'INVALIDATED_ORIGINAL_SIGNATURE_PARTS_RETAINED'
    assert _parts(result['content'])['word/_rels/document.xml.rels'] == etree.tostring(relationships)
    assert _parts(result['content'])['_xmlsignatures/synthetic-signature.xml'] == b'<signature-fixture>not a real signature</signature-fixture>'


@pytest.mark.parametrize('signal', ['both', 'relationships', 'override', 'default'])
def test_signature_identity_follows_opc_metadata_outside_the_conventional_directory(signal):
    from urllib.parse import quote

    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    original = _docx()
    parts = _parts(original)
    origin = 'customSecurity/Origin.sigs'
    signature = 'customSecurity/Signer One.sigdata'
    certificate = 'customSecurity/Signer.sigcert'
    changes = {origin: b'', signature: b'<Signature xmlns="http://www.w3.org/2000/09/xmldsig#">synthetic metadata fixture, not a verified signature</Signature>',
               certificate: b'synthetic certificate bytes, never a real credential'}
    namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    if signal in {'both', 'relationships'}:
        package_rels = etree.fromstring(parts['_rels/.rels'])
        etree.SubElement(package_rels, namespace + 'Relationship',
            Id='custom-signature-origin', Type=RT.ORIGIN, Target=origin)
        origin_rels = etree.Element(namespace + 'Relationships', nsmap={None: namespace[1:-1]})
        etree.SubElement(origin_rels, namespace + 'Relationship',
            Id='signature', Type=RT.SIGNATURE, Target=quote('Signer One.sigdata'))
        signature_rels = etree.Element(namespace + 'Relationships', nsmap={None: namespace[1:-1]})
        etree.SubElement(signature_rels, namespace + 'Relationship',
            Id='certificate', Type=RT.CERTIFICATE, Target='Signer.sigcert')
        changes.update({'_rels/.rels': etree.tostring(package_rels),
            'customSecurity/_rels/Origin.sigs.rels': etree.tostring(origin_rels),
            'customSecurity/_rels/Signer One.sigdata.rels': etree.tostring(signature_rels)})
    if signal in {'both', 'override', 'default'}:
        types = etree.fromstring(parts['[Content_Types].xml'])
        ct_namespace = '{http://schemas.openxmlformats.org/package/2006/content-types}'
        for part, content_type in [(origin, CT.OPC_DIGITAL_SIGNATURE_ORIGIN),
                (signature, CT.OPC_DIGITAL_SIGNATURE_XMLSIGNATURE), (certificate, CT.OPC_DIGITAL_SIGNATURE_CERTIFICATE)]:
            attributes = {'PartName': '/' + quote(part)} if signal != 'default' else {'Extension': part.rsplit('.', 1)[-1]}
            etree.SubElement(types, ct_namespace + ('Override' if signal != 'default' else 'Default'),
                             ContentType=content_type, **attributes)
        changes['[Content_Types].xml'] = etree.tostring(types)
    raw = _rewrite(original, changes)
    package = _package(raw, 'docx')
    assert package.projection()['signature_parts'] == sorted([origin, signature, certificate])
    target = next(item for item in package.objects if item['value'] == 'Original target text')
    with pytest.raises(ValueError, match='SIGNATURE_INVALIDATION_CONFIRMATION_REQUIRED'):
        package.edit([_command(target, 'Explicitly confirmed text')], expected_package_sha256=package.sha256)
    result = package.edit([_command(target, 'Explicitly confirmed text')], expected_package_sha256=package.sha256,
                          acknowledge_signature_invalidation=True)
    before, after = _parts(raw), _parts(result['content'])
    assert all(before[name] == after[name] for name in before if name != 'word/document.xml')
    assert result['receipt']['signature_state'] == 'INVALIDATED_ORIGINAL_SIGNATURE_PARTS_RETAINED'
    assert result['receipt']['active_content_execution'] == 'NOT_RUN'
    assert result['receipt']['native_application_acceptance'] == 'NOT_VERIFIED'


def test_unsigned_custom_parts_do_not_claim_a_signature_or_require_acknowledgement():
    original = _docx()
    parts = _parts(original)
    rels = etree.fromstring(parts['_rels/.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='ordinary-custom', Type='https://example.invalid/digital-signature/signature', Target='customSecurity/ordinary.xml')
    raw = _rewrite(original, {'_rels/.rels': etree.tostring(rels),
        'customSecurity/ordinary.xml': b'<metadata>digital-signature is ordinary text here</metadata>'})
    package = _package(raw, 'docx')
    assert package.projection()['signature_parts'] == []
    target = next(item for item in package.objects if item['value'] == 'Original target text')
    result = package.edit([_command(target, 'Ordinary permitted edit')], expected_package_sha256=package.sha256)
    assert result['receipt']['signature_state'] == 'NOT_PRESENT'
    assert _parts(result['content'])['customSecurity/ordinary.xml'] == _parts(raw)['customSecurity/ordinary.xml']


@pytest.mark.parametrize('mutation', ['zip_path', 'missing_relation', 'entity', 'active_content', 'oversized_member'])
def test_package_path_relationship_entity_active_content_and_expansion_guards(mutation):
    raw = _docx()
    if mutation == 'zip_path':
        raw = _rewrite(raw, {'../escape.xml': b'<escape/>'})
    elif mutation == 'missing_relation':
        raw = _rewrite(raw, {'word/_rels/document.xml.rels': b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="bad" Target="../../outside.xml" Type="test"/></Relationships>'})
    elif mutation == 'entity':
        raw = _rewrite(raw, {'word/document.xml': b'<!DOCTYPE document [<!ENTITY ex SYSTEM "file:///forbidden">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">&ex;</w:document>'})
    elif mutation == 'active_content':
        raw = _rewrite(raw, {'word/vbaProject.bin': b'not executed'})
    else:
        raw = _rewrite(raw, {'customXml/huge.xml': b'x' * (8 * 1024 * 1024 + 1)})
    if mutation == 'active_content':
        package = _package(raw, 'docx')
        with pytest.raises(ValueError, match='ACTIVE_CONTENT_EDIT_UNSUPPORTED'):
            package.edit([_command(package.objects[0], 'changed')], expected_package_sha256=package.sha256)
    else:
        with pytest.raises(ValueError, match='OFFICE_'):
            _package(raw, 'docx')


def test_saved_office_edit_uses_parent_cas_reopens_exact_bytes_preserves_original_and_is_idempotent(office_workspace):
    ws = office_workspace
    raw = _docx()
    path = ws['root'] / 'private-original.docx'
    path.write_bytes(raw)
    first = ws['documents'].save_private_document(path, actor=ws['actor'], scope=SCOPE, request_id='office-original-001')
    source_id = first['artifact']['artifact_id']
    projected = ws['documents'].read_office_document(source_id, actor=ws['actor'], scope=SCOPE, limit=2)
    assert projected['has_more'] and projected['next_offset'] == 2
    all_objects = ws['documents'].read_office_document(source_id, actor=ws['actor'], scope=SCOPE, limit=500)
    target = next(item for item in all_objects['objects'] if item['value'] == 'Original target text')
    args = {'actor': ws['actor'], 'scope': SCOPE, 'request_id': 'office-edit-0001', 'reason': 'Revise the selected writing',
        'expected_sha256': first['artifact']['sha256'], 'expected_package_sha256': hashlib.sha256(raw).hexdigest(),
        'commands': [_command(target, 'Saved local object edit')]}
    second = ws['documents'].edit_office_document(source_id, **args)
    assert second['document']['version'] == 2 and second['document']['parent_artifact_id'] == source_id
    assert second['document']['work_id'] == first['document']['work_id']
    reopened = ws['documents'].read_office_document(second['artifact']['artifact_id'], actor=ws['actor'], scope=SCOPE)
    assert any(item['value'] == 'Saved local object edit' for item in reopened['objects'])
    original = ws['work'].download_attachment(source_id, name=first['document']['attachments'][0]['name'], actor=ws['actor'], scope=SCOPE)
    assert original['content'] == raw
    replay = ws['documents'].edit_office_document(source_id, **args)
    assert replay['artifact']['artifact_id'] == second['artifact']['artifact_id'] and replay['replayed']
    with pytest.raises(saved_work.SavedWorkError, match='WORK_REVISION_CONFLICT'):
        ws['documents'].edit_office_document(source_id, **{**args, 'request_id': 'office-stale-0002'})
    ws['app'].config['LUMENIS_OPEN_WORKSPACE'] = False
    with pytest.raises(saved_work.SavedWorkError, match='WORK_PERMISSION_DENIED'):
        ws['documents'].read_office_document(second['artifact']['artifact_id'], actor=ws['actor'], scope=SCOPE)
