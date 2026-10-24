"""PPTX text-export regressions; library readback is not Office-app acceptance."""
import builtins
import io
import json
import os
import socket
import sqlite3
import zipfile

import pytest

from lumenis.shared.reporting.word_pdf_export import (
    ReportExporter,
    delivery_markdown,
)


def snapshot():
    long_line = "长文本片段🙂 / 中文 / café / tab\t保留。" * 220
    return {
        "schema": "lumenis.saved-work/v1",
        "title": "标题 / Research / 자료 😀\t完整保留",
        "markdown": (
            "首段：来源和范围保留。\n"
            + long_line
            + "\n한국어와中文段落、emoji 🧭、literal <tag> & 原样。"
        ),
        "kind": "research",
        "payload": {"saved_marker": "SNAPSHOT-COMPLETE-7", "long_text": long_line},
        "citations": [],
        "context": {
            "domain_id": "cth",
            "subject_id": "subject-7",
            "brand_focus_id": None,
            "global_scope": False,
            "actor_id": 19,
        },
        "work_id": "WORK-0123456789abcdef01234567",
        "version": 3,
        "request_id": "pptx-export-0001",
        "parent_artifact_id": "ART-0123456789abcdef",
    }


def source_rows():
    return [{
        "evidence_id": 71,
        "source_id": "SRC-PPTX-7",
        "source_url": "https://example.invalid/source/中文?q=editable",
        "scope": {"domain_id": "cth", "subject_id": "subject-7",
                  "brand_focus_id": None, "global_scope": False},
        "rights": "owner-approved test text; no product-fact claim",
        "quote": "引用内容：完整来源文字 😀",
    }]


def _read_sections(raw):
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(raw))
    content = {}
    for slide in presentation.slides:
        assert len(slide.shapes) == 2
        heading, body = slide.shapes[0], slide.shapes[1]
        heading_text = heading.text_frame.text
        assert heading_text.startswith("Saved document | ")
        label = heading_text.removeprefix("Saved document | ").split(" (continued ", 1)[0]
        content.setdefault(label, []).extend(p.text for p in body.text_frame.paragraphs)
    return presentation, content


def _expected_paragraphs(text):
    lines = text.splitlines()
    if not lines:
        lines = [""]
    if text.endswith(("\n", "\r", "\v", "\f", "\x1c", "\x1d", "\x1e", "\x85", "\u2028", "\u2029")):
        lines.append("")
    result = []
    for line in lines:
        if not line:
            result.append("")
        else:
            result.extend(line[start:start + 64] for start in range(0, len(line), 64))
    return result or [""]


def test_saved_work_pptx_is_editable_and_preserves_every_section_in_memory(
    monkeypatch, tmp_path,
):
    import pptx

    from lumenis.shared import database

    assert pptx.__version__ == "1.0.2"

    def forbidden(*args, **kwargs):
        raise AssertionError("PPTX memory export must not access a path, database, or network")

    original_open = builtins.open
    original_io_open = io.open

    def no_path_write(open_fn):
        def guarded(path, mode="r", *args, **kwargs):
            if any(flag in mode for flag in ("w", "a", "x", "+")):
                forbidden(path, mode, *args, **kwargs)
            return open_fn(path, mode, *args, **kwargs)
        return guarded

    monkeypatch.setattr(os, "makedirs", forbidden)
    monkeypatch.setattr(builtins, "open", no_path_write(original_open))
    monkeypatch.setattr(io, "open", no_path_write(original_io_open))
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    monkeypatch.setattr(database, "get_repository", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)

    document, sources = snapshot(), source_rows()
    before = set(tmp_path.iterdir())
    raw = ReportExporter.export_saved_document(document, sources=sources, format="pptx")
    assert set(tmp_path.iterdir()) == before
    assert raw.startswith(b"PK")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert archive.testzip() is None
        assert "ppt/presentation.xml" in archive.namelist()

    presentation, exported = _read_sections(raw)
    assert len(presentation.slides) > 1
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert 0 <= shape.left < presentation.slide_width
            assert 0 <= shape.top < presentation.slide_height
            assert shape.left + shape.width <= presentation.slide_width
            assert shape.top + shape.height <= presentation.slide_height
    expected = dict(zip(
        ("Title", "Report", "Lineage", "Sources", "Snapshot", "JSON"),
        (document["title"], document["markdown"],
         "Sources / original scope / rights",
         json.dumps(sources, ensure_ascii=False, indent=2),
         "Complete saved snapshot",
         json.dumps(document, ensure_ascii=False, sort_keys=True, indent=2)),
    ))
    assert set(exported) == set(expected)
    for label, text in expected.items():
        assert exported[label] == _expected_paragraphs(text)
    assert "标题 / Research / 자료 😀\t完整保留" in "".join(exported["Title"])
    assert "Sources / original scope / rights" in "".join(exported["Lineage"])
    assert "https://example.invalid/source/中文?q=editable" in "".join(exported["Sources"])
    assert "SNAPSHOT-COMPLETE-7" in "".join(exported["JSON"])
    assert any(shape.has_text_frame and shape.text_frame.text
               for slide in presentation.slides for shape in slide.shapes)


def test_saved_work_pptx_rejects_oversized_snapshot_without_partial_bytes():
    document = snapshot()
    document["markdown"] = "超大输出" * 250_000
    with pytest.raises(ValueError, match="^WORK_EXPORT_PPTX_TOO_LARGE$"):
        ReportExporter.export_saved_document(document, sources=source_rows(), format="pptx")


def test_saved_work_owner_exports_pptx_with_rechecks_and_content_digest(monkeypatch):
    import hashlib

    from lumenis.shared.reporting import saved_work
    from lumenis.system.security import permission

    artifact_id = "ART-0123456789abcdef"
    item = {
        "artifact": {"artifact_id": artifact_id, "sha256": "a" * 64},
        "content": b"saved-work-content",
        "document": snapshot(),
        "receipt": {"sources": source_rows()},
        "return_ref": {"scope": "existing-shared-workspace"},
    }
    reads, effects = [], []

    def read_committed_work(actual_id, *, actor, scope):
        reads.append((actual_id, actor, scope))
        return item

    def check_effect(actor, scope, operation, resource):
        effects.append((actor, scope, operation, resource))

    monkeypatch.setattr(saved_work, "read_committed_work", read_committed_work)
    monkeypatch.setattr(saved_work, "_check_effect", check_effect)
    actor, scope = object(), {"scope": "existing"}
    result = saved_work.export_work(artifact_id, actor=actor, scope=scope, format="pptx")

    assert result["filename"] == f"{artifact_id}.pptx"
    assert result["media_type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert result["sha256"] == hashlib.sha256(result["content"]).hexdigest()
    assert result["return_ref"] == item["return_ref"]
    assert len(reads) == 2
    assert all(read == (artifact_id, actor, scope) for read in reads)
    assert len(effects) == 2
    assert all(effect[:3] == (actor, scope, permission.OperationType.DATA_EXPORT) for effect in effects)
    assert effects[0][3] == effects[1][3]
    _read_sections(result["content"])


def test_saved_work_export_still_rejects_unsupported_format():
    with pytest.raises(ValueError, match="^unsupported saved document export format$"):
        ReportExporter.export_saved_document(snapshot(), sources=source_rows(), format="keynote")


def test_delivery_quotes_match_only_revalidated_source_identity():
    doc = {"title": "报告", "markdown": "正文", "citations": [
        {"evidence_id": 71, "quote": "获准原文"},
        {"evidence_id": 72, "quote": "无当前来源不得输出"},
    ]}
    result = delivery_markdown(doc, source_rows())
    assert "获准原文" in result
    assert "无当前来源不得输出" not in result
    assert "evidence_id" not in result and "source_scope" not in result


@pytest.mark.parametrize("format", ["docx", "xlsx", "pptx", "pdf"])
def test_delivery_retains_user_content_without_internal_snapshot(format):
    document = {
        "title": "季度复盘 / Review / 검토",
        "markdown": "## Findings\nConfirmed wording.\n\n| Metric | Value |\n| --- | --- |\n| Sales | 125 |\n| Missing | — |\n| Formula | =1+2 |\n\n## Next action\nKeep original records.",
        "payload": {"internal_marker": "PRIVATE-AUDIT-ONLY"},
        "context": {"actor_id": 712},
    }
    sources = [{"title": "Reviewed source", "source_url": "https://example.invalid/research",
                "quote": "A quoted passage.", "scope": {"private_marker": "INTERNAL-SCOPE-ONLY"}}]
    raw = ReportExporter.export_saved_document(document, sources=sources, format=format, purpose="delivery")
    if format == "docx":
        from docx import Document
        from docx.oxml.ns import qn
        result = Document(io.BytesIO(raw))
        assert result.styles["Normal"].font.name == "Arial"
        assert result.styles["Heading 2"].element.rPr.rFonts.get(qn("w:eastAsia")) == "Microsoft YaHei"
        font_table = next(part for part in result.part.package.iter_parts()
                          if str(part.partname) == "/word/fontTable.xml")
        assert b'Courier' not in font_table.blob
        assert any(p.style.name == "Heading 2" and p.text == "Findings" for p in result.paragraphs)
        assert [[cell.text for cell in row.cells] for row in result.tables[0].rows] == [
            ["Metric", "Value"], ["Sales", "125"], ["Missing", "—"], ["Formula", "=1+2"]]
        text = "\n".join(p.text for p in result.paragraphs)
    elif format == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw))
        assert book.sheetnames == ["Report", "Table 1"]
        assert book["Table 1"]["B4"].data_type == "s"
        assert book["Table 1"]["B4"].value == "=1+2"
        text = "\n".join(str(cell.value or "") for sheet in book for row in sheet for cell in row)
    elif format == "pptx":
        from pptx import Presentation
        result = Presentation(io.BytesIO(raw))
        text = "\n".join(shape.text for slide in result.slides for shape in slide.shapes if shape.has_text_frame)
        assert "Metric | Value" in text
    else:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        assert "Metric" in text and "125" in text
    assert "Confirmed wording." in text and "Keep original records." in text
    assert "https://example.invalid/research" in text and "A quoted passage." in text
    assert "PRIVATE-AUDIT-ONLY" not in text and "INTERNAL-SCOPE-ONLY" not in text
    assert "Complete saved snapshot" not in text


def test_delivery_purpose_is_part_of_effect_identity_and_rechecked(monkeypatch):
    from lumenis.shared.reporting import saved_work

    item = {"artifact": {"sha256": "a" * 64}, "document": snapshot(),
            "receipt": {"sources": []}, "return_ref": {"version": 3}}
    reads, resources = [], []
    def read(*args, **kwargs):
        reads.append(True)
        return item
    monkeypatch.setattr(saved_work, "read_committed_work", read)
    monkeypatch.setattr(saved_work, "_check_effect", lambda actor, scope, operation, resource: resources.append(resource))
    common = {"artifact_id": "ART-0123456789abcdef", "actor": None, "scope": {}, "format": "md"}
    delivered = saved_work.export_work(**common, purpose="delivery")
    assert delivered["content"] == item["document"]["markdown"].encode()
    saved_work.export_work(**common, purpose="audit_snapshot")
    assert len(reads) == 4
    assert resources[0] == resources[1] and resources[2] == resources[3]
    assert resources[0] != resources[2]
    with pytest.raises(saved_work.SavedWorkError, match="WORK_EXPORT_PURPOSE_INVALID"):
        saved_work.export_work(**common, purpose="ignore_permission")
    assert len(reads) == 4
