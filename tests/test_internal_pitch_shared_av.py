"""Shared continuation: actual owners, no Subject tables, isolated account/local HTTP."""
import copy
import hashlib
import io
import json
from pathlib import Path
import threading
import urllib.request
import urllib.error

import pytest
from docx import Document
from flask import Flask
from werkzeug.serving import make_server

from lumenis.shared.database import DatabaseRepository
from lumenis.shared.content_creation import internal_pitch as pitch
from lumenis.shared.reporting import saved_work as work
from lumenis.system.security import permission

from tests import test_internal_pitch_workflow as helpers
SHARED = {'scope_kind': 'shared_workspace', 'version': 1}
PREFIX = helpers.PREFIX


def account(ws, grants=('read', 'write'), product='SHARED'):
    with ws['repo'].transaction() as con:
        n = con.execute('SELECT COUNT(*) FROM v4_users').fetchone()[0] + 1
        actor_id = con.execute("INSERT INTO v4_users(username,password_hash,role,display_name,status) VALUES(?,?,'user',?,'active')", (f'shared-{n}', 'synthetic-unused-password-hash', f'Shared {n}')).lastrowid
        for action in grants:
            con.execute('INSERT INTO v4_member_authorizations(user_id,product_id,module,action,data_scope,time_range,granted_by) VALUES(?,?,?,?,NULL,NULL,NULL)', (actor_id, product, 'reports', action))
        token = f'synthetic-shared-session-{actor_id}'
        con.execute("INSERT INTO v4_sessions(user_id,token,expires_at) VALUES(?,?,'2099-01-01 00:00:00')", (actor_id, token))
    return {'id': actor_id}, {'Authorization': 'Bearer ' + token}


@pytest.fixture(params=['local', 'account'])
def shared(request, tmp_path, monkeypatch):
    import ai_task_engine, shared_artifacts
    import account.models as accounts
    from lumenis.auth import models as auth_models, service as auth
    path = tmp_path/'shared-av.sqlite3'
    monkeypatch.setenv('ZBC_DB_PATH', str(path))
    monkeypatch.setenv('ZBC_ARTIFACT_DIR', str(tmp_path/'artifacts'))
    monkeypatch.delenv('PERMISSION_GATE_TESTING', raising=False)
    monkeypatch.delenv('LUMENIS_BOOTSTRAP_OWNER_USERNAME', raising=False)
    monkeypatch.delenv('LUMENIS_BOOTSTRAP_OWNER_PASSWORD', raising=False)
    monkeypatch.setattr(permission, '_TESTING_BYPASS_ENABLED', False)
    repo = DatabaseRepository(str(path))
    for module in (ai_task_engine, shared_artifacts, accounts, auth_models, work, pitch):
        monkeypatch.setattr(module, 'get_repository', lambda: repo)
    # Auth owner performs actual token/session lookup against this isolated DB.
    monkeypatch.setattr(auth, 'get_db', lambda: repo.connect())
    accounts.init_account_db()
    monkeypatch.setattr(auth, '_initialized', True)
    with repo.transaction() as con:
        con.executescript(ai_task_engine._TASK_ENGINE_SCHEMA)
    shared_artifacts.ensure_artifact_table()
    gate = permission.PermissionGate()
    gate.grant_permission(permission.PermissionLevel.COMMIT)
    monkeypatch.setattr(permission, 'get_permission_gate', lambda: gate)
    pitch.engine.cache_clear()
    app = Flask('shared-av-isolated')
    app.config['LUMENIS_OPEN_WORKSPACE'] = request.param == 'local'
    ws = {'repo': repo, 'gate': gate, 'app': app, 'root': tmp_path, 'scope': dict(SHARED), 'mode': request.param}
    user, headers = account(ws)
    ws.update(account=user, headers={} if request.param == 'local' else headers)
    with app.test_request_context('/', headers={'Host':'localhost'}, environ_overrides={'REMOTE_ADDR':'127.0.0.1'}):
        ws['actor'] = permission.resolve_local_workspace_principal(module='reports', action='write') if request.param == 'local' else user
        yield ws
    pitch.engine.cache_clear()


@pytest.fixture
def client(shared):
    yield helpers.client.__wrapped__(shared)


def values(ws, request='shared-plan-001'):
    result = helpers.audiovisual_arguments(ws, request)
    result['shot_plan']['brief_text'] = 'Shared original brief: preserve hands and uncertainty. 不预设口播。'
    return result


def http_values(ws, request='shared-plan-001'):
    v = values(ws, request)
    v.pop('actor')
    return v


def call(ws, client, method, suffix='', **kwargs):
    kwargs.setdefault('headers', ws['headers'])
    return getattr(client, method)(PREFIX + suffix, **kwargs)


def tables(ws):
    return {r[0] for r in ws['repo'].query_all("SELECT name FROM sqlite_master WHERE type='table'")}


def state(ws):
    return (helpers.count(ws), ws['repo'].query_one('SELECT COUNT(*) FROM ai_task')[0])


def revoke(ws, action):
    with ws['repo'].transaction() as con:
        con.execute('DELETE FROM v4_member_authorizations WHERE user_id=? AND action=?', (ws['account']['id'], action))


def test_shared_http_save_rev_before_after_download_complete_plan(shared, client):
    assert not any('subject_registry' in name for name in tables(shared))
    args = http_values(shared)
    created = call(shared, client, 'post', json=args)
    assert created.status_code == 202, created.get_json()
    task = created.get_json()['result']['task_id']
    result = call(shared, client, 'post', '/' + task + '/run', json={'scope': SHARED})
    assert result.status_code == 200, result.get_json()
    saved = result.get_json()['result']
    assert saved['scope'] == SHARED
    document = saved['work']['document']
    assert document['context'] == {**SHARED, **work._context_owner(work.authorize_saved_work(shared['actor'], SHARED, action='read'))}
    snapshot = document['payload']['audiovisual_plan']
    assert snapshot['input'] == args['shot_plan']
    out = call(shared, client, 'get', '/' + task + '/export', query_string={'scope_kind':'shared_workspace','version':'1'})
    assert out.status_code == 200, out.get_json() if out.is_json else out.data[:50]
    assert hashlib.sha256(out.data).hexdigest() == out.headers['X-Artifact-SHA256']
    text = '\n'.join(p.text for p in Document(io.BytesIO(out.data)).paragraphs)
    for string in [args['manuscript']['sections'][0]['body'], args['shot_plan']['brief_text'], snapshot['input_sha256'], 'NOT_EXECUTED', 'NEEDS_MATERIAL', 'NEEDS_REVIEW', 'DECLARED_NOT_REVALIDATED', 'dot-av-v2:R10', 'same_action_chain_verified', 'Unknown / not provided', 'not material licences']:
        assert string in text, string
    for string in ['hands','tools','causality','silent-craft','USER_EXPLICIT','current_user_brief','forbidden']:
        assert string in text
    revised = http_values(shared, 'shared-plan-002')
    revised['shot_plan']['brief_text'] += ' Revised intent.'
    revised.update(work_id=document['work_id'], previous_artifact_id=saved['action_receipt']['artifact_id'], expected_draft_digest=document['draft_digest'])
    pending2 = call(shared, client, 'post', json=revised).get_json()['result']
    second = call(shared, client, 'post', '/' + pending2['task_id'] + '/run', json={'scope': SHARED}).get_json()['result']
    assert second['action_receipt']['version'] == 2
    assert second['action_receipt']['parent_sha256'] == saved['action_receipt']['sha256']
    old = call(shared, client, 'get', '/' + task, query_string={'scope_kind':'shared_workspace','version':'1'})
    assert old.get_json()['result']['work']['document'] == document
    old_bytes = call(shared, client, 'get', '/' + task + '/export', query_string={'scope_kind':'shared_workspace','version':'1'})
    assert old_bytes.data == out.data
    assert not any('subject_registry' in name for name in tables(shared))


def test_shared_actual_loopback_http_server(shared, client):
    server = make_server('127.0.0.1', 0, shared['app'])
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f'http://127.0.0.1:{server.server_port}' + PREFIX
    def send(suffix='', body=None):
        request = urllib.request.Request(base + suffix, data=json.dumps(body).encode() if body is not None else None,
            headers={**shared['headers'], **({'Content-Type':'application/json'} if body is not None else {})})
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, response.read(), response.headers
    try:
        code, data, _ = send(body=http_values(shared, 'socket-plan-001'))
        assert code == 202
        task = json.loads(data)['result']['task_id']
        assert send('/'+task+'/run', {'scope': SHARED})[0] == 200
        assert send('/'+task+'?scope_kind=shared_workspace&version=1')[0] == 200
        code, raw, headers = send('/'+task+'/export?scope_kind=shared_workspace&version=1')
        assert code == 200 and hashlib.sha256(raw).hexdigest() == headers['X-Artifact-SHA256']
        assert 'Complete audiovisual plan' in '\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    finally:
        server.shutdown(); thread.join(10); server.server_close()


@pytest.mark.parametrize('scope', [
    {'scope_kind':'shared_workspace','version':True}, {'scope_kind':'shared_workspace','version':False},
    {'scope_kind':'shared_workspace','version':'1'}, {'scope_kind':'shared_workspace','version':1.0},
    {'scope_kind':'shared_workspace','version':2}, {'scope_kind':'shared_workspace'},
    {'scope_kind':'shared_workspace','version':1,'actor_id':1},
    {'scope_kind':'shared_workspace','version':1,'domain_id':'cth'},
    {'scope_kind':'shared_workspace','version':1,'global_scope':True},
])
def test_shared_malformed_selector_does_not_create_task(shared, client, scope):
    args = http_values(shared); args['scope'] = scope
    response = call(shared, client, 'post', json=args)
    assert response.status_code == 400, response.get_json()
    assert state(shared) == (0, 0)


@pytest.mark.parametrize('query', [
    [('scope_kind','shared_workspace'),('version','1'),('version','1')],
    [('scope_kind','shared_workspace'),('scope_kind','shared_workspace'),('version','1')],
    {'scope_kind':'shared_workspace','version':'true'}, {'scope_kind':'shared_workspace','version':'01'},
    {'scope_kind':'shared_workspace','version':'1','domain_id':'cth'},
    {'scope_kind':'shared_workspace','version':'1','actor_id':'1'},
])
def test_shared_query_pollution_fails_closed(shared, client, query):
    task = pitch.prepare_internal_pitch(**values(shared))['task_id']
    for suffix in ('', '/export'):
        response = call(shared, client, 'get', '/' + task + suffix, query_string=query)
        assert response.status_code == 400, response.get_json()
    response = call(shared, client, 'post', '/' + task + '/run', json={'scope': SHARED}, query_string=query)
    assert response.status_code == 400
    assert pitch.read_internal_pitch(task, actor=shared['actor'], scope=SHARED)['state'] == 'pending'
    assert helpers.count(shared) == 0


def test_shared_duplicate_json_actor_and_scope_locations_fail_closed(shared, client):
    for patch in [{'actor': {'id':1}}, {'actor_id':1}, {'principal_ref':'workspace:local'}]:
        args=http_values(shared);args.update(patch)
        assert call(shared, client, 'post', json=args).status_code == 400
    raw=json.dumps(http_values(shared)).replace('"scope": {', '"scope": {}, "scope": {')
    assert call(shared, client, 'post', data=raw, content_type='application/json').status_code == 400
    assert call(shared, client, 'post', json=http_values(shared), query_string={'scope_kind':'shared_workspace','version':'1'}).status_code == 400
    assert state(shared) == (0,0)


def test_shared_scope_does_not_open_domain_source_profile(shared):
    args=values(shared);args.pop('shot_plan');args['profile']=pitch.SOURCE_PROFILE
    args['upstream_ref']={'kind':'saved_work','artifact_id':'forged','expected_version':1,'expected_sha256':'a'*64}
    with pytest.raises(ValueError,match='INTERNAL_PITCH_SOURCE_SCOPE_REQUIRED'):
        pitch.prepare_internal_pitch(**args)
    assert state(shared)==(0,0)


def test_shared_owner_and_cross_scope_cannot_borrow_task_or_attachment(shared):
    committed=helpers.run(shared,pitch.prepare_internal_pitch(**values(shared)))
    task=committed['task_id']; other,_=account(shared)
    for actor,scope in [(other,SHARED), ({'id':True},SHARED), (shared['actor'],{'domain_id':'cth','subject_id':None,'brand_focus_id':None,'global_scope':True})]:
        for action in (pitch.read_internal_pitch,pitch.run_internal_pitch,pitch.cancel_internal_pitch,pitch.download_internal_pitch):
            with pytest.raises((ValueError,PermissionError)):
                action(task,actor=actor,scope=scope)
    context=work.authorize_saved_work(shared['actor'],SHARED,action='read')
    with shared['repo'].connection(read_only=True,wal=False) as con:
        for actor in (None,other):
            with pytest.raises(ValueError):
                work._read_attachment(con,committed['action_receipt']['output'],context,committed['action_receipt']['work_id'],actor=actor)
    assert helpers.count(shared)==2


def test_shared_parent_digest_and_foreign_parent_are_not_overwritten(shared):
    first=helpers.run(shared,pitch.prepare_internal_pitch(**values(shared)))
    parent=first['work']['document']; count=helpers.count(shared)
    for index,changes in enumerate([{'expected_draft_digest':'f'*64}, {'previous_artifact_id':'foreign-artifact'}]):
        args=values(shared,f'stale-parent-{index}')
        args.update(work_id=parent['work_id'],previous_artifact_id=first['action_receipt']['artifact_id'],expected_draft_digest=parent['draft_digest'])
        args.update(changes)
        pending=pitch.prepare_internal_pitch(**args)
        with pytest.raises((ValueError,PermissionError,LookupError)):
            helpers.run(shared,pending)
        assert helpers.count(shared)==count
    assert pitch.read_internal_pitch(first['task_id'],actor=shared['actor'],scope=SHARED)['work']['document']==parent


def test_shared_cancel_retry_and_lost_ack_use_committed_owner(shared,monkeypatch):
    cancelled=pitch.prepare_internal_pitch(**values(shared,'cancel-shared-1'))
    assert pitch.cancel_internal_pitch(cancelled['task_id'],actor=shared['actor'],scope=SHARED)['state']=='cancelled'
    with pytest.raises(ValueError,match='NOT_RUNNABLE'):
        helpers.run(shared,cancelled)
    pending=pitch.prepare_internal_pitch(**values(shared,'ack-shared-2'))
    original=work.save_work
    def lose_ack(*args,**kwargs):
        original(*args,**kwargs)
        raise RuntimeError('SIMULATED_ACK_LOSS')
    monkeypatch.setattr(work,'save_work',lose_ack)
    with pytest.raises(RuntimeError,match='SIMULATED_ACK_LOSS'):
        helpers.run(shared,pending)
    monkeypatch.setattr(work,'save_work',original)
    import scripts.pitch_export
    monkeypatch.setattr(scripts.pitch_export,'export_pitch_docx',lambda *a,**k:pytest.fail('committed retry regenerated document'))
    restored=helpers.run(shared,pending)
    assert restored['state']=='succeeded' and helpers.count(shared)==2
    assert pitch.prepare_internal_pitch(**values(shared,'ack-shared-2'))['action_receipt']==restored['action_receipt']


def test_shared_render_cancel_rolls_back_all_outputs(shared,monkeypatch):
    pending=pitch.prepare_internal_pitch(**values(shared))
    import scripts.pitch_export
    original=scripts.pitch_export._append_audiovisual_plan
    def cancel(*args,**kwargs):
        pitch.cancel_internal_pitch(pending['task_id'],actor=shared['actor'],scope=SHARED)
        return original(*args,**kwargs)
    monkeypatch.setattr(scripts.pitch_export,'_append_audiovisual_plan',cancel)
    with pytest.raises(ValueError,match='CANCEL'):
        helpers.run(shared,pending)
    assert helpers.count(shared)==0
    assert pitch.read_internal_pitch(pending['task_id'],actor=shared['actor'],scope=SHARED)['state']=='cancelled'


@pytest.mark.parametrize('name',['prepare','run','cancel','export'])
def test_shared_effect_only_four_exact_names_keep_operation_action_and_risk(shared,name):
    resource='internal-pitch-'+name+':v1:'+'a'*64
    check=lambda **kw:permission.check_shared_workspace_effect(shared['actor'],SHARED,kw.pop('operation',permission.OperationType.DATA_EXPORT),kw.pop('resource',resource),module=kw.pop('module','reports'),action=kw.pop('action','read'),gate=shared['gate'],**kw)
    assert check().allowed
    for options in [{'operation':permission.OperationType.DB_WRITE},{'module':'content'},{'action':'write'},
                    {'resource':resource+':suffix'},{'resource':resource.replace(':v1:',':v2:')},
                    {'resource':resource[:-1]},{'resource':'internal-pitch-publish:v1:'+'a'*64}]:
        assert not check(**options).allowed
    fresh=permission.PermissionGate()
    assert not permission.check_shared_workspace_effect(shared['actor'],SHARED,permission.OperationType.DATA_EXPORT,resource,module='reports',action='read',gate=fresh).allowed


def test_shared_read_only_cannot_prepare_run_cancel_but_can_read_export(shared):
    if shared['mode']=='local':
        # Use the real member account in the same isolated store, independent of transport.
        shared['actor']=shared['account']
    first=helpers.run(shared,pitch.prepare_internal_pitch(**values(shared)))
    pending=pitch.prepare_internal_pitch(**values(shared,'revoke-write-2'))
    revoke(shared,'write')
    for name in ['prepare','run','cancel']:
        result=permission.check_shared_workspace_effect(shared['actor'],SHARED,permission.OperationType.DATA_EXPORT,'internal-pitch-'+name+':v1:'+'a'*64,module='reports',action='read',gate=shared['gate'])
        assert not result.allowed
    with pytest.raises(ValueError,match='PERMISSION'):
        helpers.run(shared,pending)
    with pytest.raises(ValueError,match='PERMISSION'):
        pitch.cancel_internal_pitch(pending['task_id'],actor=shared['actor'],scope=SHARED)
    assert pitch.download_internal_pitch(first['task_id'],actor=shared['actor'],scope=SHARED)['content']
    revoke(shared,'read')
    for action in (pitch.read_internal_pitch,pitch.download_internal_pitch):
        with pytest.raises(ValueError,match='PERMISSION'):
            action(first['task_id'],actor=shared['actor'],scope=SHARED)


def test_shared_revoke_at_export_final_check_returns_no_bytes(shared,monkeypatch):
    shared['actor']=shared['account']
    saved=helpers.run(shared,pitch.prepare_internal_pitch(**values(shared)))
    original=work._read_attachment
    reads=0
    def revoked(*args,**kwargs):
        nonlocal reads
        raw=original(*args,**kwargs)
        reads+=1
        if reads==2:
            revoke(shared,'read')
        return raw
    monkeypatch.setattr(work,'_read_attachment',revoked)
    with pytest.raises(ValueError,match='PERMISSION'):
        pitch.download_internal_pitch(saved['task_id'],actor=shared['actor'],scope=SHARED)
    assert reads==2


def test_shared_closed_workspace_and_forged_session_get_no_task(shared,client):
    shared['app'].config['LUMENIS_OPEN_WORKSPACE']=False
    for headers in [{},{'X-Actor-ID':str(shared['account']['id'])},{'Authorization':'Bearer nonexistent'}]:
        response=call(shared,client,'post',json=http_values(shared),headers=headers)
        assert response.status_code==401,response.get_json()
    assert state(shared)==(0,0)


def test_complete_docx_all_plan_fields_and_snapshot_tamper(shared,tmp_path):
    from lumenis.shared.content_creation.contracts import AudiovisualPlan,Shot,ShotPlan
    from podcast.intelligence.episode_runtime import PodcastPlan,PodcastSegment
    from scripts.pitch_export import export_pitch_docx
    args=values(shared)
    output=args['shot_plan']['outputs'][0]
    output['plan']=ShotPlan((Shot(shot_purpose='Unique visible action',duration_sec=35,subject_action='Unique hand detail'),)).to_dict()
    second=copy.deepcopy(output)
    second.update(output_id='podcast-original',medium='podcast',canvas_ratio=None,purpose='Preserve interview')
    second['sound']={'original':'required','voice':'optional','music':'forbidden'}
    second['plan']=PodcastPlan('Unique episode','Unique episode brief','en',(PodcastSegment('seg-unique','DECLARED_UNRESOLVED',0,35,'Unique original transcript','Unique edited transcript','Unique speaker','Unique segment',(),2),),revision=3).to_dict()
    args['shot_plan']['outputs'].append(second)
    args['shot_plan']['outputs'][0]['constraints'].update(nested={'key':'Unique constraint','disabled':False,'unknown':None}, empty=[])
    snapshot=AudiovisualPlan.from_dict(args['shot_plan']).snapshot()
    path=tmp_path/'full.docx'
    export_pitch_docx({},str(path),manuscript=args['manuscript'],audiovisual_plan=snapshot)
    text='\n'.join(p.text for p in Document(str(path)).paragraphs)
    def leaves(value):
        if isinstance(value,dict):
            for child in value.values():yield from leaves(child)
        elif isinstance(value,list):
            for child in value:yield from leaves(child)
        elif value is not None and type(value) is not bool and value!='':yield str(value)
    for scalar in leaves(snapshot):
        assert scalar in text,scalar
    assert 'Unique hand detail' in text and 'Unique edited transcript' in text
    assert 'No items supplied' in text and 'Unknown / not provided' in text and 'No' in text
    bad=copy.deepcopy(snapshot);bad['execution']='EXECUTED'
    with pytest.raises(ValueError,match='AV_PLAN_SNAPSHOT_INVALID'):
        export_pitch_docx({},str(tmp_path/'bad.docx'),manuscript=args['manuscript'],audiovisual_plan=bad)
    assert not (tmp_path/'bad.docx').exists()


from tests.test_creation_upstream_flow import upstream
workspace = helpers.workspace


def test_shared_signed_citation_preserves_source_scope_and_blocks_after_revocation(upstream):
    ws,research,_=upstream
    args=helpers.audiovisual_arguments(ws,'shared-signed-source')
    args['scope']=dict(SHARED)
    args['citations']=copy.deepcopy(research['document']['citations'])
    task=pitch.prepare_internal_pitch(**args)
    result=pitch.run_internal_pitch(task['task_id'],actor=ws['actor'],scope=SHARED)
    assert result['work']['receipt']['sources'][0]['scope']==ws['scope']
    assert pitch.download_internal_pitch(task['task_id'],actor=ws['actor'],scope=SHARED)['content']
    with ws['repo'].transaction() as con:
        con.execute("UPDATE lumenis_subject_registry SET status='inactive' WHERE subject_id='SUBJ-PITCH'")
    assert work.authorize_saved_work(ws['actor'],SHARED,action='read')['scope_kind']=='shared_workspace'
    for action in (pitch.read_internal_pitch,pitch.download_internal_pitch):
        with pytest.raises(ValueError,match='WORK_PERMISSION_DENIED'):
            action(task['task_id'],actor=ws['actor'],scope=SHARED)


def test_shared_current_account_scope_does_not_fallback_to_cth_or_forged_role(shared):
    actor,_=account(shared,product='CTH')
    args=values(shared);args['actor']={**actor,'role':'owner'}
    with pytest.raises(ValueError,match='WORK_PERMISSION_DENIED'):
        pitch.prepare_internal_pitch(**args)
    assert state(shared)==(0,0)


def test_shared_http_other_authenticated_account_cannot_read_run_cancel_or_export(shared,client):
    shared['app'].config['LUMENIS_OPEN_WORKSPACE']=False
    actor,headers=account(shared)
    args=http_values(shared)
    result=call(shared,client,'post',json=args,headers=headers)
    assert result.status_code==202
    task=result.get_json()['result']['task_id']
    assert call(shared,client,'post','/'+task+'/run',json={'scope':SHARED},headers=headers).status_code==200
    other,other_headers=account(shared)
    for suffix in ('','/export'):
        response=call(shared,client,'get','/'+task+suffix,query_string={'scope_kind':'shared_workspace','version':'1'},headers=other_headers)
        assert response.status_code==403,response.get_json()
        assert 'work' not in response.get_json()
    for operation in ('run','cancel'):
        assert call(shared,client,'post','/'+task+'/'+operation,json={'scope':SHARED},headers=other_headers).status_code==403
    assert helpers.count(shared)==2


@pytest.mark.parametrize('language,heading,unknown',[('zh','完整视听计划','未知／未提供'),('ko','전체 시청각 계획','알 수 없음 / 미제공')])
def test_complete_docx_language_preserves_plan_unknowns(shared,tmp_path,language,heading,unknown):
    from lumenis.shared.content_creation.contracts import AudiovisualPlan
    from scripts.pitch_export import export_pitch_docx
    args=values(shared);args['manuscript']['language']=language
    snapshot=AudiovisualPlan.from_dict(args['shot_plan']).snapshot()
    output=export_pitch_docx({},str(tmp_path/'localized.docx'),manuscript=args['manuscript'],audiovisual_plan=snapshot)
    text='\n'.join(p.text for p in Document(output).paragraphs)
    assert heading in text and unknown in text and 'NOT_EXECUTED' in text
    assert args['shot_plan']['brief_text'] in text


def test_shared_commit_time_write_revocation_rolls_back_outputs(shared,monkeypatch):
    shared['actor']=shared['account']
    task=pitch.prepare_internal_pitch(**values(shared))
    original=work.save_work
    def save_with_revoke(*args,**kwargs):
        commit=kwargs['before_commit']
        def before_commit(con,artifact,document):
            con.execute("DELETE FROM v4_member_authorizations WHERE user_id=? AND action='write'",(shared['account']['id'],))
            commit(con,artifact,document)
        kwargs['before_commit']=before_commit
        return original(*args,**kwargs)
    monkeypatch.setattr(work,'save_work',save_with_revoke)
    with pytest.raises(ValueError,match='PERMISSION'):
        helpers.run(shared,task)
    assert helpers.count(shared)==0
    # The revoked row is within the rejected artifact transaction and rolled back.
    assert work.authorize_saved_work(shared['actor'],SHARED,action='write')
    assert pitch.read_internal_pitch(task['task_id'],actor=shared['actor'],scope=SHARED)['state']=='failed'


@pytest.mark.parametrize('field', ['estimate', 'nested_number', 'nested_boolean', 'nested_list', 'predicate'])
def test_av_snapshot_comparison_rejects_bool_number_alias_at_every_depth(shared, field):
    from lumenis.shared.content_creation.contracts import AudiovisualPlan, Shot, ShotPlan
    args=values(shared)
    output=args['shot_plan']['outputs'][0]
    output['duration']={'min_seconds':1,'max_seconds':1,'tolerance_seconds':0}
    output['plan']=ShotPlan((Shot(shot_purpose='Single second plan',duration_sec=1),)).to_dict()
    output['constraints']={'nested':{'number':1,'flag':True,'list':[{'zero':0}]}}
    predicate=args['shot_plan']['rule_assessments'][0]['required_conditions'][0]
    predicate.update(value=True,provenance={'kind':'USER_EXPLICIT','locator':'current_user_brief'})
    plan=AudiovisualPlan.from_dict(args['shot_plan']);snapshot=plan.snapshot()
    assert plan.matches_snapshot(snapshot)
    if field=='estimate':snapshot['outputs'][0]['estimated_duration_seconds']=True
    elif field=='nested_number':snapshot['input']['outputs'][0]['constraints']['nested']['number']=True
    elif field=='nested_boolean':snapshot['input']['outputs'][0]['constraints']['nested']['flag']=1
    elif field=='nested_list':snapshot['input']['outputs'][0]['constraints']['nested']['list'][0]['zero']=False
    else:snapshot['input']['rule_assessments'][0]['required_conditions'][0]['value']=1
    assert not plan.matches_snapshot(snapshot)


def test_av_snapshot_comparison_keeps_finite_integer_float_compatibility(shared,tmp_path):
    from lumenis.shared.content_creation.contracts import AudiovisualPlan, Shot, ShotPlan
    from scripts.pitch_export import export_pitch_docx
    args=values(shared)
    output=args['shot_plan']['outputs'][0]
    output['duration']={'min_seconds':1,'max_seconds':1,'tolerance_seconds':0}
    output['plan']=ShotPlan((Shot(shot_purpose='Single second plan',duration_sec=1),)).to_dict()
    output['constraints']={'nested':{'number':1,'flag':True,'large_integer':10**400}}
    plan=AudiovisualPlan.from_dict(args['shot_plan']);snapshot=plan.snapshot()
    snapshot['outputs'][0]['estimated_duration_seconds']=1
    assert plan.matches_snapshot(snapshot)
    path=export_pitch_docx({},str(tmp_path/'numeric-compatible.docx'),manuscript=args['manuscript'],audiovisual_plan=snapshot)
    assert Document(path).paragraphs
    snapshot['input']['outputs'][0]['constraints']['nested']['number']=1.0
    assert plan.matches_snapshot(snapshot)
    snapshot['outputs'][0]['estimated_duration_seconds']=float('inf')
    assert not plan.matches_snapshot(snapshot)


def test_av_export_rejects_boolean_derived_estimate_and_writes_no_document(shared,tmp_path):
    from lumenis.shared.content_creation.contracts import AudiovisualPlan,Shot,ShotPlan
    from scripts.pitch_export import export_pitch_docx
    args=values(shared);output=args['shot_plan']['outputs'][0]
    output['duration']={'min_seconds':1,'max_seconds':1,'tolerance_seconds':0}
    output['plan']=ShotPlan((Shot(shot_purpose='One second',duration_sec=1),)).to_dict()
    snapshot=AudiovisualPlan.from_dict(args['shot_plan']).snapshot()
    snapshot['outputs'][0]['estimated_duration_seconds']=True
    with pytest.raises(ValueError,match='AV_PLAN_SNAPSHOT_INVALID'):
        export_pitch_docx({},str(tmp_path/'boolean-estimate.docx'),manuscript=args['manuscript'],audiovisual_plan=snapshot)
    assert not (tmp_path/'boolean-estimate.docx').exists()


def test_shared_read_rejects_boolean_nested_snapshot_after_committed_owner(shared,monkeypatch):
    from lumenis.shared.content_creation.contracts import Shot,ShotPlan
    args=values(shared);output=args['shot_plan']['outputs'][0]
    output['duration']={'min_seconds':1,'max_seconds':1,'tolerance_seconds':0}
    output['plan']=ShotPlan((Shot(shot_purpose='One second',duration_sec=1),)).to_dict()
    first=helpers.run(shared,pitch.prepare_internal_pitch(**args))
    original=work.reopen_work
    def substitute(*args,**kwargs):
        result=original(*args,**kwargs)
        result['document']['payload']['audiovisual_plan']['outputs'][0]['estimated_duration_seconds']=True
        return result
    monkeypatch.setattr(work,'reopen_work',substitute)
    with pytest.raises(ValueError,match='INTERNAL_PITCH_AV_PLAN_INTEGRITY_FAILED'):
        pitch.read_internal_pitch(first['task_id'],actor=shared['actor'],scope=SHARED)
