"""Isolated service regressions; these generated inputs are not business acceptance."""
from __future__ import annotations

import json
import hashlib
import io
import subprocess
import sys
import zipfile

import pytest
from flask import Flask

from lumenis.shared.database import DatabaseRepository
from lumenis.system.security import permission


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("ZBC_DB_PATH", str(tmp_path / "isolated.sqlite3"))
    monkeypatch.setenv("ZBC_ARTIFACT_DIR", str(tmp_path / "artifacts"))
    monkeypatch.setenv("LUMENIS_EVIDENCE_HMAC_KEY_RING", json.dumps({"test": "isolated-test-key"}))
    monkeypatch.setenv("LUMENIS_EVIDENCE_HMAC_ACTIVE_KEY_ID", "test")
    monkeypatch.setattr(permission, "_TESTING_BYPASS_ENABLED", False)
    import canonical_data
    import shared_artifacts
    from lumenis.shared.evidence import store
    from lumenis.shared.reporting import saved_work
    from lumenis.importer import documents
    repository = DatabaseRepository(str(tmp_path / "isolated.sqlite3"))
    for module in (store, shared_artifacts, saved_work, documents):
        monkeypatch.setattr(module, "get_repository", lambda: repository)
    shared_artifacts.ensure_artifact_table()
    store.init_evidence_store(str(tmp_path / "isolated.sqlite3"))
    with repository.transaction() as con:
        store.ensure_evidence_schema(con)
    canonical_data.ensure_canonical_table(str(tmp_path / "isolated.sqlite3"))
    gate = permission.PermissionGate()
    gate.grant_permission(permission.PermissionLevel.COMMIT)
    monkeypatch.setattr(permission, "get_permission_gate", lambda: gate)
    app = Flask(__name__)
    app.config["LUMENIS_OPEN_WORKSPACE"] = True
    with app.test_request_context("/", headers={"Host": "localhost"},
                                  environ_overrides={"REMOTE_ADDR": "127.0.0.1"}):
        actor = permission.resolve_local_workspace_principal(module="reports", action="write")
        scope = {"domain_id": "cth", "subject_id": None, "brand_focus_id": None, "global_scope": True}
        yield {"root": tmp_path, "repo": repository, "gate": gate, "actor": actor, "scope": scope,
               "work": saved_work, "documents": documents, "app": app}


def _source(workspace, *, filename="source.txt", source_url="https://www.swireproperties.com/"):
    path = workspace["root"] / filename
    path.write_bytes(b"Original source line\nRepeated line\nRepeated line\n")
    provenance = {"source_url": source_url, "published_at": None,
                  "captured_at": "2026-10-02T01:00:00Z", "rights": "public-reference-only",
                  "title": "Generated parser fixture"}
    args = {"actor": workspace["actor"], "scope": workspace["scope"], "provenance": provenance}
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value,
                                      resource=resource, granted_by="isolated-test")
    result = workspace["documents"].import_document(str(path), **args)
    return path, args, result


def _save(workspace, *, request_id="request-0001", **changes):
    params = {"title": "Research draft", "markdown": "A source-backed draft for review.",
              "kind": "research", "payload": {"unknowns": ["Usefulness needs user review"]},
              "citations": [], "request_id": request_id, "actor": workspace["actor"],
              "scope": workspace["scope"]}
    params.update(changes)
    return workspace["work"].save_work(**params)


_TEST_OFFICE_FORMATS = {
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "word/document.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p/></w:body></w:document>',
    ),
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xl/workbook.xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        b'<x:workbook xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><x:sheets/></x:workbook>',
    ),
    "pptx": (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "ppt/presentation.xml",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        b'<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sldIdLst/></p:presentation>',
    ),
}


def _office_attachment(package_type, *, padding=b"", extra_members=(), compression=zipfile.ZIP_DEFLATED,
                       main_xml=None):
    media_type, main_part, main_content_type, default_main_xml = _TEST_OFFICE_FORMATS[package_type]
    main_xml = main_xml or default_main_xml
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/{main_part}" ContentType="{main_content_type}"/>'
        '</Types>'
    ).encode()
    relationship = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        f'Target="{main_part}"/>'
        '</Relationships>'
    ).encode()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=compression) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationship)
        archive.writestr(main_part, main_xml)
        for name, content in extra_members:
            archive.writestr(name, content)
        if padding:
            archive.writestr("padding.bin", padding)
    return {"name": f"sample.{package_type}", "content": buffer.getvalue(), "media_type": media_type}


def test_import_cite_save_reopen_revise_export_and_retry(workspace):
    path, args, source = _source(workspace)
    again = workspace["documents"].import_document(str(path), **args)
    assert again["replayed"] and again["canonical_id"] == source["canonical_id"]
    assert again["evidence_id"] == source["evidence_id"]
    located = workspace["documents"].locate_quote(source["evidence_id"], actor=workspace["actor"],
                                                 scope=workspace["scope"], start=35, end=48)
    assert located["citation"]["quote"] == "Repeated line"
    assert located["locators"][0]["index"] == 3
    saved = _save(workspace, citations=[located["citation"]])
    initial_id = saved["artifact"]["artifact_id"]
    reopened = workspace["work"].reopen_work(initial_id, actor=workspace["actor"], scope=workspace["scope"])
    assert reopened["document"] == saved["document"]
    replayed = _save(workspace, citations=[located["citation"]])
    assert replayed["replayed"] and replayed["artifact"]["artifact_id"] == initial_id
    revision = _save(workspace, request_id="request-0002", work_id=saved["document"]["work_id"],
                     previous_artifact_id=initial_id, markdown="Revised human conclusion.", citations=[located["citation"]])
    assert revision["document"]["version"] == 2
    assert workspace["work"].reopen_work(initial_id, actor=workspace["actor"], scope=workspace["scope"])["document"]["version"] == 1
    found = workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"], query="Revised")
    assert found["items"][0]["artifact"]["artifact_id"] == revision["artifact"]["artifact_id"]
    exported = workspace["work"].export_work(revision["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"], format="json")
    assert json.loads(exported["content"])["version"] == 2
    assert exported["return_ref"]["scope"] == workspace["scope"]


def test_selected_preview_rows_project_to_canonical_children_under_source_evidence(workspace):
    from lumenis.importer.analyzer import extract_source_rows, source_rows_fingerprint
    from lumenis.importer.parser import parse_file

    path = workspace["root"] / "selected-rows.csv"
    source_bytes = (
        "标题,正文,作者,发布日期,门店\n"
        "未选行,第一行内容,作者甲,2026-01-01,广州\n"
        "选中行,第二行内容,作者乙,2026-01-02,深圳\n"
        "未选行二,第三行内容,作者丙,2026-01-03,上海\n"
    ).encode("utf-8")
    path.write_bytes(source_bytes)
    provenance = {
        "source_url": "https://www.swireproperties.com/research/source.csv",
        "published_at": None, "captured_at": "2026-10-02T01:00:00Z",
        "rights": "public-reference-only", "title": "Selected row source",
    }
    args = {"actor": workspace["actor"], "scope": workspace["scope"],
            "provenance": provenance}
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(
        permission.OperationType.DATA_IMPORT.value, resource=resource,
        granted_by="isolated-test",
    )
    parsed = parse_file(str(path))
    source_rows = extract_source_rows(parsed)
    expected_fingerprint = source_rows_fingerprint(parsed, source_rows, count=3)

    with pytest.raises(ValueError, match="DOCUMENT_PREVIEW_SELECTION_CHANGED"):
        workspace["documents"].import_document(
            str(path), **args, selected_row_indices=[1], preview_row_count=3,
            expected_preview_rows_sha256="0" * 64,
        )
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM app_evidence_records")[0] == 0
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM canonical_data")[0] == 0

    imported = workspace["documents"].import_document(
        str(path), **args, selected_row_indices=[1], preview_row_count=3,
        expected_preview_rows_sha256=expected_fingerprint,
    )
    assert imported["replayed"] is False
    parent = workspace["repo"].query_one(
        "SELECT canonical_id, source_type FROM canonical_data WHERE source_id=? AND source_type='document'",
        (imported["source_id"],),
    )
    assert parent["canonical_id"] == imported["canonical_id"]
    assert parent["source_type"] == "document"
    assert imported["selected_row_indices"] == [1]
    assert imported["canonical_rows_inserted"] == 1
    assert len(imported["canonical_row_records"]) == 1
    selected = imported["canonical_row_records"][0]
    assert selected["row_index"] == 1
    assert selected["source_locator"] == {
        "unit": "parsed_table_row", "table_index": 0, "row_index": 2,
    }

    child = workspace["repo"].query_one(
        "SELECT * FROM canonical_data WHERE canonical_id=?", (selected["canonical_id"],),
    )
    assert child["source_type"] == "document_row"
    assert child["source_id"] == imported["source_id"]
    assert child["scope_evidence_ref"] == str(imported["evidence_id"])
    assert child["title"] == "选中行"
    assert child["content"] == "第二行内容"
    assert child["author_name"] == "作者乙"
    assert child["publish_datetime"] == "2026-01-02"
    assert child["sentiment"] == "unknown" and child["risk_level"] == "unknown"
    assert child["ai_provider"] == "unknown" and child["ai_analysis"] == ""
    row_provenance = json.loads(child["provenance"])
    assert row_provenance["analysis_status"] == "NOT_RUN"
    assert row_provenance["source_row_sha256"] == selected["source_row_sha256"]
    raw_ref = json.loads(child["raw_ref"])
    assert raw_ref["evidence_id"] == imported["evidence_id"]
    assert raw_ref["source_locator"] == selected["source_locator"]

    reopened = workspace["documents"].reopen_imported_source(
        imported["evidence_id"], actor=workspace["actor"], scope=workspace["scope"],
    )
    assert reopened["content"] == source_bytes

    replayed = workspace["documents"].import_document(
        str(path), **args, selected_row_indices=[1], preview_row_count=3,
        expected_preview_rows_sha256=expected_fingerprint,
    )
    assert replayed["replayed"] is True
    assert replayed["canonical_id"] == imported["canonical_id"]
    assert replayed["canonical_rows_inserted"] == 0
    assert replayed["canonical_row_records"] == imported["canonical_row_records"]
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM canonical_data")[0] == 2

    another_selection = workspace["documents"].import_document(
        str(path), **args, selected_row_indices=[0], preview_row_count=3,
        expected_preview_rows_sha256=expected_fingerprint,
    )
    assert another_selection["source_replayed"] is True
    assert another_selection["replayed"] is False
    assert another_selection["canonical_rows_inserted"] == 1
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM app_evidence_records")[0] == 1
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM canonical_data")[0] == 3


def test_unapproved_or_revoked_import_has_no_source_side_write(workspace):
    path, args, result = _source(workspace)
    path.write_text("Changed document", encoding="utf-8")
    before = workspace["repo"].query_one("SELECT COUNT(*) FROM app_evidence_records")[0]
    with pytest.raises(ValueError, match="WORK_EFFECT_PERMISSION_DENIED"):
        workspace["documents"].import_document(str(path), **args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM app_evidence_records")[0] == before
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM canonical_data")[0] == 1


def test_missing_signer_rolls_back_evidence_and_canonical(workspace, monkeypatch):
    path, args, _ = _source(workspace)
    path.write_text("Another document", encoding="utf-8")
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value, resource=resource)
    monkeypatch.delenv("LUMENIS_EVIDENCE_HMAC_KEY_RING")
    monkeypatch.delenv("LUMENIS_EVIDENCE_HMAC_ACTIVE_KEY_ID")
    with pytest.raises(ValueError):
        workspace["documents"].import_document(str(path), **args)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM canonical_data")[0] == 1
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM app_evidence_records")[0] == 1


def test_stale_parent_and_idempotency_conflict_preserve_versions(workspace):
    first = _save(workspace)
    wid, aid = first["document"]["work_id"], first["artifact"]["artifact_id"]
    second = _save(workspace, request_id="request-0002", work_id=wid, previous_artifact_id=aid)
    with pytest.raises(ValueError, match="WORK_REVISION_CONFLICT"):
        _save(workspace, request_id="request-0003", work_id=wid, previous_artifact_id=aid)
    with pytest.raises(ValueError, match="WORK_REQUEST_REUSE_CONFLICT"):
        _save(workspace, markdown="Different content with same request ID")
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 2


def test_podcast_draft_saves_before_render_and_retains_complete_payload(workspace):
    episode = {"PodcastPlan": {"topic": "Fixture only"}, "script": "Revised words",
               "original_transcript": "Original words", "corrections": [{"start": 0, "text": "Revised"}],
               "corrected_transcript": "Revised words", "segments": [], "media_fingerprints": [],
               "chapters": [{"title": "Opening", "start": 0}], "source_refs": [], "Evidence": [],
               "revision": 2, "Timeline": {"version": 1, "clips": []},
               "rendition": {"artifact_ref": "ART-0000000000000001", "episode_revision": 1},
               "new_script_state": "awaiting_recording"}
    contract = {"id": "isolated-W4-contract-fixture", "version": "1",
                "required_payload_paths": list(episode)}
    result = _save(workspace, kind="podcast_episode", payload=episode, contract=contract,
                   task_id="fixture-podcast-task", assertion_id="fixture-save-no-render")
    assert result["document"]["payload"] == episode
    assert result["document"]["payload"]["rendition"]["episode_revision"] == 1
    assert result["artifact"]["media_type"] == "application/json"
    receipt = result["receipt"]
    assert receipt["constraint_observations"]["media_lineage"] == "W4_W6_NOT_VERIFIED"
    read = workspace["work"].read_committed_work(result["artifact"]["artifact_id"],
                                                 actor=workspace["actor"], scope=workspace["scope"])
    assert hashlib.sha256(read["content"]).hexdigest() == receipt["sha256"]
    assert len(read["content"]) == receipt["size_bytes"]
    args = {"kind": "podcast_episode", "payload": {**episode, "script": "Another revision"},
            "contract": contract, "work_id": result["document"]["work_id"],
            "previous_artifact_id": result["artifact"]["artifact_id"], "request_id": "episode-edit-0002"}
    with pytest.raises(ValueError, match="WORK_DRAFT_REVISION_CONFLICT"):
        _save(workspace, **args)
    second = _save(workspace, **args, expected_draft_digest=result["document"]["draft_digest"])
    assert _save(workspace, **args, expected_draft_digest=result["document"]["draft_digest"])["replayed"]
    assert second["document"]["payload"]["rendition"]["episode_revision"] == 1
    with pytest.raises(ValueError, match="WORK_PAYLOAD_INCOMPLETE"):
        _save(workspace, kind="podcast_episode", payload={}, contract=contract)


def test_validated_podcast_plan_saves_reopens_and_revises(workspace):
    from dataclasses import replace
    from podcast.intelligence.episode_runtime import PodcastPlan, PodcastSegment

    _, _, source = _source(workspace, filename="podcast-plan-source.txt")
    citation = workspace["documents"].locate_quote(
        source["evidence_id"], actor=workspace["actor"], scope=workspace["scope"], start=0, end=20,
    )["citation"]
    first_plan = PodcastPlan(
        title="Source-backed episode",
        description="A reviewed original episode plan.",
        language="en",
        segments=(PodcastSegment(
            segment_id="segment-1", media_asset_id="MEDIA-source", source_in=0, source_out=4,
            original_text="Original source wording", corrected_text="Reviewed episode wording",
            speaker="Host", chapter_title="Opening", evidence_refs=(f"evidence:{source['evidence_id']}",),
        ),),
    )
    contract = {"id": "lumenis.podcast-plan", "version": "1", "required_payload_paths": ["podcast_plan"]}
    first = _save(
        workspace, kind="podcast_episode", markdown=first_plan.script_text(),
        payload={"podcast_plan": first_plan.to_dict()}, citations=[citation], contract=contract,
        request_id="podcast-plan-0001",
    )
    assert first["document"]["payload"]["podcast_plan"] == first_plan.to_dict()
    assert first["receipt"]["constraint_observations"]["media_lineage"] == "W4_W6_NOT_VERIFIED"
    assert workspace["work"].reopen_work(
        first["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"],
    )["document"]["payload"]["podcast_plan"] == first_plan.to_dict()

    revised_plan = replace(first_plan, revision=2, segments=(replace(
        first_plan.segments[0], corrected_text="Reviewed revision wording",
    ),))
    second = _save(
        workspace, kind="podcast_episode", markdown=revised_plan.script_text(),
        payload={"podcast_plan": revised_plan.to_dict()}, citations=[citation], contract=contract,
        request_id="podcast-plan-0002", work_id=first["document"]["work_id"],
        previous_artifact_id=first["artifact"]["artifact_id"],
        expected_draft_digest=first["document"]["draft_digest"],
    )
    assert second["document"]["version"] == 2
    assert workspace["work"].reopen_work(
        first["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"],
    )["document"]["payload"]["podcast_plan"] == first_plan.to_dict()

    invalid = {"podcast_plan": first_plan.to_dict()}
    invalid["podcast_plan"]["segments"][0]["source_out"] = 0
    before = workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]
    with pytest.raises(ValueError, match="WORK_PODCAST_PLAN_INVALID"):
        _save(
            workspace, kind="podcast_episode", markdown="Invalid plan", payload=invalid,
            citations=[citation], contract=contract, request_id="podcast-plan-invalid",
        )
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == before


def test_corruption_and_ungrounded_citations_never_reopen_success(workspace):
    _, _, source = _source(workspace)
    with pytest.raises(ValueError, match="WORK_CITATION_UNGROUNDED"):
        _save(workspace, citations=[{"evidence_id": source["evidence_id"], "start": 0, "end": 8, "quote": "Invented"}])
    saved = _save(workspace)
    path = workspace["root"] / "artifacts" / saved["artifact"]["filename"]
    path.write_text("tampered", encoding="utf-8")
    with pytest.raises(ValueError, match="WORK_INTEGRITY_FAILED"):
        workspace["work"].reopen_work(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"])
    found = workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"])
    assert found["unverified_candidates"] == 1 and not found["items"]


def test_foreign_actor_denied_and_permission_loss_prevents_save(workspace, monkeypatch):
    first = _save(workspace)
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(first["artifact"]["artifact_id"], actor={"actor_id": None}, scope=workspace["scope"])
    workspace["app"].config["LUMENIS_OPEN_WORKSPACE"] = False
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        _save(workspace, request_id="request-0002")
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 1


def test_storage_failure_cleans_file_and_retry_recovers(workspace, monkeypatch):
    original = workspace["work"]._read
    def fail(*args, **kwargs):
        raise OSError("simulated interrupted verification before commit")
    monkeypatch.setattr(workspace["work"], "_read", fail)
    with pytest.raises(OSError):
        _save(workspace)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 0
    assert not list((workspace["root"] / "artifacts").iterdir())
    monkeypatch.setattr(workspace["work"], "_read", original)
    assert _save(workspace)["document"]["version"] == 1


def test_search_cursor_is_bound_to_query_and_scope(workspace):
    _save(workspace)
    _save(workspace, request_id="request-0002")
    first = workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"], limit=1)
    assert first["has_more"] and first["next_cursor"]
    second = workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"], limit=1, cursor=first["next_cursor"])
    assert second["items"][0]["artifact"]["artifact_id"] != first["items"][0]["artifact"]["artifact_id"]
    with pytest.raises(ValueError, match="WORK_CURSOR_INVALID"):
        workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"], query="changed", cursor=first["next_cursor"])


def test_parser_rejects_corrupt_utf8_instead_of_silent_loss(tmp_path):
    from lumenis.importer.parser import parse_file
    path = tmp_path / "bad.txt"
    path.write_bytes(b"valid\xffinvalid")
    with pytest.raises(ValueError, match="DOCUMENT_TEXT_ENCODING_OR_CONTROL_INVALID"):
        parse_file(str(path))


def test_global_sources_keep_original_focus_and_revocation_is_fresh(workspace):
    from subject_registry import add_subject
    refs = []
    for name in ("First fixture", "Second fixture"):
        subject, _ = add_subject({"domain_id": "CTH", "subject_type": "company", "name_en": name, "name_zh": name})
        focused = {**workspace["scope"], "global_scope": False, "subject_id": subject["subject_id"]}
        _, _, source = _source({**workspace, "scope": focused}, filename=name + ".txt")
        citation = workspace["documents"].locate_quote(source["evidence_id"], actor=workspace["actor"],
                                                      scope=focused, start=0, end=20)["citation"]
        refs.append(citation)
    saved = _save(workspace, citations=refs)
    assert [s["scope"]["global_scope"] for s in saved["receipt"]["sources"]] == [False, False]
    assert len({s["scope"]["subject_id"] for s in saved["receipt"]["sources"]}) == 2
    global_quote = workspace["documents"].locate_quote(refs[0]["evidence_id"], actor=workspace["actor"],
                                                      scope=workspace["scope"], start=0, end=20)
    assert global_quote["citation"]["quote"] == refs[0]["quote"]
    assert global_quote["original_scope"] == saved["receipt"]["sources"][0]["scope"]
    withdrawn = saved["receipt"]["sources"][0]["scope"]["subject_id"]
    workspace["repo"].execute("UPDATE lumenis_subject_registry SET status='inactive' WHERE subject_id=?", (withdrawn,))
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"])
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].export_work(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"])


def test_mixed_domain_missing_rights_and_forged_source_scope_fail_closed(workspace):
    from subject_registry import add_subject
    brand, _ = add_subject({"domain_id": "BMP", "subject_type": "brand", "name_en": "Fixture brand", "name_zh": "测试品牌"})
    focused = {"domain_id": "zbmp", "subject_id": None, "brand_focus_id": brand["brand_focus_id"], "global_scope": False}
    _, _, source = _source({**workspace, "scope": focused})
    cite = workspace["documents"].locate_quote(source["evidence_id"], actor=workspace["actor"], scope=focused, start=0, end=20)["citation"]
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        _save(workspace, citations=[cite])
    bad = {**cite, "source_scope": workspace["scope"]}
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        _save({**workspace, "scope": focused}, citations=[bad])
    path, args, _ = _source(workspace, filename="another.txt")
    with pytest.raises(ValueError, match="DOCUMENT_PROVENANCE_REQUIRED"):
        workspace["documents"].import_document(str(path), **{**args, "provenance": {k:v for k,v in args["provenance"].items() if k != "rights"}})


@pytest.mark.parametrize("domain", ["CTH", "BMP"])
def test_signed_source_without_original_scope_cannot_be_used_as_global(workspace, domain):
    from subject_registry import add_subject
    from lumenis.shared.evidence import store
    item, _ = add_subject({"domain_id": domain, "subject_type": "company" if domain == "CTH" else "brand",
                           "name_en": "Original scope regression", "name_zh": "逐来源范围回归"})
    key = "subject_id" if domain == "CTH" else "brand_focus_id"
    global_scope = {**workspace["scope"], "domain_id": "cth" if domain == "CTH" else "zbmp"}
    scope = {**global_scope, "global_scope": False, key: item[key]}
    _, _, imported = _source({**workspace, "scope": scope})
    quote = workspace["documents"].locate_quote(imported["evidence_id"], actor=workspace["actor"],
                                                scope=scope, start=0, end=20)["citation"]
    global_workspace = {**workspace, "scope": global_scope}
    saved = _save(global_workspace, citations=[quote])
    original_bytes = workspace["work"].read_committed_work(saved["artifact"]["artifact_id"],
                                                         actor=workspace["actor"], scope=global_scope)["content"]
    # Simulate a legacy/malformed *writer-signed* source, rather than testing
    # only a bad signature. Signature assurance cannot invent a missing scope.
    with workspace["repo"].transaction() as con:
        metadata = json.loads(con.execute("SELECT collection_metadata FROM app_evidence_records WHERE id=?",
                                          (imported["evidence_id"],)).fetchone()[0])
        del metadata["context"][key]
        con.execute("UPDATE app_evidence_records SET collection_metadata=? WHERE id=?",
                    (json.dumps(metadata), imported["evidence_id"]))
        store._sign_evidence_record(con, imported["evidence_id"])
    with workspace["repo"].connection(read_only=True, wal=False) as con:
        cursor = con.execute("SELECT * FROM app_evidence_records WHERE id=?", (imported["evidence_id"],))
        signed = dict(zip((c[0] for c in cursor.description), cursor.fetchone()))
        assert store.classify_evidence_integrity(signed)["verified"] is True
    before = workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        _save(global_workspace, request_id="missing-source-scope", citations=[quote])
    for reader in (workspace["work"].reopen_work, workspace["work"].export_work):
        with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
            reader(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=global_scope)
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        workspace["documents"].locate_quote(imported["evidence_id"], actor=workspace["actor"],
                                            scope=global_scope, start=0, end=20)
    found = workspace["work"].search_work(actor=workspace["actor"], scope=global_scope)
    assert found["items"] == [] and found["unverified_candidates"] == 1
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == before
    path = workspace["root"] / "artifacts" / saved["artifact"]["filename"]
    assert path.read_bytes() == original_bytes


def test_research_inference_keeps_identity_and_boundaries_through_selection(workspace):
    path = workspace["root"] / "derived-note.txt"
    path.write_text("A self-authored method proposal: compare publisher statements with field observations.\n")
    provenance = {"source_url": None, "source_origin": {"kind": "approved-file", "ref": "owned-method-note"},
                  "published_at": None, "captured_at": "2026-10-03T01:00:00Z", "rights": "owner-approved",
                  "title": "Own derived method, not one of the sealed researcher packages",
                  "content_origin": {"kind": "researcher-note", "claim_type": "method_proposal",
                                     "research_id": "OWN-METHOD-DEVELOPMENT-1"}}
    args = dict(actor=workspace["actor"], scope=workspace["scope"], provenance=provenance)
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value, resource=resource)
    imported = workspace["documents"].import_document(str(path), **args)
    assert imported["evidence_id"] != provenance["content_origin"]["research_id"]
    assert imported["source_boundaries"]["claim_type"] == "method_proposal"
    assert imported["source_boundaries"]["fact_verification"] == "NOT_VERIFIED"
    again = workspace["documents"].import_document(str(path), **args)
    assert again["replayed"] and again["source_boundaries"] == imported["source_boundaries"]
    canonical = workspace["repo"].query_one("SELECT source_platform FROM canonical_data WHERE canonical_id=?",
                                            (imported["canonical_id"],))
    assert canonical[0] == "approved_research_document"
    located = workspace["documents"].locate_quote(imported["evidence_id"], actor=workspace["actor"],
                                                scope=workspace["scope"], start=0, end=38)
    saved = _save(workspace, title="Method pending independent verification", citations=[located["citation"]])
    selected = workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"],
                                            query="Method pending")["items"][0]
    assert selected["artifact"]["artifact_id"] == saved["artifact"]["artifact_id"]
    reopened = workspace["work"].reopen_work(selected["artifact"]["artifact_id"], actor=workspace["actor"],
                                           scope=workspace["scope"])
    source = reopened["receipt"]["sources"][0]
    assert source["content_origin"] == provenance["content_origin"]
    assert source["claim_type"] == "method_proposal" and source["learning_admission"] == "NOT_ADMITTED"
    assert source["training_eligible"] is False and source["held_out"] is False
    assert reopened["receipt"]["constraint_observations"]["factual_verification"] == "NOT_ESTABLISHED_BY_STORAGE"
    exported = workspace["work"].export_work(saved["artifact"]["artifact_id"], actor=workspace["actor"],
                                            scope=workspace["scope"], format="md")["content"].decode()
    assert '"claim_type": "method_proposal"' in exported and '"held_out": false' in exported
    # Changing provenance is a new admission, even with identical note bytes.
    changed = {**args, "provenance": {**provenance, "content_origin": {"kind": "publisher-material", "claim_type": "source_statement"}}}
    with pytest.raises(ValueError, match="WORK_EFFECT_PERMISSION_DENIED"):
        workspace["documents"].import_document(str(path), **changed)


@pytest.mark.parametrize("url", ["https://openstax.org/details/books/psychology-2e", "https://assets.openstax.org/body.txt"])
def test_openstax_body_never_enters_decoder_or_searchable_evidence(workspace, monkeypatch, url):
    path = workspace["root"] / "restricted.txt"
    path.write_text("Synthetic prohibited-input sentinel; no OpenStax body was fetched.")
    provenance = {"source_url": url, "published_at": None, "captured_at": "2026-10-03T01:00:00Z",
                  "rights": "public-reference-only", "title": "Link-only restriction regression"}
    monkeypatch.setattr(workspace["documents"], "parse_file", lambda *_: pytest.fail("restricted text reached decoder"))
    for call in (workspace["documents"].document_import_resource, workspace["documents"].import_document):
        with pytest.raises(ValueError, match="DOCUMENT_SOURCE_LINK_ONLY_REQUIRED"):
            call(str(path), actor=workspace["actor"], scope=workspace["scope"], provenance=provenance)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM app_evidence_records")[0] == 0
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM canonical_data")[0] == 0


@pytest.mark.parametrize("origin", [
    {"kind": "researcher-note", "claim_type": "verified_fact"},
    {"kind": "researcher-note", "claim_type": "research_inference", "training_eligible": True},
    {"kind": "self-authored", "claim_type": "author_statement", "research_id": 43},
])
def test_import_provenance_cannot_claim_fact_training_or_package_admission(workspace, origin):
    path = workspace["root"] / "admission.txt"
    path.write_text("Own note, never independently verified business truth.")
    provenance = {"source_url": None, "source_origin": {"kind": "approved-file", "ref": "owned-admission-note"},
                  "published_at": None, "captured_at": "2026-10-03T01:00:00Z", "rights": "owner-approved",
                  "title": "Admission boundary", "content_origin": origin}
    with pytest.raises(ValueError, match="DOCUMENT_CONTENT_ORIGIN_INVALID"):
        workspace["documents"].import_document(str(path), actor=workspace["actor"], scope=workspace["scope"], provenance=provenance)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 0


def test_decoder_cannot_replace_admitted_nested_research_provenance(workspace, monkeypatch):
    path = workspace["root"] / "stable-provenance.txt"
    path.write_text("Own method: compare source claims with observations.")
    origin = {"kind": "researcher-note", "claim_type": "method_proposal", "research_id": "OWN-ADMITTED-METHOD"}
    provenance = {"source_url": None, "source_origin": {"kind": "approved-file", "ref": "owned-stable-note"},
                  "published_at": None, "captured_at": "2026-10-03T01:00:00Z", "rights": "owner-approved",
                  "title": "Stable research admission", "content_origin": origin}
    args = dict(actor=workspace["actor"], scope=workspace["scope"], provenance=provenance)
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value, resource=resource)
    original_decoder = workspace["documents"].parse_file
    def mutate_caller_during_decode(*values, **options):
        origin["claim_type"] = "research_inference"
        origin["research_id"] = "OWN-REPLACED-METHOD"
        provenance["source_origin"]["ref"] = "changed-file-origin"
        return original_decoder(*values, **options)
    monkeypatch.setattr(workspace["documents"], "parse_file", mutate_caller_during_decode)
    saved = workspace["documents"].import_document(str(path), **args)
    assert saved["source_boundaries"]["content_origin"] == {
        "kind": "researcher-note", "claim_type": "method_proposal", "research_id": "OWN-ADMITTED-METHOD"}
    archived = workspace["documents"].reopen_imported_source(saved["evidence_id"], actor=workspace["actor"],
                                                          scope=workspace["scope"])
    assert archived["source_origin"]["ref"] == "owned-stable-note"
    with pytest.raises(ValueError, match="WORK_EFFECT_PERMISSION_DENIED"):
        workspace["documents"].import_document(str(path), **args)


def test_approved_reference_summary_retains_restriction_and_exact_upstream_link(workspace):
    path=workspace["root"] / "worker-summary.txt"
    upstream="https://marketingscience.info/learn-with-us/commercial-research/identifying-and-prioritising-category-entry-points"
    path.write_text(json.dumps({"original_source_url":upstream,"summary":"A self-authored dated worker summary, not the publisher body."}))
    provenance={"source_url":None,"source_origin":{"kind":"approved-file","ref":"approved-reference-summary"},
                "published_at":None,"captured_at":"2026-10-03T01:00:00Z","rights":"public-reference-only",
                "title":"Approved reference summary","rights_statement":"LINK_AND_ORIGINAL_SUMMARY_ONLY; MODEL_TRAINING_NOT_AUTHORIZED",
                "content_origin":{"kind":"researcher-note","claim_type":"research_inference","research_id":"OWN-SUMMARY-LINK"}}
    args=dict(actor=workspace["actor"],scope=workspace["scope"],provenance=provenance)
    resource=workspace["documents"].document_import_resource(str(path),**args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value,resource=resource)
    source=workspace["documents"].import_document(str(path),**args)
    archived=workspace["documents"].reopen_imported_source(source["evidence_id"],actor=workspace["actor"],scope=workspace["scope"])
    assert json.loads(archived["content"])["original_source_url"]==upstream
    assert source["source_boundaries"]["rights_statement"]==provenance["rights_statement"]
    assert source["source_boundaries"]["training_eligible"] is False
    with pytest.raises(ValueError,match="DOCUMENT_USE_NOT_AUTHORIZED"):
        workspace["documents"].import_document(str(path),**{**args,"provenance":{**provenance,"intended_use":"internal_draft"}})


@pytest.mark.parametrize("domain", ["CTH", "BMP"])
def test_general_method_reference_preserves_original_scope_and_cannot_transfer_facts(workspace, domain):
    from subject_registry import add_subject
    global_scope = {**workspace["scope"], "domain_id": "cth" if domain == "CTH" else "zbmp"}
    item, _ = add_subject({"domain_id": domain, "subject_type": "company" if domain == "CTH" else "brand",
                           "name_en": "Method application scope", "name_zh": "方法选用范围"})
    key = "subject_id" if domain == "CTH" else "brand_focus_id"
    focus = {**global_scope, "global_scope": False, key: item[key]}
    path = workspace["root"] / "general-method.txt"
    path.write_text("Method proposal: compare original claims, dates and counterexamples.")
    provenance = {"source_url": None, "source_origin": {"kind": "approved-file", "ref": "owned-general-method"},
                  "published_at": None, "captured_at": "2026-10-03T01:00:00Z", "rights": "owner-approved",
                  "title": "General method", "intended_use": "research_reference",
                  "rights_statement": "Original synthesized method; MODEL_TRAIN not authorized",
                  "content_origin": {"kind": "researcher-note", "claim_type": "method_proposal", "research_id": "OWN-GENERAL-METHOD"},
                  "research_transport": {"sha256": "a"*64, "parent_manifest_sha256": "b"*64, "record_id": "OWN-GENERAL-METHOD"}}
    args = dict(actor=workspace["actor"], scope=global_scope, provenance=provenance)
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value, resource=resource)
    source = workspace["documents"].import_document(str(path), **args)
    quote = workspace["documents"].locate_quote(source["evidence_id"], actor=workspace["actor"],
                                                scope=global_scope, start=0, end=30)["citation"]
    focused_workspace = {**workspace, "scope": focus}
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        _save(focused_workspace, citations=[quote])
    method_quote = {**quote, "purpose": "method_reference"}
    saved = _save(focused_workspace, citations=[method_quote])
    reopened = workspace["work"].reopen_work(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=focus)
    original = reopened["receipt"]["sources"][0]
    assert original["scope"] == global_scope and original["purposes"] == ["method_reference"]
    assert original["rights_statement"] == provenance["rights_statement"]
    assert original["research_transport"] == provenance["research_transport"]
    assert original["fact_verification"] == "NOT_VERIFIED" and original["held_out"] is False
    with pytest.raises(ValueError, match="WORK_REQUEST_REUSE_CONFLICT"):
        _save(focused_workspace, citations=[method_quote], markdown="Another use with the committed request ID")
    foreign_domain = "zbmp" if domain == "CTH" else "cth"
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        _save({**workspace, "scope": {**workspace["scope"], "domain_id": foreign_domain}}, citations=[method_quote])
    # A general source statement cannot opt into this method-only use merely
    # by putting a purpose flag on a citation.
    _, _, statement = _source({**workspace, "scope": global_scope}, filename="statement.txt")
    source_quote = workspace["documents"].locate_quote(statement["evidence_id"], actor=workspace["actor"],
                                                       scope=global_scope, start=0, end=20)["citation"]
    with pytest.raises(ValueError, match="WORK_METHOD_REFERENCE_INVALID"):
        _save(focused_workspace, request_id="statement-is-not-method", citations=[{**source_quote, "purpose": "method_reference"}])
    # Revoke the real local transport; both source and target authorization
    # must be freshly established before reopening or exporting the result.
    workspace["app"].config["LUMENIS_OPEN_WORKSPACE"] = False
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].export_work(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=focus)
    workspace["app"].config["LUMENIS_OPEN_WORKSPACE"] = True
    assert workspace["work"].reopen_work(saved["artifact"]["artifact_id"], actor=workspace["actor"],
                                       scope=focus)["artifact"]["sha256"] == saved["artifact"]["sha256"]


def test_withdrawn_source_is_removed_by_new_revision_without_exposing_old_content(workspace):
    from subject_registry import add_subject
    subject, _ = add_subject({"domain_id":"CTH", "subject_type":"company", "name_en":"Withdrawn fixture", "name_zh":"撤销来源测试"})
    scope = {**workspace["scope"], "global_scope":False, "subject_id":subject["subject_id"]}
    _, _, source = _source({**workspace,"scope":scope})
    cite = workspace["documents"].locate_quote(source["evidence_id"],actor=workspace["actor"],scope=scope,start=0,end=20)["citation"]
    first = _save(workspace,citations=[cite])
    workspace["repo"].execute("DELETE FROM lumenis_subject_registry WHERE subject_id=?",(subject["subject_id"],))
    with pytest.raises(ValueError,match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(first["artifact"]["artifact_id"],actor=workspace["actor"],scope=workspace["scope"])
    repaired = _save(workspace,request_id="repair-withdrawn-source",work_id=first["document"]["work_id"],
                     previous_artifact_id=first["artifact"]["artifact_id"],citations=[],markdown="Withdrawn source removed; replacement evidence is required.")
    assert repaired["document"]["version"] == 2
    assert repaired["receipt"]["sources"] == []
    assert repaired["receipt"]["parent"]["sha256"] == first["artifact"]["sha256"]
    with pytest.raises(ValueError,match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(first["artifact"]["artifact_id"],actor=workspace["actor"],scope=workspace["scope"])


@pytest.mark.parametrize("format", ["docx", "pdf", "xlsx"])
def test_binary_export_contains_real_saved_snapshot(workspace, format):
    saved = _save(workspace, markdown="Actual saved words with Chinese 中文.\nSecond paragraph.")
    result = workspace["work"].export_work(saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"], format=format)
    if format == "docx":
        from docx import Document
        text = "\n".join(p.text for p in Document(io.BytesIO(result["content"])).paragraphs)
    elif format == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(result["content"])) as pdf:
            text = "\n".join(page.extract_text() for page in pdf.pages)
    else:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(result["content"]), read_only=True)
        text = "\n".join(str(row[2]) for row in book.active.values)
    assert "Actual saved words" in text and saved["document"]["work_id"] in text
    assert hashlib.sha256(result["content"]).hexdigest() == result["sha256"]


def test_permission_revoked_immediately_before_commit_rolls_back(workspace, monkeypatch):
    original = workspace["work"]._read
    def withdraw(*args, **kwargs):
        result = original(*args, **kwargs)
        workspace["app"].config["LUMENIS_OPEN_WORKSPACE"] = False
        return result
    monkeypatch.setattr(workspace["work"], "_read", withdraw)
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        _save(workspace)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 0
    assert not list((workspace["root"] / "artifacts").iterdir())


def test_delivery_preferences_and_acknowledgements_persist(workspace):
    from lumenis.shared.reporting import delivery_preferences as delivery
    initial = delivery.read_preferences(actor=workspace["actor"], scope=workspace["scope"])
    assert initial["persisted"] is False
    saved = _save(workspace)
    values = initial["preferences"]
    values["notifications_enabled"] = False
    values["quiet_hours"]["enabled"] = True
    first = delivery.save_preferences(values, actor=workspace["actor"], scope=workspace["scope"], request_id="preferences-0001")
    feed = delivery.notification_feed(actor=workspace["actor"], scope=workspace["scope"])
    assert feed["notifications"][0]["visible"]  # required human decisions cannot be suppressed
    values["acknowledged_artifact_ids"] = [saved["artifact"]["artifact_id"]]
    args = dict(actor=workspace["actor"], scope=workspace["scope"], request_id="preferences-0002",
                previous_artifact_id=first["artifact"]["artifact_id"], expected_draft_digest=first["document"]["draft_digest"])
    delivery.save_preferences(values, **args)
    assert delivery.save_preferences(values, **args)["replayed"]
    assert delivery.read_preferences(actor=workspace["actor"], scope=workspace["scope"])["version"] == 2
    assert delivery.notification_feed(actor=workspace["actor"], scope=workspace["scope"])["notifications"][0]["acknowledged"]
    values["quiet_hours"]["timezone"] = "Not/AZone"
    with pytest.raises(ValueError, match="DELIVERY_TIMEZONE_INVALID"):
        delivery.save_preferences(values, **args)


def test_fence_denial_leaves_no_persistent_write(workspace, monkeypatch):
    from lumenis.system.security.effect_fence import FenceDecision
    class DeniedFence:
        def check(self, **kwargs):
            return FenceDecision(allowed=False, action="write", fenced=True, reason="fixture policy denial")
    monkeypatch.setattr(workspace["work"], "get_effect_fence", lambda: DeniedFence())
    with pytest.raises(ValueError, match="WORK_EFFECT_FENCE_DENIED"):
        _save(workspace)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 0


_PROCESS = '''
import json, os, sys
from flask import Flask
from lumenis.system.security import permission
from lumenis.shared.reporting import saved_work
app = Flask(__name__)
app.config["LUMENIS_OPEN_WORKSPACE"] = True
permission.get_permission_gate().grant_permission(permission.PermissionLevel.COMMIT)
with app.test_request_context("/", headers={"Host":"localhost"}, environ_overrides={"REMOTE_ADDR":"127.0.0.1"}):
    actor = permission.resolve_local_workspace_principal(module="reports", action="write")
    scope = {"domain_id":"cth", "subject_id":None, "brand_focus_id":None, "global_scope":True}
    if sys.argv[1] == "read":
        item = saved_work.read_committed_work(sys.argv[2], actor=actor, scope=scope)
        print(json.dumps(item["receipt"]))
    else:
        if sys.argv[1] == "crash":
            def crash(*args, **kwargs): os._exit(71)
            saved_work._read = crash
        args = json.loads(sys.argv[2])
        try:
            item = saved_work.save_work(**args, actor=actor, scope=scope)
            print(json.dumps({"ref":item["artifact"]["artifact_id"], "version":item["document"]["version"]}))
        except ValueError as exc:
            print(str(exc))
            sys.exit(3)
'''


def test_hard_process_interrupt_retry_and_fresh_reader(workspace):
    args = dict(title="Crash regression", markdown="Committed after explicit retry", kind="report",
                payload={}, citations=[], request_id="process-crash-0001")
    crash = subprocess.run([sys.executable, "-c", _PROCESS, "crash", json.dumps(args)], capture_output=True, text=True, timeout=20)
    assert crash.returncode == 71, crash.stderr
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 0
    assert not workspace["work"].search_work(actor=workspace["actor"], scope=workspace["scope"])["items"]
    saved = workspace["work"].save_work(**args, actor=workspace["actor"], scope=workspace["scope"])
    ref = saved["artifact"]["artifact_id"]
    read = subprocess.run([sys.executable, "-c", _PROCESS, "read", ref], capture_output=True, text=True, timeout=20)
    assert read.returncode == 0, read.stderr
    assert json.loads(read.stdout)["sha256"] == saved["artifact"]["sha256"]
    assert workspace["work"].save_work(**args, actor=workspace["actor"], scope=workspace["scope"])["replayed"]


def test_concurrent_process_revisions_cannot_overwrite_same_parent(workspace):
    first = _save(workspace)
    args = dict(title="Concurrent revision", markdown="Only one competing revision commits", kind="research",
                payload={}, citations=[], work_id=first["document"]["work_id"],
                previous_artifact_id=first["artifact"]["artifact_id"])
    processes = [subprocess.Popen([sys.executable, "-c", _PROCESS, "save", json.dumps({**args,"request_id":f"parallel-request-{i}"})],
                                  stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True) for i in range(2)]
    results = [p.communicate(timeout=20) for p in processes]
    assert sorted(p.returncode for p in processes) == [0, 3], results
    assert any("WORK_REVISION_CONFLICT" in output for output, _ in results)
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0] == 2


def test_native_document_tables_and_sheet_coordinates_remain_locatable(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from lumenis.importer.parser import parse_file
    word = Document()
    word.add_paragraph("Before the table")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "真实表格来源"
    word.add_paragraph("After the table")
    path = tmp_path / "own-sample.docx"
    word.save(path)
    parsed = parse_file(str(path))
    loc = next(x for x in parsed.locators if x["unit"] == "word_table_cell")
    assert parsed.text[loc["start"]:loc["end"]] == "真实表格来源"
    assert parsed.text.index("Before") < loc["start"] < parsed.text.index("After")
    sheet = Workbook()
    sheet.active.title = "Evidence Sheet"
    sheet.active.append([0, False, "=SUM(A1)"])
    path = tmp_path / "own-sample.xlsx"
    sheet.save(path)
    parsed = parse_file(str(path))
    assert parsed.text == "0 | False | =SUM(A1)"
    assert parsed.locators[-1]["sheet"] == "Evidence Sheet"
    assert parsed.locators[-1]["column"] == 3


def test_approved_file_needs_no_fabricated_public_url_and_original_reopens(workspace):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Owned minimal sample for approved-file intake.")
    path = workspace["root"] / "owned-material.docx"
    doc.save(path)
    provenance = {"source_url":None,"published_at":None,"captured_at":"2026-10-02T01:00:00Z",
                  "rights":"owner-approved","title":"Owned input regression", "intended_use":"internal_draft",
                  "source_origin":{"kind":"approved-file","ref":"approved-upload-fixture-0001"}}
    args = dict(actor=workspace["actor"],scope=workspace["scope"],provenance=provenance)
    resource = workspace["documents"].document_import_resource(str(path),**args)
    workspace["gate"].grant_approval(permission.OperationType.DATA_IMPORT.value,resource=resource)
    source = workspace["documents"].import_document(str(path),**args)
    original = workspace["documents"].reopen_imported_source(source["evidence_id"],actor=workspace["actor"],scope=workspace["scope"])
    assert original["content"]==path.read_bytes()
    assert original["artifact_ref"]==source["source_artifact_ref"]
    path.unlink()  # original upload may be cleaned after successful intake
    assert workspace["documents"].reopen_imported_source(source["evidence_id"],actor=workspace["actor"],scope=workspace["scope"])["content"]==original["content"]


def test_unapproved_source_is_not_decoded(workspace, monkeypatch):
    path = workspace["root"] / "unapproved.html"
    path.write_text("<html>Unapproved input</html>")
    provenance = {"source_url":"https://www.nike.com/","published_at":None,"captured_at":"2026-10-02T01:00:00Z",
                  "rights":"public-reference-only","title":"Unapproved decoder regression"}
    def forbidden_decoder(*args):
        raise AssertionError("unapproved content reached decoder")
    monkeypatch.setattr(workspace["documents"],"parse_file",forbidden_decoder)
    args = dict(actor=workspace["actor"],scope=workspace["scope"],provenance=provenance)
    assert workspace["documents"].document_import_resource(str(path),**args)
    with pytest.raises(ValueError,match="WORK_EFFECT_PERMISSION_DENIED"):
        workspace["documents"].import_document(str(path),**args)


def test_original_source_corruption_prevents_report_reopen(workspace):
    _, _, source = _source(workspace)
    citation = workspace["documents"].locate_quote(source["evidence_id"],actor=workspace["actor"],scope=workspace["scope"],start=0,end=20)["citation"]
    report = _save(workspace,citations=[citation])
    filename = workspace["repo"].query_one("SELECT relative_path FROM lumenis_shared_artifacts WHERE artifact_id=?",(source["source_artifact_ref"],))[0]
    (workspace["root"] / "artifacts" / filename).write_bytes(b"Changed original input")
    with pytest.raises(ValueError,match="DOCUMENT_SOURCE_ARTIFACT_INVALID"):
        workspace["work"].reopen_work(report["artifact"]["artifact_id"],actor=workspace["actor"],scope=workspace["scope"])


def test_artifact_collision_cannot_replace_or_delete_saved_bytes(workspace, monkeypatch):
    first = _save(workspace)
    file = workspace["root"] / "artifacts" / first["artifact"]["filename"]
    original = file.read_bytes()
    class FixedUUID:
        hex = first["artifact"]["artifact_id"][4:] + "0"*16
    monkeypatch.setattr(workspace["work"].uuid,"uuid4",lambda:FixedUUID())
    with pytest.raises(ValueError,match="WORK_ARTIFACT_COLLISION"):
        _save(workspace,request_id="collision-regression-0002")
    assert file.read_bytes()==original
    assert workspace["repo"].query_one("SELECT COUNT(*) FROM lumenis_shared_artifacts")[0]==1


def test_binary_xls_uses_legacy_decoder_and_retains_zero_boolean_and_coordinates(tmp_path):
    # Generate genuine BIFF8 directly; this tests xlrd, not Calc conversion.
    # The earlier soffice exit 137 is retained in macro_pipeline/xls-safe-recheck.log.
    import xlwt
    from lumenis.importer.parser import parse_file
    pytest.importorskip("xlrd", reason="XLS runtime must be installed for XLS acceptance")
    path = tmp_path / "owned.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Owned sheet")
    sheet.write(0, 0, "Own source statement")
    sheet.write(0, 1, 0)
    sheet.write(0, 2, False)
    sheet.write(1, 1, "Exact second-row source")
    book.save(str(path))
    assert path.read_bytes().startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    result = parse_file(str(path))
    assert result.tables[0][0] == ["Own source statement", "0", "False"]
    cell = next(item for item in result.locators if item["row"] == 2 and item["column"] == 2)
    assert cell["sheet"] == "Owned sheet"
    assert result.text[cell["start"]:cell["end"]] == "Exact second-row source"
    assert result.metadata["calculation_state"] == "retained_cached_values_not_recalculated"
    assert result.metadata["formula_expressions"] == "not_exposed_by_xlrd"
    assert result.metadata["source_fingerprint"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_binary_doc_decodes_owned_original_and_preserves_source_fingerprint(tmp_path):
    import shutil
    from docx import Document
    from lumenis.importer.parser import parse_file
    if not shutil.which("soffice"):
        pytest.skip("this owned binary DOC check requires the optional local LibreOffice decoder")
    source = tmp_path / "owned.docx"
    document = Document()
    document.add_paragraph("Original owned engineering statement.")
    document.add_paragraph("Repeated source statement.")
    document.add_paragraph("Repeated source statement.")
    document.save(source)
    result = subprocess.run([shutil.which("soffice"), "--headless", "--convert-to", "doc:MS Word 97",
                             "--outdir", str(tmp_path), str(source)],
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    path = tmp_path / "owned.doc"
    assert path.read_bytes().startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    parsed = parse_file(str(path))
    assert "Original owned engineering statement." in parsed.text
    repeated = [item for item in parsed.locators if
                parsed.text[item["start"]:item["end"]] == "Repeated source statement."]
    assert len(repeated) == 2 and repeated[0]["start"] != repeated[1]["start"]
    assert parsed.metadata["source_fingerprint"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert parsed.metadata["embedded_macros"] == "disabled_by_ephemeral_profile"
    assert parsed.metadata["external_links"] == "updates_disabled_by_ephemeral_profile"
    assert all(item["unit"] == "converted_word_paragraph" for item in repeated)


def test_saved_work_roundtrips_docx_xlsx_and_pptx_with_exact_staging_types(workspace):
    attachments = [_office_attachment(kind) for kind in ("docx", "xlsx", "pptx")]
    saved = _save(workspace, attachments=attachments)
    bindings = saved["document"]["attachments"]
    assert [binding["name"] for binding in bindings] == [item["name"] for item in attachments]

    reopened = workspace["work"].reopen_work(
        saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"],
    )
    assert reopened["document"]["attachments"] == bindings
    for binding, original in zip(bindings, attachments):
        record = workspace["repo"].query_one(
            "SELECT relative_path, media_type, sha256, size_bytes FROM lumenis_shared_artifacts WHERE artifact_id=?",
            (binding["artifact_id"],),
        )
        extension = "." + original["name"].rsplit(".", 1)[1]
        assert record["relative_path"] == binding["artifact_id"] + extension
        assert record["media_type"] == original["media_type"]
        assert record["sha256"] == binding["sha256"]
        assert record["size_bytes"] == len(original["content"])
        assert (workspace["root"] / "artifacts" / record["relative_path"]).read_bytes() == original["content"]

    replayed = _save(workspace, attachments=attachments)
    assert replayed["replayed"] is True
    assert replayed["artifact"]["artifact_id"] == saved["artifact"]["artifact_id"]
    assert workspace["repo"].query_one(
        "SELECT COUNT(*) FROM lumenis_shared_artifacts WHERE kind='saved_document_attachment'",
    )[0] == 3


def test_saved_report_artifact_download_reopens_the_existing_saved_work(workspace, monkeypatch):
    from lumenis.shared.reporting import export_api

    app = workspace["app"]
    app.register_blueprint(export_api.export_bp)
    monkeypatch.setattr(
        export_api, "_request_report_scope",
        lambda action="read": (workspace["actor"], workspace["scope"], [], None),
    )
    saved = _save(workspace, kind="report", title="Scoped business report",
                  payload={"report": {"summary": "Saved report content", "metrics": {"records": 2},
                                      "source_use_bindings": []}})

    with app.test_client() as client:
        response = client.get(f"/api/export/file/{saved['artifact']['artifact_id']}.json")

    assert response.status_code == 200
    assert response.headers["X-Lumenis-Saved"] == "true"
    assert response.headers["X-Lumenis-Content-SHA256"] == saved["artifact"]["sha256"]
    assert json.loads(response.data)["payload"]["report"]["summary"] == "Saved report content"
    reopened = workspace["work"].reopen_work(
        saved["artifact"]["artifact_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert reopened["document"]["version"] == 1


def test_saved_report_download_stops_when_current_source_use_is_revoked(workspace, monkeypatch):
    from lumenis.shared.reporting import export_api

    app = workspace["app"]
    app.register_blueprint(export_api.export_bp)
    monkeypatch.setattr(
        export_api, "_request_report_scope",
        lambda action="read": (workspace["actor"], workspace["scope"], [], None),
    )
    saved = _save(workspace, kind="report", title="Scoped business report",
                  payload={"report": {"summary": "Saved report content", "metrics": {"records": 2},
                                      "source_use_bindings": []}})
    checks = []

    def reject_current_source_use(_actor, _scope, _item):
        checks.append("revalidated")
        raise PermissionError("REPORT_SOURCE_USE_UNVERIFIED")

    monkeypatch.setattr(export_api, "_verify_saved_report_source_use", reject_current_source_use)

    with app.test_client() as client:
        response = client.get(f"/api/export/file/{saved['artifact']['artifact_id']}.json")

    assert checks == ["revalidated"]
    assert response.status_code == 403
    assert response.get_json()["code"] == "REPORT_SOURCE_USE_UNVERIFIED"


def test_saved_report_source_use_rechecks_current_canonical_and_declared_rights(monkeypatch):
    from contextlib import nullcontext

    import database
    from lumenis.shared import database as shared_database
    from lumenis.shared.reporting import business_brief, operational_brief, saved_work
    from lumenis.shared.reporting.export_api import _verify_saved_report_source_use

    scope = {"domain_id": "cth", "subject_id": "SUBJECT-A",
             "brand_focus_id": None, "global_scope": False}
    binding = {"source_type": "note", "source_id": "NOTE-A", "scope": scope,
               "scope_evidence_ref": "scope-evidence-v1", "content_evidence_id": 2,
               "content_evidence_ref": "app_evidence_records:2"}
    item = {"document": {"kind": "report", "citations": [], "payload": {"report": {
        "source_use_bindings": [binding],
        "source_details": [{"source_type": "note", "source_id": "NOTE-A", **scope,
                            "scope_evidence_ref": "scope-evidence-v1"}],
        "event_details": [],
    }}}}
    source_row = {"source_id": "NOTE-A", "source_type": "note",
                  "scope_evidence_ref": "scope-evidence-v1", "source_platform": "x",
                  "publish_date": None, "captured_at": "2026-10-14T00:00:00Z"}

    class Query:
        def fetchall(self):
            return [("canonical-raw-ref", "2026-10-14T00:00:00Z")]

    class Connection:
        def execute(self, _sql, _params):
            return Query()

    class Repository:
        def connection(self, **_kwargs):
            return nullcontext(Connection())

    monkeypatch.setattr(shared_database, "get_repository", lambda: Repository())
    monkeypatch.setattr(database, "_verified_event_source_refs", lambda *_args: (
        [], {("notes", "note", "NOTE-A"): dict(source_row)}
    ))
    monkeypatch.setattr(operational_brief, "_content_evidence_ref", lambda *_args: "app_evidence_records:2")
    monkeypatch.setattr(saved_work, "authorize", lambda *_args, **_kwargs: {"scope": scope})
    monkeypatch.setattr(saved_work, "_verified_evidence", lambda *_args: {
        "raw_content": "Exact source body", "asset_id": "NOTE-A",
    })
    uses = []

    def allow_source_use(material, actor, *, connection):
        uses.append((material, actor, connection))
        return [{"evidence_ref": "app_evidence_records:2"}]

    monkeypatch.setattr(business_brief, "_source_use_material", allow_source_use)
    assert _verify_saved_report_source_use(object(), scope, item) == 1
    assert uses and uses[-1][0]["scope"] == scope
    assert uses[-1][0]["sources"][0]["text"] == "Exact source body"

    def reject_revoked_source(_material, _actor, *, connection):
        raise PermissionError("REPORT_SOURCE_USE_UNVERIFIED")

    monkeypatch.setattr(business_brief, "_source_use_material", reject_revoked_source)
    with pytest.raises(PermissionError, match="REPORT_SOURCE_USE_UNVERIFIED"):
        _verify_saved_report_source_use(object(), scope, item)


def test_export_generate_saves_exact_rendition_and_returns_legacy_download_contract(workspace, monkeypatch):
    from lumenis.shared.reporting import export_api

    app = workspace["app"]
    app.register_blueprint(export_api.export_bp)
    monkeypatch.setattr(
        export_api, "_request_report_scope",
        lambda action="read": (workspace["actor"], workspace["scope"], [], None),
    )
    report = {
        "brand_name": "Scoped business",
        "period_label": "日报",
        "start_date": "2026-10-02",
        "end_date": "2026-10-02",
        "executive_summary": "A saved scoped report.",
        "key_metrics": [{"name": "Records", "value": 2, "unit": "items"}],
        "source_use_bindings": [],
    }
    monkeypatch.setattr(
        export_api, "_build_report",
        lambda kind, start, end, **_kwargs: (kind, start, end, report, 0, [], []),
    )
    monkeypatch.setattr(export_api, "_report_citations", lambda *_args: [])
    monkeypatch.setattr(export_api, "_verify_report_source_use", lambda *_args: 0)
    rendition = _office_attachment("xlsx")
    monkeypatch.setattr(export_api, "_render_report", lambda _report, fmt: (
        rendition["content"], rendition["media_type"]
    ) if fmt == "xlsx" else pytest.fail("unexpected report format"))

    with app.test_client() as client:
        generated = client.post("/api/export/generate?type=daily&date=2026-10-02&format=xlsx")
        body = generated.get_json()
        assert generated.status_code == 200, body
        assert body["ok"] is True and body["saved"] is True
        assert body["download_url"].endswith(f"{body['artifact_id']}.xlsx")
        assert body["size"] == len(rendition["content"])
        assert body["artifact_hash"] == hashlib.sha256(rendition["content"]).hexdigest()
        assert body["persistence_state"] == "SAVEDWORK_COMMITTED_AND_REOPENED"

        downloaded = client.get(body["download_url"])
        listed = client.get("/api/export/list")

    assert downloaded.status_code == 200
    assert downloaded.data == rendition["content"]
    assert downloaded.headers["X-Lumenis-Saved"] == "true"
    assert listed.status_code == 200
    assert any(item["artifact"]["artifact_id"] == body["artifact_id"]
               for item in listed.get_json()["items"])
    reopened = workspace["work"].reopen_work(
        body["artifact_id"], actor=workspace["actor"], scope=workspace["scope"])
    assert reopened["document"]["payload"]["report"] == report
    assert reopened["document"]["attachments"][0]["sha256"] == body["artifact_hash"]


def test_report_spreadsheets_preserve_all_fields_and_render_formulas_as_text():
    import csv

    from openpyxl import load_workbook

    from lumenis.shared.reporting.export_api import _render_report

    report = {
        "report_id": "REPORT-SAFE-1",
        "scope": {"domain_id": "cth", "subject_id": "SUBJECT-A"},
        "sources": [{"title": "=1+1", "content": "+SUM(A1:A3)",
                     "url": "https://example.invalid/source"}],
        "risk_events": [{"summary": "@user text", "note": "-cmd"}],
        "appendix": {"unknown_dimension": None, "collected_zero": 0,
                     "negative_integer": -1, "negative_fraction": -0.5},
    }

    raw_csv, _ = _render_report(report, "csv")
    rows = list(csv.reader(io.StringIO(raw_csv.decode("utf-8-sig"))))
    csv_cells = {row[1]: row[2] for row in rows[1:] if len(row) == 3}
    assert csv_cells["sources[0].title"] == "'=1+1"
    assert csv_cells["sources[0].content"] == "'+SUM(A1:A3)"
    assert csv_cells["sources[0].url"] == "https://example.invalid/source"
    assert csv_cells["unknown_dimension"] == "null"
    assert csv_cells["collected_zero"] == "0"
    assert csv_cells["negative_integer"] == "-1"
    assert csv_cells["negative_fraction"] == "-0.5"

    raw_xlsx, _ = _render_report(report, "xlsx")
    sheet = load_workbook(io.BytesIO(raw_xlsx), data_only=False).active
    xlsx_cells = {sheet.cell(row, 2).value: sheet.cell(row, 3)
                  for row in range(2, sheet.max_row + 1)}
    assert xlsx_cells["sources[0].title"].data_type == "s"
    assert xlsx_cells["sources[0].title"].value == "'=1+1"
    assert xlsx_cells["sources[0].content"].data_type == "s"
    assert xlsx_cells["sources[0].content"].value == "'+SUM(A1:A3)"
    assert xlsx_cells["sources[0].url"].value == "https://example.invalid/source"
    assert xlsx_cells["appendix.unknown_dimension"].value == "null"
    assert xlsx_cells["appendix.collected_zero"].value == 0
    assert xlsx_cells["appendix.negative_integer"].value == -1
    assert xlsx_cells["appendix.negative_fraction"].value == -0.5


def test_report_xlsx_retains_verified_source_and_event_detail_sheets():
    from openpyxl import load_workbook

    from lumenis.shared.reporting.export_api import (
        _attach_report_details,
        _render_report,
    )

    report = _attach_report_details(
        {"key_metrics": [{"name": "Records", "value": 1}]},
        [{"source_type": "note", "source_id": "N-1", "title": "Observed source",
          "publish_date": "2026-10-01", "views": None, "likes": 3,
          "comments": 2, "collects": 3, "shares": 4}],
        [{"id": "E-1", "title": "Observed event", "risk_level": "medium",
          "mention_count": 1, "summary": "Verified event summary",
          "sources": [{"source_type": "note", "source_id": "N-1", "title": "Observed source"}]}],
    )
    raw, _ = _render_report(report, "xlsx")
    book = load_workbook(io.BytesIO(raw), data_only=False)
    assert {"完整报告", "指标", "来源明细", "事件明细"}.issubset(book.sheetnames)

    source_sheet = book["来源明细"]
    source_values = [cell.value for row in source_sheet.iter_rows() for cell in row]
    assert "N-1" in source_values and "Observed source" in source_values
    assert "1.views" in source_values and "1.likes" in source_values
    assert "1.engagement" in source_values and 12 in source_values

    event_sheet = book["事件明细"]
    event_values = [cell.value for row in event_sheet.iter_rows() for cell in row]
    assert "E-1" in event_values and "Observed event" in event_values
    assert "Verified event summary" in event_values and "1.mention_count" in event_values


def test_report_error_exposes_only_known_codes_and_never_unknown_exception_text():
    from lumenis.shared.reporting.export_api import _report_error

    secret_path = "D:/private/customer-production.db"
    with Flask(__name__).app_context():
        safe_response, safe_status = _report_error(ValueError("SOURCE_SCOPE_EVIDENCE_REQUIRED"))
        assert safe_status == 409
        assert safe_response.get_json()["code"] == "SOURCE_SCOPE_EVIDENCE_REQUIRED"

        unsafe_response, unsafe_status = _report_error(RuntimeError(f"failed while reading {secret_path}"))
        assert unsafe_status == 500
        assert unsafe_response.get_json()["code"] == "REPORT_EXPORT_FAILED"
        assert secret_path not in unsafe_response.get_data(as_text=True)


def test_saved_work_rejects_mislabeled_malformed_and_unsafe_office_attachments(workspace):
    valid_docx = _office_attachment("docx")
    wrong_mime = {**valid_docx, "media_type": _TEST_OFFICE_FORMATS["xlsx"][0]}
    wrong_suffix = {**valid_docx, "name": "sample.xlsx"}
    not_a_zip = {**valid_docx, "content": b"not an OOXML ZIP"}
    wrong_main = _office_attachment("docx", main_xml=b"<document/>")
    traversal = _office_attachment("docx", extra_members=(("../outside.xml", b"<part/>"),))
    bomb = _office_attachment("docx", padding=b"A" * (workspace["work"]._LIMIT + 1))
    assert len(bomb["content"]) < workspace["work"]._LIMIT

    encrypted = _office_attachment("xlsx")
    encrypted_bytes = bytearray(encrypted["content"])
    local_header = encrypted_bytes.find(b"PK\x03\x04")
    central_header = encrypted_bytes.find(b"PK\x01\x02")
    assert local_header >= 0 and central_header >= 0
    encrypted_bytes[local_header + 6] |= 1
    encrypted_bytes[central_header + 8] |= 1
    encrypted["content"] = bytes(encrypted_bytes)

    corrupt_crc = _office_attachment("pptx", compression=zipfile.ZIP_STORED)
    main_xml = _TEST_OFFICE_FORMATS["pptx"][3]
    corrupt_bytes = bytearray(corrupt_crc["content"])
    main_offset = corrupt_bytes.find(main_xml)
    assert main_offset >= 0
    corrupt_bytes[main_offset + 10] ^= 1
    corrupt_crc["content"] = bytes(corrupt_bytes)

    unsafe_name = {**valid_docx, "name": "../sample.docx"}
    for index, item in enumerate((wrong_mime, wrong_suffix, not_a_zip, wrong_main, traversal,
                                  bomb, encrypted, corrupt_crc, unsafe_name)):
        with pytest.raises(ValueError, match="WORK_ATTACHMENT_INVALID"):
            _save(workspace, request_id=f"invalid-{index:02d}", attachments=[item])
    assert workspace["repo"].query_one(
        "SELECT COUNT(*) FROM lumenis_shared_artifacts WHERE kind IN ('saved_document', 'saved_document_attachment')",
    )[0] == 0
    artifact_root = workspace["root"] / "artifacts"
    assert not artifact_root.exists() or not list(artifact_root.glob("ART-*"))


def test_saved_work_retains_attachment_count_and_aggregate_size_limits(workspace):
    nine = [_office_attachment("docx") | {"name": f"sample-{index}.docx"} for index in range(9)]
    with pytest.raises(ValueError, match="WORK_ATTACHMENT_INVALID"):
        _save(workspace, attachments=nine)

    large = [_office_attachment("xlsx", padding=b"x" * 4_300_000, compression=zipfile.ZIP_STORED)
             | {"name": f"large-{index}.xlsx"} for index in range(2)]
    assert all(len(item["content"]) < workspace["work"]._LIMIT for item in large)
    assert sum(len(item["content"]) for item in large) > workspace["work"]._LIMIT
    with pytest.raises(ValueError, match="WORK_ATTACHMENT_INVALID"):
        _save(workspace, request_id="aggregate-limit-0001", attachments=large)
    assert workspace["repo"].query_one(
        "SELECT COUNT(*) FROM lumenis_shared_artifacts WHERE kind IN ('saved_document', 'saved_document_attachment')",
    )[0] == 0


def test_attachment_staging_rolls_back_all_artifacts_on_precommit_failure(workspace):
    attachment = _office_attachment("pptx")

    def fail_before_commit(*_args):
        raise RuntimeError("isolated commit interruption")

    with pytest.raises(RuntimeError, match="isolated commit interruption"):
        _save(workspace, attachments=[attachment], before_commit=fail_before_commit)
    assert workspace["repo"].query_one(
        "SELECT COUNT(*) FROM lumenis_shared_artifacts WHERE kind IN ('saved_document', 'saved_document_attachment')",
    )[0] == 0
    artifact_root = workspace["root"] / "artifacts"
    assert not artifact_root.exists() or not list(artifact_root.iterdir())
