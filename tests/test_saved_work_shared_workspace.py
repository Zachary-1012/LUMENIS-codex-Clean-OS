"""Isolated regressions for actor-owned SavedWork in the Shared workspace.

These tests exercise service, Permission, account-table and Artifact continuity.
They are not production or real-user E2E evidence.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json

import pytest
from flask import Flask

from lumenis.shared.database import DatabaseRepository
from lumenis.system.security import permission


SHARED_SCOPE = {"scope_kind": "shared_workspace", "version": 1}
_SHARED_GRANTS = ("SHARED", "reports")
_EXPORT_OPERATION = permission.OperationType.DATA_EXPORT.value


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    """Use one disposable canonical SQLite repository and an actual Gate."""
    db_path = tmp_path / "isolated-shared-work.sqlite3"
    artifact_dir = tmp_path / "artifacts"
    monkeypatch.setenv("ZBC_DB_PATH", str(db_path))
    monkeypatch.setenv("ZBC_ARTIFACT_DIR", str(artifact_dir))
    monkeypatch.setenv(
        "LUMENIS_EVIDENCE_HMAC_KEY_RING",
        json.dumps({"isolated-test": "isolated-test-evidence-key"}),
    )
    monkeypatch.setenv("LUMENIS_EVIDENCE_HMAC_ACTIVE_KEY_ID", "isolated-test")
    monkeypatch.delenv("PERMISSION_GATE_TESTING", raising=False)
    monkeypatch.delenv("LUMENIS_BOOTSTRAP_OWNER_USERNAME", raising=False)
    monkeypatch.delenv("LUMENIS_BOOTSTRAP_OWNER_PASSWORD", raising=False)
    monkeypatch.setattr(permission, "_TESTING_BYPASS_ENABLED", False)

    import account.models as account_models
    import canonical_data
    import shared_artifacts
    import subject_registry
    from lumenis.importer import documents
    from lumenis.shared.evidence import store
    from lumenis.shared.reporting import saved_work

    repository = DatabaseRepository(str(db_path))
    for module in (account_models, shared_artifacts, store, documents, saved_work):
        monkeypatch.setattr(module, "get_repository", lambda repo=repository: repo)
    if hasattr(subject_registry, "get_repository"):
        monkeypatch.setattr(subject_registry, "get_repository", lambda: repository)

    shared_artifacts.ensure_artifact_table()
    store.init_evidence_store(str(db_path))
    with repository.transaction() as con:
        store.ensure_evidence_schema(con)
    canonical_data.ensure_canonical_table(str(db_path))
    account_models.init_account_db()

    gate = permission.PermissionGate()
    gate.grant_permission(permission.PermissionLevel.COMMIT)
    monkeypatch.setattr(permission, "get_permission_gate", lambda: gate)

    app = Flask("saved-work-shared-workspace-tests")
    app.config["LUMENIS_OPEN_WORKSPACE"] = True
    with app.test_request_context(
        "/",
        headers={"Host": "localhost"},
        environ_overrides={"REMOTE_ADDR": "127.0.0.1"},
    ):
        local_actor = permission.resolve_local_workspace_principal(
            module="reports", action="write",
        )
        assert type(local_actor) is permission.LocalWorkspacePrincipal
        yield {
            "root": tmp_path,
            "artifact_dir": artifact_dir,
            "db_path": db_path,
            "repo": repository,
            "gate": gate,
            "app": app,
            "local_actor": local_actor,
            "documents": documents,
            "work": saved_work,
            "account_models": account_models,
            "actor_counter": 0,
        }


def _new_account(workspace, *, actions=("read", "write")):
    """Create an active member and actual Shared reports grants in the test DB."""
    workspace["actor_counter"] += 1
    username = f"shared-member-{workspace['actor_counter']}"
    with workspace["repo"].transaction() as con:
        cursor = con.execute(
            "INSERT INTO v4_users (username,password_hash,role,display_name,status) "
            "VALUES (?,?,?,?,?)",
            (username, "test-only-password-hash", "user", username, "active"),
        )
        actor_id = int(cursor.lastrowid)
        for action in actions:
            con.execute(
                "INSERT INTO v4_member_authorizations "
                "(user_id,product_id,module,action,data_scope,time_range,granted_by) "
                "VALUES (?,?,?,?,?,?,?)",
                (actor_id, *_SHARED_GRANTS, action, None, None, None),
            )
    return {"id": actor_id}


def _save_args(
    workspace,
    *,
    actor=None,
    scope=None,
    request_id="shared-save-request-0001",
    title="Workspace research note",
    markdown="A user-authored note for isolated Shared SavedWork coverage.",
    payload=None,
    citations=None,
    work_id=None,
    previous_artifact_id=None,
    expected_draft_digest=None,
    contract=None,
    task_id=None,
    assertion_id=None,
    attachments=None,
):
    return {
        "title": title,
        "markdown": markdown,
        "kind": "research",
        "payload": {"origin": "direct_user_brief"} if payload is None else payload,
        "citations": [] if citations is None else citations,
        "request_id": request_id,
        "actor": workspace["local_actor"] if actor is None else actor,
        "scope": SHARED_SCOPE if scope is None else scope,
        "work_id": work_id,
        "previous_artifact_id": previous_artifact_id,
        "expected_draft_digest": expected_draft_digest,
        "contract": contract,
        "task_id": task_id,
        "assertion_id": assertion_id,
        "attachments": attachments,
    }


def _grant_exact_save_export(workspace, args):
    resource = workspace["work"].work_write_resource(**args)
    workspace["gate"].grant_approval(
        _EXPORT_OPERATION,
        resource=resource,
        granted_by="isolated-shared-work-human-approval",
    )
    return resource


def _grant_exact_work_export(workspace, saved, *, format="json"):
    descriptor = {
        "id": saved["artifact"]["artifact_id"],
        "sha256": saved["artifact"]["sha256"],
        "format": format,
    }
    raw = json.dumps(
        descriptor, ensure_ascii=False, allow_nan=False,
        sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")
    resource = "saved-work-export:v1:" + hashlib.sha256(raw).hexdigest()
    workspace["gate"].grant_approval(
        _EXPORT_OPERATION,
        resource=resource,
        granted_by="isolated-shared-work-human-approval",
    )
    return resource


def _revision_save_args(
    workspace, saved, *, request_id, markdown, reason, attachments=None,
):
    document = saved["document"]
    artifact = saved["artifact"]
    edit = {
        "pointer": "/markdown",
        "expected": document["markdown"],
        "value": markdown,
    }
    payload = copy.deepcopy(document["payload"])
    payload["revision_note"] = {
        "schema": "lumenis.user-work-local-revision/v1",
        "reason": reason,
        "parent": {
            "artifact_id": artifact["artifact_id"],
            "sha256": artifact["sha256"],
            "version": document["version"],
        },
        "edits": [edit],
        "source_adoptions": [],
        "offset_unit": "unicode_codepoint",
        "source_relation": "explicit user selection; not an independent correction of source Evidence",
    }
    return _save_args(
        workspace,
        actor=workspace["actor_for_revision"],
        scope=SHARED_SCOPE,
        request_id=request_id,
        title=document["title"],
        markdown=markdown,
        payload=payload,
        citations=document["citations"],
        work_id=document["work_id"],
        previous_artifact_id=artifact["artifact_id"],
        expected_draft_digest=document["draft_digest"],
        contract=document["contract"],
        task_id=document.get("task_id"),
        assertion_id=document.get("assertion_id"),
        attachments=attachments,
    )


def _artifact_state(workspace):
    with workspace["repo"].connection(read_only=True, wal=False) as con:
        rows = con.execute(
            "SELECT artifact_id,kind,product_id,domain_id,source_execution_id,evidence_id,note_id,"
            "relative_path,sha256,status FROM lumenis_shared_artifacts ORDER BY artifact_id"
        ).fetchall()
    files = ()
    if workspace["artifact_dir"].exists():
        files = tuple(sorted(path.name for path in workspace["artifact_dir"].iterdir()))
    return tuple(tuple(row) for row in rows), files


def _source(workspace, *, scope, actor=None):
    content = "Synthetic source header.\nQuoted source line for permission checks.\n"
    path = workspace["root"] / "synthetic-source.txt"
    path.write_text(content, encoding="utf-8")
    provenance = {
        "source_url": None,
        "published_at": None,
        "captured_at": "2026-10-11T00:00:00Z",
        "rights": "owner-approved",
        "title": "Synthetic isolated test source",
        "intended_use": "research_reference",
        "source_origin": {
            "kind": "approved-file",
            "ref": "shared-workspace-source-" + hashlib.sha256(
                content.encode("utf-8"),
            ).hexdigest(),
        },
        "content_origin": {"kind": "self-authored", "claim_type": "synthetic_sample"},
        "rights_statement": "Self-authored synthetic test input; no external-source claim.",
    }
    source_actor = workspace["local_actor"] if actor is None else actor
    args = {"actor": source_actor, "scope": scope, "provenance": provenance}
    resource = workspace["documents"].document_import_resource(str(path), **args)
    workspace["gate"].grant_approval(
        permission.OperationType.DATA_IMPORT.value,
        resource=resource,
        granted_by="isolated-source-import-approval",
    )
    imported = workspace["documents"].import_document(str(path), **args)
    quote = "Quoted source line for permission checks."
    start = content.index(quote)
    located = workspace["documents"].locate_quote(
        imported["evidence_id"], actor=source_actor, scope=scope,
        start=start, end=start + len(quote),
    )
    return path, imported, located


def _focused_cth_scope(workspace):
    from subject_registry import add_subject

    subject, _created = add_subject({
        "domain_id": "CTH",
        "subject_type": "company",
        "name_en": "Isolated Shared workspace citation subject",
        "name_zh": "隔离共享工作区引用主体",
    })
    return {
        "domain_id": "cth",
        "subject_id": subject["subject_id"],
        "brand_focus_id": None,
        "global_scope": False,
    }


def _grant_account_scope(workspace, actor, *, product_id, scope_key, scope_id):
    data_scope = json.dumps(
        {scope_key: [scope_id]}, ensure_ascii=False, sort_keys=True,
        separators=(",", ":"),
    )
    with workspace["repo"].transaction() as con:
        for action in ("read", "write"):
            con.execute(
                "INSERT INTO v4_member_authorizations "
                "(user_id,product_id,module,action,data_scope,time_range,granted_by) "
                "VALUES (?,?,?,?,?,?,?)",
                (actor["id"], product_id, "reports", action, data_scope, None, None),
            )


def _grant_global_account_scope(workspace, actor, *, product_id):
    data_scope = json.dumps(
        {"global": True}, ensure_ascii=False, sort_keys=True,
        separators=(",", ":"),
    )
    with workspace["repo"].transaction() as con:
        for action in ("read", "write"):
            con.execute(
                "INSERT INTO v4_member_authorizations "
                "(user_id,product_id,module,action,data_scope,time_range,granted_by) "
                "VALUES (?,?,?,?,?,?,?)",
                (actor["id"], product_id, "reports", action, data_scope, None, None),
            )


def _seed_approved_source_use_fixture(workspace, monkeypatch, actor):
    """Seed test-only SourceUse inputs while retaining the real validators."""
    import canonical_data
    import subject_registry
    from canonical_data import CanonicalRecord, persist_canonical
    from lumenis.shared.database.business_schema import apply_business_schema
    from lumenis.shared.evidence import store

    subject, _created = subject_registry.add_subject({
        "domain_id": "CTH",
        "subject_type": "company",
        "name_en": "Synthetic Shared export denial source",
        "name_zh": "共享成果来源导出拒绝合成主体",
    })
    source_scope = {
        "domain_id": "cth", "subject_id": subject["subject_id"],
        "brand_focus_id": None, "global_scope": False,
    }
    _grant_account_scope(
        workspace, actor, product_id="CTH", scope_key="subject_ids",
        scope_id=source_scope["subject_id"],
    )
    context = workspace["work"].authorize(actor, source_scope, action="read")
    owner = {
        key: value for key, value in context.items()
        if key not in {"domain_id", "subject_id", "brand_focus_id", "global_scope"}
    }
    source_id = "shared-source-use-export-test-20261011"
    canonical_id = "canonical-shared-source-use-export-test-20261011"
    source_url = "https://example.invalid/shared-source-use-export-test"
    body = (
        "Synthetic source text includes the exact phrase internal-use citation "
        "for a permission regression."
    )
    raw = body.encode("utf-8")
    body_sha256 = hashlib.sha256(raw).hexdigest()

    with workspace["repo"].transaction() as con:
        apply_business_schema(con)
        canonical_data.ensure_canonical_table(connection=con)
        scope_evidence_id = store.record_evidence(
            connection=con,
            source="isolated-test-canonical-scope",
            project="LUMENIS",
            product="CTH",
            feature_id="SHARED_SOURCE_USE_SCOPE_FIXTURE",
            asset_id=source_id,
            version="shared-source-scope:test-v1",
            test_name="shared-source-use-export-scope-fixture",
            result="PASS",
            status="verified",
            inputs={"source_ids": [source_id]},
            outputs={"scope_fixture": "test-only"},
            domain_id="cth",
            subject_id=source_scope["subject_id"],
            global_scope=False,
        )
        artifact, artifact_path = workspace["work"]._stage_bytes(
            con,
            raw=raw,
            context=context,
            kind="imported_document_source",
            source_execution_id=source_id,
            extension=".txt",
            media_type="text/plain; charset=utf-8",
        )
        metadata = {
            "context": context,
            "owner": owner,
            "source_url": source_url,
            "source_artifact_ref": artifact["artifact_id"],
            "source_fingerprint": body_sha256,
            "parsed_text": body,
            "locators": [{
                "unit": "fixture-original-text",
                "start": 0,
                "end": len(body),
                "offset_unit": "unicode-codepoint",
            }],
            "provenance": {
                "source_url": source_url,
                "published_at": "UNKNOWN",
                "captured_at": "UNKNOWN",
                "rights": "UNKNOWN",
                "intended_use": "research_reference",
                "source_kind": "original_source_evidence",
                "content_origin": {
                    "kind": "publisher-material",
                    "claim_type": "source_statement",
                },
            },
        }
        content_evidence_id = store.record_evidence(
            connection=con,
            source="isolated-test-approved-source",
            project="LUMENIS",
            product="CTH",
            feature_id="SHARED_SOURCE_USE_CONTENT_FIXTURE",
            asset_id=source_id,
            version="shared-source-content:test-v1",
            test_name="shared-source-use-export-content-fixture",
            result="PASS",
            status="verified",
            inputs={"source_ids": [source_id]},
            outputs={"source_artifact_ref": artifact["artifact_id"]},
            raw_content=body,
            collection_metadata=metadata,
            domain_id="cth",
            subject_id=source_scope["subject_id"],
            global_scope=False,
        )
        con.execute(
            "UPDATE lumenis_shared_artifacts SET evidence_id=? WHERE artifact_id=?",
            (str(content_evidence_id), artifact["artifact_id"]),
        )
        content_locator = f"app_evidence_records:{content_evidence_id}"
        scope_locator = f"app_evidence_records:{scope_evidence_id}"
        persist_canonical(con, CanonicalRecord(
            canonical_id=canonical_id,
            product_id="CTH",
            domain_id="cth",
            subject_id=source_scope["subject_id"],
            scope_status="VERIFIED",
            scope_evidence_ref=str(scope_evidence_id),
            source_id=source_id,
            external_id=source_id,
            source_type="collected_note",
            source_table="app_evidence_records",
            title="Synthetic internal source-use fixture",
            content=body,
            url=source_url,
            canonical_url=source_url,
            source_platform="UNKNOWN",
            publish_datetime="",
            captured_at="",
            raw_ref=json.dumps({
                "source_id": source_id,
                "scope_evidence_ref": scope_locator,
                "content_evidence_ref": content_locator,
                "source_artifact_ref": artifact["artifact_id"],
            }, ensure_ascii=False, sort_keys=True),
        ))

    entry = {
        "receipt_id": "PO-20261011-APPROVED-INTERNAL-SOURCE-USE",
        "source_id": source_id,
        "source_url": source_url,
        "canonical_id": canonical_id,
        "domain_id": "cth",
        "subject_id": source_scope["subject_id"],
        "brand_focus_id": None,
        "scope_evidence_ref": scope_locator,
        "content_evidence_ref": content_locator,
        "body_sha256": body_sha256,
        "snapshot_kind": "existing_source_evidence_snapshot",
        "allowed_purposes": ["internal_research", "short_quote", "cth_brief"],
        "allowed_operations": ["derive_reference", "read", "save", "revise", "reopen"],
        "rights_state": "UNKNOWN",
        "commercial_state": "UNKNOWN",
        "publication_state": "NOT_APPROVED",
        "training_state": "NOT_APPROVED",
    }
    receipt = {
        "schema": "lumenis.po-approval-receipt/v1",
        "source_use_schema": "lumenis.approved-internal-source-use/v1",
        "approval_id": entry["receipt_id"],
        "approval_source_type": "direct_product_owner_instruction",
        "approved_by": "Product Owner",
        "approved_at": "2026-10-11",
        "status": "APPROVED",
        "receipt_id": entry["receipt_id"],
        "constraints": ["test-only synthetic input in an isolated unit test"],
        "scope_manifest": {"source_id": source_id},
        "sources": [entry],
        "quote_limit": {
            "max_unique_unicode_codepoints_per_source_per_saved_work": 160,
            "measurement": "unicode-codepoints",
            "scope": "this synthetic test receipt only",
        },
    }

    def load_test_receipt(receipt_id):
        assert receipt_id == receipt["receipt_id"]
        return receipt

    monkeypatch.setattr(permission, "_load_source_use_receipt", load_test_receipt)
    return {
        "source_id": source_id,
        "canonical_id": canonical_id,
        "source_url": source_url,
        "body": body,
        "body_sha256": body_sha256,
        "scope_evidence_id": scope_evidence_id,
        "scope_locator": scope_locator,
        "content_evidence_id": content_evidence_id,
        "content_locator": content_locator,
        "artifact": artifact,
        "artifact_path": artifact_path,
        "scope": source_scope,
        "entry": entry,
        "receipt": receipt,
    }


def _confirm_and_create_source_reference(workspace, source, actor, request_id):
    request = {
        "source_id": source["source_id"],
        "canonical_id": source["canonical_id"],
        "expected_body_sha256": source["body_sha256"],
        "request_id": request_id,
    }
    pending = permission.prepare_approved_source_reference_confirmation(
        actor=actor, scope=source["scope"], request=request, gate=workspace["gate"],
    )
    assert pending["allowed"] is True
    assert pending["code"] == "APPROVED_SOURCE_REFERENCE_CONFIRMATION_PENDING"
    confirmed = permission.confirm_approved_source_reference_confirmation(
        actor=actor, scope=source["scope"], confirmation_id=pending["confirmation_id"],
        confirmed=True, gate=workspace["gate"],
    )
    assert confirmed["allowed"] is True
    assert confirmed["code"] == "APPROVED_SOURCE_REFERENCE_READY_FOR_RETRY"
    assert confirmed["resource"] == pending["resource"]
    assert "approval_token" not in confirmed
    created = workspace["documents"].create_approved_source_reference(
        source_id=source["source_id"],
        canonical_id=source["canonical_id"],
        expected_body_sha256=source["body_sha256"],
        request_id=request_id,
        actor=actor,
        scope=source["scope"],
    )
    return pending, confirmed, created


def _actual_ooxml_attachment(kind):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    output = io.BytesIO()
    if kind == "docx":
        document = Document()
        document.add_paragraph("Synthetic editable test attachment.")
        document.save(output)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif kind == "xlsx":
        workbook = Workbook()
        workbook.active["A1"] = "Synthetic editable test attachment"
        workbook.save(output)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif kind == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, 4_000_000, 1_000_000)
        shape.text = "Synthetic editable test attachment"
        presentation.save(output)
        media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    else:
        raise AssertionError(f"unsupported test package: {kind}")
    return {
        "name": f"shared-test-{kind}.{kind}",
        "media_type": media_type,
        "content": output.getvalue(),
    }


def _attachment_bytes(workspace, saved, *, actor):
    with workspace["repo"].connection(read_only=True, wal=False) as con:
        return [
            workspace["work"]._read_attachment(
                con, binding, saved["document"]["context"], saved["document"]["work_id"],
                actor=actor,
            )
            for binding in saved["document"].get("attachments", [])
        ]


def test_shared_saved_work_save_replay_read_revise_search_and_exact_export(workspace):
    actor = _new_account(workspace)
    workspace["actor_for_revision"] = actor
    args = _save_args(workspace, actor=actor)
    resource = _grant_exact_save_export(workspace, args)

    first = workspace["work"].save_work(**args)
    assert first["replayed"] is False
    expected_context = {
        "scope_kind": "shared_workspace", "version": 1, "actor_id": actor["id"],
    }
    assert workspace["work"].authorize_saved_work(
        actor, SHARED_SCOPE, action="read",
    ) == expected_context
    assert first["document"]["context"] == expected_context
    assert first["return_ref"]["scope"] == SHARED_SCOPE
    assert resource.startswith("saved-work:v1:")

    row = workspace["repo"].query_one(
        "SELECT kind,product_id,domain_id,source_execution_id,evidence_id,note_id "
        "FROM lumenis_shared_artifacts WHERE artifact_id=?",
        (first["artifact"]["artifact_id"],),
    )
    assert tuple(row) == (
        "saved_document", "SHARED", "content", first["document"]["work_id"], "", "",
    )
    before_replay = _artifact_state(workspace)
    replay = workspace["work"].save_work(**args)
    assert replay["replayed"] is True
    assert replay["artifact"]["artifact_id"] == first["artifact"]["artifact_id"]
    assert _artifact_state(workspace) == before_replay

    read = workspace["work"].read_committed_work(
        first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
    )
    reopened = workspace["work"].reopen_work(
        first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
    )
    assert read["document"] == first["document"]
    assert reopened["document"] == first["document"]
    assert read["return_ref"]["scope"] == SHARED_SCOPE

    revised_markdown = "The user explicitly revised this Shared workspace note."
    reason = "Correct the user-authored draft after review."
    revised_args = _revision_save_args(
        workspace, first, request_id="shared-revision-request-0002",
        markdown=revised_markdown, reason=reason,
    )
    _grant_exact_save_export(workspace, revised_args)
    revised = workspace["work"].revise_work(
        first["artifact"]["artifact_id"],
        actor=actor,
        scope=SHARED_SCOPE,
        request_id="shared-revision-request-0002",
        expected_sha256=first["artifact"]["sha256"],
        reason=reason,
        edits=[{
            "pointer": "/markdown",
            "expected": first["document"]["markdown"],
            "value": revised_markdown,
        }],
    )
    assert revised["document"]["version"] == 2
    assert revised["document"]["parent_artifact_id"] == first["artifact"]["artifact_id"]
    assert revised["document"]["markdown"] == revised_markdown
    assert revised["document"]["payload"]["revision_note"]["reason"] == reason
    assert revised["return_ref"]["scope"] == SHARED_SCOPE

    found = workspace["work"].search_work(
        actor=actor, scope=SHARED_SCOPE, query="explicitly revised",
    )
    assert [item["artifact"]["artifact_id"] for item in found["items"]] == [
        revised["artifact"]["artifact_id"],
    ]

    _grant_exact_work_export(workspace, revised)
    exported = workspace["work"].export_work(
        revised["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE, format="json",
    )
    assert json.loads(exported["content"])["version"] == 2
    assert exported["return_ref"]["scope"] == SHARED_SCOPE
    assert hashlib.sha256(exported["content"]).hexdigest() == exported["sha256"]


def test_shared_saved_work_is_account_owned_and_rejects_foreign_cursor_and_work_id(workspace):
    first_actor = _new_account(workspace)
    second_actor = _new_account(workspace)
    first_args = _save_args(workspace, actor=first_actor, request_id="shared-owner-a-first")
    first_resource = _grant_exact_save_export(workspace, first_args)
    first = workspace["work"].save_work(**first_args)
    assert first_resource.startswith("saved-work:v1:")

    second_args = _save_args(
        workspace, actor=first_actor, request_id="shared-owner-a-second",
        title="Second actor-owned workspace note",
    )
    _grant_exact_save_export(workspace, second_args)
    second = workspace["work"].save_work(**second_args)

    assert workspace["work"].authorize_saved_work(
        second_actor, SHARED_SCOPE, action="read",
    ) == {
        "scope_kind": "shared_workspace", "version": 1, "actor_id": second_actor["id"],
    }
    before_foreign_access = _artifact_state(workspace)
    with pytest.raises(ValueError, match="WORK_CONTEXT_MISMATCH"):
        workspace["work"].reopen_work(
            first["artifact"]["artifact_id"], actor=second_actor, scope=SHARED_SCOPE,
        )
    foreign_search = workspace["work"].search_work(
        actor=second_actor, scope=SHARED_SCOPE, query="workspace note",
    )
    assert foreign_search["items"] == []
    assert _artifact_state(workspace) == before_foreign_access

    first_page = workspace["work"].search_work(
        actor=first_actor, scope=SHARED_SCOPE, limit=1,
    )
    assert first_page["has_more"] is True
    assert first_page["next_cursor"]
    with pytest.raises(ValueError, match="WORK_CURSOR_INVALID"):
        workspace["work"].search_work(
            actor=second_actor, scope=SHARED_SCOPE, limit=1,
            cursor=first_page["next_cursor"],
        )

    stolen_work_args = _save_args(
        workspace,
        actor=second_actor,
        request_id="shared-owner-b-cross-work-id",
        work_id=first["document"]["work_id"],
        previous_artifact_id=first["artifact"]["artifact_id"],
        expected_draft_digest=first["document"]["draft_digest"],
    )
    _grant_exact_save_export(workspace, stolen_work_args)
    before_cross_owner_write = _artifact_state(workspace)
    with pytest.raises(ValueError, match="WORK_CONTEXT_MISMATCH"):
        workspace["work"].save_work(**stolen_work_args)
    assert _artifact_state(workspace) == before_cross_owner_write
    assert second["document"]["context"]["actor_id"] == first_actor["id"]


@pytest.mark.parametrize(
    "bad_scope",
    [
        None,
        [],
        {},
        {"scope_kind": "shared_workspace"},
        {"scope_kind": "shared_workspace", "version": True},
        {"scope_kind": "shared_workspace", "version": 1.0},
        {"scope_kind": "shared_workspace", "version": 2},
        {"scope_kind": "cth", "version": 1},
        {"scope_kind": "shared_workspace", "version": 1, "actor_id": 901},
        {"scope_kind": "shared_workspace", "version": 1, "principal_ref": "client"},
        {"scope_kind": "shared_workspace", "version": 1, "workspace_id": "client"},
        {"scope_kind": "shared_workspace", "version": 1, "domain_id": "cth"},
    ],
    ids=[
        "none", "list", "empty", "missing-version", "bool-version", "float-version",
        "wrong-version", "wrong-kind", "client-actor", "client-principal",
        "client-workspace", "client-domain",
    ],
)
def test_shared_saved_work_rejects_noncanonical_scope_before_artifacts(workspace, bad_scope):
    actor = _new_account(workspace)
    before = _artifact_state(workspace)
    with pytest.raises(ValueError):
        workspace["work"].authorize_saved_work(actor, bad_scope, action="write")

    args = _save_args(workspace, actor=actor)
    args["scope"] = bad_scope
    with pytest.raises(ValueError):
        workspace["work"].work_write_resource(**args)
    assert _artifact_state(workspace) == before


def test_legacy_saved_work_and_document_importers_reject_shared_scope(workspace):
    actor = _new_account(workspace)
    before = _artifact_state(workspace)
    with pytest.raises(ValueError, match="WORK_SCOPE_INVALID"):
        workspace["work"].authorize(actor, SHARED_SCOPE, action="read")

    path = workspace["root"] / "must-not-import.txt"
    path.write_text("This content must not enter a Shared import path.", encoding="utf-8")
    provenance = {
        "source_url": "https://example.invalid/legacy-shared-import-check",
        "published_at": None,
        "captured_at": "2026-10-11T00:00:00Z",
        "rights": "public-reference-only",
        "title": "Synthetic legacy importer boundary",
    }
    with pytest.raises(ValueError, match="WORK_SCOPE_INVALID"):
        workspace["documents"].document_import_resource(
            str(path), actor=actor, scope=SHARED_SCOPE, provenance=provenance,
        )
    with pytest.raises(ValueError, match="WORK_SCOPE_INVALID"):
        workspace["documents"].import_document(
            str(path), actor=actor, scope=SHARED_SCOPE, provenance=provenance,
        )
    assert _artifact_state(workspace) == before


@pytest.mark.parametrize("actions", [("read",), ("write",)], ids=["read-only", "write-only"])
def test_shared_saved_work_requires_both_read_and_write_account_grants(workspace, actions):
    actor = _new_account(workspace, actions=actions)
    args = _save_args(workspace, actor=actor)
    before = _artifact_state(workspace)

    if "write" in actions:
        _grant_exact_save_export(workspace, args)
    with pytest.raises(ValueError):
        workspace["work"].save_work(**args)
    assert _artifact_state(workspace) == before


def test_shared_saved_work_requires_exact_gate_approval_and_retries_after_rollback(workspace):
    actor = _new_account(workspace)
    args = _save_args(workspace, actor=actor)
    # The canonical DATA_EXPORT spec is auto-allowed at MEDIUM risk. Configure
    # this real Gate instance with an explicit approval requirement so the
    # wrong-resource versus exact-resource behavior below is exercised through
    # PermissionGate.check rather than bypassed by the default policy.
    workspace["gate"].register_operation(
        permission.OperationType.DATA_EXPORT,
        permission.OperationSpec(
            permission.PermissionLevel.VERIFY,
            permission.RiskLevel.HIGH,
            "Isolated exact-resource SavedWork approval regression",
            auto_allow=False,
        ),
    )
    export_spec = workspace["gate"].get_operation_spec(
        permission.OperationType.DATA_EXPORT,
    )
    assert export_spec.risk is permission.RiskLevel.HIGH and export_spec.auto_allow is False
    resource = workspace["work"].work_write_resource(**args)
    workspace["gate"].grant_approval(
        _EXPORT_OPERATION,
        resource="saved-work:v1:unrelated-resource",
        granted_by="isolated-unrelated-approval",
    )
    before = _artifact_state(workspace)
    with pytest.raises(ValueError, match="WORK_EFFECT_PERMISSION_DENIED"):
        workspace["work"].save_work(**args)
    assert _artifact_state(workspace) == before

    workspace["gate"].grant_approval(
        _EXPORT_OPERATION,
        resource=resource,
        granted_by="isolated-exact-approval",
    )

    def interrupt_before_commit(*_args):
        raise RuntimeError("isolated precommit interruption")

    with pytest.raises(RuntimeError, match="isolated precommit interruption"):
        workspace["work"].save_work(**args, before_commit=interrupt_before_commit)
    assert _artifact_state(workspace) == before

    retried = workspace["work"].save_work(**args)
    assert retried["replayed"] is False
    assert retried["document"]["version"] == 1
    assert len(_artifact_state(workspace)[0]) == len(before[0]) + 1


def test_shared_saved_work_rechecks_final_write_grant_and_rolls_back(workspace):
    actor = _new_account(workspace)
    args = _save_args(workspace, actor=actor)
    _grant_exact_save_export(workspace, args)
    before = _artifact_state(workspace)

    def revoke_write_in_transaction(con, _record, _document):
        con.execute(
            "DELETE FROM v4_member_authorizations "
            "WHERE user_id=? AND product_id='SHARED' AND module='reports' AND action='write'",
            (actor["id"],),
        )

    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].save_work(**args, before_commit=revoke_write_in_transaction)
    assert _artifact_state(workspace) == before
    assert workspace["repo"].query_one(
        "SELECT COUNT(*) FROM v4_member_authorizations "
        "WHERE user_id=? AND product_id='SHARED' AND module='reports' AND action='write'",
        (actor["id"],),
    )[0] == 1

    retry = workspace["work"].save_work(**args)
    assert retry["document"]["version"] == 1


def test_shared_saved_work_rechecks_read_permission_after_commit(workspace):
    actor = _new_account(workspace)
    args = _save_args(workspace, actor=actor)
    _grant_exact_save_export(workspace, args)
    saved = workspace["work"].save_work(**args)
    before = _artifact_state(workspace)

    workspace["repo"].execute(
        "DELETE FROM v4_member_authorizations "
        "WHERE user_id=? AND product_id='SHARED' AND module='reports' AND action='read'",
        (actor["id"],),
    )
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(
            saved["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
        )
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].search_work(actor=actor, scope=SHARED_SCOPE)
    assert _artifact_state(workspace) == before

    workspace["repo"].execute(
        "INSERT INTO v4_member_authorizations "
        "(user_id,product_id,module,action,data_scope,time_range,granted_by) "
        "VALUES (?,?,?,?,?,?,?)",
        (actor["id"], "SHARED", "reports", "read", None, None, None),
    )
    reopened = workspace["work"].reopen_work(
        saved["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
    )
    assert reopened["document"] == saved["document"]


def test_shared_container_permission_does_not_authorize_original_source_scope(workspace):
    source_scope = {
        "domain_id": "cth", "subject_id": None,
        "brand_focus_id": None, "global_scope": True,
    }
    shared_actor = _new_account(workspace)
    _grant_global_account_scope(workspace, shared_actor, product_id="CTH")
    _path, imported, located = _source(
        workspace, scope=source_scope, actor=shared_actor,
    )
    assert workspace["work"].authorize(
        shared_actor, source_scope, action="read",
    )["actor_id"] == shared_actor["id"]
    shared_context = workspace["work"].authorize_saved_work(
        shared_actor, SHARED_SCOPE, action="write",
    )
    assert shared_context["actor_id"] == shared_actor["id"]
    assert located["original_scope"] == source_scope

    args = _save_args(
        workspace, actor=shared_actor, citations=[located["citation"]],
        payload={"source_evidence_id": imported["evidence_id"]},
    )
    _grant_exact_save_export(workspace, args)
    before = _artifact_state(workspace)

    workspace["repo"].execute(
        "DELETE FROM v4_member_authorizations "
        "WHERE user_id=? AND product_id='CTH' AND module='reports' AND action='read'",
        (shared_actor["id"],),
    )
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].authorize(shared_actor, source_scope, action="read")
    assert workspace["work"].authorize(
        shared_actor, source_scope, action="write",
    )["actor_id"] == shared_actor["id"]
    assert workspace["work"].authorize_saved_work(
        shared_actor, SHARED_SCOPE, action="read",
    )["actor_id"] == shared_actor["id"]
    assert workspace["work"].authorize_saved_work(
        shared_actor, SHARED_SCOPE, action="write",
    )["actor_id"] == shared_actor["id"]
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].save_work(**args)
    assert _artifact_state(workspace) == before


def test_shared_container_rejects_foreign_original_source_owner(workspace):
    source_scope = {
        "domain_id": "cth", "subject_id": None,
        "brand_focus_id": None, "global_scope": True,
    }
    _path, imported, located = _source(workspace, scope=source_scope)
    shared_actor = _new_account(workspace)
    _grant_global_account_scope(workspace, shared_actor, product_id="CTH")
    assert workspace["work"].authorize_saved_work(
        shared_actor, SHARED_SCOPE, action="write",
    )["actor_id"] == shared_actor["id"]
    args = _save_args(
        workspace,
        actor=shared_actor,
        citations=[located["citation"]],
        payload={"source_evidence_id": imported["evidence_id"]},
    )
    _grant_exact_save_export(workspace, args)
    before = _artifact_state(workspace)
    with pytest.raises(ValueError, match="WORK_EVIDENCE_CONTEXT_MISMATCH"):
        workspace["work"].save_work(**args)
    assert _artifact_state(workspace) == before


def test_shared_saved_work_retains_original_scope_and_withdrawal_allows_source_free_revision(workspace):
    source_scope = _focused_cth_scope(workspace)
    _path, imported, located = _source(workspace, scope=source_scope)
    actor = workspace["local_actor"]
    assert workspace["work"].authorize_saved_work(
        actor, SHARED_SCOPE, action="write",
    ) == {
        "scope_kind": "shared_workspace", "version": 1,
        "principal_kind": actor.principal_kind, "principal_ref": actor.principal_ref,
    }
    args = _save_args(
        workspace, actor=actor, citations=[located["citation"]],
        payload={"source_evidence_id": imported["evidence_id"]},
    )
    _grant_exact_save_export(workspace, args)
    first = workspace["work"].save_work(**args)
    assert first["document"]["citations"]
    assert first["receipt"]["sources"][0]["scope"] == source_scope
    assert first["return_ref"]["scope"] == SHARED_SCOPE

    workspace["repo"].execute(
        "DELETE FROM lumenis_subject_registry WHERE subject_id=?",
        (source_scope["subject_id"],),
    )
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(
            first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
        )

    revised_args = _save_args(
        workspace,
        actor=actor,
        request_id="source-withdrawal-revision-0002",
        title=first["document"]["title"],
        markdown="The user revised the note and removed its withdrawn source citation.",
        payload={"source_removed_after_withdrawal": True},
        citations=[],
        work_id=first["document"]["work_id"],
        previous_artifact_id=first["artifact"]["artifact_id"],
        expected_draft_digest=first["document"]["draft_digest"],
        contract=first["document"]["contract"],
    )
    _grant_exact_save_export(workspace, revised_args)
    revised = workspace["work"].save_work(**revised_args)
    assert revised["document"]["version"] == 2
    assert revised["document"]["parent_artifact_id"] == first["artifact"]["artifact_id"]
    assert revised["document"]["citations"] == []
    assert revised["receipt"]["sources"] == []
    assert revised["return_ref"]["scope"] == SHARED_SCOPE
    assert workspace["work"].reopen_work(
        revised["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
    )["document"] == revised["document"]
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].reopen_work(
            first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
        )


def test_shared_saved_work_roundtrips_revises_and_recovers_real_ooxml_attachments(workspace):
    actor = _new_account(workspace)
    workspace["actor_for_revision"] = actor
    attachments = [_actual_ooxml_attachment(kind) for kind in ("docx", "xlsx", "pptx")]
    args = _save_args(workspace, actor=actor, attachments=attachments)
    _grant_exact_save_export(workspace, args)
    first = workspace["work"].save_work(**args)

    bindings = first["document"]["attachments"]
    assert [binding["name"] for binding in bindings] == [item["name"] for item in attachments]
    assert _attachment_bytes(workspace, first, actor=actor) == [item["content"] for item in attachments]
    with workspace["repo"].connection(read_only=True, wal=False) as con:
        rows = con.execute(
            "SELECT artifact_id,kind,product_id,domain_id,sha256,size_bytes FROM lumenis_shared_artifacts "
            "WHERE source_execution_id=? ORDER BY kind,artifact_id",
            (first["document"]["work_id"],),
        ).fetchall()
    artifact_rows = {row["artifact_id"]: dict(row) for row in rows}
    assert set(artifact_rows) == {
        first["artifact"]["artifact_id"], *(binding["artifact_id"] for binding in bindings),
    }
    assert all(
        (row["product_id"], row["domain_id"]) == ("SHARED", "content")
        for row in artifact_rows.values()
    )
    assert artifact_rows[first["artifact"]["artifact_id"]]["kind"] == "saved_document"
    for binding, original in zip(bindings, attachments):
        row = artifact_rows[binding["artifact_id"]]
        assert row["kind"] == "saved_document_attachment"
        assert row["sha256"] == hashlib.sha256(original["content"]).hexdigest()
        assert row["size_bytes"] == len(original["content"])

    tampered_binding = bindings[0]
    artifact_row = workspace["repo"].query_one(
        "SELECT relative_path FROM lumenis_shared_artifacts WHERE artifact_id=?",
        (tampered_binding["artifact_id"],),
    )
    attachment_path = workspace["artifact_dir"] / artifact_row["relative_path"]
    original_bytes = attachment_path.read_bytes()
    artifact_before_tamper = _artifact_state(workspace)
    attachment_path.write_bytes(b"tampered test attachment bytes")
    with pytest.raises(ValueError, match="WORK_ATTACHMENT_INTEGRITY_FAILED"):
        workspace["work"].reopen_work(
            first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
        )
    revision_markdown = "The user revised this document while preserving its office files."
    reason = "Retain the submitted editable materials during a user revision."
    with pytest.raises(ValueError, match="WORK_ATTACHMENT_INTEGRITY_FAILED"):
        workspace["work"].revise_work(
            first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
            request_id="shared-attachment-revision-0002",
            expected_sha256=first["artifact"]["sha256"], reason=reason,
            edits=[{
                "pointer": "/markdown", "expected": first["document"]["markdown"],
                "value": revision_markdown,
            }],
        )
    assert _artifact_state(workspace) == artifact_before_tamper

    attachment_path.write_bytes(original_bytes)
    recovered = workspace["work"].reopen_work(
        first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
    )
    assert recovered["document"]["attachments"] == bindings
    assert _attachment_bytes(workspace, recovered, actor=actor) == [item["content"] for item in attachments]

    revised_args = _revision_save_args(
        workspace, first, request_id="shared-attachment-revision-0002",
        markdown=revision_markdown, reason=reason, attachments=attachments,
    )
    _grant_exact_save_export(workspace, revised_args)
    revised = workspace["work"].revise_work(
        first["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
        request_id="shared-attachment-revision-0002",
        expected_sha256=first["artifact"]["sha256"], reason=reason,
        edits=[{
            "pointer": "/markdown", "expected": first["document"]["markdown"],
            "value": revision_markdown,
        }],
    )
    assert revised["document"]["version"] == 2
    assert [item["name"] for item in revised["document"]["attachments"]] == [
        item["name"] for item in attachments
    ]
    assert [item["sha256"] for item in revised["document"]["attachments"]] == [
        item["sha256"] for item in bindings
    ]
    assert [item["artifact_id"] for item in revised["document"]["attachments"]] != [
        item["artifact_id"] for item in bindings
    ]
    assert _attachment_bytes(workspace, revised, actor=actor) == [item["content"] for item in attachments]


def test_shared_saved_work_rejects_legacy_binary_attachment_bytes_before_staging(workspace):
    actor = _new_account(workspace)
    legacy_binary = bytes.fromhex("d0cf11e0a1b11ae1") + b"isolated legacy binary fixture"
    args = _save_args(workspace, actor=actor, attachments=[{
        "name": "legacy-office.docx",
        "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "content": legacy_binary,
    }])
    before = _artifact_state(workspace)
    with pytest.raises(ValueError, match="WORK_ATTACHMENT_INVALID"):
        workspace["work"].work_write_resource(**args)
    assert _artifact_state(workspace) == before


def test_shared_saved_work_exports_pptx_with_complete_readback_and_scope_ref(workspace):
    from pptx import Presentation

    actor = _new_account(workspace)
    args = _save_args(
        workspace, actor=actor,
        request_id="shared-pptx-export-save-0001",
        title="Shared PPTX export roundtrip",
        markdown="A complete saved result is serialized into editable slides.",
    )
    _grant_exact_save_export(workspace, args)
    saved = workspace["work"].save_work(**args)
    _grant_exact_work_export(workspace, saved, format="pptx")

    exported = workspace["work"].export_work(
        saved["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
        format="pptx",
    )
    presentation = Presentation(io.BytesIO(exported["content"]))
    sections = {}
    for slide in presentation.slides:
        heading = slide.shapes[0].text_frame.text
        label = heading.removeprefix("Saved document | ").split(" (continued ", 1)[0]
        body = slide.shapes[1].text_frame
        sections.setdefault(label, []).extend(paragraph.text for paragraph in body.paragraphs)

    assert "Title" in sections and "Report" in sections and "Lineage" in sections
    assert "Sources" in sections and "Snapshot" in sections and "JSON" in sections
    assert "".join(sections["Title"]) == saved["document"]["title"]
    assert "".join(sections["Report"]) == saved["document"]["markdown"]
    assert "".join(sections["Lineage"]) == "Sources / original scope / rights"
    assert "".join(sections["Sources"]) == "[]"
    assert "".join(sections["Snapshot"]) == "Complete saved snapshot"
    serialized_document = json.dumps(
        saved["document"], ensure_ascii=False, sort_keys=True, indent=2,
    )
    expected_json_segments = []
    for line in serialized_document.splitlines():
        if line:
            expected_json_segments.extend(
                line[start:start + 64] for start in range(0, len(line), 64)
            )
        else:
            expected_json_segments.append("")
    assert sections["JSON"] == expected_json_segments

    assert exported["media_type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert exported["filename"] == saved["artifact"]["artifact_id"] + ".pptx"
    assert exported["sha256"] == hashlib.sha256(exported["content"]).hexdigest()
    assert exported["return_ref"]["scope"] == SHARED_SCOPE


def test_shared_pptx_export_rechecks_account_read_grant_after_conversion(
    workspace, monkeypatch,
):
    from pptx import Presentation
    from lumenis.shared.reporting.word_pdf_export import ReportExporter

    actor = _new_account(workspace)
    args = _save_args(
        workspace, actor=actor,
        request_id="shared-pptx-revoked-read-save-0001",
        title="Revoked account read after conversion",
        markdown="The final account check must deny this converted slide deck.",
    )
    _grant_exact_save_export(workspace, args)
    saved = workspace["work"].save_work(**args)
    _grant_exact_work_export(workspace, saved, format="pptx")
    before = _artifact_state(workspace)

    real_export = ReportExporter.export_saved_document
    converted = []

    def convert_then_revoke(document, *, sources, format):
        pptx_bytes = real_export(document, sources=sources, format=format)
        converted.append(pptx_bytes)
        with workspace["repo"].transaction() as con:
            deleted = con.execute(
                "DELETE FROM v4_member_authorizations "
                "WHERE user_id=? AND product_id='SHARED' AND module='reports' AND action='read'",
                (actor["id"],),
            )
            assert deleted.rowcount == 1
        return pptx_bytes

    monkeypatch.setattr(
        ReportExporter, "export_saved_document", staticmethod(convert_then_revoke),
    )
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].export_work(
            saved["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
            format="pptx",
        )

    assert len(converted) == 1
    presentation = Presentation(io.BytesIO(converted[0]))
    extracted = "\n".join(
        shape.text_frame.text
        for slide in presentation.slides for shape in slide.shapes
        if shape.has_text_frame
    )
    assert saved["document"]["title"] in extracted
    assert saved["document"]["markdown"] in extracted
    assert workspace["repo"].query_one(
        "SELECT COUNT(*) FROM v4_member_authorizations "
        "WHERE user_id=? AND product_id='SHARED' AND module='reports' AND action='read'",
        (actor["id"],),
    )[0] == 0
    assert _artifact_state(workspace) == before


def test_shared_saved_work_source_use_fixture_denies_export_operation(workspace, monkeypatch):
    from lumenis.shared.evidence import store

    actor = _new_account(workspace)
    source = _seed_approved_source_use_fixture(workspace, monkeypatch, actor)
    entry = source["entry"]
    receipt = source["receipt"]
    assert entry["allowed_purposes"] == ["internal_research", "short_quote", "cth_brief"]
    assert entry["allowed_operations"] == ["derive_reference", "read", "save", "revise", "reopen"]
    assert "export" not in entry["allowed_operations"]
    assert entry["rights_state"] == "UNKNOWN"
    assert receipt["quote_limit"]["max_unique_unicode_codepoints_per_source_per_saved_work"] == 160

    expected_signed_scope = {
        "product_id": "CTH", "domain_id": "cth",
        "subject_id": source["scope"]["subject_id"], "global": False,
    }
    for evidence_id in (source["scope_evidence_id"], source["content_evidence_id"]):
        evidence = dict(workspace["repo"].query_one(
            "SELECT * FROM app_evidence_records WHERE id=?", (evidence_id,),
        ))
        assert store.verify_evidence_record_integrity(evidence)
        assert json.loads(evidence["inputs_json"])["_lumenis_scope"] == expected_signed_scope

    request_id = "shared-source-use-export-derive-0001"
    _pending, confirmed, derived = _confirm_and_create_source_reference(
        workspace, source, actor, request_id,
    )
    assert confirmed["allowed"] is True
    assert derived["rights_state"] == "UNKNOWN"
    quote = "internal-use citation"
    start = source["body"].index(quote)
    located = workspace["documents"].locate_quote(
        derived["evidence_id"], actor=actor, scope=source["scope"],
        start=start, end=start + len(quote),
    )
    assert len(located["citation"]["quote"]) <= 160
    assert located["original_scope"] == source["scope"]

    args = _save_args(
        workspace, actor=actor, scope=SHARED_SCOPE,
        request_id="shared-source-use-export-save-0001",
        title="Shared note with restricted internal source use",
        markdown=f"The internal citation remains restricted: {quote}",
        payload={"source_id": source["source_id"]},
        citations=[located["citation"]],
    )
    _grant_exact_save_export(workspace, args)
    saved = workspace["work"].save_work(**args)
    assert saved["receipt"]["sources"][0]["scope"] == source["scope"]
    assert saved["receipt"]["sources"][0]["rights"] == "UNKNOWN"
    assert saved["return_ref"]["scope"] == SHARED_SCOPE
    before_export = _artifact_state(workspace)

    with pytest.raises(ValueError, match="SOURCE_USE_PURPOSE_NOT_ALLOWED"):
        workspace["work"].export_work(
            saved["artifact"]["artifact_id"], actor=actor, scope=SHARED_SCOPE,
            format="pptx",
        )
    assert _artifact_state(workspace) == before_export


def _postcommit_saved_document_record(workspace, artifact_id):
    """Prove a separate connection can see the Artifact before revoking access."""
    row = workspace["repo"].query_one(
        "SELECT artifact_id,kind,source_execution_id,sha256,relative_path,status "
        "FROM lumenis_shared_artifacts WHERE artifact_id=? AND kind='saved_document'",
        (artifact_id,),
    )
    assert row is not None
    record = tuple(row)
    raw = (workspace["artifact_dir"] / record[4]).read_bytes()
    file_sha256 = hashlib.sha256(raw).hexdigest()
    assert file_sha256 == record[3]
    return record, file_sha256


def _assert_postcommit_artifact_unchanged(workspace, state, record, file_sha256):
    matching = [row for row in state[0] if row[0] == record[0]]
    assert len(matching) == 1
    assert matching[0][1] == "saved_document"
    assert matching[0][4] == record[2]
    assert matching[0][7] == record[4]
    assert matching[0][8] == record[3]
    artifact_path = workspace["artifact_dir"] / record[4]
    assert hashlib.sha256(artifact_path.read_bytes()).hexdigest() == file_sha256 == record[3]


def _assert_committed_readback_error(error, *, request_id, replayed, result):
    assert isinstance(error, ValueError)
    assert str(error) == "WORK_COMMITTED_REOPEN_UNAVAILABLE"
    status = error.commit_status
    assert set(status) == {
        "schema", "state", "request_id", "replayed", "readback_error_code", "result",
    }
    assert status["schema"] == "lumenis.saved-work-commit-status/v1"
    assert status["state"] == "committed_readback_unavailable"
    assert status["request_id"] == request_id
    assert type(status["replayed"]) is bool and status["replayed"] is replayed
    assert status["readback_error_code"] == "WORK_PERMISSION_DENIED"
    assert status["result"] == result
    assert not {
        "artifact_id", "document", "title", "markdown", "payload", "citations",
        "receipt", "content", "raw_content", "raw_bytes",
    }.intersection(status)
    return status


def test_shared_saved_work_reports_committed_source_readback_withdrawal_and_recovers(
    workspace, monkeypatch,
):
    source_scope = {
        "domain_id": "cth", "subject_id": None,
        "brand_focus_id": None, "global_scope": True,
    }
    actor = _new_account(workspace)
    _grant_global_account_scope(workspace, actor, product_id="CTH")
    _path, imported, located = _source(
        workspace, scope=source_scope, actor=actor,
    )
    args = _save_args(
        workspace,
        actor=actor,
        request_id="shared-postcommit-source-read-0001",
        citations=[located["citation"]],
        payload={"source_evidence_id": imported["evidence_id"]},
    )
    _grant_exact_save_export(workspace, args)
    before = _artifact_state(workspace)
    real_reopen = workspace["work"].reopen_work
    reopen_calls = []
    committed = {}

    def revoke_source_read_after_commit(artifact_id, *, actor, scope, connection=None):
        reopen_calls.append(artifact_id)
        if len(reopen_calls) == 1:
            # The independent SELECT proves COMMIT completed before this revoke.
            record, file_sha256 = _postcommit_saved_document_record(workspace, artifact_id)
            committed.update(record=record, file_sha256=file_sha256)
            with workspace["repo"].transaction() as con:
                deleted = con.execute(
                    "DELETE FROM v4_member_authorizations "
                    "WHERE user_id=? AND product_id='CTH' AND module='reports' AND action='read'",
                    (actor["id"],),
                )
                assert deleted.rowcount == 1
        return real_reopen(
            artifact_id, actor=actor, scope=scope, connection=connection,
        )

    monkeypatch.setattr(workspace["work"], "reopen_work", revoke_source_read_after_commit)
    with pytest.raises(workspace["work"].SavedWorkCommittedReadbackError) as first_error:
        workspace["work"].save_work(**args)

    record = committed["record"]
    expected_draft_digest = hashlib.sha256(json.dumps(
        args["payload"], ensure_ascii=False, allow_nan=False,
        sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    expected_result = {
        "artifact_id": record[0],
        "sha256": record[3],
        "work_id": record[2],
        "version": 1,
        "draft_digest": expected_draft_digest,
        "scope": SHARED_SCOPE,
    }
    _assert_committed_readback_error(
        first_error.value, request_id=args["request_id"], replayed=False,
        result=expected_result,
    )
    assert workspace["work"].authorize(
        actor, source_scope, action="write",
    )["actor_id"] == actor["id"]
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].authorize(actor, source_scope, action="read")
    assert workspace["work"].authorize_saved_work(
        actor, SHARED_SCOPE, action="read",
    )["actor_id"] == actor["id"]

    committed_state = _artifact_state(workspace)
    assert len(committed_state[0]) == len(before[0]) + 1
    assert len(committed_state[1]) == len(before[1]) + 1
    _assert_postcommit_artifact_unchanged(
        workspace, committed_state, record, committed["file_sha256"],
    )

    with pytest.raises(workspace["work"].SavedWorkCommittedReadbackError) as denied_replay:
        workspace["work"].save_work(**args)
    denied_status = _assert_committed_readback_error(
        denied_replay.value, request_id=args["request_id"], replayed=True,
        result=expected_result,
    )
    assert denied_status["result"] == expected_result
    assert len(reopen_calls) == 2 and reopen_calls == [record[0], record[0]]
    assert _artifact_state(workspace) == committed_state
    _assert_postcommit_artifact_unchanged(
        workspace, committed_state, record, committed["file_sha256"],
    )

    global_scope = json.dumps(
        {"global": True}, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    )
    with workspace["repo"].transaction() as con:
        con.execute(
            "INSERT INTO v4_member_authorizations "
            "(user_id,product_id,module,action,data_scope,time_range,granted_by) "
            "VALUES (?,?,?,?,?,?,?)",
            (actor["id"], "CTH", "reports", "read", global_scope, None, None),
        )
    replayed = workspace["work"].save_work(**args)
    assert replayed["replayed"] is True
    assert replayed["artifact"]["artifact_id"] == record[0]
    assert replayed["artifact"]["sha256"] == record[3]
    assert replayed["document"]["citations"][0]["quote"] == located["citation"]["quote"]
    assert replayed["receipt"]["sources"][0]["scope"] == source_scope
    assert len(reopen_calls) == 3 and reopen_calls == [record[0], record[0], record[0]]
    assert _artifact_state(workspace) == committed_state
    _assert_postcommit_artifact_unchanged(
        workspace, committed_state, record, committed["file_sha256"],
    )


def test_shared_saved_work_withholds_committed_pointer_when_container_read_is_revoked(
    workspace, monkeypatch,
):
    source_scope = {
        "domain_id": "cth", "subject_id": None,
        "brand_focus_id": None, "global_scope": True,
    }
    actor = _new_account(workspace)
    _grant_global_account_scope(workspace, actor, product_id="CTH")
    _path, imported, located = _source(
        workspace, scope=source_scope, actor=actor,
    )
    args = _save_args(
        workspace,
        actor=actor,
        request_id="shared-postcommit-container-read-0001",
        citations=[located["citation"]],
        payload={"source_evidence_id": imported["evidence_id"]},
    )
    _grant_exact_save_export(workspace, args)
    before = _artifact_state(workspace)
    real_reopen = workspace["work"].reopen_work
    reopen_calls = []
    committed = {}

    def revoke_container_read_after_commit(artifact_id, *, actor, scope, connection=None):
        reopen_calls.append(artifact_id)
        if len(reopen_calls) == 1:
            record, file_sha256 = _postcommit_saved_document_record(workspace, artifact_id)
            committed.update(record=record, file_sha256=file_sha256)
            with workspace["repo"].transaction() as con:
                deleted = con.execute(
                    "DELETE FROM v4_member_authorizations "
                    "WHERE user_id=? AND product_id='SHARED' AND module='reports' AND action='read'",
                    (actor["id"],),
                )
                assert deleted.rowcount == 1
            # The original source grant is independent of the Shared container grant.
            assert workspace["work"].authorize(
                actor, source_scope, action="read",
            )["actor_id"] == actor["id"]
        return real_reopen(
            artifact_id, actor=actor, scope=scope, connection=connection,
        )

    monkeypatch.setattr(workspace["work"], "reopen_work", revoke_container_read_after_commit)
    with pytest.raises(workspace["work"].SavedWorkCommittedReadbackError) as first_error:
        workspace["work"].save_work(**args)

    _assert_committed_readback_error(
        first_error.value, request_id=args["request_id"], replayed=False, result=None,
    )
    assert workspace["work"].authorize(
        actor, source_scope, action="read",
    )["actor_id"] == actor["id"]
    assert workspace["work"].authorize_saved_work(
        actor, SHARED_SCOPE, action="write",
    )["actor_id"] == actor["id"]
    with pytest.raises(ValueError, match="WORK_PERMISSION_DENIED"):
        workspace["work"].authorize_saved_work(
            actor, SHARED_SCOPE, action="read",
        )
    committed_state = _artifact_state(workspace)
    record = committed["record"]
    assert len(committed_state[0]) == len(before[0]) + 1
    assert len(committed_state[1]) == len(before[1]) + 1
    _assert_postcommit_artifact_unchanged(
        workspace, committed_state, record, committed["file_sha256"],
    )

    with pytest.raises(workspace["work"].SavedWorkCommittedReadbackError) as denied_replay:
        workspace["work"].save_work(**args)
    _assert_committed_readback_error(
        denied_replay.value, request_id=args["request_id"], replayed=True, result=None,
    )
    # Denied export/read is fenced before the real reopen, so no source/content reader runs.
    assert len(reopen_calls) == 1 and reopen_calls == [record[0]]
    assert _artifact_state(workspace) == committed_state
    _assert_postcommit_artifact_unchanged(
        workspace, committed_state, record, committed["file_sha256"],
    )

    with workspace["repo"].transaction() as con:
        con.execute(
            "INSERT INTO v4_member_authorizations "
            "(user_id,product_id,module,action,data_scope,time_range,granted_by) "
            "VALUES (?,?,?,?,?,?,?)",
            (actor["id"], "SHARED", "reports", "read", None, None, None),
        )
    replayed = workspace["work"].save_work(**args)
    assert replayed["replayed"] is True
    assert replayed["artifact"]["artifact_id"] == record[0]
    assert replayed["artifact"]["sha256"] == record[3]
    assert replayed["document"]["citations"][0]["quote"] == located["citation"]["quote"]
    assert replayed["receipt"]["sources"][0]["scope"] == source_scope
    assert len(reopen_calls) == 2 and reopen_calls == [record[0], record[0]]
    assert _artifact_state(workspace) == committed_state
    _assert_postcommit_artifact_unchanged(
        workspace, committed_state, record, committed["file_sha256"],
    )
