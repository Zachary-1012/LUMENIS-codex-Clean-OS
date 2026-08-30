# -*- coding: utf-8 -*-
"""读取三个基线文档的内容"""
import docx
import os

docs = [
    (r'D:\wps云空间\WPS Cloud\WPSDrive\280397037\WPS云盘\Z-Brand-Core_全局代码级基线_V1.0_豆包执行版.docx', 'V1.0'),
    (r'D:\wps云空间\WPS Cloud\WPSDrive\280397037\WPS云盘\Z-Brand-Core_全局代码级基线_V2.0_豆包执行版.docx', 'V2.0'),
    (r'D:\wps云空间\WPS Cloud\WPSDrive\280397037\WPS云盘\Z-Brand-Core_全局代码级基线_V3.0_豆包执行版.docx', 'V3.0'),
]

for path, version in docs:
    print(f'\n{"="*70}')
    print(f'  {version} 文档: {os.path.basename(path)}')
    print(f'{"="*70}')
    if not os.path.exists(path):
        print(f'  文件不存在: {path}')
        continue
    try:
        doc = docx.Document(path)
        print(f'  段落数: {len(doc.paragraphs)}')
        print(f'  表格数: {len(doc.tables)}')
        print()
        # 输出所有标题和前几行内容
        heading_count = 0
        content_count = 0
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ''
            # 标题
            if 'Heading' in style or '标题' in style:
                print(f'  [{style}] {text[:100]}')
                heading_count += 1
            elif content_count < 50:  # 只输出前50行正文
                print(f'    {text[:120]}')
                content_count += 1
        print(f'\n  标题数: {heading_count}, 输出正文行数: {content_count}')
        # 输出表格概要
        for ti, table in enumerate(doc.tables[:5]):
            print(f'\n  表格{ti+1}: {len(table.rows)}行 x {len(table.columns)}列')
            for ri, row in enumerate(table.rows[:3]):
                cells = [cell.text.strip()[:30] for cell in row.cells]
                print(f'    行{ri}: {cells}')
    except Exception as e:
        print(f'  读取失败: {e}')
